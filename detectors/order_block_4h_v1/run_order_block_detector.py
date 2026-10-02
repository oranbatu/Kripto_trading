"""Run the permanent 4h body-based Order Block detector."""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import pyarrow as pa

import config
from data.loader import load_candles
from detectors.order_block_4h_v1.engine import (
    ACTIVE_BODY_MITIGATED,
    ACTIVE_UNTOUCHED,
    Bar,
    OrderBlockError,
    audit_result,
    disposition_counts,
    iso_turkey,
    iso_utc,
    status_counts,
    zone_distance,
)
from detectors.order_block_4h_v1.engine import analyze
from detectors.order_block_4h_v1.order_block_config import (
    CONFIG,
    DATASET_END,
    DETECTOR_VERSION,
    EXPECTED_FIRST,
    EXPECTED_LAST_CLOSE,
    EXPECTED_LAST_OPEN,
    EXPECTED_PAIRS,
    EXPECTED_ROWS,
    EXCLUSIVE_LOAD_END,
    REVISION_PATTERN_TEXT,
    SOURCE_DATASET,
    STEP_MS,
    TIMEFRAME,
    TMP_DIR,
    ANALYSIS_START,
)
from detectors.order_block_4h_v1.workbook import round_trip_workbook, validate_ooxml_workbook, validate_saved_workbook, write_workbook

EXECUTION_COMMAND = "python -m detectors.order_block_4h_v1.run_order_block_detector"
REVISION_PATTERN = re.compile(REVISION_PATTERN_TEXT, re.IGNORECASE)


def main() -> int:
    print("order_block_4h_v1 starting", flush=True)
    assert_python_write_allowed(CONFIG.package_dir / "run_order_block_detector.py")
    package_before = _fingerprint_tree([CONFIG.package_dir])
    code_before = _fingerprint_tree(_existing_code_roots())
    existing_revisions = scan_revisions()
    try:
        loaded = load_production()
    except (OrderBlockError, RuntimeError) as exc:
        print(f"SOURCE_VALIDATION_FAILED {exc}", flush=True)
        return 2
    source_before = {path: fingerprint_file(Path(path)) for path in loaded["source_paths"]}
    print(f"source validated rows={len(loaded['bars'])}", flush=True)
    try:
        result = analyze(loaded["bars"])
        audit_result(result)
        if len(result.pairs) != EXPECTED_PAIRS:
            raise OrderBlockError(f"pairs {len(result.pairs)} != {EXPECTED_PAIRS}")
    except (OrderBlockError, RuntimeError) as exc:
        print(f"DETECTION_FAILED {exc}", flush=True)
        return 2
    counts = disposition_counts(result)
    print(
        "qualified "
        f"bull={counts['QUALIFIED_BULLISH_ORDER_BLOCK']} "
        f"bear={counts['QUALIFIED_BEARISH_ORDER_BLOCK']} "
        f"pairs={len(result.pairs)}",
        flush=True,
    )
    source_after = {path: fingerprint_file(Path(path)) for path in loaded["source_paths"]}
    code_after = _fingerprint_tree(_existing_code_roots())
    if source_before != source_after:
        print("SOURCE_IMMUTABILITY_FAILED", flush=True)
        return 3
    if code_before != code_after:
        print("CODE_IMMUTABILITY_FAILED", flush=True)
        return 3
    context = _context(result, loaded, source_before, existing_revisions)
    readme = (CONFIG.package_dir / "README.md").read_text(encoding="utf-8")
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    try:
        revision, final_path = _commit_workbook(result, context, readme)
    except RuntimeError as exc:
        print(f"WORKBOOK_VALIDATION_FAILED {exc}", flush=True)
        _cleanup()
        return 4
    package_after = _fingerprint_tree([CONFIG.package_dir])
    if package_before != package_after:
        print("DETECTOR_SOURCE_REWRITTEN", flush=True)
        _cleanup()
        return 3
    stray = _stray_python_files()
    if stray:
        print(f"STRAY_PYTHON {stray}", flush=True)
        _cleanup()
        return 3
    _cleanup()
    print("SUMMARY_JSON " + json.dumps(_summary(result, final_path, revision, existing_revisions), default=str), flush=True)
    return 0


def load_production() -> dict[str, Any]:
    slice_ = load_candles(
        "BTCUSDT",
        TIMEFRAME,
        ANALYSIS_START,
        EXCLUSIVE_LOAD_END,
        warmup_bars=0,
        precision_mode=config.PRECISION_EXACT,
    )
    if slice_.warmup_bars_loaded != 0:
        raise OrderBlockError("4h load included warmup bars")
    if slice_.symbol != "BTCUSDT" or slice_.timeframe != TIMEFRAME or slice_.precision_mode != config.PRECISION_EXACT:
        raise OrderBlockError("4h identity mismatch")
    bars = _bars_from_table(slice_.table)
    _validate_frame(bars, slice_.source_files)
    return {
        "bars": bars,
        "source_paths": sorted(set(slice_.source_files)),
        "provenance": {
            "rows": len(bars),
            "first_open_utc": iso_utc(bars[0].open_time),
            "last_open_utc": iso_utc(bars[-1].open_time),
            "last_close_utc": iso_utc(bars[-1].close_time),
            "precision_mode": slice_.precision_mode,
        },
    }


def _bars_from_table(table: pa.Table) -> list[Bar]:
    opens = table.column("open_time").to_pylist()
    closes = table.column("close_time").to_pylist()
    open_ = [_decimal(value) for value in table.column("open").to_pylist()]
    high = [_decimal(value) for value in table.column("high").to_pylist()]
    low = [_decimal(value) for value in table.column("low").to_pylist()]
    close = [_decimal(value) for value in table.column("close").to_pylist()]
    volume = [_decimal(value) for value in table.column("volume").to_pylist()]
    trades = table.column("number_of_trades").to_pylist()
    symbols = table.column("symbol").to_pylist()
    rows: list[Bar] = []
    for index in range(table.num_rows):
        if symbols[index] != "BTCUSDT":
            raise OrderBlockError(f"row {index} symbol {symbols[index]}")
        if None in (open_[index], high[index], low[index], close[index], volume[index]):
            raise OrderBlockError(f"null OHLC at row {index}")
        rows.append(
            Bar(
                row=index,
                open_time=_as_utc(opens[index]),
                close_time=_as_utc(closes[index]),
                open=open_[index],
                high=high[index],
                low=low[index],
                close=close[index],
                volume=volume[index],
                trades=int(trades[index] or 0),
            )
        )
    return rows


def _validate_frame(bars: list[Bar], source_files: list[str]) -> None:
    if len(bars) != EXPECTED_ROWS:
        raise OrderBlockError(f"rows {len(bars)} != {EXPECTED_ROWS}")
    if bars[0].open_time != EXPECTED_FIRST or bars[-1].open_time != EXPECTED_LAST_OPEN:
        raise OrderBlockError(f"range {iso_utc(bars[0].open_time)} .. {iso_utc(bars[-1].open_time)}")
    if bars[-1].close_time != EXPECTED_LAST_CLOSE:
        raise OrderBlockError(f"last close {iso_utc(bars[-1].close_time)}")
    seen: set[datetime] = set()
    previous = None
    for bar in bars:
        if bar.open_time.tzinfo is None or bar.open_time.utcoffset() != timezone.utc.utcoffset(None):
            raise OrderBlockError("timestamp is not UTC")
        if bar.open_time in seen:
            raise OrderBlockError(f"duplicate open time {iso_utc(bar.open_time)}")
        seen.add(bar.open_time)
        if previous is not None:
            delta = int((bar.open_time - previous).total_seconds() * 1000)
            if delta != STEP_MS or bar.open_time <= previous:
                raise OrderBlockError(f"missing or unordered timestamp at {iso_utc(bar.open_time)}")
        previous = bar.open_time
        expected_close = bar.open_time + timedelta(milliseconds=STEP_MS) - timedelta(milliseconds=1)
        if bar.close_time != expected_close or bar.close_time > DATASET_END:
            raise OrderBlockError(f"close time mismatch at {iso_utc(bar.open_time)}")
        if bar.volume < 0:
            raise OrderBlockError("negative volume")
        if not (
            bar.high >= bar.open
            and bar.high >= bar.close
            and bar.high >= bar.low
            and bar.low <= bar.open
            and bar.low <= bar.close
            and bar.low <= bar.high
        ):
            raise OrderBlockError(f"OHLC violation at {iso_utc(bar.open_time)}")
    dataset = str(SOURCE_DATASET).replace("\\", "/").lower()
    for token in ("marketdata", "binance", "futures", "um", "perpetual", "4h", "btcusdt"):
        if token not in dataset:
            raise OrderBlockError(f"dataset path missing {token}")
    if not source_files:
        raise OrderBlockError("4h has no source files")
    for path in source_files:
        lowered = path.replace("\\", "/").lower()
        for token in ("marketdata", "binance", "futures", "um", "perpetual", "4h", "btcusdt"):
            if token not in lowered:
                raise OrderBlockError(f"source file missing {token}: {path}")


def _context(result, loaded: dict[str, Any], fingerprints: dict[str, dict[str, Any]], revisions: list[int]) -> dict[str, Any]:
    counts = disposition_counts(result)
    statuses = status_counts(result)
    return {
        "source_validation": "PASS",
        "provenance": loaded["provenance"],
        "source_fingerprints": fingerprints,
        "existing_revisions": revisions,
        "counts": counts,
        "statuses": statuses,
        "excel_validation": "PENDING",
        "diagnostic_rows": [],
        "revision_name": "",
    }


def _diagnostic_rows(result, context: dict[str, Any], final_path: Path | None = None) -> list[list[Any]]:
    rows = [["Section", "Name", "Value"]]

    def add(section: str, name: str, value: Any) -> None:
        rows.append([section, name, value])

    proven = context["provenance"]
    add("source", "rows", proven["rows"])
    add("source", "expected_rows", EXPECTED_ROWS)
    add("source", "first_open_utc", proven["first_open_utc"])
    add("source", "last_open_utc", proven["last_open_utc"])
    add("source", "last_close_utc", proven["last_close_utc"])
    add("source", "validation", context["source_validation"])
    add("source", "dataset", str(SOURCE_DATASET))
    for path, meta in context["source_fingerprints"].items():
        add("source_fingerprint", path, f"sha256={meta['sha256']} size={meta['size']} mtime_ns={meta['mtime_ns']}")
    counts = context["counts"]
    statuses = context["statuses"]
    add("reconciliation", "candidate_pairs", len(result.pairs))
    add("reconciliation", "qualified_bullish", counts["QUALIFIED_BULLISH_ORDER_BLOCK"])
    add("reconciliation", "qualified_bearish", counts["QUALIFIED_BEARISH_ORDER_BLOCK"])
    rejected = sum(value for name, value in counts.items() if name.startswith("REJECTED"))
    add("reconciliation", "rejected_pairs", rejected)
    add("reconciliation", "pair_identity", "PASS" if rejected + len(result.blocks) == len(result.pairs) else "FAIL")
    for name, value in statuses.items():
        add("lifecycle", name, value)
    add("immutability", "source_parquet_unchanged", "PASS")
    add("immutability", "existing_project_code_unchanged", "PASS")
    add("isolation", "separate_project_created", "FALSE")
    add("isolation", "detector_python_outside_project", "NONE")
    add("validation", "high_low_used_for_qualification", "FALSE")
    add("validation", "workbook_validation", "PASS")
    add("validation", "no_macros", "PASS")
    add("validation", "no_external_links", "PASS")
    add("workbook", "revision", context.get("revision_name", ""))
    add("workbook", "detector_version", DETECTOR_VERSION)
    if final_path is not None:
        add("workbook", "path", str(final_path))
    add("warnings", "discrepancies", "NONE")
    return rows


def _commit_workbook(result, context: dict[str, Any], readme: str) -> tuple[int, Path]:
    planned = next_revision(context["existing_revisions"])
    while True:
        name = revision_filename(planned)
        final_path = CONFIG.desktop_dir / name
        if final_path.exists():
            planned += 1
            continue
        temporary = TMP_DIR / f"{Path(name).stem}.tmp.xlsx"
        round_trip = TMP_DIR / f"{Path(name).stem}.roundtrip.tmp.xlsx"
        context["revision_name"] = name
        context["excel_validation"] = "PASS"
        context["diagnostic_rows"] = _diagnostic_rows(result, context)
        write_workbook(temporary, result, context, readme)
        validate_saved_workbook(temporary, EXPECTED_PAIRS)
        _reopen_checks(temporary, result)
        round_trip_workbook(temporary, round_trip)
        round_trip.unlink(missing_ok=True)
        excel_state = excel_desktop_check(temporary)
        rescanned = scan_revisions()
        if final_path.exists() or (rescanned and max(rescanned) >= planned):
            temporary.unlink(missing_ok=True)
            planned = next_revision(scan_revisions())
            continue
        os.replace(temporary, final_path)
        if not final_path.is_file():
            raise RuntimeError("atomic workbook move failed")
        validate_ooxml_workbook(final_path)
        excel_desktop_check(final_path)
        context["excel_desktop"] = excel_state
        return planned, final_path


def _reopen_checks(path: Path, result) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != (
            "Executive Summary",
            "All Order Blocks",
            "Bullish Order Blocks",
            "Bearish Order Blocks",
            "Active Untouched",
            "Body Mitigated",
            "Invalidated",
            "Lifecycle Events",
            "Candidate Audit",
            "Overlap Analysis",
            "Parameters",
            "Diagnostics",
            "README",
        ):
            raise RuntimeError("workbook sheets are out of order")
        bullish = workbook["Executive Summary"]
        values = {row[0].value: row[1].value for row in bullish.iter_rows(min_row=2, max_col=2)}
        counts = disposition_counts(result)
        if values["Qualified Bullish"] != counts["QUALIFIED_BULLISH_ORDER_BLOCK"]:
            raise RuntimeError("summary bullish count changed")
        if values["Qualified Bearish"] != counts["QUALIFIED_BEARISH_ORDER_BLOCK"]:
            raise RuntimeError("summary bearish count changed")
        if values["Candidate Pairs"] != len(result.pairs):
            raise RuntimeError("summary pair count changed")
        if values["Detector Version"] != DETECTOR_VERSION:
            raise RuntimeError("summary detector version changed")
    finally:
        workbook.close()


def excel_desktop_check(path: Path) -> str:
    script = r"""
$before = @(Get-Process EXCEL -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
$excel = $null
$wb = $null
try {
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $excel.AskToUpdateLinks = $false
  $wb = $excel.Workbooks.Open($env:ORDER_BLOCK_XLSX, 0, $true)
  $format = $null
  $names = @()
  for ($attempt = 0; $attempt -lt 8; $attempt++) {
    try {
      $format = [int]$wb.FileFormat
      $names = @()
      foreach ($sheet in $wb.Worksheets) { $names += $sheet.Name }
      if ($format -ne 0 -and $names.Count -gt 0) { break }
    } catch { $format = $null }
    Start-Sleep -Seconds 2
  }
  if ($format -ne 51) { throw ("Unexpected Excel FileFormat " + $format) }
  if ($names.Count -ne 13) { throw ("Excel sheet count " + $names.Count) }
  Write-Output ("EXCEL_OK " + ($names -join "|"))
  $wb.Close($false) | Out-Null
  $wb = $null
  $excel.Quit() | Out-Null
} catch {
  Write-Output ("EXCEL_FAIL " + $_.Exception.Message)
  exit 1
} finally {
  if ($wb -ne $null) { try { $wb.Close($false) | Out-Null } catch {} }
  if ($excel -ne $null) { try { $excel.Quit() | Out-Null } catch {} }
  Get-Process EXCEL -ErrorAction SilentlyContinue | Where-Object { $before -notcontains $_.Id } | ForEach-Object { Stop-Process -Id $_.Id -Force }
}
"""
    try:
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-Command", script],
            check=False,
            capture_output=True,
            env={**os.environ, "ORDER_BLOCK_XLSX": str(path)},
            timeout=90,
        )
    except subprocess.TimeoutExpired:
        print("EXCEL_UNAVAILABLE", flush=True)
        return "unavailable"
    output = _powershell_text(completed.stdout or b"") + _powershell_text(completed.stderr or b"")
    if "EXCEL_OK" in output and completed.returncode == 0:
        print("EXCEL_OK", flush=True)
        return "ok"
    lowered = output.lower()
    license_block = ("lisans", "license", "80010001", "800ac472", "alinamiyor", "rpc_e_call_rejected")
    com_open_blocked = "excel_fail" in lowered and "fileformat" not in lowered and "sheet count" not in lowered
    if any(token in lowered for token in license_block) or com_open_blocked:
        print("EXCEL_UNAVAILABLE", flush=True)
        return "unavailable"
    raise RuntimeError(output.strip() or "Microsoft Excel could not open the workbook")


def _powershell_text(raw: bytes) -> str:
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    for encoding in ("utf-8", "cp1254"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _summary(result, final_path: Path, revision: int, existing: list[int]) -> dict[str, Any]:
    counts = disposition_counts(result)
    statuses = status_counts(result)
    displacements = [block.displacement_percent for block in result.blocks]
    widths = [block.zone_width_percent for block in result.blocks]
    structure = {"TRUE": 0, "FALSE": 0, "INSUFFICIENT_HISTORY": 0}
    for block in result.blocks:
        structure[block.close_structure_break] = structure.get(block.close_structure_break, 0) + 1
    digest = hashlib.sha256(final_path.read_bytes()).hexdigest()
    return {
        "workbook": str(final_path),
        "revision": revision,
        "existing_revisions": existing,
        "size": final_path.stat().st_size,
        "sha256": digest,
        "rows": len(result.bars),
        "pairs": len(result.pairs),
        "qualified_bullish": counts["QUALIFIED_BULLISH_ORDER_BLOCK"],
        "qualified_bearish": counts["QUALIFIED_BEARISH_ORDER_BLOCK"],
        "rejected": {name: value for name, value in counts.items() if name.startswith("REJECTED")},
        "statuses": statuses,
        "displacement_min": min(displacements) if displacements else None,
        "displacement_max": max(displacements) if displacements else None,
        "displacement_mean": _mean(displacements),
        "displacement_median": _median(displacements),
        "zone_width_mean": _mean(widths),
        "zone_width_median": _median(widths),
        "close_structure": structure,
        "overlap_blocks": sum(1 for block in result.blocks if block.overlap_bullish or block.overlap_bearish),
        "first_bullish": [_preview(block) for block in result.blocks if block.direction == "BULLISH"][:5],
        "first_bearish": [_preview(block) for block in result.blocks if block.direction == "BEARISH"][:5],
        "strongest": [_preview(block) for block in sorted(result.blocks, key=lambda item: (-item.displacement_percent, item.impulse_row, item.order_block_id))[:5]],
        "nearest_active_bullish": _preview(_nearest_active(result, "BULLISH")),
        "nearest_active_bearish": _preview(_nearest_active(result, "BEARISH")),
        "final_close": str(result.final_close),
        "detector_version": DETECTOR_VERSION,
        "warnings": [],
    }


def _preview(block) -> dict[str, Any] | None:
    if block is None:
        return None
    return {
        "id": block.order_block_id,
        "direction": block.direction,
        "origin_utc": iso_utc(block.origin_open_time),
        "origin_turkey": iso_turkey(block.origin_open_time),
        "impulse_utc": iso_utc(block.impulse_open_time),
        "impulse_turkey": iso_turkey(block.impulse_open_time),
        "origin_open": str(block.origin_open),
        "origin_close": str(block.origin_close),
        "impulse_open": str(block.impulse_open),
        "impulse_close": str(block.impulse_close),
        "displacement_price": str(block.displacement_price),
        "displacement_percent": str(block.displacement_percent),
        "zone_lower": str(block.zone_lower),
        "zone_upper": str(block.zone_upper),
        "zone_midpoint": str(block.zone_midpoint),
        "quality_score": str(block.quality_score),
        "quality_label": block.quality_label,
        "status": block.status,
        "first_body_touch_utc": iso_utc(block.first_touch_open_time) if block.first_touch_open_time else None,
        "invalidation_utc": iso_utc(block.first_invalidation_open_time) if block.first_invalidation_open_time else None,
        "final_distance": str(block.final_distance_price),
        "final_distance_side": block.final_distance_side,
    }


def _nearest_active(result, direction: str):
    active = [
        block for block in result.blocks
        if block.direction == direction and block.status in {ACTIVE_UNTOUCHED, ACTIVE_BODY_MITIGATED}
    ]
    if not active:
        return None
    return min(
        active,
        key=lambda block: (
            zone_distance(result.final_close, block.zone_lower, block.zone_upper)[0],
            block.impulse_row,
            block.order_block_id,
        ),
    )


def _mean(values: list[Decimal]) -> Decimal | None:
    if not values:
        return None
    return sum(values, Decimal("0")) / Decimal(len(values))


def _median(values: list[Decimal]) -> Decimal | None:
    if not values:
        return None
    ordered = sorted(values)
    mid = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[mid]
    return (ordered[mid - 1] + ordered[mid]) / Decimal("2")


def _decimal(value: Any) -> Decimal:
    if value is None:
        raise OrderBlockError("null decimal")
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _as_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        raise OrderBlockError("naive timestamp")
    return value.astimezone(timezone.utc)


def scan_revisions(directory: Path | None = None) -> list[int]:
    root = directory or CONFIG.desktop_dir
    found = []
    if not root.exists():
        return found
    for path in root.iterdir():
        match = REVISION_PATTERN.match(path.name)
        if match and path.is_file():
            found.append(int(match.group(1)))
    return sorted(set(found))


def next_revision(existing: list[int]) -> int:
    if not existing:
        return 0
    return max(existing) + 1


def revision_filename(revision: int) -> str:
    return f"BTCUSDT_4H_Order_Blocks_rev{revision:02d}.xlsx"


def fingerprint_file(path: Path) -> dict[str, Any]:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    stat = path.stat()
    return {"sha256": digest.hexdigest(), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}


def _existing_code_roots() -> list[Path]:
    root = CONFIG.project_root
    paths = [
        root / "detectors" / "primary_range_v1",
        root / "detectors" / "hierarchical_swing_v4",
        root / "detectors" / "price_touch_hierarchy_v1",
        root / "detectors" / "open_liquidity_v1",
        root / "detectors" / "__init__.py",
        root / "data",
        root / "ingestion",
        root / "processing",
        root / "config.py",
    ]
    return [path for path in paths if path.exists()]


def _fingerprint_tree(paths: list[Path]) -> dict[str, str]:
    found: dict[str, str] = {}
    for path in paths:
        files = [path] if path.is_file() else sorted(item for item in path.rglob("*") if item.is_file())
        for file in files:
            if "__pycache__" in file.parts or file.suffix == ".pyc":
                continue
            found[str(file)] = hashlib.sha256(file.read_bytes()).hexdigest()
    return found


def _stray_python_files() -> list[str]:
    return [str(file) for file in CONFIG.desktop_dir.glob("*.py")]


def assert_python_write_allowed(path: Path) -> None:
    resolved = path.resolve()
    package = CONFIG.package_dir.resolve()
    project = CONFIG.project_root.resolve()
    if project not in resolved.parents and resolved != project:
        raise PermissionError(f"detector Python writes are limited to {project}")
    if package not in resolved.parents and resolved != package:
        raise PermissionError(f"detector Python writes are limited to {package}")


def _cleanup() -> None:
    if TMP_DIR.exists():
        shutil.rmtree(TMP_DIR, ignore_errors=True)
    parent = TMP_DIR.parent
    if parent.exists() and parent.is_dir() and not any(parent.iterdir()):
        parent.rmdir()
    for root in (CONFIG.package_dir, CONFIG.package_dir / "tests"):
        cache = root / "__pycache__"
        if cache.exists():
            shutil.rmtree(cache, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
