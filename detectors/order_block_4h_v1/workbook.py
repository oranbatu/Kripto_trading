"""Conservative Open XML workbook for the 4h Order Block detector."""
from __future__ import annotations

import math
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Sequence
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.order_block_4h_v1.engine import (
    ACTIVE_BODY_MITIGATED,
    ACTIVE_UNTOUCHED,
    FORWARD_METRIC_CLASS,
    INVALIDATED,
    QUALIFICATION_PRICE_MODE,
    STRUCTURE_METRIC,
    TERMINAL,
    ZONE_DEFINITION,
    AnalysisResult,
    OrderBlock,
    disposition_counts,
    iso_turkey,
    iso_utc,
    status_counts,
    zone_distance,
)
from detectors.order_block_4h_v1.engine import _mean as mean_decimal
from detectors.order_block_4h_v1.engine import _median as median_decimal
from detectors.order_block_4h_v1.order_block_config import (
    DETECTOR_VERSION,
    PLAIN_RANGE_SHEETS,
    TABLE_NAMES,
    WORKBOOK_SHEETS,
)

PRICE = "0.00000000"
PERCENT_POINTS = '0.00000000"%"'
EXCEL_MAX_TEXT = 32767
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
OOXML_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_EMPTY_TYPED_CELL = re.compile(br'<c\b[^>]*\bt="(?:n|b|s|str|inlineStr|d|e)"[^>]*/>')


def _solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


HEADER_FILL = _solid("1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
HEADER_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
BULL_FILL = _solid("D9EAD3")
BEAR_FILL = _solid("F4CCCC")
ACTIVE_FILL = _solid("C9DAF8")
MITIGATED_FILL = _solid("FCE5CD")
INVALID_FILL = _solid("E6B8B7")
TERMINAL_FILL = _solid("D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="center")


def write_workbook(path: Path, result: AnalysisResult, context: dict[str, Any], readme: str) -> None:
    workbook = Workbook()
    sheets = {
        "Executive Summary": _executive(result, context),
        "All Order Blocks": _blocks(result.blocks),
        "Bullish Order Blocks": _blocks([block for block in result.blocks if block.direction == "BULLISH"]),
        "Bearish Order Blocks": _blocks([block for block in result.blocks if block.direction == "BEARISH"]),
        "Active Untouched": _blocks([block for block in result.blocks if block.status == ACTIVE_UNTOUCHED]),
        "Body Mitigated": _blocks([block for block in result.blocks if block.status == ACTIVE_BODY_MITIGATED]),
        "Invalidated": _blocks([block for block in result.blocks if block.status == INVALIDATED]),
        "Lifecycle Events": _events(result),
        "Candidate Audit": _audit(result),
        "Overlap Analysis": _overlaps(result.blocks),
        "Parameters": _parameters(),
        "Diagnostics": context["diagnostic_rows"],
        "README": _readme_rows(readme),
    }
    first = workbook.active
    first.title = WORKBOOK_SHEETS[0]
    for name in WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(name)
    for name in WORKBOOK_SHEETS:
        _write_sheet(workbook[name], sheets[name], TABLE_NAMES.get(name))
    _highlight(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def validate_saved_workbook(path: Path, expected_pairs: int) -> None:
    validate_ooxml_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = workbook["Candidate Audit"].max_row - 1
        if rows != expected_pairs:
            raise RuntimeError(f"candidate audit rows {rows} != {expected_pairs}")
    finally:
        workbook.close()


def validate_ooxml_workbook(path: Path) -> None:
    payload = path.read_bytes()
    if payload[:4] != b"PK\x03\x04":
        raise RuntimeError("workbook does not begin with the ZIP signature")
    if path.suffix.lower() != ".xlsx":
        raise RuntimeError("workbook extension is not .xlsx")
    if not zipfile.is_zipfile(path):
        raise RuntimeError("zipfile.is_zipfile returned false")
    required = (
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/worksheets/sheet1.xml",
    )
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("ZipFile.testzip found a corrupt member")
        names = set(archive.namelist())
        for member in required:
            if member not in names:
                raise RuntimeError(f"missing OOXML member {member}")
        if any("vba" in name.lower() for name in names):
            raise RuntimeError("workbook contains a VBA part")
        if any("externalLink" in name for name in names):
            raise RuntimeError("workbook contains an external link")
        for name in sorted(names):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                sheet_payload = archive.read(name)
                if _EMPTY_TYPED_CELL.search(sheet_payload):
                    raise RuntimeError(f"{name} contains a typed cell with no value")
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise RuntimeError(f"XML parse failed for {name}: {exc}") from exc
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                    _validate_sheet_xml(name, root)
        _validate_workbook_relationships(archive)
    _validate_openpyxl_workbook(path)


def _validate_sheet_xml(name: str, root: ET.Element) -> None:
    dimension = root.find("m:dimension", OOXML_NS)
    auto_filter = root.find("m:autoFilter", OOXML_NS)
    if auto_filter is not None and dimension is not None:
        if auto_filter.attrib.get("ref", "") != dimension.attrib.get("ref", ""):
            raise RuntimeError(f"{name} autoFilter does not match the used range")


def _validate_workbook_relationships(archive: zipfile.ZipFile) -> None:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.attrib.get("Id"): rel.attrib.get("Target", "") for rel in relationships}
    for sheet in root.findall("m:sheets/m:sheet", OOXML_NS):
        title = sheet.attrib.get("name", "")
        if not title or len(title) > 31:
            raise RuntimeError(f"invalid sheet name {title!r}")
        relation = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = targets.get(relation, "").lstrip("/")
        member = target if target.startswith("xl/") else "xl/" + target
        if member not in archive.namelist():
            raise RuntimeError(f"sheet {title} references missing part {member}")


def _validate_openpyxl_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError(f"sheet order {workbook.sheetnames}")
        if workbook.vba_archive is not None:
            raise RuntimeError("workbook contains macros")
        if getattr(workbook, "_external_links", []):
            raise RuntimeError("workbook contains external links")
        names = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > EXCEL_MAX_ROWS or worksheet.max_column > EXCEL_MAX_COLUMNS:
                raise RuntimeError(f"{worksheet.title} exceeds an Excel worksheet limit")
            if worksheet.title in PLAIN_RANGE_SHEETS and worksheet.tables:
                raise RuntimeError(f"{worksheet.title} must stay a plain range")
            expected = TABLE_NAMES.get(worksheet.title)
            if expected is not None:
                if worksheet["A2"].value == "NO_RESULTS":
                    if worksheet.tables:
                        raise RuntimeError(f"{worksheet.title} created a table for NO_RESULTS")
                elif not worksheet.tables:
                    raise RuntimeError(f"{worksheet.title} is missing its Excel table")
            for table in worksheet.tables.values():
                names.append(table.name)
                _validate_table(worksheet, table)
        if len(names) != len(set(names)):
            raise RuntimeError(f"duplicate table names {names}")
        for defined_name in workbook.defined_names.values():
            attr = getattr(defined_name, "attr_text", "") or ""
            if "#REF!" in attr or "#NAME?" in attr:
                raise RuntimeError(f"invalid defined name {defined_name.name}")
    finally:
        workbook.close()
    streamed = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if tuple(streamed.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError("read-only sheet order mismatch")
        for worksheet in streamed.worksheets:
            for _row in worksheet.iter_rows():
                pass
    finally:
        streamed.close()


def _validate_table(worksheet, table) -> None:
    name = table.name
    if not name or not (name[0].isalpha() or name[0] == "_") or not all(ch.isalnum() or ch == "_" for ch in name):
        raise RuntimeError(f"invalid table display name {name}")
    start, end = table.ref.split(":")
    from openpyxl.utils import coordinate_to_tuple

    start_row, start_col = coordinate_to_tuple(start)
    end_row, end_col = coordinate_to_tuple(end)
    if start_row < 1 or end_row > worksheet.max_row or start_col < 1 or end_col > worksheet.max_column or end_row <= start_row:
        raise RuntimeError(f"table {name} reference {table.ref} is invalid")
    headers = [worksheet.cell(start_row, col).value for col in range(start_col, end_col + 1)]
    if any(value in (None, "") for value in headers) or len(headers) != len(set(headers)):
        raise RuntimeError(f"table {name} headers are blank or duplicated")


def round_trip_workbook(source: Path, copy_path: Path) -> None:
    workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    try:
        original_sheets = tuple(workbook.sheetnames)
        original_summary = [tuple(row) for row in workbook["Executive Summary"].iter_rows(values_only=True)]
        original_dimensions = {worksheet.title: worksheet.dimensions for worksheet in workbook.worksheets}
        original_headers = {worksheet.title: tuple(cell.value for cell in worksheet[1]) for worksheet in workbook.worksheets}
        workbook.save(copy_path)
    finally:
        workbook.close()
    reopened = load_workbook(copy_path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(reopened.sheetnames) != original_sheets:
            raise RuntimeError("round trip changed the sheet order")
        summary = [tuple(row) for row in reopened["Executive Summary"].iter_rows(values_only=True)]
        if summary != original_summary:
            raise RuntimeError("round trip changed the executive summary")
        for worksheet in reopened.worksheets:
            if worksheet.dimensions != original_dimensions[worksheet.title]:
                raise RuntimeError(f"round trip changed dimensions for {worksheet.title}")
            headers = tuple(cell.value for cell in worksheet[1])
            if headers != original_headers[worksheet.title]:
                raise RuntimeError(f"round trip changed headers for {worksheet.title}")
    finally:
        reopened.close()


def _write_sheet(ws, rows: Sequence[Sequence[Any]], table_name: str | None) -> None:
    prepared = _prepare_rows(rows)
    for row in prepared:
        ws.append(row)
    if not prepared:
        return
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = HEADER_BORDER
    width = len(prepared[0])
    for index, header in enumerate(prepared[0], start=1):
        label = header if isinstance(header, str) else ""
        ws.column_dimensions[get_column_letter(index)].width = min(42, max(12, len(label) + 2))
    ws.row_dimensions[1].height = 30
    if ws.max_row > EXCEL_MAX_ROWS or ws.max_column > EXCEL_MAX_COLUMNS:
        raise RuntimeError(f"{ws.title} exceeds an Excel worksheet limit")
    if len(prepared) >= 2:
        ref = f"A1:{get_column_letter(width)}{len(prepared)}"
        ws.auto_filter.ref = ref
        if table_name and prepared[1][0] != "NO_RESULTS":
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    _apply_number_formats(ws)


def _prepare_rows(rows: Sequence[Sequence[Any]]) -> list[list[Any]]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    prepared = []
    for row in rows:
        padded = list(row) + [None] * (width - len(row))
        prepared.append([excel_value(value) for value in padded])
    if len(prepared) == 1:
        prepared.append(["NO_RESULTS"] + [None] * (width - 1))
    return prepared


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        return float(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, datetime):
        if value.tzinfo is not None and value.tzinfo.utcoffset(value) is not None:
            return iso_utc(value)
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        text = _clean_text(value)
        return text[:EXCEL_MAX_TEXT] or None
    text = _clean_text(str(value))
    return text[:EXCEL_MAX_TEXT] or None


def _clean_text(text: str) -> str:
    return "".join(char if char in "\t\n\r" or ord(char) >= 32 else " " for char in text)


def _apply_number_formats(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    for index, name in enumerate(headers, start=1):
        if not isinstance(name, str) or "UTC" in name or "Turkey" in name:
            continue
        if "Percent" in name:
            kind = "percent"
        elif any(token in name for token in ("Open", "Close", "High", "Low", "Price", "Lower", "Upper", "Midpoint", "Distance", "Overlap", "Extreme", "Excursion")):
            kind = "price"
        else:
            continue
        number_format = PERCENT_POINTS if kind == "percent" else PRICE
        for cell in ws.iter_rows(min_row=2, min_col=index, max_col=index, max_row=ws.max_row):
            if isinstance(cell[0].value, (int, float)):
                cell[0].number_format = number_format


def _highlight(workbook: Workbook) -> None:
    fills = {
        "BULLISH": BULL_FILL,
        "BEARISH": BEAR_FILL,
        "QUALIFIED_BULLISH_ORDER_BLOCK": BULL_FILL,
        "QUALIFIED_BEARISH_ORDER_BLOCK": BEAR_FILL,
        ACTIVE_UNTOUCHED: ACTIVE_FILL,
        ACTIVE_BODY_MITIGATED: MITIGATED_FILL,
        INVALIDATED: INVALID_FILL,
        TERMINAL: TERMINAL_FILL,
    }
    for worksheet in workbook.worksheets:
        if worksheet.max_row < 2 or worksheet.max_row > 8000:
            continue
        headers = {cell.value: cell.column for cell in worksheet[1]}
        for header in ("Direction", "Lifecycle Status", "Primary Disposition"):
            column = headers.get(header)
            if not column:
                continue
            for cell in worksheet.iter_rows(min_row=2, min_col=column, max_col=column, max_row=worksheet.max_row):
                fill = fills.get(cell[0].value)
                if fill is not None:
                    cell[0].fill = fill


def _time(value: datetime | None, turkey: bool = False) -> str | None:
    if value is None:
        return None
    return iso_turkey(value) if turkey else iso_utc(value)


def _text_decimal(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return format(value, "f")


def _blocks(blocks: Sequence[OrderBlock]) -> list[list[Any]]:
    rows = [[name for name, _getter in BLOCK_COLUMNS]]
    for block in blocks:
        rows.append([getter(block) for _name, getter in BLOCK_COLUMNS])
    return rows


def _events(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Order Block ID", "Direction", "Event Type", "Event Row", "Open UTC", "Open Turkey",
        "Close UTC", "Close Turkey", "Candle Open", "Candle Close", "Bars From Confirmation",
        "Hours From Confirmation", "Detail", "Forward Metric Class", "Detector Version",
    ]]
    for event in result.events:
        rows.append([
            event.order_block_id, event.direction, event.event_type, event.row,
            iso_utc(event.open_time), iso_turkey(event.open_time), iso_utc(event.close_time), iso_turkey(event.close_time),
            event.open, event.close, event.bars_from_confirmation, event.bars_from_confirmation * 4,
            event.detail, FORWARD_METRIC_CLASS, DETECTOR_VERSION,
        ])
    return rows


def _audit(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Candidate Pair ID", "Origin Row", "Impulse Row", "Origin Open UTC", "Origin Open Turkey",
        "Impulse Open UTC", "Impulse Open Turkey", "Origin Open", "Origin Close", "Impulse Open",
        "Impulse Close", "Origin Direction", "Impulse Direction", "Bullish Displacement Percent",
        "Bearish Displacement Percent", "Unrounded Bullish Displacement Percent",
        "Unrounded Bearish Displacement Percent", "Bullish Qualification", "Bearish Qualification",
        "Primary Disposition", "Qualified Order Block ID", "Rejection Reason", "Confirmation UTC",
        "Confirmation Turkey",
    ]]
    for pair in result.pairs:
        rows.append([
            pair.pair_id, pair.origin_row, pair.impulse_row, iso_utc(pair.origin_open_time), iso_turkey(pair.origin_open_time),
            iso_utc(pair.impulse_open_time), iso_turkey(pair.impulse_open_time), pair.origin_open, pair.origin_close,
            pair.impulse_open, pair.impulse_close, pair.origin_direction, pair.impulse_direction,
            pair.bullish_displacement_percent, pair.bearish_displacement_percent,
            _text_decimal(pair.bullish_displacement_percent), _text_decimal(pair.bearish_displacement_percent),
            pair.bullish_qualification, pair.bearish_qualification, pair.disposition, pair.order_block_id,
            "NOT_APPLICABLE" if pair.order_block_id else pair.disposition,
            iso_utc(pair.impulse_close_time), iso_turkey(pair.impulse_close_time),
        ])
    return rows


def _overlaps(blocks: Sequence[OrderBlock]) -> list[list[Any]]:
    rows = [[
        "Order Block ID", "Direction", "Zone Lower", "Zone Upper", "Overlapping Bullish Count",
        "Overlapping Bearish Count", "Same Direction Overlap Count", "Opposite Direction Overlap Count",
        "Max Overlap Price", "Max Overlap Percent", "Overlap Cluster ID", "Detector Version",
    ]]
    for block in blocks:
        rows.append([
            block.order_block_id, block.direction, block.zone_lower, block.zone_upper, block.overlap_bullish,
            block.overlap_bearish, block.same_direction_overlap, block.opposite_direction_overlap,
            block.max_overlap_price, block.max_overlap_percent, block.cluster_id, DETECTOR_VERSION,
        ])
    return rows


def _executive(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    counts = disposition_counts(result)
    statuses = status_counts(result)
    displacements = [block.displacement_percent for block in result.blocks]
    widths = [block.zone_width_percent for block in result.blocks]
    structure = {"TRUE": 0, "FALSE": 0, "INSUFFICIENT_HISTORY": 0}
    for block in result.blocks:
        structure[block.close_structure_break] = structure.get(block.close_structure_break, 0) + 1
    bullish = [block for block in result.blocks if block.direction == "BULLISH"]
    bearish = [block for block in result.blocks if block.direction == "BEARISH"]
    nearest_bull = _nearest(bullish, result.final_close)
    nearest_bear = _nearest(bearish, result.final_close)
    overlapping = sum(1 for block in result.blocks if block.overlap_bullish or block.overlap_bearish)
    rows = [
        ["Metric", "Value", "Notes"],
        ["Symbol", "BTCUSDT", "BINANCE:BTCUSDT.P"],
        ["Market", "Binance USD-M PERPETUAL", "futures/um"],
        ["Timeframe", "4h", "completed candles only"],
        ["Detector Version", DETECTOR_VERSION, QUALIFICATION_PRICE_MODE],
        ["Workbook Revision", context.get("revision_name", ""), ""],
        ["Source Rows", len(result.bars), context.get("source_validation", "")],
        ["Candidate Pairs", len(result.pairs), "one disposition each"],
        ["Qualified Bullish", counts["QUALIFIED_BULLISH_ORDER_BLOCK"], ""],
        ["Qualified Bearish", counts["QUALIFIED_BEARISH_ORDER_BLOCK"], ""],
        ["Rejected Same Direction", counts["REJECTED_SAME_DIRECTION"], ""],
        ["Rejected Origin Doji", counts["REJECTED_ORIGIN_DOJI"], ""],
        ["Rejected Impulse Doji", counts["REJECTED_IMPULSE_DOJI"], ""],
        ["Rejected Bullish Below 1 Percent", counts["REJECTED_BULLISH_DISPLACEMENT_BELOW_1_PERCENT"], ""],
        ["Rejected Bearish Below 1 Percent", counts["REJECTED_BEARISH_DISPLACEMENT_BELOW_1_PERCENT"], ""],
        ["Rejected Direction Mismatch", counts["REJECTED_DIRECTION_MISMATCH"], ""],
        ["Active Untouched", statuses[ACTIVE_UNTOUCHED], ""],
        ["Body Mitigated", statuses[ACTIVE_BODY_MITIGATED], ""],
        ["Invalidated By Body Close", statuses[INVALIDATED], ""],
        ["Terminal At Dataset End", statuses[TERMINAL], ""],
        ["Minimum Qualified Displacement Percent", min(displacements) if displacements else None, "unrounded"],
        ["Maximum Qualified Displacement Percent", max(displacements) if displacements else None, "unrounded"],
        ["Average Displacement Percent", mean_decimal(displacements), ""],
        ["Median Displacement Percent", median_decimal(displacements), ""],
        ["Average Zone Width Percent", mean_decimal(widths), "origin body"],
        ["Median Zone Width Percent", median_decimal(widths), "origin body"],
        ["Close Structure True", structure["TRUE"], STRUCTURE_METRIC],
        ["Close Structure False", structure["FALSE"], "descriptive only"],
        ["Close Structure Insufficient History", structure["INSUFFICIENT_HISTORY"], ""],
        ["Blocks With Zone Overlap", overlapping, "blocks are not merged"],
        ["Final Close", result.final_close, iso_utc(result.final_close_time)],
        ["Nearest Active Bullish", nearest_bull.order_block_id if nearest_bull else "NONE", _nearest_note(nearest_bull, result.final_close)],
        ["Nearest Active Bearish", nearest_bear.order_block_id if nearest_bear else "NONE", _nearest_note(nearest_bear, result.final_close)],
        ["Qualification Price Mode", QUALIFICATION_PRICE_MODE, "high and low do not qualify"],
        ["Zone Definition", ZONE_DEFINITION, ""],
        ["Excel Validation", context.get("excel_validation", ""), ""],
    ]
    return rows


def _nearest(blocks: Sequence[OrderBlock], final_close: Decimal) -> OrderBlock | None:
    active = [block for block in blocks if block.status in {ACTIVE_UNTOUCHED, ACTIVE_BODY_MITIGATED}]
    if not active:
        return None
    return min(active, key=lambda block: (zone_distance(final_close, block.zone_lower, block.zone_upper)[0], block.impulse_row, block.order_block_id))


def _nearest_note(block: OrderBlock | None, final_close: Decimal) -> str:
    if block is None:
        return "no active zone"
    distance, side = zone_distance(final_close, block.zone_lower, block.zone_upper)
    return f"{side} distance {format(distance, 'f')} zone {format(block.zone_lower, 'f')} to {format(block.zone_upper, 'f')}"


def _parameters() -> list[list[Any]]:
    values = [
        ("symbol", "BTCUSDT"),
        ("tradingview_symbol", "BINANCE:BTCUSDT.P"),
        ("market", "Binance USD-M PERPETUAL"),
        ("timeframe", "4h"),
        ("qualification_price_mode", QUALIFICATION_PRICE_MODE),
        ("origin_impulse_relationship", "IMMEDIATELY_ADJACENT_CANDLES"),
        ("bullish_origin", "CLOSE_LESS_THAN_OPEN"),
        ("bullish_impulse", "CLOSE_GREATER_THAN_OPEN"),
        ("bearish_origin", "CLOSE_GREATER_THAN_OPEN"),
        ("bearish_impulse", "CLOSE_LESS_THAN_OPEN"),
        ("minimum_displacement_percent", "1.00"),
        ("bullish_displacement_formula", "(IMPULSE_CLOSE - ORIGIN_OPEN) / ORIGIN_OPEN x 100"),
        ("bearish_displacement_formula", "(ORIGIN_OPEN - IMPULSE_CLOSE) / ORIGIN_OPEN x 100"),
        ("exactly_one_percent_passes", "TRUE"),
        ("values_below_one_percent_fail", "TRUE"),
        ("zone_definition", ZONE_DEFINITION),
        ("zone_lower", "MIN(ORIGIN_OPEN, ORIGIN_CLOSE)"),
        ("zone_upper", "MAX(ORIGIN_OPEN, ORIGIN_CLOSE)"),
        ("high_low_used_for_qualification", "FALSE"),
        ("wick_used_for_qualification", "FALSE"),
        ("atr_used_for_qualification", "FALSE"),
        ("fvg_required", "FALSE"),
        ("bos_required", "FALSE"),
        ("body_baseline_lookback", "20"),
        ("close_structure_lookback", "20"),
        ("close_structure_break_required", "FALSE"),
        ("lifecycle_starts_after_impulse", "TRUE"),
        ("primary_mitigation_mode", "LATER_CANDLE_BODY_INTERSECTION"),
        ("bullish_invalidation", "LATER_CLOSE_STRICTLY_BELOW_ZONE_LOWER"),
        ("bearish_invalidation", "LATER_CLOSE_STRICTLY_ABOVE_ZONE_UPPER"),
        ("breaker_block_conversion", "FALSE"),
        ("computational_timezone", "UTC"),
        ("display_timezone", "Europe/Istanbul"),
        ("timestamp_excel_storage", "ISO_8601_TEXT"),
        ("excel_format", "XLSX_OFFICE_OPEN_XML"),
        ("excel_writer", "OPENPYXL_3_1_5"),
        ("excel_macro_enabled", "FALSE"),
        ("excel_external_links", "FALSE"),
        ("detector_version", DETECTOR_VERSION),
    ]
    return [["Parameter", "Value"], *([name, value] for name, value in values)]


def _readme_rows(readme: str) -> list[list[Any]]:
    lines = [line for line in readme.splitlines() if line.strip()]
    return [["README Line"], *([line] for line in lines)]


def _optional_time(value: datetime | None, turkey: bool = False) -> str | None:
    return _time(value, turkey)


BLOCK_COLUMNS: list[tuple[str, Callable[[OrderBlock], Any]]] = [
    ("Order Block ID", lambda block: block.order_block_id),
    ("Direction", lambda block: block.direction),
    ("Sequence", lambda block: block.sequence),
    ("Origin Row", lambda block: block.origin_row),
    ("Impulse Row", lambda block: block.impulse_row),
    ("Origin Open UTC", lambda block: iso_utc(block.origin_open_time)),
    ("Origin Open Turkey", lambda block: iso_turkey(block.origin_open_time)),
    ("Origin Close UTC", lambda block: iso_utc(block.origin_close_time)),
    ("Origin Close Turkey", lambda block: iso_turkey(block.origin_close_time)),
    ("Impulse Open UTC", lambda block: iso_utc(block.impulse_open_time)),
    ("Impulse Open Turkey", lambda block: iso_turkey(block.impulse_open_time)),
    ("Impulse Close UTC", lambda block: iso_utc(block.impulse_close_time)),
    ("Impulse Close Turkey", lambda block: iso_turkey(block.impulse_close_time)),
    ("Confirmed At UTC", lambda block: iso_utc(block.confirmed_at)),
    ("Confirmed At Turkey", lambda block: iso_turkey(block.confirmed_at)),
    ("Origin Open", lambda block: block.origin_open),
    ("Origin Close", lambda block: block.origin_close),
    ("Origin High Audit", lambda block: block.origin_high),
    ("Origin Low Audit", lambda block: block.origin_low),
    ("Impulse Open", lambda block: block.impulse_open),
    ("Impulse Close", lambda block: block.impulse_close),
    ("Impulse High Audit", lambda block: block.impulse_high),
    ("Impulse Low Audit", lambda block: block.impulse_low),
    ("Origin Direction", lambda block: block.origin_direction),
    ("Impulse Direction", lambda block: block.impulse_direction),
    ("Origin Body Price", lambda block: block.origin_body_price),
    ("Origin Body Percent", lambda block: block.origin_body_percent),
    ("Impulse Body Price", lambda block: block.impulse_body_price),
    ("Impulse Body Percent", lambda block: block.impulse_body_percent),
    ("Impulse To Origin Body Ratio", lambda block: block.impulse_to_origin_body_ratio),
    ("Displacement Price", lambda block: block.displacement_price),
    ("Displacement Percent", lambda block: block.displacement_percent),
    ("Unrounded Displacement Percent", lambda block: _text_decimal(block.displacement_percent)),
    ("Minimum Required Displacement Percent", lambda block: block.minimum_displacement_percent),
    ("Displacement Pass", lambda block: block.displacement_pass),
    ("Zone Lower", lambda block: block.zone_lower),
    ("Zone Upper", lambda block: block.zone_upper),
    ("Zone Midpoint", lambda block: block.zone_midpoint),
    ("Zone Width Price", lambda block: block.zone_width_price),
    ("Zone Width Percent", lambda block: block.zone_width_percent),
    ("Rolling Median Body Price", lambda block: block.rolling_median_body),
    ("Rolling Mean Body Price", lambda block: block.rolling_mean_body),
    ("Impulse Body To Median Ratio", lambda block: block.impulse_body_to_median_ratio),
    ("Displacement To Median Body Ratio", lambda block: block.displacement_to_median_body_ratio),
    ("Close Structure Break", lambda block: block.close_structure_break),
    ("Prior Close Extreme", lambda block: block.prior_close_extreme),
    ("Break Distance Price", lambda block: block.break_distance),
    ("Break Percent", lambda block: block.break_percent),
    ("Quality Score", lambda block: block.quality_score),
    ("Quality Label", lambda block: block.quality_label),
    ("Lifecycle Status", lambda block: block.status),
    ("Ever Body Touched", lambda block: block.ever_body_touched),
    ("First Body Touch Row", lambda block: block.first_touch_row),
    ("First Body Touch UTC", lambda block: _optional_time(block.first_touch_open_time)),
    ("First Body Touch Turkey", lambda block: _optional_time(block.first_touch_open_time, True)),
    ("Bars To Body Touch", lambda block: block.bars_to_touch),
    ("Hours To Body Touch", lambda block: block.hours_to_touch),
    ("Days To Body Touch", lambda block: block.days_to_touch),
    ("Body Overlap Price", lambda block: block.body_overlap_price),
    ("Body Overlap Percent", lambda block: block.body_overlap_percent),
    ("Retrace Percent Display", lambda block: block.retrace_percent_display),
    ("Midpoint Reached", lambda block: block.ever_midpoint),
    ("Full Zone Traversed", lambda block: block.ever_full_traversal),
    ("First Invalidation Row", lambda block: block.first_invalidation_row),
    ("First Invalidation UTC", lambda block: _optional_time(block.first_invalidation_open_time)),
    ("Bars To Invalidation", lambda block: block.bars_to_invalidation),
    ("Body Revisit Count", lambda block: block.body_revisit_count),
    ("Touch Episode Count", lambda block: block.touch_episode_count),
    ("Final Close", lambda block: block.final_close),
    ("Final Distance Price", lambda block: block.final_distance_price),
    ("Final Distance Percent", lambda block: block.final_distance_percent),
    ("Final Distance Side", lambda block: block.final_distance_side),
    ("Right Censored", lambda block: block.right_censored),
    ("Censoring Reason", lambda block: block.censoring_reason),
    ("Open At Dataset End", lambda block: block.open_at_dataset_end),
    ("Overlapping Bullish Count", lambda block: block.overlap_bullish),
    ("Overlapping Bearish Count", lambda block: block.overlap_bearish),
    ("Same Direction Overlap Count", lambda block: block.same_direction_overlap),
    ("Opposite Direction Overlap Count", lambda block: block.opposite_direction_overlap),
    ("Max Overlap Price", lambda block: block.max_overlap_price),
    ("Max Overlap Percent", lambda block: block.max_overlap_percent),
    ("Overlap Cluster ID", lambda block: block.cluster_id),
    ("Wick Only Contact Diagnostic", lambda block: block.wick_only_contact),
    ("First Event", lambda block: block.first_event),
    ("Latest Event", lambda block: block.latest_event),
    ("Qualification Price Mode", lambda block: QUALIFICATION_PRICE_MODE),
    ("Zone Definition", lambda block: ZONE_DEFINITION),
    ("Structure Metric", lambda block: STRUCTURE_METRIC),
    ("Forward Metric Class", lambda block: FORWARD_METRIC_CLASS),
    ("Detector Version", lambda block: DETECTOR_VERSION),
]
