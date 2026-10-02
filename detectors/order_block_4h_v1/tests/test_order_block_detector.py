"""Specification tests for the 4h body-based Order Block detector."""
from __future__ import annotations

import hashlib
import math
import os
import shutil
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from detectors.order_block_4h_v1.engine import (
    ACTIVE_BODY_MITIGATED,
    ACTIVE_UNTOUCHED,
    INVALIDATED,
    QUALIFIED_BEARISH,
    QUALIFIED_BULLISH,
    REJECTED_BEARISH_BELOW,
    REJECTED_BULLISH_BELOW,
    REJECTED_DIRECTION_MISMATCH,
    REJECTED_IMPULSE_DOJI,
    REJECTED_ORIGIN_DOJI,
    REJECTED_SAME_DIRECTION,
    TERMINAL,
    Bar,
    analyze,
    audit_result,
    iso_turkey,
    iso_utc,
)
from detectors.order_block_4h_v1.order_block_config import CONFIG, DETECTOR_VERSION, PROJECT_ROOT, WORKBOOK_SHEETS
from detectors.order_block_4h_v1.run_order_block_detector import (
    assert_python_write_allowed,
    fingerprint_file,
    next_revision,
    revision_filename,
    scan_revisions,
)
from detectors.order_block_4h_v1.workbook import excel_value, validate_ooxml_workbook, write_workbook

D = Decimal


def bar(row: int, open_, close, high=None, low=None) -> Bar:
    open_ = D(str(open_))
    close = D(str(close))
    high = D(str(max(open_, close) if high is None else high))
    low = D(str(min(open_, close) if low is None else low))
    open_time = datetime(2024, 1, 1, tzinfo=timezone.utc) + timedelta(hours=4 * row)
    close_time = open_time + timedelta(hours=4) - timedelta(milliseconds=1)
    return Bar(row, open_time, close_time, open_, high, low, close, D("1"), 1)


def bullish_pair(extra=None) -> list[Bar]:
    rows = [bar(0, "100", "99"), bar(1, "99", "101")]
    if extra:
        rows.extend(extra)
    return rows


class OrderBlockDetectorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = PROJECT_ROOT / "tmp" / "order_block_4h_v1_unit"
        if cls.temp_dir.exists():
            shutil.rmtree(cls.temp_dir)
        cls.temp_dir.mkdir(parents=True)
        bars = [
            bar(0, "100", "99"),
            bar(1, "99", "101"),
            bar(2, "100", "101"),
            bar(3, "102", "104"),
            bar(4, "104", "100"),
        ]
        cls.result = analyze(bars)
        audit_result(cls.result)
        cls.workbook_path = cls.temp_dir / "sample.xlsx"
        write_workbook(
            cls.workbook_path,
            cls.result,
            {"revision_name": "BTCUSDT_4H_Order_Blocks_rev00.xlsx", "source_validation": "PASS", "excel_validation": "PASS", "diagnostic_rows": [["Section", "Name", "Value"], ["test", "sample", "PASS"]]},
            "Order Block detector\nBody rules only\n",
        )
        validate_ooxml_workbook(cls.workbook_path)

    @classmethod
    def tearDownClass(cls) -> None:
        if cls.temp_dir.exists():
            shutil.rmtree(cls.temp_dir, ignore_errors=True)

    def test_01_bullish_pattern(self) -> None:
        result = analyze(bullish_pair())
        self.assertEqual(result.pairs[0].disposition, QUALIFIED_BULLISH)
        self.assertEqual(result.blocks[0].direction, "BULLISH")

    def test_02_bearish_pattern(self) -> None:
        result = analyze([bar(0, "100", "102"), bar(1, "102", "98")])
        self.assertEqual(result.pairs[0].disposition, QUALIFIED_BEARISH)
        self.assertEqual(result.blocks[0].direction, "BEARISH")

    def test_03_exact_one_percent_bullish_passes(self) -> None:
        result = analyze([bar(0, "100", "99"), bar(1, "100", "101")])
        self.assertEqual(result.blocks[0].displacement_percent, D("1"))
        self.assertEqual(result.pairs[0].disposition, QUALIFIED_BULLISH)

    def test_04_below_one_percent_bullish_fails(self) -> None:
        result = analyze([bar(0, "100", "99"), bar(1, "100", "100.99999999")])
        self.assertEqual(result.pairs[0].disposition, REJECTED_BULLISH_BELOW)
        self.assertEqual(result.blocks, [])
        self.assertLess(result.pairs[0].bullish_displacement_percent, D("1"))

    def test_05_exact_one_percent_bearish_passes(self) -> None:
        result = analyze([bar(0, "100", "101"), bar(1, "100", "99")])
        self.assertEqual(result.blocks[0].displacement_percent, D("1"))
        self.assertEqual(result.pairs[0].disposition, QUALIFIED_BEARISH)

    def test_06_below_one_percent_bearish_fails(self) -> None:
        result = analyze([bar(0, "100", "101"), bar(1, "100", "99.00000001")])
        self.assertEqual(result.pairs[0].disposition, REJECTED_BEARISH_BELOW)
        self.assertEqual(result.blocks, [])

    def test_07_high_low_do_not_affect_qualification(self) -> None:
        first = analyze([bar(0, "100", "99", high="130", low="70"), bar(1, "99", "101", high="140", low="60"), bar(2, "110", "111", high="112", low="109")])
        second = analyze([bar(0, "100", "99", high="101", low="98"), bar(1, "99", "101", high="102", low="98"), bar(2, "110", "111", high="160", low="50")])
        self.assertEqual(first.pairs[0].disposition, second.pairs[0].disposition)
        self.assertEqual(first.blocks[0].zone_lower, second.blocks[0].zone_lower)
        self.assertEqual(first.blocks[0].zone_upper, second.blocks[0].zone_upper)
        self.assertEqual(first.blocks[0].status, second.blocks[0].status)
        self.assertEqual(first.blocks[0].displacement_percent, second.blocks[0].displacement_percent)
        self.assertNotEqual(first.blocks[0].origin_high, second.blocks[0].origin_high)

    def test_08_doji_origin_rejected(self) -> None:
        result = analyze([bar(0, "100", "100"), bar(1, "100", "102")])
        self.assertEqual(result.pairs[0].disposition, REJECTED_ORIGIN_DOJI)

    def test_09_doji_impulse_rejected(self) -> None:
        result = analyze([bar(0, "100", "99"), bar(1, "100", "100")])
        self.assertEqual(result.pairs[0].disposition, REJECTED_IMPULSE_DOJI)

    def test_10_same_direction_rejected(self) -> None:
        result = analyze([bar(0, "100", "101"), bar(1, "101", "104")])
        self.assertEqual(result.pairs[0].disposition, REJECTED_SAME_DIRECTION)
        self.assertIn(REJECTED_DIRECTION_MISMATCH, (
            "REJECTED_DIRECTION_MISMATCH",
        ))

    def test_11_non_adjacent_candle_cannot_qualify(self) -> None:
        result = analyze([bar(0, "100", "99"), bar(1, "99", "99.2"), bar(2, "99.2", "120")])
        self.assertEqual(result.blocks, [])
        self.assertEqual(result.pairs[0].disposition, REJECTED_BULLISH_BELOW)
        self.assertEqual(result.pairs[1].disposition, REJECTED_SAME_DIRECTION)

    def test_12_impulse_cannot_mitigate_itself(self) -> None:
        result = analyze([bar(0, "100", "90"), bar(1, "95", "101")])
        block = result.blocks[0]
        self.assertEqual(block.status, TERMINAL)
        self.assertFalse(block.ever_body_touched)
        self.assertIsNone(block.first_touch_row)

    def test_13_lifecycle_starts_after_impulse(self) -> None:
        result = analyze([bar(0, "100", "90"), bar(1, "95", "101"), bar(2, "110", "112")])
        block = result.blocks[0]
        self.assertEqual(block.status, ACTIVE_UNTOUCHED)
        self.assertFalse(block.ever_body_touched)
        self.assertGreater(result.bars[2].row, block.impulse_row)

    def test_14_bullish_body_zone(self) -> None:
        block = analyze(bullish_pair()).blocks[0]
        self.assertEqual(block.zone_upper, D("100"))
        self.assertEqual(block.zone_lower, D("99"))
        self.assertEqual(block.zone_midpoint, D("99.5"))

    def test_15_bearish_body_zone(self) -> None:
        block = analyze([bar(0, "100", "104"), bar(1, "104", "98")]).blocks[0]
        self.assertEqual(block.zone_lower, D("100"))
        self.assertEqual(block.zone_upper, D("104"))
        self.assertEqual(block.zone_midpoint, D("102"))

    def test_16_body_touch_at_near_edge(self) -> None:
        result = analyze(bullish_pair([bar(2, "100", "101")]))
        block = result.blocks[0]
        self.assertEqual(block.status, ACTIVE_BODY_MITIGATED)
        self.assertEqual(block.retrace_percent_display, D("0"))
        self.assertFalse(block.ever_invalidated)

    def test_17_body_touch_at_midpoint(self) -> None:
        result = analyze([bar(0, "100", "98"), bar(1, "98", "101"), bar(2, "99", "100")])
        block = result.blocks[0]
        self.assertEqual(block.retrace_percent_display, D("50"))
        self.assertTrue(block.ever_midpoint)
        self.assertEqual(block.status, ACTIVE_BODY_MITIGATED)

    def test_18_body_touch_at_far_edge(self) -> None:
        result = analyze([bar(0, "100", "98"), bar(1, "98", "101"), bar(2, "98", "99")])
        block = result.blocks[0]
        self.assertEqual(block.retrace_percent_display, D("100"))
        self.assertEqual(block.first_touch_close, D("99"))
        self.assertEqual(block.status, ACTIVE_BODY_MITIGATED)

    def test_19_wick_only_contact_is_not_body_mitigation(self) -> None:
        result = analyze(bullish_pair([bar(2, "102", "103", high="104", low="99.5")]))
        block = result.blocks[0]
        self.assertEqual(block.status, ACTIVE_UNTOUCHED)
        self.assertFalse(block.ever_body_touched)
        self.assertTrue(block.wick_only_contact)

    def test_20_bullish_close_at_zone_lower_does_not_invalidate(self) -> None:
        result = analyze([bar(0, "100", "98"), bar(1, "98", "101"), bar(2, "99", "98")])
        block = result.blocks[0]
        self.assertEqual(block.first_touch_close, D("98"))
        self.assertEqual(block.status, ACTIVE_BODY_MITIGATED)
        self.assertFalse(block.ever_invalidated)

    def test_21_bullish_close_below_zone_lower_invalidates(self) -> None:
        result = analyze([bar(0, "100", "98"), bar(1, "98", "101"), bar(2, "99", "97.9")])
        block = result.blocks[0]
        self.assertEqual(block.status, INVALIDATED)
        self.assertTrue(block.ever_body_touched)
        self.assertEqual(block.first_invalidation_row, 2)

    def test_22_bearish_close_at_zone_upper_does_not_invalidate(self) -> None:
        result = analyze([bar(0, "100", "102"), bar(1, "102", "99"), bar(2, "101", "102")])
        block = result.blocks[0]
        self.assertEqual(block.status, ACTIVE_BODY_MITIGATED)
        self.assertFalse(block.ever_invalidated)

    def test_23_bearish_close_above_zone_upper_invalidates(self) -> None:
        result = analyze([bar(0, "100", "102"), bar(1, "102", "99"), bar(2, "101", "102.1")])
        self.assertEqual(result.blocks[0].status, INVALIDATED)

    def test_24_gap_behavior(self) -> None:
        inside = analyze([bar(0, "100", "99"), bar(1, "99", "101"), bar(2, "98", "99.5")])
        through = analyze([bar(0, "100", "99"), bar(1, "99", "101"), bar(2, "98", "97")])
        self.assertTrue(inside.blocks[0].ever_gap)
        self.assertEqual(inside.blocks[0].status, ACTIVE_BODY_MITIGATED)
        self.assertTrue(through.blocks[0].ever_gap)
        self.assertEqual(through.blocks[0].status, INVALIDATED)

    def test_25_first_body_touch_is_selected(self) -> None:
        result = analyze(bullish_pair([bar(2, "110", "111"), bar(3, "100", "101"), bar(4, "99.5", "100")]))
        self.assertEqual(result.blocks[0].first_touch_row, 3)

    def test_26_first_invalidation_is_selected(self) -> None:
        result = analyze([
            bar(0, "100", "98"),
            bar(1, "98", "101"),
            bar(2, "110", "111"),
            bar(3, "99", "97"),
            bar(4, "96", "95"),
        ])
        self.assertEqual(result.blocks[0].first_invalidation_row, 3)

    def test_27_mitigated_then_invalidated_history_is_preserved(self) -> None:
        result = analyze([bar(0, "100", "98"), bar(1, "98", "101"), bar(2, "99", "99.5"), bar(3, "99", "97")])
        block = result.blocks[0]
        self.assertEqual(block.first_touch_row, 2)
        self.assertEqual(block.first_invalidation_row, 3)
        self.assertTrue(block.ever_body_touched)
        self.assertEqual(block.status, INVALIDATED)

    def test_28_active_untouched(self) -> None:
        result = analyze(bullish_pair([bar(2, "110", "112"), bar(3, "112", "113")]))
        self.assertEqual(result.blocks[0].status, ACTIVE_UNTOUCHED)
        self.assertGreater(len(result.bars), result.blocks[0].impulse_row + 1)

    def test_29_terminal_classification(self) -> None:
        result = analyze(bullish_pair())
        self.assertEqual(result.blocks[0].status, TERMINAL)
        self.assertEqual(result.blocks[0].impulse_row, len(result.bars) - 1)

    def test_30_candidate_pair_count_is_5933(self) -> None:
        price = D("100")
        bars = []
        for index in range(5934):
            close = price + D("0.01") if index % 2 == 0 else price - D("0.01")
            bars.append(bar(index, price, close))
            price = close
        result = analyze(bars)
        self.assertEqual(len(result.pairs), 5933)
        self.assertEqual(len(result.bars) - 1, 5933)

    def test_31_candidate_reconciliation(self) -> None:
        result = analyze(bullish_pair([bar(2, "110", "111"), bar(3, "111", "109")]))
        audit_result(result)
        qualified = sum(1 for pair in result.pairs if pair.disposition.startswith("QUALIFIED"))
        rejected = sum(1 for pair in result.pairs if pair.disposition.startswith("REJECTED"))
        self.assertEqual(qualified + rejected, len(result.pairs))
        self.assertEqual(qualified, len(result.blocks))

    def test_32_duplicate_qualified_keys_are_zero(self) -> None:
        result = analyze([bar(0, "100", "99"), bar(1, "99", "102"), bar(2, "102", "104"), bar(3, "104", "100")])
        audit_result(result)
        keys = [(block.origin_row, block.impulse_row, block.direction) for block in result.blocks]
        self.assertEqual(len(keys), len(set(keys)))

    def test_33_ids_are_deterministic(self) -> None:
        bars = [bar(0, "100", "99"), bar(1, "99", "102"), bar(2, "102", "104"), bar(3, "104", "100")]
        first = analyze(bars)
        second = analyze(bars)
        self.assertEqual([block.order_block_id for block in first.blocks], [block.order_block_id for block in second.blocks])
        self.assertTrue(first.blocks[0].order_block_id.startswith("OB4H-BULL-"))
        self.assertTrue(any(block.order_block_id.startswith("OB4H-BEAR-") for block in first.blocks))

    def test_34_close_structure_is_descriptive_only(self) -> None:
        bars = [bar(index, "150", "149") for index in range(20)]
        bars.append(bar(20, "100", "99"))
        bars.append(bar(21, "99", "101"))
        result = analyze(bars)
        block = result.blocks[0]
        self.assertEqual(block.close_structure_break, "FALSE")
        self.assertEqual(block.direction, "BULLISH")
        self.assertGreaterEqual(block.displacement_percent, D("1"))

    def test_35_quality_score_does_not_reject_qualified_blocks(self) -> None:
        result = analyze(bullish_pair())
        block = result.blocks[0]
        self.assertLess(block.quality_score, D("30"))
        self.assertEqual(len(result.blocks), 1)

    def test_36_confirmation_has_no_lookahead(self) -> None:
        early = analyze(bullish_pair())
        later = analyze(bullish_pair([bar(2, "99.2", "99.4")]))
        self.assertEqual(early.blocks[0].confirmed_at, later.blocks[0].confirmed_at)
        self.assertEqual(early.blocks[0].displacement_percent, later.blocks[0].displacement_percent)
        self.assertEqual(early.blocks[0].zone_lower, later.blocks[0].zone_lower)
        self.assertNotEqual(early.blocks[0].status, later.blocks[0].status)

    def test_37_utc_and_turkey_conversion(self) -> None:
        moment = datetime(2026, 9, 15, 20, tzinfo=timezone.utc)
        self.assertEqual(iso_utc(moment), "2026-09-15T20:00:00Z")
        self.assertEqual(iso_turkey(moment), "2026-09-15T23:00:00+03:00")

    def test_38_dynamic_revision_allocation(self) -> None:
        directory = self.temp_dir / "revisions"
        directory.mkdir()
        (directory / "BTCUSDT_4H_Order_Blocks_rev00.xlsx").write_bytes(b"PK")
        (directory / "BTCUSDT_4H_Order_Blocks_rev02.xlsx").write_bytes(b"PK")
        (directory / "BTCUSDT_MTF_Open_Liquidity_rev09.xlsx").write_bytes(b"PK")
        found = scan_revisions(directory)
        self.assertEqual(found, [0, 2])
        self.assertEqual(next_revision([]), 0)
        self.assertEqual(next_revision(found), 3)
        self.assertEqual(revision_filename(3), "BTCUSDT_4H_Order_Blocks_rev03.xlsx")

    def test_39_existing_workbook_is_not_overwritten(self) -> None:
        directory = self.temp_dir / "keep"
        directory.mkdir()
        path = directory / "BTCUSDT_4H_Order_Blocks_rev00.xlsx"
        path.write_bytes(b"keep-me")
        found = scan_revisions(directory)
        self.assertNotEqual(revision_filename(next_revision(found)), path.name)
        self.assertEqual(path.read_bytes(), b"keep-me")

    def test_40_genuine_xlsx(self) -> None:
        self.assertTrue(zipfile.is_zipfile(self.workbook_path))
        self.assertEqual(self.workbook_path.read_bytes()[:4], b"PK\x03\x04")
        self.assertIsNone(zipfile.ZipFile(self.workbook_path).testzip())

    def test_41_required_ooxml_members(self) -> None:
        names = set(zipfile.ZipFile(self.workbook_path).namelist())
        for member in ("[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "xl/theme/theme1.xml"):
            self.assertIn(member, names)

    def test_42_xml_parses(self) -> None:
        import xml.etree.ElementTree as ET

        with zipfile.ZipFile(self.workbook_path) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ET.fromstring(archive.read(name))

    def test_43_normal_openpyxl_reopen(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=False, data_only=False, keep_links=False)
        try:
            self.assertEqual(len(workbook.sheetnames), 13)
        finally:
            workbook.close()

    def test_44_streaming_reopen(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=True, keep_links=False)
        try:
            seen = 0
            for worksheet in workbook.worksheets:
                for _row in worksheet.iter_rows():
                    seen += 1
            self.assertGreater(seen, 13)
        finally:
            workbook.close()

    def test_45_thirteen_sheets_in_order(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=True)
        try:
            self.assertEqual(tuple(workbook.sheetnames), WORKBOOK_SHEETS)
        finally:
            workbook.close()

    def test_46_excel_safe_timestamps(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=False)
        try:
            value = workbook["All Order Blocks"]["F2"].value
            self.assertIsInstance(value, str)
            self.assertTrue(value.endswith("Z"))
            self.assertNotIn("+00:00", value)
        finally:
            workbook.close()

    def test_47_decimal_conversion_is_excel_safe(self) -> None:
        self.assertEqual(excel_value(D("101.25")), 101.25)
        self.assertIsInstance(excel_value(D("101.25")), float)

    def test_48_nan_and_infinity_become_blank(self) -> None:
        self.assertIsNone(excel_value(float("nan")))
        self.assertIsNone(excel_value(float("inf")))
        self.assertIsNone(excel_value(float("-inf")))
        self.assertIsNone(excel_value(D("NaN")))
        self.assertEqual(excel_value("a\x00b"), "a b")
        self.assertEqual(len(excel_value("x" * 40000)), 32767)

    def test_49_table_names_are_unique(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=False, data_only=False)
        try:
            names = [table.name for worksheet in workbook.worksheets for table in worksheet.tables.values()]
            self.assertEqual(len(names), len(set(names)))
            self.assertIn("tblBullishOrderBlocks", names)
            self.assertIn("tblCandidateAudit", names)
        finally:
            workbook.close()

    def test_50_table_ranges_are_valid_and_empty_sheets_have_no_table(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=False, data_only=False)
        try:
            untouched = workbook["Active Untouched"]
            self.assertEqual(untouched["A2"].value, "NO_RESULTS")
            self.assertEqual(len(untouched.tables), 0)
            self.assertTrue(untouched.auto_filter.ref)
            for worksheet in workbook.worksheets:
                for table in worksheet.tables.values():
                    self.assertIn(":", table.ref)
            self.assertEqual(len(workbook["Executive Summary"].tables), 0)
            self.assertEqual(len(workbook["Diagnostics"].tables), 0)
            self.assertEqual(len(workbook["README"].tables), 0)
        finally:
            workbook.close()

    def test_51_no_macros(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=False, data_only=False)
        try:
            self.assertIsNone(workbook.vba_archive)
        finally:
            workbook.close()
        self.assertFalse(any("vba" in name.lower() for name in zipfile.ZipFile(self.workbook_path).namelist()))

    def test_52_no_external_links(self) -> None:
        workbook = load_workbook(self.workbook_path, read_only=False, data_only=False)
        try:
            self.assertEqual(list(workbook._external_links), [])
        finally:
            workbook.close()
        self.assertFalse(any("externalLink" in name for name in zipfile.ZipFile(self.workbook_path).namelist()))

    def test_53_final_file_reopens_after_atomic_move(self) -> None:
        destination = self.temp_dir / "moved.xlsx"
        os.replace(self.workbook_path, destination)
        validate_ooxml_workbook(destination)
        os.replace(destination, self.workbook_path)

    def test_54_source_fingerprint_is_stable(self) -> None:
        path = self.temp_dir / "source.parquet"
        path.write_bytes(b"parquet-bytes")
        self.assertEqual(fingerprint_file(path), fingerprint_file(path))
        analyze(bullish_pair())
        self.assertEqual(path.read_bytes(), b"parquet-bytes")

    def test_55_existing_detector_files_are_outside_this_package(self) -> None:
        package = CONFIG.package_dir.resolve()
        for name in ("primary_range_v1", "hierarchical_swing_v4", "open_liquidity_v1", "price_touch_hierarchy_v1"):
            other = PROJECT_ROOT / "detectors" / name
            self.assertTrue(other.exists())
            self.assertFalse(str(other).startswith(str(package)))
        target = PROJECT_ROOT / "detectors" / "open_liquidity_v1" / "__init__.py"
        before = hashlib.sha256(target.read_bytes()).hexdigest()
        analyze(bullish_pair())
        self.assertEqual(hashlib.sha256(target.read_bytes()).hexdigest(), before)

    def test_56_no_separate_project(self) -> None:
        self.assertEqual(CONFIG.package_dir.parent.name, "detectors")
        self.assertEqual(CONFIG.package_dir.parents[1], PROJECT_ROOT)
        self.assertTrue(str(CONFIG.package_dir).startswith(str(PROJECT_ROOT)))

    def test_57_python_writes_stay_inside_the_package(self) -> None:
        with self.assertRaises(PermissionError):
            assert_python_write_allowed(Path(r"C:\Users\oranb\Desktop\order_block.py"))
        assert_python_write_allowed(CONFIG.package_dir / "engine.py")

    def test_58_two_runs_are_deterministic(self) -> None:
        bars = [bar(0, "100", "99"), bar(1, "99", "103"), bar(2, "103", "104"), bar(3, "104", "101"), bar(4, "100", "99.5")]
        first = analyze(bars)
        second = analyze(bars)
        audit_result(first)
        self.assertEqual(
            [(block.order_block_id, block.status, str(block.displacement_percent), str(block.zone_lower), str(block.zone_upper)) for block in first.blocks],
            [(block.order_block_id, block.status, str(block.displacement_percent), str(block.zone_lower), str(block.zone_upper)) for block in second.blocks],
        )
        self.assertEqual(DETECTOR_VERSION, "BTCUSDT_4H_BODY_ORDER_BLOCK_V1")


if __name__ == "__main__":
    unittest.main()
