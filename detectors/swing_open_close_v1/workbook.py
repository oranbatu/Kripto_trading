"""Conservative Open XML workbook for the Swing Open/Close detector."""
from __future__ import annotations

import math
import re
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.swing_open_close_v1.engine import (
    COMPACT_SWING,
    CONFIRMED_COMPACT,
    CONFIRMED_STANDARD,
    CROSSED_RETURN,
    EXACT_RETURN,
    STANDARD_SWING,
    SWING_HIGH,
    SWING_LOW,
    AnalysisResult,
    Swing,
    iso_turkey,
    iso_utc,
)
from detectors.swing_open_close_v1.swing_config import (
    COMPACT_WIDTH,
    DETECTOR_VERSION,
    DIRECTION_PRIORITY,
    EQUALITY_TOLERANCE_PERCENT,
    PLAIN_RANGE_SHEETS,
    STANDARD_WIDTH,
    TABLE_NAMES,
    WORKBOOK_SHEETS,
)

PRICE = "0.00000000"
PERCENT_POINTS = '0.00000000"%"'
EXCEL_MAX_TEXT = 32767
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
OOXML_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_EMPTY_TYPED_CELL = re.compile(br'<c\b[^>]*\bt="(?:n|b|s|str|inlineStr|d|e)"[^>]*/>')
LARGE_SHEETS = frozenset({"Formation Candles", "Candidate Audit"})


def _solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


HEADER_FILL = _solid("1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
HEADER_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
HIGH_FILL = _solid("F4CCCC")
LOW_FILL = _solid("D9EAD3")
COMPACT_FILL = _solid("C9DAF8")
STANDARD_FILL = _solid("FCE5CD")
EXACT_FILL = _solid("FFF2CC")
CROSS_FILL = _solid("D0E2F3")
FAIL_FILL = _solid("E6B8B7")
WRAP = Alignment(wrap_text=True, vertical="center")


def write_workbook(path: Path, result: AnalysisResult, context: dict[str, Any], readme: str) -> None:
    workbook = Workbook()
    sheets = {
        "Executive Summary": _executive(result, context),
        "Confirmed Swings": _swings(result.swings),
        "Swing Highs": _swings([item for item in result.swings if item.direction == SWING_HIGH]),
        "Swing Lows": _swings([item for item in result.swings if item.direction == SWING_LOW]),
        "Compact Swings": _swings([item for item in result.swings if item.formation_class == COMPACT_SWING]),
        "Standard Swings": _swings([item for item in result.swings if item.formation_class == STANDARD_SWING]),
        "Formation Candles": _formation(result),
        "Candidate Audit": _audit(result),
        "Overlap Analysis": _overlaps(result.swings),
        "Forward Evaluation": _forward(result.swings),
        "Parameters": _parameters(),
        "Diagnostics": context["diagnostic_rows"],
        "README": _readme_rows(readme),
    }
    active = workbook.active
    active.title = WORKBOOK_SHEETS[0]
    for name in WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(name)
    for name in WORKBOOK_SHEETS:
        _write_sheet(workbook[name], sheets[name], None if name in PLAIN_RANGE_SHEETS else TABLE_NAMES.get(name))
    _highlight(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def validate_saved_workbook(path: Path, expected_attempts: int) -> None:
    validate_ooxml_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = workbook["Candidate Audit"].max_row - 1
        if rows != expected_attempts:
            raise RuntimeError(f"candidate audit rows {rows} != {expected_attempts}")
    finally:
        workbook.close()


def validate_ooxml_workbook(path: Path) -> None:
    payload = path.read_bytes()
    if payload[:4] != b"PK\x03\x04":
        raise RuntimeError("workbook does not begin with the ZIP signature")
    if path.suffix.lower() != ".xlsx":
        raise RuntimeError("workbook extension is not .xlsx")
    if not zipfile.is_zipfile(path):
        raise RuntimeError("zipfile.is_zipfile returned false")
    required = (
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/worksheets/sheet1.xml",
    )
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("ZipFile.testzip found a corrupt member")
        names = set(archive.namelist())
        for member in required:
            if member not in names:
                raise RuntimeError(f"missing OOXML member {member}")
        if any("vba" in name.lower() for name in names):
            raise RuntimeError("workbook contains a VBA part")
        if any("externalLink" in name for name in names):
            raise RuntimeError("workbook contains an external link")
        for name in sorted(names):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                sheet_payload = archive.read(name)
                if _EMPTY_TYPED_CELL.search(sheet_payload):
                    raise RuntimeError(f"{name} contains a typed cell with no value")
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise RuntimeError(f"XML parse failed for {name}: {exc}") from exc
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                    _validate_sheet_xml(name, root)
        _validate_relationships(archive)
    _validate_openpyxl_workbook(path)


def _validate_sheet_xml(name: str, root: ET.Element) -> None:
    dimension = root.find("m:dimension", OOXML_NS)
    auto_filter = root.find("m:autoFilter", OOXML_NS)
    if auto_filter is not None and dimension is not None:
        if auto_filter.attrib.get("ref", "") != dimension.attrib.get("ref", ""):
            raise RuntimeError(f"{name} autoFilter does not match the used range")


def _validate_relationships(archive: zipfile.ZipFile) -> None:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib.get("Id"): item.attrib.get("Target", "") for item in relationships}
    for sheet in root.findall("m:sheets/m:sheet", OOXML_NS):
        title = sheet.attrib.get("name", "")
        if not title or len(title) > 31:
            raise RuntimeError(f"invalid sheet name {title!r}")
        relation = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = targets.get(relation, "").lstrip("/")
        member = target if target.startswith("xl/") else "xl/" + target
        if member not in archive.namelist():
            raise RuntimeError(f"sheet {title} references missing part {member}")


def _validate_openpyxl_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError(f"sheet order {workbook.sheetnames}")
        if workbook.vba_archive is not None:
            raise RuntimeError("workbook contains macros")
        if getattr(workbook, "_external_links", []):
            raise RuntimeError("workbook contains external links")
        names: list[str] = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > EXCEL_MAX_ROWS or worksheet.max_column > EXCEL_MAX_COLUMNS:
                raise RuntimeError(f"{worksheet.title} exceeds an Excel worksheet limit")
            if worksheet.title in LARGE_SHEETS or worksheet.title in PLAIN_RANGE_SHEETS:
                if worksheet.tables:
                    raise RuntimeError(f"{worksheet.title} must stay a plain filtered range")
            expected = TABLE_NAMES.get(worksheet.title)
            if expected is not None:
                if worksheet["A2"].value == "NO_RESULTS":
                    if worksheet.tables:
                        raise RuntimeError(f"{worksheet.title} created a table for NO_RESULTS")
                elif not worksheet.tables:
                    raise RuntimeError(f"{worksheet.title} is missing its Excel table")
            for table in worksheet.tables.values():
                names.append(table.name)
                _validate_table(worksheet, table)
        if len(names) != len(set(names)):
            raise RuntimeError(f"duplicate table names {names}")
    finally:
        workbook.close()
    streamed = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for worksheet in streamed.worksheets:
            for _row in worksheet.iter_rows():
                pass
    finally:
        streamed.close()


def _validate_table(worksheet, table) -> None:
    from openpyxl.utils import coordinate_to_tuple

    if not table.name or not (table.name[0].isalpha() or table.name[0] == "_"):
        raise RuntimeError(f"invalid table name {table.name}")
    start, end = table.ref.split(":")
    start_row, start_col = coordinate_to_tuple(start)
    end_row, end_col = coordinate_to_tuple(end)
    if end_row <= start_row or end_row > worksheet.max_row or end_col > worksheet.max_column:
        raise RuntimeError(f"table {table.name} reference {table.ref} is invalid")
    headers = [worksheet.cell(start_row, col).value for col in range(start_col, end_col + 1)]
    if any(value in (None, "") for value in headers) or len(headers) != len(set(headers)):
        raise RuntimeError(f"table {table.name} headers are blank or duplicated")


def round_trip_workbook(source: Path, copy_path: Path) -> None:
    workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    try:
        sheets = tuple(workbook.sheetnames)
        summary = [tuple(row) for row in workbook["Executive Summary"].iter_rows(values_only=True)]
        dimensions = {worksheet.title: worksheet.dimensions for worksheet in workbook.worksheets}
        headers = {worksheet.title: tuple(cell.value for cell in worksheet[1]) for worksheet in workbook.worksheets}
        workbook.save(copy_path)
    finally:
        workbook.close()
    reopened = load_workbook(copy_path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(reopened.sheetnames) != sheets:
            raise RuntimeError("round trip changed the sheet order")
        if [tuple(row) for row in reopened["Executive Summary"].iter_rows(values_only=True)] != summary:
            raise RuntimeError("round trip changed the executive summary")
        for worksheet in reopened.worksheets:
            if worksheet.dimensions != dimensions[worksheet.title]:
                raise RuntimeError(f"round trip changed dimensions for {worksheet.title}")
            if tuple(cell.value for cell in worksheet[1]) != headers[worksheet.title]:
                raise RuntimeError(f"round trip changed headers for {worksheet.title}")
    finally:
        reopened.close()


def _write_sheet(ws, rows: Sequence[Sequence[Any]], table_name: str | None) -> None:
    prepared = _prepare_rows(rows)
    for row in prepared:
        ws.append(row)
    if not prepared:
        return
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = HEADER_BORDER
    for index, header in enumerate(prepared[0], start=1):
        label = header if isinstance(header, str) else ""
        ws.column_dimensions[get_column_letter(index)].width = min(42, max(12, len(label) + 2))
    ws.row_dimensions[1].height = 30
    if len(prepared) >= 2:
        ref = f"A1:{get_column_letter(len(prepared[0]))}{len(prepared)}"
        ws.auto_filter.ref = ref
        if table_name and prepared[1][0] != "NO_RESULTS":
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    if ws.max_row <= 20000:
        _apply_number_formats(ws)


def _prepare_rows(rows: Sequence[Sequence[Any]]) -> list[list[Any]]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    prepared = []
    for row in rows:
        padded = list(row) + [None] * (width - len(row))
        prepared.append([excel_value(value) for value in padded])
    if len(prepared) == 1:
        prepared.append(["NO_RESULTS"] + [None] * (width - 1))
    return prepared


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        return float(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, datetime):
        return iso_utc(value)
    text = _clean_text(value if isinstance(value, str) else str(value))
    return text[:EXCEL_MAX_TEXT] or None


def _clean_text(text: str) -> str:
    return "".join(char if char in "\t\n\r" or ord(char) >= 32 else " " for char in text)


def _apply_number_formats(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    for index, name in enumerate(headers, start=1):
        if not isinstance(name, str):
            continue
        if any(token in name for token in ("UTC", "Turkey", "Row", "Count", "Label", "Class", "Type", "Status", "Flag", "Version", "Bars", "Hours", "Days", "Score", "Trades", "Position", "Sequence", "View", "Reason", "Role", "ID", "Id")):
            continue
        if "Percent" in name:
            number_format = PERCENT_POINTS
        elif any(token in name for token in ("Price", "Open", "High", "Low", "Close", "Reference", "Extreme", "Volume", "Overshoot", "Movement", "Excursion")):
            number_format = PRICE
        else:
            continue
        for cell in ws.iter_rows(min_row=2, min_col=index, max_col=index, max_row=ws.max_row):
            if isinstance(cell[0].value, (int, float)) and not isinstance(cell[0].value, bool):
                cell[0].number_format = number_format


def _highlight(workbook: Workbook) -> None:
    fills = {
        SWING_HIGH: HIGH_FILL,
        SWING_LOW: LOW_FILL,
        COMPACT_SWING: COMPACT_FILL,
        STANDARD_SWING: STANDARD_FILL,
        EXACT_RETURN: EXACT_FILL,
        CROSSED_RETURN: CROSS_FILL,
        "REJECTED_WIDTH_BELOW_3_50": FAIL_FILL,
        "REJECTED_WIDTH_BELOW_2_00": FAIL_FILL,
    }
    for worksheet in workbook.worksheets:
        if worksheet.title in LARGE_SHEETS or worksheet.max_row < 2 or worksheet.max_row > 2500:
            continue
        headers = {cell.value: cell.column for cell in worksheet[1]}
        for header in ("Swing Type", "Formation Class", "Completion Type", "Status", "Overlap Classification"):
            column = headers.get(header)
            if not column:
                continue
            for cell in worksheet.iter_rows(min_row=2, min_col=column, max_col=column, max_row=worksheet.max_row):
                fill = fills.get(cell[0].value)
                if fill is not None:
                    cell[0].fill = fill


def _text(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return format(value, "f")


def _time(value: datetime | None, turkey: bool = False) -> str | None:
    if value is None:
        return None
    return iso_turkey(value) if turkey else iso_utc(value)


def _swings(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[name for name, _getter in SWING_COLUMNS]]
    for swing in swings:
        rows.append([getter(swing) for _name, getter in SWING_COLUMNS])
    return rows


def _formation(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Swing ID", "Swing Type", "Formation Class", "Candle Role", "Formation Position",
        "Interior Position", "Open UTC", "Open Turkey", "Close UTC", "Close Turkey",
        "Open", "High", "Low", "Close", "Volume", "Quote Volume", "Trades",
        "Taker Buy Base", "Taker Buy Quote", "Is Swing Open", "Is Extreme", "Is Swing Close",
        "Is Plateau Member", "Is Representative Extreme", "Distance From Reference",
        "Distance From Extreme", "Known At UTC", "Timing Class",
    ]]
    for swing in result.swings:
        for offset, row in enumerate(range(swing.open_row, swing.close_row + 1)):
            bar = result.bars[row]
            interior_position = None if row in (swing.open_row, swing.close_row) else row - swing.open_row
            if row == swing.open_row:
                role = "SWING_OPEN"
            elif row == swing.close_row:
                role = "SWING_CLOSE"
            elif row == swing.extreme_row:
                role = "EXTREME"
            else:
                role = "INTERIOR"
            rows.append([
                swing.swing_id, swing.direction, swing.formation_class, role, offset + 1, interior_position,
                iso_utc(bar.open_time), iso_turkey(bar.open_time), iso_utc(bar.close_time), iso_turkey(bar.close_time),
                bar.open, bar.high, bar.low, bar.close, bar.volume, bar.quote_volume, bar.trades,
                bar.taker_base, bar.taker_quote, row == swing.open_row, row in swing.equal_extreme_rows,
                row == swing.close_row, swing.plateau_start <= row <= swing.plateau_end, row == swing.extreme_row,
                bar.close - swing.reference, bar.close - swing.extreme_price, iso_utc(bar.close_time),
                "KNOWN_AT_SWING_CONFIRMATION",
            ])
    return rows


def _audit(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Attempt ID", "Direction", "Swing Open Row", "Swing Open UTC", "Swing Open Turkey",
        "Reference Price", "Width Achieved Row", "Width Achieved UTC", "Width Achieved Turkey",
        "First Return Row", "First Return UTC", "First Return Turkey", "Extreme Row", "Extreme Price",
        "Interior Count", "Formation Class", "Required Width Percent", "Open To Extreme Percent",
        "Close To Extreme Percent", "Unrounded Open To Extreme Percent", "Unrounded Close To Extreme Percent",
        "Status", "Confirmed Swing ID", "Rejection Reason", "Terminal Flag", "Conflict Flag",
        "Conflict Resolution", "Detector Version",
    ]]
    for attempt in result.attempts:
        rows.append([
            attempt.attempt_id, attempt.direction, attempt.open_row, iso_utc(attempt.open_time), iso_turkey(attempt.open_time),
            attempt.reference, attempt.width_achieved_row, _time(attempt.width_achieved_time), _time(attempt.width_achieved_time, True),
            attempt.close_row, _time(attempt.close_time), _time(attempt.close_time, True), attempt.extreme_row, attempt.extreme_price,
            attempt.interior_count, attempt.formation_class, attempt.required_width, attempt.open_to_extreme_percent,
            attempt.close_to_extreme_percent, _text(attempt.open_to_extreme_percent), _text(attempt.close_to_extreme_percent),
            attempt.status, attempt.swing_id if attempt.status in {CONFIRMED_COMPACT, CONFIRMED_STANDARD} else None,
            "NOT_APPLICABLE" if attempt.status in {CONFIRMED_COMPACT, CONFIRMED_STANDARD} else attempt.status,
            attempt.terminal, attempt.conflict, attempt.conflict_resolution, DETECTOR_VERSION,
        ])
    return rows


def _overlaps(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[
        "Swing ID", "Swing Type", "Overlap Classification", "Overlap Count", "Same Type Overlap Count",
        "Opposite Type Overlap Count", "Contained Swing Count", "Parent Swing IDs", "Child Swing IDs",
        "Shared Candle Count", "Shared Candle Percent", "Timing Class",
    ]]
    for swing in swings:
        rows.append([
            swing.swing_id, swing.direction, swing.overlap_class, swing.overlap_count, swing.same_type_overlap,
            swing.opposite_type_overlap, swing.contained_count, swing.parent_ids, swing.child_ids,
            swing.shared_candles, swing.shared_percent, "EX_POST_OVERLAP_METRIC",
        ])
    return rows


def _forward(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[
        "Swing ID", "Swing Type", "Next Same Type", "Next Opposite Type", "Bars To Next Swing",
        "Favorable Excursion", "Adverse Excursion", "First Break Row", "First Close Beyond Row",
        "Retest Extreme Row", "Retest Reference Row", "Forward Data Censored", "Forward Censoring Reason",
        "Timing Class",
    ]]
    for swing in swings:
        rows.append([
            swing.swing_id, swing.direction, swing.next_same_id, swing.next_opposite_id, swing.bars_to_next,
            swing.favorable_excursion, swing.adverse_excursion, swing.first_break_row, swing.first_close_beyond_row,
            swing.retest_extreme_row, swing.retest_reference_row, swing.forward_censored, swing.forward_censor_reason,
            "EX_POST_FORWARD_METRIC",
        ])
    return rows


def _executive(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    highs = [item for item in result.swings if item.direction == SWING_HIGH]
    lows = [item for item in result.swings if item.direction == SWING_LOW]
    compact = [item for item in result.swings if item.formation_class == COMPACT_SWING]
    standard = [item for item in result.swings if item.formation_class == STANDARD_SWING]
    statuses: dict[str, int] = {}
    for attempt in result.attempts:
        statuses[attempt.status] = statuses.get(attempt.status, 0) + 1
    widths = [min(item.open_to_extreme_percent, item.close_to_extreme_percent) for item in result.swings]
    rows = [
        ["Metric", "Value", "Notes"],
        ["Symbol", "BTCUSDT", "BINANCE:BTCUSDT.P"],
        ["Market", "Binance USD-M PERPETUAL", "futures/um"],
        ["Timeframe", "1h", "completed candles only"],
        ["Detector Version", DETECTOR_VERSION, "SWING_OPEN_CANDLE_OPEN"],
        ["Workbook Revision", context.get("revision_name", ""), ""],
        ["Source Rows", len(result.bars), context.get("source_validation", "")],
        ["Directional Attempts", len(result.attempts), "one row per open and direction"],
        ["Confirmed Swings", len(result.swings), ""],
        ["Swing Highs", len(highs), ""],
        ["Swing Lows", len(lows), ""],
        ["Compact Swings", len(compact), COMPACT_WIDTH],
        ["Standard Swings", len(standard), STANDARD_WIDTH],
        ["Exact Returns", sum(1 for item in result.swings if item.completion_type == EXACT_RETURN), ""],
        ["Crossed Returns", sum(1 for item in result.swings if item.completion_type == CROSSED_RETURN), ""],
        ["Expired Attempts", statuses.get("EXPIRED_NO_RETURN_WITHIN_MAX_DURATION", 0), ""],
        ["Terminal Attempts", statuses.get("TERMINAL_DATASET_END", 0), ""],
        ["Minimum Binding Width Percent", min(widths) if widths else None, "unrounded gate"],
        ["Maximum Binding Width Percent", max(widths) if widths else None, ""],
        ["Plateau Swings", sum(1 for item in result.swings if item.plateau_length > 1), ""],
        ["Overlapping Swings", sum(1 for item in result.swings if item.overlap_count), "retained"],
        ["Excel Validation", context.get("excel_validation", ""), ""],
    ]
    for status, count in sorted(statuses.items()):
        rows.append([f"Status {status}", count, ""])
    return rows


def _parameters() -> list[list[Any]]:
    values = [
        ("symbol", "BTCUSDT"),
        ("tradingview_symbol", "BINANCE:BTCUSDT.P"),
        ("market", "Binance USD-M PERPETUAL"),
        ("timeframe", "1h"),
        ("swing_reference", "SWING_OPEN_CANDLE_OPEN"),
        ("swing_high_extreme", "MAXIMUM_INTERIOR_HIGH"),
        ("swing_low_extreme", "MINIMUM_INTERIOR_LOW"),
        ("extreme_must_be_interior", "TRUE"),
        ("extreme_must_be_central", "FALSE"),
        ("fixed_left_bars", "NONE"),
        ("fixed_right_bars", "NONE"),
        ("compact_minimum_interior_candles", "1"),
        ("compact_maximum_interior_candles", "5"),
        ("compact_minimum_width_percent", COMPACT_WIDTH),
        ("standard_minimum_interior_candles", "6"),
        ("standard_maximum_interior_candles", "15"),
        ("standard_minimum_width_percent", STANDARD_WIDTH),
        ("both_width_measurements_must_pass", "TRUE"),
        ("swing_high_completion", "SWING_CLOSE_CLOSE_LESS_THAN_OR_EQUAL_TO_REFERENCE"),
        ("swing_low_completion", "SWING_CLOSE_CLOSE_GREATER_THAN_OR_EQUAL_TO_REFERENCE"),
        ("exact_reference_close_passes", "TRUE"),
        ("crossed_reference_close_passes", "TRUE"),
        ("first_return_close_is_binding", "TRUE"),
        ("later_close_cherry_picking", "FALSE"),
        ("plateau_representative", "LAST_EXTREME_OCCURRENCE"),
        ("overlapping_swings_retained", "TRUE"),
        ("primary_nested_filter", "FALSE"),
        ("strict_alternation_required", "FALSE"),
        ("equality_tolerance_percent", EQUALITY_TOLERANCE_PERCENT),
        ("direction_conflict_priority", DIRECTION_PRIORITY),
        ("conflict_resolution_order", "EARLIER_WIDTH_THEN_EARLIER_RETURN_THEN_LARGER_WIDTH_THEN_PRIORITY"),
        ("computational_timezone", "UTC"),
        ("display_timezone", "Europe/Istanbul"),
        ("timestamp_excel_storage", "ISO_8601_TEXT"),
        ("excel_format", "XLSX_OFFICE_OPEN_XML"),
        ("excel_writer", "OPENPYXL_3_1_5"),
        ("excel_macro_enabled", "FALSE"),
        ("excel_external_links", "FALSE"),
        ("large_sheet_mode", "PLAIN_RANGE_WITH_AUTOFILTER"),
        ("formation_candles_formal_excel_table", "FALSE"),
        ("candidate_audit_formal_excel_table", "FALSE"),
        ("path_efficiency", "CLAMPED_NET_MOVE_DIVIDED_BY_CLOSE_PATH"),
        ("detector_version", DETECTOR_VERSION),
        ("metric_timing_confirmation", "KNOWN_AT_SWING_CONFIRMATION"),
        ("metric_timing_overlap", "EX_POST_OVERLAP_METRIC"),
        ("metric_timing_forward", "EX_POST_FORWARD_METRIC"),
        ("metric_timing_width", "KNOWN_AT_WIDTH_ACHIEVEMENT"),
        ("metric_timing_open", "KNOWN_AT_SWING_OPEN_CLOSE"),
    ]
    return [["Parameter", "Value"], *([name, value] for name, value in values)]


def _readme_rows(readme: str) -> list[list[Any]]:
    lines = [line for line in readme.splitlines() if line.strip()]
    return [["README Line"], *([line] for line in lines)]


SWING_COLUMNS: list[tuple[str, Any]] = [
    ("Swing ID", lambda swing: swing.swing_id),
    ("Sequence", lambda swing: swing.sequence),
    ("Swing Type", lambda swing: swing.direction),
    ("Structure Label", lambda swing: swing.structure_label),
    ("Formation Class", lambda swing: swing.formation_class),
    ("Swing Open Row", lambda swing: swing.open_row),
    ("Extreme Row", lambda swing: swing.extreme_row),
    ("Swing Close Row", lambda swing: swing.close_row),
    ("Swing Open UTC", lambda swing: iso_utc(swing.open_time)),
    ("Swing Open Turkey", lambda swing: iso_turkey(swing.open_time)),
    ("Swing Open Close UTC", lambda swing: iso_utc(swing.open_close_time)),
    ("Swing Open Close Turkey", lambda swing: iso_turkey(swing.open_close_time)),
    ("Extreme UTC", lambda swing: iso_utc(swing.extreme_open_time)),
    ("Extreme Turkey", lambda swing: iso_turkey(swing.extreme_open_time)),
    ("Extreme Close UTC", lambda swing: iso_utc(swing.extreme_close_time)),
    ("Extreme Close Turkey", lambda swing: iso_turkey(swing.extreme_close_time)),
    ("Swing Close UTC", lambda swing: iso_utc(swing.close_open_time)),
    ("Swing Close Turkey", lambda swing: iso_turkey(swing.close_open_time)),
    ("Swing Close Close UTC", lambda swing: iso_utc(swing.confirmed_at)),
    ("Swing Close Close Turkey", lambda swing: iso_turkey(swing.confirmed_at)),
    ("Swing Open High", lambda swing: swing.open_candle_high),
    ("Swing Open Low", lambda swing: swing.open_candle_low),
    ("Swing Open Candle Close", lambda swing: swing.open_candle_close),
    ("Extreme Candle Open", lambda swing: swing.extreme_candle_open),
    ("Extreme Candle High", lambda swing: swing.extreme_candle_high),
    ("Extreme Candle Low", lambda swing: swing.extreme_candle_low),
    ("Extreme Candle Close", lambda swing: swing.extreme_candle_close),
    ("Swing Close Candle Open", lambda swing: swing.close_candle_open),
    ("Swing Close Candle High", lambda swing: swing.close_candle_high),
    ("Swing Close Candle Low", lambda swing: swing.close_candle_low),
    ("Confirmed At UTC", lambda swing: iso_utc(swing.confirmed_at)),
    ("Confirmed At Turkey", lambda swing: iso_turkey(swing.confirmed_at)),
    ("Width Achieved Row", lambda swing: swing.width_achieved_row),
    ("Reference Price", lambda swing: swing.reference),
    ("Extreme Price", lambda swing: swing.extreme_price),
    ("Swing Close Price", lambda swing: swing.close_price),
    ("Open To Extreme Price", lambda swing: swing.open_to_extreme_price),
    ("Open To Extreme Percent", lambda swing: swing.open_to_extreme_percent),
    ("Close To Extreme Price", lambda swing: swing.close_to_extreme_price),
    ("Close To Extreme Percent", lambda swing: swing.close_to_extreme_percent),
    ("Unrounded Open To Extreme Percent", lambda swing: _text(swing.open_to_extreme_percent)),
    ("Unrounded Close To Extreme Percent", lambda swing: _text(swing.close_to_extreme_percent)),
    ("Applicable Minimum Width Percent", lambda swing: swing.required_width),
    ("Both Width Rules Passed", lambda swing: True),
    ("Completion Type", lambda swing: swing.completion_type),
    ("Completion Difference Price", lambda swing: swing.completion_difference_price),
    ("Completion Difference Percent", lambda swing: swing.completion_difference_percent),
    ("Completion Overshoot Price", lambda swing: swing.overshoot_price),
    ("Completion Overshoot Percent", lambda swing: swing.overshoot_percent),
    ("Interior Candle Count", lambda swing: swing.interior_count),
    ("Total Formation Candles", lambda swing: swing.total_candles),
    ("Formation Left Bars", lambda swing: swing.left_bars),
    ("Formation Right Bars", lambda swing: swing.right_bars),
    ("Extreme Position", lambda swing: swing.extreme_position),
    ("Plateau Start Row", lambda swing: swing.plateau_start),
    ("Plateau End Row", lambda swing: swing.plateau_end),
    ("Plateau Length", lambda swing: swing.plateau_length),
    ("Equal Extreme Count", lambda swing: swing.equal_extreme_count),
    ("Duration Hours", lambda swing: swing.duration_hours),
    ("Duration Days", lambda swing: swing.duration_days),
    ("Bullish Candles", lambda swing: swing.bullish_candles),
    ("Bearish Candles", lambda swing: swing.bearish_candles),
    ("Doji Candles", lambda swing: swing.doji_candles),
    ("Path Length", lambda swing: swing.path_length),
    ("Net Close Movement", lambda swing: swing.net_close_movement),
    ("Base Volume", lambda swing: swing.base_volume),
    ("Quote Volume", lambda swing: swing.quote_volume),
    ("Trades", lambda swing: swing.trades),
    ("Quality Score", lambda swing: swing.quality_score),
    ("Quality Label", lambda swing: swing.quality_label),
    ("Overlap Classification", lambda swing: swing.overlap_class),
    ("Contained Swing Count", lambda swing: swing.contained_count),
    ("Alternating View", lambda swing: swing.alternating_view),
    ("Forward Data Censored", lambda swing: swing.forward_censored),
    ("Detector Version", lambda swing: DETECTOR_VERSION),
]


