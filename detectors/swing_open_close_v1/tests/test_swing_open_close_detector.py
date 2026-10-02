"""Required Swing Open/Close detector tests."""
from __future__ import annotations

import hashlib
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from detectors.swing_open_close_v1.engine import (
    COMPACT_THRESHOLD,
    CONFIRMED_COMPACT,
    CONFIRMED_STANDARD,
    CROSSED_RETURN,
    EXACT_RETURN,
    EXPIRED,
    REJECTED_CONFLICT,
    REJECTED_EXTREME_ON_CLOSE,
    REJECTED_EXTREME_ON_OPEN,
    REJECTED_NO_INTERIOR,
    REJECTED_WIDTH_200,
    REJECTED_WIDTH_350,
    STANDARD_THRESHOLD,
    SWING_HIGH,
    SWING_LOW,
    TERMINAL,
    Bar,
    analyze,
    audit_result,
    both_widths_pass,
    duration_rejection,
    extreme_position_rejection,
    iso_turkey,
    iso_utc,
)
from detectors.swing_open_close_v1.run_swing_open_close_detector import (
    next_revision,
    revision_filename,
    scan_revisions,
)
from detectors.swing_open_close_v1.swing_config import (
    CONFIG,
    DETECTOR_VERSION,
    PROJECT_ROOT,
    SOURCE_DATASET,
    WORKBOOK_SHEETS,
)
from detectors.swing_open_close_v1.workbook import validate_saved_workbook, write_workbook

UNIT_DIR = PROJECT_ROOT / "tmp" / "swing_open_close_v1_unit"


def bar(row: int, open_: str, high: str, low: str, close: str) -> Bar:
    start = datetime(2024, 1, 1, tzinfo=timezone.utc) + timedelta(hours=row)
    return Bar(
        row=row,
        open_time=start,
        close_time=start + timedelta(hours=1) - timedelta(milliseconds=1),
        open=Decimal(open_),
        high=Decimal(high),
        low=Decimal(low),
        close=Decimal(close),
        volume=Decimal("10"),
        quote_volume=Decimal("1000"),
        trades=3,
        taker_base=Decimal("4"),
        taker_quote=Decimal("400"),
    )


def attempt_of(result, open_row: int, direction: str):
    return next(item for item in result.attempts if item.open_row == open_row and item.direction == direction)


def highs(result):
    return [item for item in result.swings if item.direction == SWING_HIGH]


def lows(result):
    return [item for item in result.swings if item.direction == SWING_LOW]


def compact_high(extreme: str = "103.5", close: str = "100") -> list[Bar]:
    return [
        bar(0, "100", "101", "99", "100"),
        bar(1, "100", extreme, "99.5", "102"),
        bar(2, "102", "104", "98", close),
    ]


def compact_low(extreme: str = "96.5", close: str = "100") -> list[Bar]:
    return [
        bar(0, "100", "101", "99", "100"),
        bar(1, "100", "100.5", extreme, "98"),
        bar(2, "98", "101", "96", close),
    ]


def hold_until(last_row: int, extreme_high: str, final_close: str, extreme_row: int = 1) -> list[Bar]:
    rows = [bar(0, "100", "101", "99", "100")]
    for row in range(1, last_row):
        high = extreme_high if row == extreme_row else "101.2"
        rows.append(bar(row, "101", high, "99.5", "101"))
    rows.append(bar(last_row, "101", "102", "50", final_close))
    return rows


class SwingOpenCloseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        UNIT_DIR.mkdir(parents=True, exist_ok=True)
        cls.sample_bars = compact_high() + [
            bar(3, "100", "101", "99", "100"),
            bar(4, "100", "100.5", "96.5", "97"),
            bar(5, "97", "100", "96", "100"),
        ]
        for index, item in enumerate(cls.sample_bars):
            item.row = index
            item.open_time = datetime(2024, 1, 1, tzinfo=timezone.utc) + timedelta(hours=index)
            item.close_time = item.open_time + timedelta(hours=1) - timedelta(milliseconds=1)
        cls.sample = analyze(cls.sample_bars)
        audit_result(cls.sample)
        cls.sample_path = UNIT_DIR / "sample.xlsx"
        context = {
            "source_validation": "PASS",
            "provenance": {"rows": 6, "first_open_utc": "2024-01-01T00:00:00Z", "last_open_utc": "2024-01-01T05:00:00Z", "last_close_utc": "2024-01-01T05:59:59.999Z"},
            "source_fingerprints": {},
            "existing_revisions": [],
            "excel_validation": "PASS",
            "diagnostic_rows": [["Section", "Name", "Value"], ["sample", "rows", 6]],
            "revision_name": "sample",
        }
        readme = (CONFIG.package_dir / "README.md").read_text(encoding="utf-8")
        write_workbook(cls.sample_path, cls.sample, context, readme)
        validate_saved_workbook(cls.sample_path, 12)

    @classmethod
    def tearDownClass(cls) -> None:
        if UNIT_DIR.exists():
            for path in UNIT_DIR.glob("*"):
                path.unlink(missing_ok=True)

    def test_01_compact_high_one_interior(self) -> None:
        result = analyze(compact_high())
        audit_result(result)
        swing = highs(result)[0]
        self.assertEqual(swing.interior_count, 1)
        self.assertEqual(swing.formation_class, "COMPACT_SWING")
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, CONFIRMED_COMPACT)

    def test_02_compact_low_one_interior(self) -> None:
        result = analyze(compact_low())
        audit_result(result)
        swing = lows(result)[0]
        self.assertEqual(swing.interior_count, 1)
        self.assertEqual(swing.extreme_price, Decimal("96.5"))

    def test_03_compact_five_interiors(self) -> None:
        result = analyze(hold_until(6, "103.5", "100"))
        audit_result(result)
        swing = highs(result)[0]
        self.assertEqual(swing.interior_count, 5)
        self.assertEqual(swing.formation_class, "COMPACT_SWING")

    def test_04_standard_six_interiors(self) -> None:
        result = analyze(hold_until(7, "102", "100"))
        audit_result(result)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.interior_count, 6)
        self.assertEqual(swing.formation_class, "STANDARD_SWING")

    def test_05_standard_fifteen_interiors(self) -> None:
        result = analyze(hold_until(16, "102", "100"))
        audit_result(result)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.interior_count, 15)
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, CONFIRMED_STANDARD)

    def test_06_zero_interior_rejected(self) -> None:
        self.assertEqual(duration_rejection(0), REJECTED_NO_INTERIOR)
        result = analyze([bar(0, "100", "110", "90", "100"), bar(1, "100", "110", "90", "100")])
        self.assertEqual(result.swings, [])
        self.assertTrue(all(item.status == TERMINAL for item in result.attempts))

    def test_07_sixteen_interiors_rejected(self) -> None:
        self.assertEqual(duration_rejection(16), EXPIRED)
        rows = hold_until(17, "110", "100")
        rows[-1] = bar(17, "101", "102", "99", "105")
        result = analyze(rows)
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, EXPIRED)
        self.assertFalse(any(item.open_row == 0 and item.interior_count > 15 for item in result.swings))

    def test_08_compact_exact_boundary_passes(self) -> None:
        result = analyze(compact_high("103.5", "100"))
        swing = highs(result)[0]
        self.assertEqual(swing.open_to_extreme_percent, Decimal("3.5"))
        self.assertGreaterEqual(swing.close_to_extreme_percent, COMPACT_THRESHOLD)
        audit_result(result)

    def test_09_compact_below_boundary_fails(self) -> None:
        result = analyze(compact_high("103.49999999", "100"))
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, REJECTED_WIDTH_350)
        self.assertFalse(any(item.open_row == 0 for item in highs(result)))
        five = analyze(hold_until(6, "102", "100"))
        self.assertEqual(attempt_of(five, 0, SWING_HIGH).status, REJECTED_WIDTH_350)

    def test_10_standard_exact_boundary_passes(self) -> None:
        result = analyze(hold_until(7, "102", "100"))
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.open_to_extreme_percent, Decimal("2"))
        self.assertGreaterEqual(swing.open_to_extreme_percent, STANDARD_THRESHOLD)
        audit_result(result)

    def test_11_standard_below_boundary_fails(self) -> None:
        result = analyze(hold_until(7, "101.99999999", "100"))
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, REJECTED_WIDTH_200)
        self.assertFalse(any(item.open_row == 0 for item in result.swings))

    def test_12_open_width_pass_close_width_fail(self) -> None:
        self.assertFalse(both_widths_pass(Decimal("4"), Decimal("3"), COMPACT_THRESHOLD))
        self.assertTrue(both_widths_pass(Decimal("4"), Decimal("4"), COMPACT_THRESHOLD))

    def test_13_close_width_pass_open_width_fail(self) -> None:
        result = analyze(compact_high("103", "90"))
        attempt = attempt_of(result, 0, SWING_HIGH)
        self.assertEqual(attempt.status, REJECTED_WIDTH_350)
        self.assertGreaterEqual(attempt.close_to_extreme_percent, COMPACT_THRESHOLD)
        self.assertLess(attempt.open_to_extreme_percent, COMPACT_THRESHOLD)
        self.assertFalse(any(item.open_row == 0 and item.direction == SWING_HIGH for item in result.swings))

    def test_14_both_widths_required(self) -> None:
        self.assertFalse(both_widths_pass(Decimal("3.50"), Decimal("3.49999999"), COMPACT_THRESHOLD))
        self.assertFalse(both_widths_pass(Decimal("1.99999999"), Decimal("2.00"), STANDARD_THRESHOLD))
        self.assertTrue(both_widths_pass(Decimal("2.00000000"), Decimal("2.00000000"), STANDARD_THRESHOLD))

    def test_15_high_uses_maximum_interior_high(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "99", "101"),
            bar(2, "101", "105", "99", "108"),
            bar(3, "108", "109", "99", "100"),
        ]
        result = analyze(rows)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_price, Decimal("110"))
        self.assertEqual(swing.extreme_row, 1)
        audit_result(result)

    def test_16_low_uses_minimum_interior_low(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "101", "90", "99"),
            bar(2, "99", "101", "95", "92"),
            bar(3, "92", "101", "91", "100"),
        ]
        result = analyze(rows)
        swing = next(item for item in lows(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_price, Decimal("90"))
        self.assertEqual(swing.extreme_row, 1)

    def test_17_extreme_need_not_be_central(self) -> None:
        result = analyze(hold_until(6, "110", "100", extreme_row=1))
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.left_bars, 1)
        self.assertEqual(swing.right_bars, 5)
        self.assertFalse(swing.symmetric)

    def test_18_extreme_at_first_interior(self) -> None:
        result = analyze(hold_until(4, "110", "100", extreme_row=1))
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_row, swing.open_row + 1)

    def test_19_extreme_at_last_interior(self) -> None:
        result = analyze(hold_until(4, "110", "100", extreme_row=3))
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_row, swing.close_row - 1)
        self.assertEqual(swing.right_bars, 1)

    def test_20_extreme_on_open_rejected(self) -> None:
        self.assertEqual(extreme_position_rejection(0, 0, 2), REJECTED_EXTREME_ON_OPEN)
        result = analyze(compact_high())
        swing = highs(result)[0]
        self.assertGreater(swing.extreme_row, swing.open_row)
        self.assertLess(bars_high(result, 0), swing.extreme_price)

    def test_21_extreme_on_close_rejected(self) -> None:
        self.assertEqual(extreme_position_rejection(0, 2, 2), REJECTED_EXTREME_ON_CLOSE)
        result = analyze(compact_high())
        swing = highs(result)[0]
        self.assertLess(swing.extreme_row, swing.close_row)

    def test_22_equal_high_plateau(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "99", "105"),
            bar(2, "105", "110", "99", "106"),
            bar(3, "106", "104", "99", "100"),
        ]
        result = analyze(rows)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_row, 2)
        self.assertEqual(swing.plateau_start, 1)
        self.assertEqual(swing.plateau_end, 2)
        self.assertEqual(swing.plateau_length, 2)

    def test_23_equal_low_plateau(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "101", "90", "95"),
            bar(2, "95", "101", "90", "94"),
            bar(3, "94", "101", "93", "100"),
        ]
        result = analyze(rows)
        swing = next(item for item in lows(result) if item.open_row == 0)
        self.assertEqual(swing.extreme_row, 2)
        self.assertEqual(swing.plateau_length, 2)

    def test_24_last_extreme_occurrence_selected(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "99", "105"),
            bar(2, "105", "104", "99", "106"),
            bar(3, "106", "110", "99", "107"),
            bar(4, "107", "108", "99", "100"),
        ]
        result = analyze(rows)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.equal_extreme_count, 2)
        self.assertEqual(swing.extreme_row, 3)
        self.assertEqual(swing.plateau_length, 1)

    def test_25_exact_reference_confirms_high(self) -> None:
        result = analyze(compact_high("110", "100"))
        swing = highs(result)[0]
        self.assertEqual(swing.completion_type, EXACT_RETURN)
        self.assertEqual(swing.close_price, swing.reference)

    def test_26_downward_cross_confirms_high(self) -> None:
        result = analyze(compact_high("110", "99"))
        swing = highs(result)[0]
        self.assertEqual(swing.completion_type, CROSSED_RETURN)
        self.assertLess(swing.close_price, swing.reference)

    def test_27_close_above_reference_does_not_confirm_high(self) -> None:
        rows = [bar(0, "100", "101", "99", "100"), bar(1, "100", "110", "99", "105"), bar(2, "105", "111", "104", "106")]
        result = analyze(rows)
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, TERMINAL)
        self.assertFalse(any(item.open_row == 0 for item in highs(result)))

    def test_28_exact_reference_confirms_low(self) -> None:
        result = analyze(compact_low("90", "100"))
        swing = lows(result)[0]
        self.assertEqual(swing.completion_type, EXACT_RETURN)

    def test_29_upward_cross_confirms_low(self) -> None:
        result = analyze(compact_low("90", "101"))
        swing = lows(result)[0]
        self.assertEqual(swing.completion_type, CROSSED_RETURN)
        self.assertGreater(swing.close_price, swing.reference)

    def test_30_close_below_reference_does_not_confirm_low(self) -> None:
        rows = [bar(0, "100", "101", "99", "100"), bar(1, "100", "101", "90", "95"), bar(2, "95", "96", "89", "94")]
        result = analyze(rows)
        self.assertEqual(attempt_of(result, 0, SWING_LOW).status, TERMINAL)
        self.assertFalse(any(item.open_row == 0 for item in lows(result)))

    def test_31_first_return_is_binding(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "101", "99", "101"),
            bar(2, "101", "102", "98", "99"),
            bar(3, "99", "150", "98", "140"),
            bar(4, "140", "151", "99", "100"),
        ]
        result = analyze(rows)
        attempt = attempt_of(result, 0, SWING_HIGH)
        self.assertEqual(attempt.close_row, 2)
        self.assertEqual(attempt.status, REJECTED_WIDTH_350)

    def test_32_failed_first_return_cannot_be_skipped(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "101", "99", "101"),
            bar(2, "101", "102", "98", "99"),
            bar(3, "99", "150", "98", "140"),
            bar(4, "140", "151", "99", "100"),
        ]
        result = analyze(rows)
        self.assertFalse(any(item.open_row == 0 and item.direction == SWING_HIGH for item in result.swings))

    def test_33_no_return_within_fifteen_expires(self) -> None:
        rows = [bar(0, "100", "101", "99", "100")]
        for row in range(1, 18):
            rows.append(bar(row, "101", "110" if row == 1 else "102", "99.5", "105"))
        result = analyze(rows)
        attempt = attempt_of(result, 0, SWING_HIGH)
        self.assertEqual(attempt.status, EXPIRED)
        self.assertEqual(attempt.interior_count, 15)
        self.assertFalse(any(item.open_row == 0 for item in highs(result)))

    def test_34_confirmation_is_swing_close_completion(self) -> None:
        rows = compact_high()
        result = analyze(rows)
        swing = highs(result)[0]
        self.assertEqual(swing.confirmed_at, rows[swing.close_row].close_time)
        self.assertNotEqual(swing.confirmed_at, rows[swing.extreme_row].close_time)
        self.assertNotEqual(swing.confirmed_at, rows[swing.open_row].close_time)

    def test_35_no_lookahead(self) -> None:
        base = analyze(compact_high())
        extended = compact_high() + [bar(3, "100", "180", "20", "100")]
        later = analyze(extended)
        original = highs(base)[0]
        same = next(item for item in highs(later) if item.open_row == 0)
        self.assertEqual(same.close_row, original.close_row)
        self.assertEqual(same.extreme_price, original.extreme_price)
        self.assertEqual(same.confirmed_at, original.confirmed_at)
        self.assertEqual(same.open_to_extreme_percent, original.open_to_extreme_percent)

    def test_36_width_achievement_does_not_confirm(self) -> None:
        rows = [bar(0, "100", "101", "99", "100")]
        for row in range(1, 18):
            rows.append(bar(row, "101", "110" if row == 1 else "102", "99.5", "105"))
        result = analyze(rows)
        attempt = attempt_of(result, 0, SWING_HIGH)
        self.assertEqual(attempt.width_achieved_row, 1)
        self.assertEqual(attempt.status, EXPIRED)
        self.assertEqual(highs(result), [])

    def test_37_directional_states_are_independent(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "101", "99", "101"),
            bar(2, "101", "102", "98", "99"),
            bar(3, "99", "101", "96", "97"),
            bar(4, "97", "102", "95", "101"),
        ]
        result = analyze(rows)
        audit_result(result)
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, REJECTED_WIDTH_350)
        self.assertEqual(attempt_of(result, 0, SWING_LOW).status, CONFIRMED_COMPACT)
        self.assertTrue(any(item.open_row == 0 for item in lows(result)))

    def test_38_direction_conflict_is_deterministic(self) -> None:
        equal = analyze([
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "90", "100"),
            bar(2, "100", "101", "99", "100"),
        ])
        self.assertEqual(attempt_of(equal, 0, SWING_HIGH).status, CONFIRMED_COMPACT)
        self.assertEqual(attempt_of(equal, 0, SWING_LOW).status, REJECTED_CONFLICT)
        self.assertIn("SWING_HIGH_THEN_SWING_LOW", attempt_of(equal, 0, SWING_HIGH).conflict_resolution)
        wider_low = analyze([
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "104", "90", "100"),
            bar(2, "100", "101", "99", "100"),
        ])
        self.assertEqual(attempt_of(wider_low, 0, SWING_LOW).status, CONFIRMED_COMPACT)
        self.assertEqual(attempt_of(wider_low, 0, SWING_HIGH).status, REJECTED_CONFLICT)
        self.assertIn("LARGER_WIDTH", attempt_of(wider_low, 0, SWING_LOW).conflict_resolution)
        again = analyze([
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "90", "100"),
            bar(2, "100", "101", "99", "100"),
        ])
        self.assertEqual([item.status for item in again.attempts], [item.status for item in equal.attempts])

    def test_39_overlapping_swings_retained(self) -> None:
        rows = [
            bar(0, "100", "101", "99", "100"),
            bar(1, "100", "110", "99.5", "105"),
            bar(2, "100", "101", "99", "100"),
            bar(3, "100", "110", "99.5", "105"),
            bar(4, "105", "106", "99", "100"),
        ]
        result = analyze(rows)
        audit_result(result)
        owned = [item for item in highs(result) if item.open_row in (0, 2)]
        self.assertEqual(len(owned), 2)
        self.assertTrue(all(item.overlap_count > 0 for item in owned))
        self.assertEqual(len({(item.direction, item.open_row, item.extreme_row, item.close_row) for item in result.swings}), len(result.swings))

    def test_40_duplicate_confirmed_key_prohibited(self) -> None:
        result = analyze(compact_high())
        audit_result(result)
        keys = [(item.direction, item.open_row, item.extreme_row, item.close_row) for item in result.swings]
        self.assertEqual(len(keys), len(set(keys)))

    def test_41_ids_are_deterministic(self) -> None:
        first = analyze(compact_high())
        second = analyze(compact_high())
        self.assertEqual([item.swing_id for item in first.swings], [item.swing_id for item in second.swings])
        self.assertEqual([item.attempt_id for item in first.attempts], [item.attempt_id for item in second.attempts])
        self.assertTrue(first.swings[0].swing_id.startswith("SWH-"))

    def test_42_utc_and_turkey_conversion(self) -> None:
        moment = datetime(2026, 9, 15, 23, tzinfo=timezone.utc)
        self.assertEqual(iso_utc(moment), "2026-09-15T23:00:00Z")
        self.assertEqual(iso_turkey(moment), "2026-09-16T02:00:00+03:00")
        result = analyze(compact_high())
        swing = highs(result)[0]
        self.assertTrue(iso_utc(swing.open_time).endswith("Z"))
        self.assertIn("+03:00", iso_turkey(swing.open_time))

    def test_43_forward_metrics_do_not_change_confirmation(self) -> None:
        base = analyze(compact_high())
        extended = analyze(compact_high() + [bar(3, "100", "200", "10", "40")])
        original = highs(base)[0]
        same = next(item for item in highs(extended) if item.open_row == original.open_row)
        self.assertEqual(same.confirmed_at, original.confirmed_at)
        self.assertEqual(same.close_row, original.close_row)
        self.assertTrue(same.forward_censored)
        self.assertEqual(same.forward_censor_reason, "DATASET_END")

    def test_44_quality_score_does_not_filter(self) -> None:
        result = analyze(hold_until(16, "102", "50"))
        audit_result(result)
        swing = next(item for item in highs(result) if item.open_row == 0)
        self.assertEqual(swing.quality_label, "BASIC_CONFIRMED")
        self.assertLess(swing.quality_score, Decimal("30"))
        self.assertEqual(attempt_of(result, 0, SWING_HIGH).status, CONFIRMED_STANDARD)

    def test_45_one_audit_row_per_open_and_direction(self) -> None:
        result = analyze(compact_high())
        pairs = [(item.open_row, item.direction) for item in result.attempts]
        self.assertEqual(len(result.attempts), len(result.bars) * 2)
        self.assertEqual(len(pairs), len(set(pairs)))

    def test_46_dynamic_revision(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.assertEqual(next_revision(scan_revisions(root)), 0)
            (root / "BTCUSDT_1H_Swing_Structure_rev00.xlsx").write_bytes(b"a")
            (root / "BTCUSDT_1H_Swing_Structure_rev02.xlsx").write_bytes(b"b")
            (root / "BTCUSDT_4H_Order_Blocks_rev09.xlsx").write_bytes(b"c")
            (root / "BTCUSDT_MTF_Open_Liquidity_rev03.xlsx").write_bytes(b"d")
            self.assertEqual(scan_revisions(root), [0, 2])
            self.assertEqual(next_revision(scan_revisions(root)), 3)
            self.assertEqual(revision_filename(3), "BTCUSDT_1H_Swing_Structure_rev03.xlsx")

    def test_47_existing_workbook_not_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target = root / revision_filename(0)
            target.write_bytes(b"keep")
            planned = next_revision(scan_revisions(root))
            self.assertNotEqual(root / revision_filename(planned), target)
            self.assertEqual(target.read_bytes(), b"keep")

    def test_48_genuine_xlsx(self) -> None:
        self.assertEqual(self.sample_path.read_bytes()[:2], b"PK")
        self.assertTrue(zipfile.is_zipfile(self.sample_path))
        self.assertIsNone(zipfile.ZipFile(self.sample_path).testzip())

    def test_49_ooxml_validation(self) -> None:
        names = set(zipfile.ZipFile(self.sample_path).namelist())
        for member in ("[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "xl/theme/theme1.xml"):
            self.assertIn(member, names)

    def test_50_xml_parsing(self) -> None:
        import xml.etree.ElementTree as ET

        with zipfile.ZipFile(self.sample_path) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ET.fromstring(archive.read(name))

    def test_51_normal_openpyxl_reopen(self) -> None:
        workbook = load_workbook(self.sample_path, read_only=False, data_only=False, keep_links=False)
        try:
            self.assertEqual(tuple(workbook.sheetnames), WORKBOOK_SHEETS)
            self.assertIsNone(workbook.vba_archive)
        finally:
            workbook.close()

    def test_52_streaming_reopen(self) -> None:
        workbook = load_workbook(self.sample_path, read_only=True, data_only=True)
        try:
            count = 0
            for worksheet in workbook.worksheets:
                for _row in worksheet.iter_rows():
                    count += 1
            self.assertGreater(count, 13)
        finally:
            workbook.close()

    def test_53_thirteen_sheets_in_order(self) -> None:
        workbook = load_workbook(self.sample_path, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, list(WORKBOOK_SHEETS))
        finally:
            workbook.close()

    def test_54_large_sheets_use_plain_filtered_ranges(self) -> None:
        workbook = load_workbook(self.sample_path, read_only=False, data_only=False)
        try:
            for name in ("Formation Candles", "Candidate Audit", "Executive Summary", "Diagnostics", "README"):
                worksheet = workbook[name]
                self.assertEqual(list(worksheet.tables), [])
                self.assertIsNotNone(worksheet.auto_filter.ref)
            self.assertEqual(workbook["Candidate Audit"].max_row - 1, 12)
            self.assertEqual(workbook["Standard Swings"]["A2"].value, "NO_RESULTS")
            self.assertEqual(list(workbook["Standard Swings"].tables), [])
            headers = [cell.value for cell in workbook["Confirmed Swings"][1]]
            open_col = headers.index("Swing Open UTC") + 1
            confirmed_col = headers.index("Confirmed At UTC") + 1
            self.assertNotEqual(workbook["Confirmed Swings"].cell(2, open_col).value, workbook["Confirmed Swings"].cell(2, confirmed_col).value)
        finally:
            workbook.close()

    def test_55_no_macros(self) -> None:
        names = " ".join(zipfile.ZipFile(self.sample_path).namelist()).lower()
        self.assertNotIn("vba", names)

    def test_56_no_external_links(self) -> None:
        names = " ".join(zipfile.ZipFile(self.sample_path).namelist())
        self.assertNotIn("externalLink", names)

    def test_57_final_file_reopens(self) -> None:
        workbook = load_workbook(self.sample_path, read_only=False, data_only=False, keep_links=False)
        try:
            parameters = {row[0].value: row[1].value for row in workbook["Parameters"].iter_rows(min_row=2)}
            self.assertEqual(parameters["symbol"], "BTCUSDT")
            self.assertEqual(parameters["detector_version"], DETECTOR_VERSION)
            self.assertEqual(parameters["timeframe"], "1h")
        finally:
            workbook.close()

    def test_58_source_files_unchanged_by_detection(self) -> None:
        if not SOURCE_DATASET.exists():
            self.skipTest("source dataset is not mounted")
        before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in SOURCE_DATASET.rglob("*.parquet")}
        analyze(compact_high())
        after = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in SOURCE_DATASET.rglob("*.parquet")}
        self.assertEqual(before, after)

    def test_59_existing_detectors_unchanged(self) -> None:
        roots = [
            PROJECT_ROOT / "detectors" / "order_block_4h_v1",
            PROJECT_ROOT / "detectors" / "open_liquidity_v1",
            PROJECT_ROOT / "detectors" / "primary_range_v1",
        ]
        before = _tree_hash(roots)
        analyze(compact_high())
        self.assertEqual(_tree_hash(roots), before)

    def test_60_no_separate_project(self) -> None:
        self.assertEqual(CONFIG.package_dir.parent.parent, PROJECT_ROOT)
        self.assertTrue(str(CONFIG.package_dir).endswith("detectors\\swing_open_close_v1") or str(CONFIG.package_dir).endswith("detectors/swing_open_close_v1"))

    def test_61_no_python_file_outside_project(self) -> None:
        for path in CONFIG.package_dir.rglob("*.py"):
            self.assertTrue(str(path.resolve()).startswith(str(PROJECT_ROOT.resolve())))
        self.assertFalse((CONFIG.desktop_dir / "run_swing_open_close_detector.py").exists())

    def test_62_two_run_determinism(self) -> None:
        first = analyze(hold_until(7, "102", "100"))
        second = analyze(hold_until(7, "102", "100"))
        self.assertEqual(
            [(item.swing_id, item.open_row, item.extreme_row, item.close_row, str(item.extreme_price)) for item in first.swings],
            [(item.swing_id, item.open_row, item.extreme_row, item.close_row, str(item.extreme_price)) for item in second.swings],
        )
        self.assertEqual([item.status for item in first.attempts], [item.status for item in second.attempts])


def bars_high(result, open_row: int) -> Decimal:
    swing = next(item for item in highs(result) if item.open_row == open_row)
    return result.bars[swing.open_row].high


def _tree_hash(roots: list[Path]) -> dict[str, str]:
    found = {}
    for root in roots:
        if not root.exists():
            continue
        for path in sorted(root.rglob("*")):
            if path.is_file() and "__pycache__" not in path.parts:
                found[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    return found


if __name__ == "__main__":
    unittest.main()
