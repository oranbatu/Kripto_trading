#!/usr/bin/env python3
r"""Isolated Binance USD-M perpetual 1h PRIMARY range detector (Range V1 / CONFIG REV05).

Market-structure analysis only. Not a trading strategy. Not auto-registered.
Invoke:

    python C:\Users\oranb\Desktop\backtest_system\detectors\primary_range_v1\run_primary_range_detector.py --symbol BTCUSDT
"""
from __future__ import annotations

import argparse
import bisect
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from decimal import Decimal, getcontext
from pathlib import Path
from zoneinfo import ZoneInfo

os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
sys.dont_write_bytecode = True
getcontext().prec = 50

DETECTOR_DIR = Path(__file__).resolve().parent
PROJECT = DETECTOR_DIR.parents[1]
if str(PROJECT) not in sys.path:
    sys.path.insert(0, str(PROJECT))

import pyarrow as pa
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import config
from data.catalog import DataCatalog
from data.exceptions import PartitionNotFoundError, UnsupportedSymbolError
from data.loader import load_candles

DETECTOR_VERSION = "UM_PERP_1H_PRIMARY_RANGE_V1"
CONFIGURATION_VERSION = "PRIMARY_RANGE_CONFIG_REV05"
HIERARCHY_MODE = "POST_DETECTION"
TIMEFRAME = "1h"
MARKET_LABEL = "Binance USD-M USDT-margined PERPETUAL"
DEFAULT_DATA_ROOT = Path(r"C:\MarketData")
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\oranb\Desktop")
ALLOWED_SYMBOLS = tuple(config.ALLOWED_SYMBOLS)
ISTANBUL = ZoneInfo("Europe/Istanbul")
UTC = timezone.utc
ZERO = Decimal("0")
ONE = Decimal("1")
DUMMY_TIME = datetime(2024, 1, 1, 0, 0, tzinfo=UTC)

ATR_PERIOD = 14
PIVOT_LEFT = 3
PIVOT_RIGHT = 3
TOUCH_TOL_ATR = Decimal("0.25")
MIN_RANGE_BARS = 500
MIN_UPPER_TOUCHES = 2
MIN_LOWER_TOUCHES = 2
MIN_ANCHOR_TOUCHES = 2
MIN_ADDITIONAL_ELIGIBLE = 0
MIN_PAIR_SEP = 30
TOUCH_SCORE_DENOM = 4
MIN_EQ_CROSS = 2
MIN_INSIDE = Decimal("0.97")
MAX_ER = Decimal("0.30")
MAX_SLOPE = Decimal("0.35")
MIN_WIDTH_ATR = Decimal("2.0")
MAX_WIDTH_PCT = Decimal("13.0")
NORMAL_BO_ATR = Decimal("0.25")
STRONG_BO_ATR = Decimal("0.50")
DUP_OVERLAP = Decimal("0.70")
DUP_BOUND_ATR = Decimal("0.50")
MERGE_GAP = 6
MERGE_BOUND_ATR = Decimal("0.25")
CLUSTER_JOIN_GAP = 1440
CANDIDATE_EXPIRE_BARS = 2160
AWAY_ATR = Decimal("1.0")
AWAY_CLOSE_ATR = Decimal("0.5")
HOUR_MS = 3_600_000
PRIMARY_REJECTION_ORDER = [
    "duration", "upper_touches", "lower_touches", "upper_pair", "lower_pair",
    "independent", "containment", "slope", "efficiency", "eq", "width", "width_pct",
    "causality", "expired", "invalidated", "other",
]
PRICE_FMT = "0.00000000"
PCT_FMT = "0.00%"
SHEET_ORDER = [
    "High Quality Ranges", "Range Details", "Boundary Touches", "EQ Crossings",
    "Moderate Ranges", "Parameters", "Diagnostics", "README",
]
FILENAME_TEMPLATE = "{symbol}_1H_Range_Detection_{revision}.xlsx"


class PrimaryRangeDetectorError(Exception):
    """User-facing detector failure that must not write a workbook."""


# Compatibility aliases used by synthetic helpers copied from the validated runner.
FIRST_OPEN = DUMMY_TIME
LAST_OPEN = DUMMY_TIME

def iso_z(dt: datetime | None) -> str | None:
    if dt is None:
        return None
    return dt.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def turkey(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    return dt.astimezone(ISTANBUL)


def excel_clock(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    aware = dt.astimezone(dt.tzinfo)
    return datetime(aware.year, aware.month, aware.day, aware.hour, aware.minute, aware.second)


def median_dec(values: list[Decimal]) -> Decimal:
    if not values:
        raise ValueError("median of empty")
    s = sorted(values)
    n = len(s)
    mid = n // 2
    if n % 2:
        return s[mid]
    return (s[mid - 1] + s[mid]) / Decimal(2)


def mad_dec(values: list[Decimal]) -> Decimal:
    if not values:
        return ZERO
    med = median_dec(values)
    return median_dec([abs(v - med) for v in values])


def clamp01(x: Decimal) -> Decimal:
    if x < ZERO:
        return ZERO
    if x > ONE:
        return ONE
    return x


def fingerprint(path: Path) -> dict:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        while True:
            chunk = fh.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
    st = path.stat()
    return {"path": str(path), "sha256": h.hexdigest(), "size": st.st_size, "mtime_ns": st.st_mtime_ns}


def fps_equal(a: list[dict], b: list[dict]) -> bool:
    ka = [(x["path"], x["sha256"], x["size"], x["mtime_ns"]) for x in a]
    kb = [(x["path"], x["sha256"], x["size"], x["mtime_ns"]) for x in b]
    return ka == kb


def default_data_root() -> Path:
    return Path(config.data_root())


def detector_script_path() -> Path:
    return Path(__file__).resolve()


def script_sha256() -> str:
    return fingerprint(detector_script_path())["sha256"]


def new_run_id() -> str:
    return datetime.now(tz=UTC).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:8]


def normalize_symbol(symbol: str) -> str:
    if not isinstance(symbol, str) or not symbol.strip():
        raise PrimaryRangeDetectorError("symbol is required, for example BTCUSDT")
    return symbol.strip().upper()


def validate_symbol(symbol: str) -> str:
    symbol = normalize_symbol(symbol)
    if symbol not in ALLOWED_SYMBOLS:
        raise PrimaryRangeDetectorError(
            f"Unsupported symbol {symbol!r}. Pass an exact Binance USD-M perpetual symbol "
            f"such as BTCUSDT. Allowed: {', '.join(ALLOWED_SYMBOLS)}"
        )
    return symbol


def validate_timeframe(timeframe: str) -> str:
    if timeframe != TIMEFRAME:
        raise PrimaryRangeDetectorError(
            f"This detector version supports only timeframe {TIMEFRAME!r}, got {timeframe!r}"
        )
    return timeframe


def derived_1h_symbol_root(data_root: Path, symbol: str) -> Path:
    return Path(data_root) / "derived" / "binance" / "futures" / "um" / "perpetual" / "1h" / f"symbol={symbol}"


def revision_filename_pattern(symbol: str) -> re.Pattern[str]:
    return re.compile(rf"^{re.escape(symbol)}_1H_Range_Detection_rev([0-9]+)\.xlsx$", re.IGNORECASE)


def is_ignored_output_file(path: Path) -> bool:
    name = path.name
    if name.startswith("~$"):
        return True
    lowered = name.lower()
    if lowered.endswith(".tmp") or lowered.endswith(".part") or lowered.endswith(".bak"):
        return True
    if ".tmp." in lowered or lowered.endswith(".xlsx.tmp") or lowered.endswith(".partial"):
        return True
    return False


def format_revision(revision_number: int) -> str:
    if revision_number < 1:
        raise PrimaryRangeDetectorError(f"revision_number must be >= 1, got {revision_number}")
    return f"rev{revision_number:02d}"


def scan_existing_revisions(output_dir: Path, symbol: str) -> list[int]:
    pat = revision_filename_pattern(symbol)
    found: list[int] = []
    if not output_dir.is_dir():
        return found
    for path in output_dir.iterdir():
        if not path.is_file() or is_ignored_output_file(path):
            continue
        if path.suffix.lower() != ".xlsx":
            continue
        match = pat.fullmatch(path.name)
        if match:
            found.append(int(match.group(1), 10))
    return found


def next_revision_number(output_dir: Path, symbol: str) -> int:
    existing = scan_existing_revisions(output_dir, symbol)
    return 1 if not existing else max(existing) + 1


def workbook_path_for(output_dir: Path, symbol: str, revision_number: int) -> Path:
    return Path(output_dir) / FILENAME_TEMPLATE.format(symbol=symbol, revision=format_revision(revision_number))


def allocate_output_path(output_dir: Path, symbol: str) -> tuple[int, str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    n = next_revision_number(output_dir, symbol)
    path = workbook_path_for(output_dir, symbol, n)
    while path.exists():
        n += 1
        path = workbook_path_for(output_dir, symbol, n)
    return n, format_revision(n), path


def fingerprint_unrelated_project_code() -> list[dict]:
    out: list[dict] = []
    skip_parts = {".venv", "site-packages", "__pycache__", "detectors"}
    for p in sorted(PROJECT.rglob("*.py")):
        if any(part in skip_parts for part in p.parts):
            continue
        out.append(fingerprint(p))
    return out


def fingerprint_matching_workbooks(output_dir: Path, symbol: str) -> dict[str, dict]:
    pat = revision_filename_pattern(symbol)
    out: dict[str, dict] = {}
    if not Path(output_dir).is_dir():
        return out
    for path in Path(output_dir).iterdir():
        if not path.is_file() or is_ignored_output_file(path):
            continue
        if path.suffix.lower() != ".xlsx":
            continue
        if pat.fullmatch(path.name) or path.name.lower() == f"{symbol.lower()}_1h_range_detection.xlsx":
            out[str(path.resolve())] = fingerprint(path)
    return out


def active_thresholds() -> dict:
    return {
        "minimum_range_bars": MIN_RANGE_BARS,
        "minimum_upper_anchor_touches": MIN_ANCHOR_TOUCHES,
        "minimum_lower_anchor_touches": MIN_ANCHOR_TOUCHES,
        "minimum_total_eligible_upper_visits": MIN_UPPER_TOUCHES,
        "minimum_total_eligible_lower_visits": MIN_LOWER_TOUCHES,
        "minimum_additional_eligible_upper_touches": MIN_ADDITIONAL_ELIGIBLE,
        "minimum_additional_eligible_lower_touches": MIN_ADDITIONAL_ELIGIBLE,
        "minimum_qualifying_touch_pair_separation": MIN_PAIR_SEP,
        "minimum_eq_crossings": MIN_EQ_CROSS,
        "minimum_inside_close_ratio": str(MIN_INSIDE),
        "maximum_efficiency_ratio": str(MAX_ER),
        "maximum_normalized_slope": str(MAX_SLOPE),
        "minimum_normalized_width_atr": str(MIN_WIDTH_ATR),
        "maximum_range_width_percent": str(MAX_WIDTH_PCT),
        "required_final_structure_type": "PRIMARY",
        "timeframe": TIMEFRAME,
        "hierarchy_classification_mode": HIERARCHY_MODE,
        "detector_version": DETECTOR_VERSION,
        "configuration_version": CONFIGURATION_VERSION,
    }


def inspect_source_inventory(data_root: Path, symbol: str) -> dict:
    root = derived_1h_symbol_root(data_root, symbol)
    if not root.is_dir():
        raise PrimaryRangeDetectorError(
            f"No derived 1h Parquet partition tree for {symbol} at {root}. "
            "The detector will not create a workbook."
        )
    files = sorted(p for p in root.rglob("data.parquet") if p.is_file())
    if not files:
        raise PrimaryRangeDetectorError(
            f"Symbol {symbol} has a 1h directory but no data.parquet partitions under {root}."
        )
    catalog = DataCatalog(data_root)
    try:
        parts = catalog.list_partitions(TIMEFRAME, symbol=symbol)
    except UnsupportedSymbolError as exc:
        raise PrimaryRangeDetectorError(str(exc)) from exc
    if not parts:
        raise PrimaryRangeDetectorError(
            f"Catalog found no {TIMEFRAME} partitions for {symbol} under {data_root}."
        )
    start = min(p.start_utc for p in parts)
    end = max(p.end_utc for p in parts)
    fps = [fingerprint(p) for p in files]
    return {
        "symbol": symbol,
        "timeframe": TIMEFRAME,
        "dataset_root": str(root),
        "partition_count": len(parts),
        "parquet_files": fps,
        "range_start": start,
        "range_end_exclusive": end,
        "data_root": str(Path(data_root)),
    }


def load_and_verify(symbol: str, data_root: Path) -> tuple:
    inventory = inspect_source_inventory(data_root, symbol)
    start = inventory["range_start"]
    end = inventory["range_end_exclusive"]
    sl = load_candles(symbol, TIMEFRAME, start, end, precision_mode="exact", data_root=data_root)
    n = sl.row_count
    problems: list[str] = []
    if n == 0:
        problems.append("loaded zero candles")
        integrity = {"rows": 0, "problems": problems, "source_files": inventory["parquet_files"]}
        return None, integrity, sl, inventory
    opens = sl.open_times_ms()
    first = sl.table.column("open_time")[0].as_py()
    last = sl.table.column("open_time")[-1].as_py()
    if sl.table.schema.field("open").type != pa.decimal128(38, 8):
        problems.append("open is not decimal128(38, 8)")
    tz = sl.table.schema.field("open_time").type.tz
    if str(tz) not in ("UTC", "utc"):
        problems.append(f"open_time tz {tz}")
    o = sl.table.column("open").to_pylist()
    h = sl.table.column("high").to_pylist()
    l = sl.table.column("low").to_pylist()
    c = sl.table.column("close").to_pylist()
    times = [sl.table.column("open_time")[i].as_py() for i in range(n)]
    nulls = ohlc_v = 0
    for i in range(n):
        if any(v is None for v in (o[i], h[i], l[i], c[i], times[i])):
            nulls += 1
            continue
        if not isinstance(o[i], Decimal):
            problems.append(f"open[{i}] type {type(o[i])}")
            break
        if h[i] < o[i] or h[i] < c[i] or l[i] > o[i] or l[i] > c[i] or h[i] < l[i]:
            ohlc_v += 1
    dups = ooo = missing = 0
    for i in range(1, n):
        dlt = opens[i] - opens[i - 1]
        if dlt == 0:
            dups += 1
        elif dlt < 0:
            ooo += 1
        elif dlt != HOUR_MS:
            missing += 1
    if nulls:
        problems.append(f"null OHLC {nulls}")
    if ohlc_v:
        problems.append(f"OHLC violations {ohlc_v}")
    if dups:
        problems.append(f"duplicate opens {dups}")
    if ooo:
        problems.append(f"out-of-order {ooo}")
    if missing:
        problems.append(f"missing hours {missing}")
    sym_token = f"symbol={symbol.lower()}"
    for p in sl.source_files:
        pl = str(p).replace("/", "\\").lower()
        posix = str(p).replace("\\", "/").lower()
        if "futures\\um\\perpetual" not in pl and "futures/um/perpetual" not in posix:
            problems.append(f"source is not USD-M perpetual: {p}")
        if sym_token not in pl:
            problems.append(f"source symbol path mismatch: {p}")
        if "\\1h\\" not in pl and "/1h/" not in posix:
            problems.append(f"source timeframe path mismatch: {p}")
        if "\\spot\\" in pl or "/spot/" in posix:
            problems.append(f"Spot data loaded: {p}")
        if "\\delivery\\" in pl or "/delivery/" in posix:
            problems.append(f"delivery data loaded: {p}")
    fps = [fingerprint(Path(p)) for p in sl.source_files]
    integrity = {
        "rows": n, "first": iso_z(first), "last": iso_z(last), "nulls": nulls,
        "ohlc_violations": ohlc_v, "duplicates": dups, "out_of_order": ooo,
        "missing": missing, "source_files": fps, "schema_version": sl.schema_version,
        "problems": problems, "symbol": symbol, "timeframe": TIMEFRAME,
        "dataset_root": inventory["dataset_root"],
        "partition_count": inventory["partition_count"],
    }
    if problems:
        return None, integrity, sl, inventory
    return (times, o, h, l, c, opens), integrity, sl, inventory



@dataclass
class Pivot:
    idx: int
    time: datetime
    confirmed_idx: int
    confirmed_time: datetime
    price: Decimal
    atr: Decimal
    kind: str


@dataclass
class Cluster:
    kind: str
    members: list[Pivot] = field(default_factory=list)

    @property
    def prices(self) -> list[Decimal]:
        return [m.price for m in self.members]

    @property
    def center(self) -> Decimal:
        return median_dec(self.prices)

    @property
    def first_idx(self) -> int:
        return min(m.idx for m in self.members)

    @property
    def last_idx(self) -> int:
        return max(m.idx for m in self.members)

    @property
    def last_confirmed(self) -> int:
        return max(m.confirmed_idx for m in self.members)


@dataclass
class Touch:
    range_id: str = ""
    boundary: str = ""
    classification: str = "ADDITIONAL_ELIGIBLE_TOUCH"
    visit_number: int | None = None
    anchor_position: str = ""
    pair_id: str = ""
    pivot_time: datetime | None = None
    pivot_idx: int = 0
    pivot_confirmed_at: datetime | None = None
    pivot_confirmed_idx: int = 0
    pivot_price: Decimal = ZERO
    frozen: Decimal = ZERO
    abs_dist: Decimal = ZERO
    atr: Decimal = ZERO
    dist_atr: Decimal = ZERO
    high: Decimal = ZERO
    low: Decimal = ZERO
    close: Decimal = ZERO
    is_distinct: bool = False
    is_anchor: bool = False
    is_additional: bool = False
    counterpart_time: datetime | None = None
    counterpart_idx: int | None = None
    pair_sep: int | None = None
    passes_pair: bool = False
    moved_away: bool = True
    reason: str = ""


@dataclass
class Crossing:
    range_id: str = ""
    number: int = 0
    time: datetime | None = None
    prev_close: Decimal = ZERO
    close: Decimal = ZERO
    eq: Decimal = ZERO
    direction: str = ""


@dataclass
class RangeRec:
    range_id: str = ""
    parent_range_id: str = "none"
    structure_type: str = "PRIMARY"
    status: str = ""
    start_idx: int = 0
    confirmed_idx: int = 0
    end_idx: int = 0
    first_outside_idx: int | None = None
    breakout_idx: int | None = None
    start: datetime | None = None
    confirmed_at: datetime | None = None
    end: datetime | None = None
    first_outside: datetime | None = None
    breakout_at: datetime | None = None
    duration_bars: int = 0
    high: Decimal = ZERO
    low: Decimal = ZERO
    eq: Decimal = ZERO
    obs_max: Decimal = ZERO
    obs_min: Decimal = ZERO
    width: Decimal = ZERO
    width_pct: Decimal = ZERO
    width_atr: Decimal = ZERO
    median_atr: Decimal = ZERO
    upper_touches: int = 0
    lower_touches: int = 0
    u_pair_exists: bool = False
    l_pair_exists: bool = False
    u_first_time: datetime | None = None
    u_second_time: datetime | None = None
    l_first_time: datetime | None = None
    l_second_time: datetime | None = None
    u_first_idx: int | None = None
    u_second_idx: int | None = None
    l_first_idx: int | None = None
    l_second_idx: int | None = None
    u_pair_sep: int | None = None
    l_pair_sep: int | None = None
    add_upper: int = 0
    add_lower: int = 0
    u_anchor_count: int = 0
    l_anchor_count: int = 0
    first_add_u_time: datetime | None = None
    first_add_l_time: datetime | None = None
    grouped_upper: int = 0
    grouped_lower: int = 0
    inelig_upper: int = 0
    inelig_lower: int = 0
    eq_cross: int = 0
    inside_ratio: Decimal = ZERO
    closes_inside: int = 0
    closes_above: int = 0
    closes_below: int = 0
    slope: Decimal = ZERO
    er: Decimal = ZERO
    wick_up: int = 0
    wick_dn: int = 0
    false_breaks: int = 0
    breakout_dir: str = ""
    breakout_close: Decimal | None = None
    breakout_atr: Decimal | None = None
    buffer_type: str = ""
    quality: Decimal = ZERO
    label: str = ""
    merge_applied: str = "NO"
    merged_ids: str = "none"
    dup_removed: int = 0
    pre_conf_bars: int = 0
    post_conf_bars: int = 0
    conf_progress: Decimal = ZERO
    u_second_conf_idx: int | None = None
    l_second_conf_idx: int | None = None
    pass_500: bool = False
    pass_u2: bool = False
    pass_l2: bool = False
    pass_u_pair: bool = False
    pass_l_pair: bool = False
    pass_indep: bool = True
    pass_icr: bool = False
    pass_eq: bool = False
    pass_er: bool = False
    pass_slope: bool = False
    pass_width: bool = False
    pass_max_width_pct: bool = False
    pass_primary: bool = False
    pass_causal: bool = False
    hierarchy_mode: str = HIERARCHY_MODE
    confirmed_at_original: datetime | None = None
    symbol: str = ""
    market: str = MARKET_LABEL
    timeframe: str = TIMEFRAME
    detector_version: str = DETECTOR_VERSION
    configuration_version: str = CONFIGURATION_VERSION
    output_revision: str = ""
    script_sha256: str = ""
    run_id: str = ""
    creation_time_utc: str = ""
    touches: list[Touch] = field(default_factory=list)
    crossings: list[Crossing] = field(default_factory=list)
    components: list[str] = field(default_factory=list)


def compute_atr(h, l, c) -> list[Decimal | None]:
    n = len(c)
    tr = [h[0] - l[0]]
    for i in range(1, n):
        tr.append(max(h[i] - l[i], abs(h[i] - c[i - 1]), abs(l[i] - c[i - 1])))
    atr: list[Decimal | None] = [None] * n
    if n < ATR_PERIOD:
        return atr
    seed = sum(tr[:ATR_PERIOD]) / Decimal(ATR_PERIOD)
    atr[ATR_PERIOD - 1] = seed
    k = Decimal(ATR_PERIOD - 1)
    p = Decimal(ATR_PERIOD)
    prev = seed
    for i in range(ATR_PERIOD, n):
        prev = (prev * k + tr[i]) / p
        atr[i] = prev
    return atr


def find_pivots(times, h, l, atr) -> tuple[list[Pivot], list[Pivot]]:
    n = len(h)
    highs: list[Pivot] = []
    lows: list[Pivot] = []
    last = n - PIVOT_RIGHT
    for i in range(PIVOT_LEFT, last):
        if atr[i + PIVOT_RIGHT] is None:
            continue
        hi, lo = h[i], l[i]
        is_ph = is_pl = True
        for k in range(1, PIVOT_LEFT + 1):
            if not (hi > h[i - k]):
                is_ph = False
            if not (lo < l[i - k]):
                is_pl = False
        for k in range(1, PIVOT_RIGHT + 1):
            if not (hi >= h[i + k]):
                is_ph = False
            if not (lo <= l[i + k]):
                is_pl = False
        conf = i + PIVOT_RIGHT
        a = atr[conf]
        if a is None or a <= ZERO:
            continue
        if is_ph:
            highs.append(Pivot(i, times[i], conf, times[conf], hi, a, "HIGH"))
        if is_pl:
            lows.append(Pivot(i, times[i], conf, times[conf], lo, a, "LOW"))
    return highs, lows


def cluster_pivots(pivots: list[Pivot], kind: str) -> list[Cluster]:
    clusters: list[Cluster] = []
    for p in pivots:
        joined = False
        for c in reversed(clusters):
            if p.idx - c.last_idx > CLUSTER_JOIN_GAP:
                continue
            if abs(p.price - c.center) <= TOUCH_TOL_ATR * p.atr:
                c.members.append(p)
                joined = True
                break
        if not joined:
            clusters.append(Cluster(kind, [p]))
    return clusters


def linreg_slope(closes: list[Decimal]) -> Decimal:
    n = len(closes)
    if n < 2:
        return ZERO
    nd = Decimal(n)
    sum_x = Decimal(n * (n - 1) // 2)
    sum_xx = Decimal((n - 1) * n * (2 * n - 1) // 6)
    sum_y = sum(closes)
    sum_xy = sum(Decimal(i) * y for i, y in enumerate(closes))
    den = nd * sum_xx - sum_x * sum_x
    if den == 0:
        return ZERO
    return (nd * sum_xy - sum_x * sum_y) / den


def efficiency_ratio(closes: list[Decimal]) -> Decimal:
    if len(closes) < 2:
        return ZERO
    num = abs(closes[-1] - closes[0])
    den = sum(abs(closes[i] - closes[i - 1]) for i in range(1, len(closes)))
    if den == 0:
        return ZERO
    return num / den


def eq_crossings(closes: list[Decimal], eq: Decimal, times: list[datetime]) -> list[Crossing]:
    out: list[Crossing] = []
    last_side = None
    for i, cl in enumerate(closes):
        if cl > eq:
            side = "ABOVE"
        elif cl < eq:
            side = "BELOW"
        else:
            continue
        if last_side is not None and side != last_side:
            direction = "ABOVE_TO_BELOW" if last_side == "ABOVE" else "BELOW_TO_ABOVE"
            out.append(Crossing(number=len(out) + 1, time=times[i], prev_close=closes[i - 1], close=cl, eq=eq, direction=direction))
        last_side = side
    return out


def moved_away(prev_idx: int, curr_idx: int, side: str, boundary: Decimal, highs, lows, closes, atr) -> bool:
    for x in range(prev_idx + 1, curr_idx):
        ax = atr[x] or atr[curr_idx]
        if ax is None:
            continue
        if side == "UPPER" and (highs[x] < boundary - AWAY_ATR * ax or closes[x] <= boundary - AWAY_CLOSE_ATR * ax):
            return True
        if side == "LOWER" and (lows[x] > boundary + AWAY_ATR * ax or closes[x] >= boundary + AWAY_CLOSE_ATR * ax):
            return True
    return False


def select_anchor_pair(visits: list[Pivot]) -> tuple[int, int] | None:
    n = len(visits)
    best = best_key = None
    for i in range(n):
        for j in range(i + 1, n):
            sep = visits[j].idx - visits[i].idx
            if sep < MIN_PAIR_SEP:
                continue
            key = (visits[j].confirmed_idx, visits[i].idx, visits[j].idx, -sep, visits[i].idx, visits[j].idx)
            if best_key is None or key < best_key:
                best_key = key
                best = (i, j)
    return best


def synthetic_pair_test() -> dict:
    rows = [100, 110, 125, 135]
    dummy = [Pivot(i, FIRST_OPEN, i + 3, FIRST_OPEN, ZERO, ONE, "HIGH") for i in rows]
    pair = select_anchor_pair(dummy)
    ok = pair == (0, 3) and dummy[3].idx - dummy[0].idx == 35
    adjacent_all_below = all(rows[k + 1] - rows[k] < 30 for k in range(len(rows) - 1))
    return {
        "ok": bool(ok and adjacent_all_below),
        "rows": rows,
        "selected": None if pair is None else (dummy[pair[0]].idx, dummy[pair[1]].idx),
        "sep": None if pair is None else dummy[pair[1]].idx - dummy[pair[0]].idx,
        "adjacent_gaps": [rows[k + 1] - rows[k] for k in range(len(rows) - 1)],
        "note": "pair (100,135) must pass even though adjacent gaps are 10, 15, 10",
    }


def structural_width_pct(high: Decimal, low: Decimal) -> Decimal:
    eq = (high + low) / Decimal(2)
    if eq == ZERO:
        return ZERO
    return ((high - low) / eq) * Decimal(100)


def synthetic_width_tests() -> dict:
    exact_h, exact_l = Decimal("106.5"), Decimal("93.5")
    exact_pct = structural_width_pct(exact_h, exact_l)
    exact_pass = exact_pct == MAX_WIDTH_PCT and exact_pct <= MAX_WIDTH_PCT
    gt_pct = structural_width_pct(Decimal("106.5000005"), exact_l)
    round_pct = Decimal("13.004")
    visually_13 = format(round_pct, ".2f") == "13.00"
    wick_unused = structural_width_pct(exact_h, exact_l) != structural_width_pct(exact_h + Decimal("2500"), exact_l - Decimal("2500"))
    return {
        "ok": bool(exact_pass and gt_pct > MAX_WIDTH_PCT and visually_13 and round_pct > MAX_WIDTH_PCT and wick_unused),
        "exact_13_passes": exact_pass,
        "exact_13_value": format(exact_pct, "f"),
        "gt_13_fails": gt_pct > MAX_WIDTH_PCT,
        "gt_13_value": format(gt_pct, "f"),
        "rounds_to_13_but_fails": visually_13 and round_pct > MAX_WIDTH_PCT,
        "rounds_value": format(round_pct, "f"),
        "wick_not_used": wick_unused,
    }


def synthetic_zero_additional_test() -> dict:
    two = [100, 135]
    dummy = [Pivot(i, FIRST_OPEN, i + 3, FIRST_OPEN, ZERO, ONE, "HIGH") for i in two]
    pair = select_anchor_pair(dummy)
    n_add = 0
    if pair is not None:
        n_add = len([v for i, v in enumerate(dummy) if i not in pair])
    upper_ok = pair is not None and dummy[pair[1]].idx - dummy[pair[0]].idx >= MIN_PAIR_SEP and len(dummy) >= MIN_UPPER_TOUCHES and n_add >= MIN_ADDITIONAL_ELIGIBLE
    lower_dummy = [Pivot(i, FIRST_OPEN, i + 3, FIRST_OPEN, ZERO, ONE, "LOW") for i in [200, 240]]
    lpair = select_anchor_pair(lower_dummy)
    l_add = 0 if lpair is None else len([v for i, v in enumerate(lower_dummy) if i not in lpair])
    lower_ok = lpair is not None and lower_dummy[lpair[1]].idx - lower_dummy[lpair[0]].idx >= MIN_PAIR_SEP and len(lower_dummy) >= MIN_LOWER_TOUCHES and l_add >= MIN_ADDITIONAL_ELIGIBLE
    ok = bool(upper_ok and lower_ok and n_add == 0 and l_add == 0)
    return {
        "ok": ok,
        "upper_rows": two,
        "upper_pair": None if pair is None else (dummy[pair[0]].idx, dummy[pair[1]].idx),
        "upper_additional": n_add,
        "lower_rows": [200, 240],
        "lower_pair": None if lpair is None else (lower_dummy[lpair[0]].idx, lower_dummy[lpair[1]].idx),
        "lower_additional": l_add,
        "note": "exactly 2+2 visits with 30-bar pairs and zero additional eligible touches must pass",
    }


def make_touch(p: Pivot, side: str, boundary: Decimal, highs, lows, closes, atr, **kw) -> Touch:
    a = atr[p.idx] if atr[p.idx] is not None else (p.atr or ZERO)
    price = highs[p.idx] if side == "UPPER" else lows[p.idx]
    dist = abs(price - boundary)
    return Touch(
        boundary=side, pivot_time=p.time, pivot_idx=p.idx, pivot_confirmed_at=p.confirmed_time,
        pivot_confirmed_idx=p.confirmed_idx, pivot_price=price, frozen=boundary, abs_dist=dist,
        atr=a or ZERO, dist_atr=(dist / a) if a else ZERO, high=highs[p.idx], low=lows[p.idx],
        close=closes[p.idx], **kw,
    )


def side_events(s, e, boundary, side, highs, lows, closes, atr, times, pivots, causal_t):
    lo = bisect.bisect_left(pivots, s, key=lambda p: p.idx)
    hi = bisect.bisect_right(pivots, e, key=lambda p: p.idx)
    in_tol: list[Pivot] = []
    out_tol: list[Pivot] = []
    for p in pivots[lo:hi]:
        if p.confirmed_idx > causal_t:
            continue
        a = atr[p.idx] if atr[p.idx] is not None else p.atr
        if a is None or a <= ZERO:
            continue
        px = highs[p.idx] if side == "UPPER" else lows[p.idx]
        if abs(px - boundary) <= TOUCH_TOL_ATR * a:
            in_tol.append(p)
        else:
            out_tol.append(p)
    groups: list[list[Pivot]] = []
    cur: list[Pivot] = []
    for p in in_tol:
        if not cur:
            cur = [p]
            continue
        if moved_away(cur[-1].idx, p.idx, side, boundary, highs, lows, closes, atr):
            groups.append(cur)
            cur = [p]
        else:
            cur.append(p)
    if cur:
        groups.append(cur)
    visits = [g[0] for g in groups]
    pair = select_anchor_pair(visits)
    events: list[Touch] = []
    pair_id = f"{side}_QPAIR"
    a0 = a1 = None
    sep = None
    if pair is not None:
        a0, a1 = visits[pair[0]], visits[pair[1]]
        sep = a1.idx - a0.idx
    visit_no = {id(v): n for n, v in enumerate(visits, start=1)}
    for g in groups:
        rep = g[0]
        is_first = a0 is not None and rep.idx == a0.idx and rep.confirmed_idx == a0.confirmed_idx
        is_second = a1 is not None and rep.idx == a1.idx and rep.confirmed_idx == a1.confirmed_idx
        if is_first or is_second:
            klass = "QUALIFYING_ANCHOR_TOUCH"
            pos = "FIRST_ANCHOR" if is_first else "SECOND_ANCHOR"
            counterpart = a1 if is_first else a0
        else:
            klass = "ADDITIONAL_ELIGIBLE_TOUCH"
            pos = ""
            counterpart = None
        events.append(make_touch(
            rep, side, boundary, highs, lows, closes, atr,
            classification=klass, visit_number=visit_no[id(rep)], anchor_position=pos,
            pair_id=pair_id if (is_first or is_second) else "", is_distinct=True,
            is_anchor=is_first or is_second,
            is_additional=(klass == "ADDITIONAL_ELIGIBLE_TOUCH" and pair is not None),
            counterpart_time=None if counterpart is None else counterpart.time,
            counterpart_idx=None if counterpart is None else counterpart.idx,
            pair_sep=sep if (is_first or is_second) else None,
            passes_pair=bool(is_first or is_second), moved_away=True, reason="DISTINCT_ELIGIBLE_VISIT",
        ))
        for extra in g[1:]:
            events.append(make_touch(
                extra, side, boundary, highs, lows, closes, atr,
                classification="GROUPED_SAME_VISIT_CONTACT", visit_number=visit_no[id(rep)],
                is_distinct=False, is_anchor=False,
                moved_away=moved_away(rep.idx, extra.idx, side, boundary, highs, lows, closes, atr),
                reason="GROUPED_SAME_CONTINUOUS_VISIT",
            ))
    for p in out_tol:
        events.append(make_touch(
            p, side, boundary, highs, lows, closes, atr,
            classification="INELIGIBLE_CONTACT", is_distinct=False, is_anchor=False,
            reason="OUTSIDE_TOUCH_TOLERANCE",
        ))
    first_add = None
    n_additional = 0
    if pair is not None:
        add_visits = [v for i, v in enumerate(visits) if i not in pair]
        n_additional = len(add_visits)
        if add_visits:
            first_add = min(add_visits, key=lambda v: (v.confirmed_idx, v.idx))
    return {
        "visits": visits, "pair": pair, "a0": a0, "a1": a1, "sep": sep, "events": events,
        "n_grouped": sum(max(len(g) - 1, 0) for g in groups), "n_inelig": len(out_tol),
        "n_additional": n_additional, "first_add": first_add,
    }


def window_stats(s, e, H, L, highs, lows, closes, atr, times, ph, pl, causal_t=None, need_full=True):
    if causal_t is None:
        causal_t = e
    bars = e - s + 1
    eq = (H + L) / Decimal(2)
    width = H - L
    u = side_events(s, e, H, "UPPER", highs, lows, closes, atr, times, ph, causal_t)
    l = side_events(s, e, L, "LOWER", highs, lows, closes, atr, times, pl, causal_t)
    fails = []
    if bars < MIN_RANGE_BARS:
        fails.append("duration")
    if len(u["visits"]) < MIN_UPPER_TOUCHES:
        fails.append("upper_touches")
    if len(l["visits"]) < MIN_LOWER_TOUCHES:
        fails.append("lower_touches")
    if u["pair"] is None:
        fails.append("upper_pair")
    if l["pair"] is None:
        fails.append("lower_pair")
    causal_ok = True
    if e < s + MIN_RANGE_BARS - 1:
        causal_ok = False
    if u["a1"] is None or l["a1"] is None:
        causal_ok = False
    elif u["a1"].confirmed_idx > causal_t or l["a1"].confirmed_idx > causal_t:
        causal_ok = False
    if not causal_ok:
        fails.append("causality")
    width_pct = structural_width_pct(H, L)
    if width_pct > MAX_WIDTH_PCT:
        fails.append("width_pct")
    if fails and not need_full:
        return {
            "bars": bars, "eq": eq, "width": width, "width_pct": width_pct, "width_atr": ZERO, "med_atr": ZERO,
            "inside": 0, "above": 0, "below": 0, "ratio": ZERO, "wick_up": 0, "wick_dn": 0,
            "slope": Decimal("999"), "er": ONE, "crosses": [],
            "u": u, "l": l, "obs_max": H, "obs_min": L, "fails": fails, "ok": False,
        }
    atr_slice = [a for a in atr[s : e + 1] if a is not None]
    med_atr = median_dec(atr_slice) if atr_slice else ZERO
    width_atr = (width / med_atr) if med_atr else ZERO
    cs = closes[s : e + 1]
    ts = times[s : e + 1]
    inside = sum(1 for v in cs if L <= v <= H)
    above = sum(1 for v in cs if v > H)
    below = sum(1 for v in cs if v < L)
    ratio = Decimal(inside) / Decimal(bars) if bars else ZERO
    wick_up = sum(1 for i in range(s, e + 1) if highs[i] > H and closes[i] <= H)
    wick_dn = sum(1 for i in range(s, e + 1) if lows[i] < L and closes[i] >= L)
    slope = linreg_slope(cs)
    nrm_slope = abs(slope * Decimal(bars)) / width if width else Decimal("999")
    er = efficiency_ratio(cs)
    crosses = eq_crossings(cs, eq, ts)
    if ratio < MIN_INSIDE:
        fails.append("containment")
    if nrm_slope > MAX_SLOPE:
        fails.append("slope")
    if er > MAX_ER:
        fails.append("efficiency")
    if len(crosses) < MIN_EQ_CROSS:
        fails.append("eq")
    if width <= ZERO or width_atr < MIN_WIDTH_ATR:
        fails.append("width")
    if width_pct > MAX_WIDTH_PCT and "width_pct" not in fails:
        fails.append("width_pct")
    return {
        "bars": bars, "eq": eq, "width": width, "width_pct": width_pct, "width_atr": width_atr, "med_atr": med_atr,
        "inside": inside, "above": above, "below": below, "ratio": ratio, "wick_up": wick_up, "wick_dn": wick_dn,
        "slope": nrm_slope, "er": er, "crosses": crosses,
        "u": u, "l": l, "obs_max": max(highs[s : e + 1]), "obs_min": min(lows[s : e + 1]),
        "fails": fails, "ok": not fails,
    }


def count_false_breaks(s, e, H, L, closes, atr):
    n = 0
    i = s
    while i <= e:
        a = atr[i] or ZERO
        up = closes[i] > H
        dn = closes[i] < L
        if not up and not dn:
            i += 1
            continue
        strong = double = False
        if i + 1 <= e:
            a2 = atr[i + 1] or a
            if up and closes[i] > H + STRONG_BO_ATR * a:
                strong = True
            if dn and closes[i] < L - STRONG_BO_ATR * a:
                strong = True
            if up and closes[i] > H + NORMAL_BO_ATR * a and closes[i + 1] > H + NORMAL_BO_ATR * a2:
                double = True
            if dn and closes[i] < L - NORMAL_BO_ATR * a and closes[i + 1] < L - NORMAL_BO_ATR * a2:
                double = True
        else:
            if up and closes[i] > H + STRONG_BO_ATR * a:
                strong = True
            if dn and closes[i] < L - STRONG_BO_ATR * a:
                strong = True
        if not strong and not double:
            n += 1
            i += 1
        else:
            break
    return n


def detect_breakout(conf, n, H, L, closes, atr):
    for i in range(conf + 1, n):
        a = atr[i] or ZERO
        if closes[i] > H + STRONG_BO_ATR * a:
            return "UP", i, i, "strong", closes[i], a
        if closes[i] < L - STRONG_BO_ATR * a:
            return "DOWN", i, i, "strong", closes[i], a
        if i + 1 < n:
            a2 = atr[i + 1] or a
            if closes[i] > H + NORMAL_BO_ATR * a and closes[i + 1] > H + NORMAL_BO_ATR * a2:
                return "UP", i, i + 1, "normal", closes[i + 1], a2
            if closes[i] < L - NORMAL_BO_ATR * a and closes[i + 1] < L - NORMAL_BO_ATR * a2:
                return "DOWN", i, i + 1, "normal", closes[i + 1], a2
    return None, None, None, "", None, None


def quality_score(st, ut_prices, lt_prices, ut_tols, lt_tols) -> tuple[Decimal, str]:
    cont = clamp01((st["ratio"] - MIN_INSIDE) / (ONE - MIN_INSIDE))
    ut_c = min(Decimal(len(st["u"]["visits"])) / Decimal(TOUCH_SCORE_DENOM), ONE)
    lt_c = min(Decimal(len(st["l"]["visits"])) / Decimal(TOUCH_SCORE_DENOM), ONE)
    touch = (ut_c + lt_c) / Decimal(2)
    osc = min(Decimal(len(st["crosses"])) / Decimal(6), ONE)
    slope_c = ONE - clamp01(st["slope"] / MAX_SLOPE)
    er_c = ONE - clamp01(st["er"] / MAX_ER)
    flat = (slope_c + er_c) / Decimal(2)
    dur = min(Decimal(st["bars"]) / Decimal(168), ONE)
    umad = mad_dec(ut_prices) if ut_prices else ZERO
    lmad = mad_dec(lt_prices) if lt_prices else ZERO
    utol = median_dec(ut_tols) if ut_tols else ZERO
    ltol = median_dec(lt_tols) if lt_tols else ZERO
    ur = ZERO if utol == ZERO and umad == ZERO else (ONE if utol == ZERO else umad / utol)
    lr = ZERO if ltol == ZERO and lmad == ZERO else (ONE if ltol == ZERO else lmad / ltol)
    bound = ONE - Decimal("0.5") * (clamp01(ur) + clamp01(lr))
    score = Decimal(100) * (
        Decimal("0.30") * cont + Decimal("0.20") * touch + Decimal("0.15") * osc
        + Decimal("0.15") * flat + Decimal("0.10") * dur + Decimal("0.10") * bound
    )
    if score >= Decimal("85"):
        label = "Very High Quality"
    elif score >= Decimal("70"):
        label = "High Quality"
    elif score >= Decimal("60"):
        label = "Moderate Quality"
    else:
        label = "Low Quality"
    return score, label


def build_range(s, conf, end_idx, status, H, L, times, highs, lows, closes, atr, ph, pl, bo, causal_t=None):
    if causal_t is None:
        causal_t = end_idx
    st = window_stats(s, end_idx, H, L, highs, lows, closes, atr, times, ph, pl, causal_t=causal_t)
    u, lside = st["u"], st["l"]
    ut_prices = [v.price for v in u["visits"]]
    lt_prices = [v.price for v in lside["visits"]]
    ut_tols = [TOUCH_TOL_ATR * (t.atr or ZERO) for t in u["events"] if t.is_distinct and t.atr]
    lt_tols = [TOUCH_TOL_ATR * (t.atr or ZERO) for t in lside["events"] if t.is_distinct and t.atr]
    score, label = quality_score(st, ut_prices, lt_prices, ut_tols, lt_tols)
    fb = count_false_breaks(s, end_idx, H, L, closes, atr)
    pre = conf - s + 1
    post = end_idx - conf
    progress = Decimal(pre) / Decimal(st["bars"]) if st["bars"] else ZERO
    rec = RangeRec(
        start_idx=s, confirmed_idx=conf, end_idx=end_idx,
        first_outside_idx=bo.get("first_out"), breakout_idx=bo.get("conf_idx"),
        start=times[s], confirmed_at=times[conf], end=times[end_idx],
        first_outside=times[bo["first_out"]] if bo.get("first_out") is not None else None,
        breakout_at=times[bo["conf_idx"]] if bo.get("conf_idx") is not None else None,
        duration_bars=st["bars"], high=H, low=L, eq=st["eq"],
        obs_max=st["obs_max"], obs_min=st["obs_min"], width=st["width"],
        width_pct=st["width_pct"], width_atr=st["width_atr"], median_atr=st["med_atr"],
        upper_touches=len(u["visits"]), lower_touches=len(lside["visits"]),
        u_pair_exists=u["pair"] is not None, l_pair_exists=lside["pair"] is not None,
        u_first_time=None if u["a0"] is None else u["a0"].time,
        u_second_time=None if u["a1"] is None else u["a1"].time,
        l_first_time=None if lside["a0"] is None else lside["a0"].time,
        l_second_time=None if lside["a1"] is None else lside["a1"].time,
        u_first_idx=None if u["a0"] is None else u["a0"].idx,
        u_second_idx=None if u["a1"] is None else u["a1"].idx,
        l_first_idx=None if lside["a0"] is None else lside["a0"].idx,
        l_second_idx=None if lside["a1"] is None else lside["a1"].idx,
        u_pair_sep=u["sep"], l_pair_sep=lside["sep"],
        add_upper=u["n_additional"], add_lower=lside["n_additional"],
        u_anchor_count=MIN_ANCHOR_TOUCHES if u["pair"] is not None else 0,
        l_anchor_count=MIN_ANCHOR_TOUCHES if lside["pair"] is not None else 0,
        first_add_u_time=None if u["first_add"] is None else u["first_add"].time,
        first_add_l_time=None if lside["first_add"] is None else lside["first_add"].time,
        grouped_upper=u["n_grouped"], grouped_lower=lside["n_grouped"],
        inelig_upper=u["n_inelig"], inelig_lower=lside["n_inelig"],
        eq_cross=len(st["crosses"]), inside_ratio=st["ratio"],
        closes_inside=st["inside"], closes_above=st["above"], closes_below=st["below"],
        slope=st["slope"], er=st["er"], wick_up=st["wick_up"], wick_dn=st["wick_dn"],
        false_breaks=fb, breakout_dir=bo.get("dir") or "",
        breakout_close=bo.get("close"), breakout_atr=bo.get("atr"), buffer_type=bo.get("buffer") or "",
        quality=score, label=label, status=status,
        touches=u["events"] + lside["events"], crossings=st["crosses"],
        pre_conf_bars=pre, post_conf_bars=post, conf_progress=progress,
        u_second_conf_idx=None if u["a1"] is None else u["a1"].confirmed_idx,
        l_second_conf_idx=None if lside["a1"] is None else lside["a1"].confirmed_idx,
        pass_500=st["bars"] >= MIN_RANGE_BARS,
        pass_u2=len(u["visits"]) >= MIN_UPPER_TOUCHES,
        pass_l2=len(lside["visits"]) >= MIN_LOWER_TOUCHES,
        pass_u_pair=u["pair"] is not None and (u["sep"] or 0) >= MIN_PAIR_SEP,
        pass_l_pair=lside["pair"] is not None and (lside["sep"] or 0) >= MIN_PAIR_SEP,
        pass_indep=True, pass_icr=st["ratio"] >= MIN_INSIDE,
        pass_eq=len(st["crosses"]) >= MIN_EQ_CROSS, pass_er=st["er"] <= MAX_ER,
        pass_slope=st["slope"] <= MAX_SLOPE,
        pass_width=st["width"] > ZERO and st["width_atr"] >= MIN_WIDTH_ATR,
        pass_max_width_pct=st["width_pct"] <= MAX_WIDTH_PCT,
        pass_causal=(
            conf >= s + MIN_RANGE_BARS - 1 and u["a1"] is not None and lside["a1"] is not None
            and conf >= u["a1"].confirmed_idx and conf >= lside["a1"].confirmed_idx
        ),
        confirmed_at_original=times[conf],
        hierarchy_mode=HIERARCHY_MODE,
    )
    return rec, st


FAIL_MAP = {
    "duration": "fail_duration", "upper_touches": "fail_upper_touches", "lower_touches": "fail_lower_touches",
    "upper_pair": "fail_upper_pair", "lower_pair": "fail_lower_pair", "independent": "fail_independent",
    "containment": "fail_containment", "slope": "fail_slope", "efficiency": "fail_efficiency",
    "eq": "fail_eq", "width": "fail_width", "width_pct": "fail_width_pct", "causality": "fail_causality",
}


def empty_fail_diag() -> dict:
    d = {k: 0 for k in (
        "pivot_highs", "pivot_lows", "high_clusters", "low_clusters",
        "fail_duration", "fail_upper_touches", "fail_lower_touches", "fail_touch_counts_either",
        "fail_upper_pair", "fail_lower_pair", "fail_u_ge2_no_pair", "fail_l_ge2_no_pair",
        "fail_independent", "fail_containment", "fail_slope", "fail_efficiency", "fail_eq",
        "fail_width", "fail_width_pct", "width_pct_exact_13", "pass_width_pct",
        "fail_causality", "invalidated", "raw_candidates", "expired",
        "cand_pot_u", "cand_pot_l", "cand_vis_u", "cand_vis_l", "cand_anc_u", "cand_anc_l",
        "cand_add_u", "cand_add_l", "cand_grp_u", "cand_grp_l", "cand_inelig_u", "cand_inelig_l",
        "cand_zero_add_u", "cand_zero_add_l",
    )}
    d["exclusive"] = {k: 0 for k in PRIMARY_REJECTION_ORDER}
    return d


def record_rejection(diag, st, *, expired=False, invalidated=False):
    fails = list(st["fails"]) if st else []
    if st:
        u, l = st["u"], st["l"]
        diag["cand_pot_u"] += len(u["events"])
        diag["cand_pot_l"] += len(l["events"])
        diag["cand_vis_u"] += len(u["visits"])
        diag["cand_vis_l"] += len(l["visits"])
        diag["cand_anc_u"] += 2 if u["pair"] is not None else 0
        diag["cand_anc_l"] += 2 if l["pair"] is not None else 0
        diag["cand_add_u"] += u["n_additional"]
        diag["cand_add_l"] += l["n_additional"]
        diag["cand_grp_u"] += u["n_grouped"]
        diag["cand_grp_l"] += l["n_grouped"]
        diag["cand_inelig_u"] += u["n_inelig"]
        diag["cand_inelig_l"] += l["n_inelig"]
        if len(u["visits"]) >= 2 and u["pair"] is None:
            diag["fail_u_ge2_no_pair"] += 1
        if len(l["visits"]) >= 2 and l["pair"] is None:
            diag["fail_l_ge2_no_pair"] += 1
        if u["pair"] is not None and u["n_additional"] == 0:
            diag["cand_zero_add_u"] += 1
        if l["pair"] is not None and l["n_additional"] == 0:
            diag["cand_zero_add_l"] += 1
        wp = st.get("width_pct")
        if wp is not None:
            if wp == MAX_WIDTH_PCT:
                diag["width_pct_exact_13"] += 1
            if wp <= MAX_WIDTH_PCT:
                diag["pass_width_pct"] += 1
    for f in fails:
        key = FAIL_MAP.get(f)
        if key:
            diag[key] += 1
    if "upper_touches" in fails or "lower_touches" in fails:
        diag["fail_touch_counts_either"] += 1
    if expired:
        diag["expired"] += 1
    if invalidated:
        diag["invalidated"] += 1
    primary = None
    for k in PRIMARY_REJECTION_ORDER:
        if k in fails:
            primary = k
            break
    if primary is None:
        primary = "expired" if expired else ("invalidated" if invalidated else "other")
    diag["exclusive"][primary] += 1


def detect(times, o, h, l, c, atr):
    ph, pl = find_pivots(times, h, l, atr)
    h_cl = [x for x in cluster_pivots(ph, "HIGH") if len(x.members) >= 2]
    l_cl = [x for x in cluster_pivots(pl, "LOW") if len(x.members) >= 2]
    diag = empty_fail_diag()
    diag["pivot_highs"] = len(ph)
    diag["pivot_lows"] = len(pl)
    diag["high_clusters"] = len(h_cl)
    diag["low_clusters"] = len(l_cl)
    print(f"clusters high={len(h_cl)} low={len(l_cl)} pivots H/L={len(ph)}/{len(pl)}", flush=True)
    n = len(c)
    raw: list[RangeRec] = []
    seen_keys: set[tuple] = set()
    pair_i = 0
    for hc in h_cl:
        for lc in l_cl:
            H0, L0 = hc.center, lc.center
            if H0 <= L0:
                continue
            if hc.last_idx < lc.first_idx - MERGE_GAP or lc.last_idx < hc.first_idx - MERGE_GAP:
                continue
            s = min(hc.first_idx, lc.first_idx)
            seed_t = max(hc.last_confirmed, lc.last_confirmed, s + MIN_RANGE_BARS - 1)
            if seed_t >= n:
                continue
            members_h = list(hc.members)
            members_l = list(lc.members)
            key = (s, round(float(H0), 2), round(float(L0), 2))
            if key in seen_keys:
                continue
            seen_keys.add(key)
            diag["raw_candidates"] += 1
            pair_i += 1
            if pair_i % 100 == 0:
                print(f"evaluated {pair_i} pairs, kept {len(raw)}", flush=True)
            confirmed = freeze_h = freeze_l = None
            freeze_u1 = freeze_l1 = None
            last_st = None
            t_limit = min(n - 1, s + CANDIDATE_EXPIRE_BARS)
            hi_iter = lo_iter = 0
            extra_h = [p for p in ph if p.confirmed_idx >= seed_t]
            extra_l = [p for p in pl if p.confirmed_idx >= seed_t]
            last_eval = -1
            for t in range(seed_t, t_limit + 1):
                joined = False
                while hi_iter < len(extra_h) and extra_h[hi_iter].confirmed_idx <= t:
                    p = extra_h[hi_iter]
                    hi_iter += 1
                    if p.idx < s:
                        continue
                    med = median_dec([m.price for m in members_h])
                    if abs(p.price - med) <= TOUCH_TOL_ATR * p.atr and p.idx - max(m.idx for m in members_h) <= CLUSTER_JOIN_GAP:
                        members_h.append(p)
                        joined = True
                while lo_iter < len(extra_l) and extra_l[lo_iter].confirmed_idx <= t:
                    p = extra_l[lo_iter]
                    lo_iter += 1
                    if p.idx < s:
                        continue
                    med = median_dec([m.price for m in members_l])
                    if abs(p.price - med) <= TOUCH_TOL_ATR * p.atr and p.idx - max(m.idx for m in members_l) <= CLUSTER_JOIN_GAP:
                        members_l.append(p)
                        joined = True
                if (
                    last_st is not None
                    and "width_pct" in last_st["fails"]
                    and "duration" not in last_st["fails"]
                    and not joined
                ):
                    continue
                should_eval = joined or t == seed_t or t == t_limit or (t - last_eval) >= 6 or (t - s + 1) == MIN_RANGE_BARS
                if not should_eval:
                    continue
                last_eval = t
                H = median_dec([m.price for m in members_h])
                L = median_dec([m.price for m in members_l])
                if H <= L:
                    continue
                st = window_stats(s, t, H, L, h, l, c, atr, times, ph, pl, causal_t=t, need_full=False)
                last_st = st
                if st["ok"]:
                    confirmed = t
                    freeze_h, freeze_l = H, L
                    freeze_u1 = st["u"]["a1"]
                    freeze_l1 = st["l"]["a1"]
                    break
            else:
                record_rejection(diag, last_st, expired=True)
                continue
            if confirmed is None:
                record_rejection(diag, last_st, invalidated=True)
                continue
            H, L = freeze_h, freeze_l
            direction, first_out, bo_conf, buf, bo_close, bo_atr = detect_breakout(confirmed, n, H, L, c, atr)
            if direction is None:
                st_full = window_stats(s, n - 1, H, L, h, l, c, atr, times, ph, pl, causal_t=n - 1)
                if st_full["ratio"] < MIN_INSIDE:
                    record_rejection(diag, st_full, invalidated=True)
                    continue
                status = "OPEN_AT_DATASET_END"
                end_idx = n - 1
                bo = {}
            else:
                range_end = first_out - 1
                if range_end < confirmed:
                    range_end = confirmed
                st_chk = window_stats(s, range_end, H, L, h, l, c, atr, times, ph, pl, causal_t=n - 1)
                if st_chk["ratio"] < MIN_INSIDE:
                    record_rejection(diag, st_chk, invalidated=True)
                    continue
                status = "BROKEN_UP" if direction == "UP" else "BROKEN_DOWN"
                end_idx = range_end
                bo = {"dir": direction, "first_out": first_out, "conf_idx": bo_conf, "buffer": buf, "close": bo_close, "atr": bo_atr}
            rec, stf = build_range(s, confirmed, end_idx, status, H, L, times, h, l, c, atr, ph, pl, bo, causal_t=n - 1)
            if freeze_u1 is not None:
                rec.u_second_conf_idx = freeze_u1.confirmed_idx
            if freeze_l1 is not None:
                rec.l_second_conf_idx = freeze_l1.confirmed_idx
            rec.pass_causal = (
                rec.confirmed_idx >= rec.start_idx + MIN_RANGE_BARS - 1
                and rec.u_second_conf_idx is not None and rec.l_second_conf_idx is not None
                and rec.confirmed_idx >= rec.u_second_conf_idx and rec.confirmed_idx >= rec.l_second_conf_idx
            )
            hard_ok = (
                rec.pass_500 and rec.pass_u2 and rec.pass_l2 and rec.pass_u_pair and rec.pass_l_pair
                and rec.pass_indep and rec.pass_icr and rec.pass_eq and rec.pass_er and rec.pass_slope
                and rec.pass_width and rec.pass_max_width_pct and rec.pass_causal
            )
            if not hard_ok:
                record_rejection(diag, stf, invalidated=True)
                continue
            rec.confirmed_at_original = rec.confirmed_at
            raw.append(rec)
    return raw, diag, ph, pl


def interval_overlap(a: RangeRec, b: RangeRec) -> Decimal:
    lo = max(a.start_idx, b.start_idx)
    hi = min(a.end_idx, b.end_idx)
    if hi < lo:
        return ZERO
    inter = hi - lo + 1
    return Decimal(inter) / Decimal(min(a.end_idx - a.start_idx + 1, b.end_idx - b.start_idx + 1))


def contained(inner: RangeRec, outer: RangeRec) -> bool:
    return inner.start_idx >= outer.start_idx and inner.end_idx <= outer.end_idx and (
        inner.end_idx - inner.start_idx < outer.end_idx - outer.start_idx
    )


def dedupe_and_merge(rows, atr, times, h, l, c, ph, pl):
    rows = sorted(rows, key=lambda r: (-r.quality, r.confirmed_idx, -r.duration_bars, r.start_idx))
    kept: list[RangeRec] = []
    removed = 0
    for cand in rows:
        dup = False
        for k in kept:
            ov = interval_overlap(cand, k)
            a_ref = cand.median_atr if cand.median_atr else k.median_atr
            if ov >= DUP_OVERLAP and abs(cand.high - k.high) <= DUP_BOUND_ATR * a_ref and abs(cand.low - k.low) <= DUP_BOUND_ATR * a_ref:
                k.dup_removed += 1
                dup = True
                removed += 1
                break
        if not dup:
            kept.append(cand)
    merged_ops = 0
    changed = True
    while changed:
        changed = False
        kept.sort(key=lambda r: (r.start_idx, r.confirmed_idx))
        i = 0
        new: list[RangeRec] = []
        while i < len(kept):
            if i + 1 < len(kept):
                a, b = kept[i], kept[i + 1]
                gap = b.start_idx - a.end_idx - 1
                a_ref = a.median_atr or b.median_atr
                if 0 <= gap <= MERGE_GAP and abs(a.high - b.high) <= MERGE_BOUND_ATR * a_ref and abs(a.low - b.low) <= MERGE_BOUND_ATR * a_ref:
                    s = min(a.start_idx, b.start_idx)
                    e = max(a.end_idx, b.end_idx)
                    H = median_dec([a.high, b.high])
                    L = median_dec([a.low, b.low])
                    st = window_stats(s, e, H, L, h, l, c, atr, times, ph, pl)
                    if st["ok"]:
                        status = a.status if a.end_idx >= b.end_idx else b.status
                        bo = {
                            "dir": a.breakout_dir or b.breakout_dir,
                            "first_out": a.first_outside_idx if a.end_idx >= b.end_idx else b.first_outside_idx,
                            "conf_idx": a.breakout_idx if a.end_idx >= b.end_idx else b.breakout_idx,
                            "buffer": a.buffer_type or b.buffer_type,
                            "close": a.breakout_close if a.end_idx >= b.end_idx else b.breakout_close,
                            "atr": a.breakout_atr if a.end_idx >= b.end_idx else b.breakout_atr,
                        }
                        rec, _ = build_range(s, min(a.confirmed_idx, b.confirmed_idx), e, status, H, L, times, h, l, c, atr, ph, pl, bo)
                        rec.merge_applied = "YES"
                        rec.merged_ids = ",".join(x for x in (a.range_id or "tmpA", b.range_id or "tmpB"))
                        rec.dup_removed = a.dup_removed + b.dup_removed
                        rec.confirmed_at_original = rec.confirmed_at
                        new.append(rec)
                        new.extend(kept[i + 2 :])
                        merged_ops += 1
                        changed = True
                        kept = new
                        break
            new.append(kept[i])
            i += 1
        if not changed:
            kept = new if new else kept
    return kept, removed, merged_ops


def classify_hierarchy(rows: list[RangeRec]) -> None:
    """Unchanged Range V1 nested rule. Post-detection; does not alter confirmed_at."""
    for r in rows:
        original = r.confirmed_at
        parent = None
        span = None
        for outer in rows:
            if outer is r:
                continue
            if contained(r, outer):
                a_ref = r.median_atr or outer.median_atr
                distinct = abs(r.high - outer.high) > DUP_BOUND_ATR * a_ref or abs(r.low - outer.low) > DUP_BOUND_ATR * a_ref
                if not distinct:
                    continue
                sp = outer.end_idx - outer.start_idx
                if parent is None or sp < span:
                    parent = outer
                    span = sp
        if parent is not None:
            r.structure_type = "NESTED"
            r.parent_range_id = parent.range_id
            r.pass_primary = False
        else:
            r.structure_type = "PRIMARY"
            r.parent_range_id = "none"
            r.pass_primary = True
        if r.confirmed_at != original:
            raise RuntimeError("hierarchy classification altered confirmed_at")
        r.hierarchy_mode = HIERARCHY_MODE


def assign_temp_ids(rows: list[RangeRec]) -> None:
    rows.sort(key=lambda r: (r.start, r.confirmed_at, r.start_idx))
    for i, r in enumerate(rows, start=1):
        r.range_id = f"TMP{i:03d}"


def assign_primary_ids(rows: list[RangeRec]) -> None:
    rows.sort(key=lambda r: (r.start, r.confirmed_at, r.start_idx))
    for i, r in enumerate(rows, start=1):
        r.range_id = f"R{i:03d}"
        r.parent_range_id = "none"
        r.structure_type = "PRIMARY"
        r.pass_primary = True
        for t in r.touches:
            t.range_id = r.range_id
        for x in r.crossings:
            x.range_id = r.range_id


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"), right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"), bottom=Side(style="thin", color="B0B0B0"),
)
GREEN_DK = PatternFill("solid", fgColor="1B7A3D")
GREEN_LT = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F4CCCC")
BLUE = PatternFill("solid", fgColor="D0E2F3")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
PURPLE = PatternFill("solid", fgColor="E2D5F1")
GRAY = PatternFill("solid", fgColor="D9D9D9")
NOTE_FILL = PatternFill("solid", fgColor="FFF8E1")

HQ_COLS = [
    "Range ID", "Parent Range ID", "Structure Type", "Symbol", "Market", "Timeframe", "Status",
    "Range Start UTC", "Range Start Turkey", "Confirmed At UTC", "Confirmed At Turkey",
    "Range End UTC", "Range End Turkey", "Breakout Confirmed At UTC",
    "Duration Bars", "Duration Days", "Pre-Confirmation Bars", "Post-Confirmation Bars", "Confirmation Progress Ratio",
    "Range High", "Range Low", "Range EQ", "Observed Maximum High", "Observed Minimum Low",
    "Range Width", "Range Width Percent", "Range Width Percent Formatted", "Normalized Width ATR",
    "Upper Touch Count", "Lower Touch Count",
    "Upper Anchor Touch Count", "Lower Anchor Touch Count",
    "Additional Eligible Upper Touch Count", "Additional Eligible Lower Touch Count",
    "Upper Qualifying Pair Exists", "Lower Qualifying Pair Exists",
    "Upper First Anchor Time", "Upper Second Anchor Time",
    "Lower First Anchor Time", "Lower Second Anchor Time",
    "Upper Qualifying Pair Separation Bars", "Lower Qualifying Pair Separation Bars",
    "First Additional Eligible Upper Touch Time", "First Additional Eligible Lower Touch Time",
    "EQ Crossing Count", "Inside Close Ratio", "Inside Close Percentage",
    "Normalized Slope", "Efficiency Ratio", "False Break Count",
    "Breakout Direction", "Breakout Close", "Quality Score", "Quality Label",
    "Hierarchy Classification Mode", "Detector Version", "Configuration Version",
    "Output Revision", "Script SHA-256", "Run ID", "Creation Time UTC",
]
DETAIL_COLS = [
    "Range ID", "Parent Range ID", "Structure Type", "Symbol", "Market", "Timeframe", "Status",
    "Range Start UTC", "Range Start Turkey", "Confirmed At UTC", "Confirmed At Turkey",
    "Range End UTC", "Range End Turkey", "First Outside Close UTC", "First Outside Close Turkey",
    "Breakout Confirmed At UTC", "Breakout Confirmed At Turkey",
    "Duration Bars", "Duration Hours", "Duration Days",
    "Pre-Confirmation Bars", "Post-Confirmation Bars", "Confirmation Progress Ratio",
    "Range High", "Range Low", "Range EQ", "Range High Exact", "Range Low Exact", "Range EQ Exact",
    "Observed Maximum High", "Observed Minimum Low", "Range Width", "Range Width Exact",
    "Range Width Percent", "Range Width Percent Exact", "Range Width Percent Formatted",
    "Normalized Width ATR", "Median ATR During Range",
    "Upper Touch Count", "Lower Touch Count",
    "Upper Anchor Touch Count", "Lower Anchor Touch Count",
    "Additional Eligible Upper Touch Count", "Additional Eligible Lower Touch Count",
    "Upper Qualifying Pair Exists", "Lower Qualifying Pair Exists",
    "Upper First Anchor Time", "Upper Second Anchor Time", "Lower First Anchor Time", "Lower Second Anchor Time",
    "Upper First Anchor Row Index", "Upper Second Anchor Row Index", "Lower First Anchor Row Index", "Lower Second Anchor Row Index",
    "Upper Qualifying Pair Separation Bars", "Lower Qualifying Pair Separation Bars",
    "First Additional Eligible Upper Touch Time", "First Additional Eligible Lower Touch Time",
    "Grouped Upper Contact Count", "Grouped Lower Contact Count",
    "Ineligible Upper Contact Count", "Ineligible Lower Contact Count",
    "EQ Crossing Count", "Inside Close Ratio", "Inside Close Percentage", "Inside Close Ratio Exact",
    "Closes Inside", "Closes Above Range High", "Closes Below Range Low",
    "Normalized Slope", "Efficiency Ratio", "Wick-Only Upper Excursions", "Wick-Only Lower Excursions",
    "False Break Count", "Breakout Direction", "Breakout Close", "Breakout Close Exact",
    "Breakout ATR", "Breakout Buffer Type", "Quality Score", "Quality Label",
    "Merge Applied", "Merged Component IDs", "Duplicate Candidates Removed",
    "Passes Minimum 500 Bars",
    "Passes Upper Anchor Pair", "Passes Lower Anchor Pair",
    "Passes Minimum Two Total Upper Visits", "Passes Minimum Two Total Lower Visits",
    "Passes Independent Boundary Evaluation", "Passes Inside Close Ratio 0.97",
    "Passes Minimum EQ Crossings", "Passes Maximum Efficiency Ratio", "Passes Maximum Normalized Slope",
    "Passes Minimum Normalized Width", "Passes Maximum Range Width Percent 13",
    "Passes PRIMARY-Only Policy", "Passes Causality Validation",
    "Hierarchy Classification Mode", "Detector Version", "Configuration Version",
    "Output Revision", "Script SHA-256", "Run ID", "Creation Time UTC",
]
TOUCH_COLS = [
    "Range ID", "Boundary Type", "Event Classification", "Visit Number", "Anchor Position", "Qualifying Pair ID",
    "Pivot Time UTC", "Pivot Time Turkey", "Pivot Row Index",
    "Pivot Confirmed At UTC", "Pivot Confirmed At Turkey",
    "Pivot Price", "Frozen Boundary Price", "Absolute Distance From Boundary",
    "ATR At Touch", "Distance In ATR", "Candle High", "Candle Low", "Candle Close",
    "Is Distinct Eligible Visit", "Is Qualifying Anchor", "Is Additional Eligible Touch",
    "Qualifying Counterpart Time UTC", "Qualifying Counterpart Row Index",
    "Qualifying Pair Separation Bars", "Passes 30-Bar Pair Separation",
    "Price Moved Meaningfully Away Before Return", "Eligibility or Grouping Reason",
]
EQ_COLS = [
    "Range ID", "Crossing Number", "Crossing Time UTC", "Crossing Time Turkey",
    "Previous Close", "Current Close", "Range EQ", "Direction",
]


def style_header(ws, ncols):
    ws.freeze_panes = "A2"
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 32


def add_table(ws, name, rows, cols):
    if rows < 2:
        return
    ref = f"A1:{get_column_letter(cols)}{rows}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    ws.auto_filter.ref = ref


def autosize(ws, min_w=12, max_w=42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            if cell.value is None:
                continue
            length = max(length, min(len(str(cell.value)), 60))
        ws.column_dimensions[letter].width = min(max(length + 2, min_w), max_w)


def status_fill(status: str):
    return {"BROKEN_UP": GREEN_LT, "BROKEN_DOWN": RED, "OPEN_AT_DATASET_END": BLUE, "PRIMARY": PURPLE}.get(status)


def class_fill(klass: str):
    return {
        "QUALIFYING_ANCHOR_TOUCH": GREEN_LT, "ADDITIONAL_ELIGIBLE_TOUCH": BLUE,
        "GROUPED_SAME_VISIT_CONTACT": ORANGE, "INELIGIBLE_CONTACT": GRAY, "PRIMARY": PURPLE,
    }.get(klass)


def row_summary(r: RangeRec) -> list:
    return [
        r.range_id, r.parent_range_id, r.structure_type, r.symbol, r.market, r.timeframe, r.status,
        excel_clock(r.start), excel_clock(turkey(r.start)),
        excel_clock(r.confirmed_at), excel_clock(turkey(r.confirmed_at)),
        excel_clock(r.end), excel_clock(turkey(r.end)), excel_clock(r.breakout_at),
        r.duration_bars, float(Decimal(r.duration_bars) / Decimal(24)),
        r.pre_conf_bars, r.post_conf_bars, float(r.conf_progress),
        float(r.high), float(r.low), float(r.eq), float(r.obs_max), float(r.obs_min),
        float(r.width), float(r.width_pct), float(r.width_pct / Decimal(100)), float(r.width_atr),
        r.upper_touches, r.lower_touches,
        r.u_anchor_count, r.l_anchor_count, r.add_upper, r.add_lower,
        r.u_pair_exists, r.l_pair_exists,
        excel_clock(r.u_first_time), excel_clock(r.u_second_time),
        excel_clock(r.l_first_time), excel_clock(r.l_second_time),
        r.u_pair_sep, r.l_pair_sep,
        excel_clock(r.first_add_u_time), excel_clock(r.first_add_l_time),
        r.eq_cross, float(r.inside_ratio), float(r.inside_ratio),
        float(r.slope), float(r.er), r.false_breaks,
        r.breakout_dir or "none", float(r.breakout_close) if r.breakout_close is not None else None,
        float(r.quality.quantize(Decimal("0.01"))), r.label,
        r.hierarchy_mode, DETECTOR_VERSION, r.configuration_version,
        r.output_revision, r.script_sha256, r.run_id, r.creation_time_utc,
    ]


def row_detail(r: RangeRec) -> list:
    return [
        r.range_id, r.parent_range_id, r.structure_type, r.symbol, r.market, r.timeframe, r.status,
        excel_clock(r.start), excel_clock(turkey(r.start)),
        excel_clock(r.confirmed_at), excel_clock(turkey(r.confirmed_at)),
        excel_clock(r.end), excel_clock(turkey(r.end)),
        excel_clock(r.first_outside), excel_clock(turkey(r.first_outside)),
        excel_clock(r.breakout_at), excel_clock(turkey(r.breakout_at)),
        r.duration_bars, r.duration_bars, float(Decimal(r.duration_bars) / Decimal(24)),
        r.pre_conf_bars, r.post_conf_bars, float(r.conf_progress),
        float(r.high), float(r.low), float(r.eq), format(r.high, "f"), format(r.low, "f"), format(r.eq, "f"),
        float(r.obs_max), float(r.obs_min), float(r.width), format(r.width, "f"),
        float(r.width_pct), format(r.width_pct, "f"), float(r.width_pct / Decimal(100)),
        float(r.width_atr), float(r.median_atr),
        r.upper_touches, r.lower_touches,
        r.u_anchor_count, r.l_anchor_count, r.add_upper, r.add_lower,
        r.u_pair_exists, r.l_pair_exists,
        excel_clock(r.u_first_time), excel_clock(r.u_second_time), excel_clock(r.l_first_time), excel_clock(r.l_second_time),
        r.u_first_idx, r.u_second_idx, r.l_first_idx, r.l_second_idx,
        r.u_pair_sep, r.l_pair_sep,
        excel_clock(r.first_add_u_time), excel_clock(r.first_add_l_time),
        r.grouped_upper, r.grouped_lower, r.inelig_upper, r.inelig_lower,
        r.eq_cross, float(r.inside_ratio), float(r.inside_ratio), format(r.inside_ratio, "f"),
        r.closes_inside, r.closes_above, r.closes_below,
        float(r.slope), float(r.er), r.wick_up, r.wick_dn,
        r.false_breaks, r.breakout_dir or "none",
        float(r.breakout_close) if r.breakout_close is not None else None,
        format(r.breakout_close, "f") if r.breakout_close is not None else "",
        float(r.breakout_atr) if r.breakout_atr is not None else None,
        r.buffer_type or "none",
        float(r.quality.quantize(Decimal("0.01"))), r.label,
        r.merge_applied, r.merged_ids or "none", r.dup_removed,
        r.pass_500, r.pass_u_pair, r.pass_l_pair, r.pass_u2, r.pass_l2, r.pass_indep, r.pass_icr,
        r.pass_eq, r.pass_er, r.pass_slope, r.pass_width, r.pass_max_width_pct, r.pass_primary, r.pass_causal,
        r.hierarchy_mode, DETECTOR_VERSION, r.configuration_version,
        r.output_revision, r.script_sha256, r.run_id, r.creation_time_utc,
    ]


def _stamp_output_revision(path: Path, revision: str, workbook_path: str) -> None:
    wb = load_workbook(path)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        if "Output Revision" in headers:
            col = headers.index("Output Revision") + 1
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 1).value is None:
                    continue
                if str(ws.cell(r, 1).value).startswith("No "):
                    continue
                ws.cell(r, col).value = revision
        for r in range(2, ws.max_row + 1):
            key = str(ws.cell(r, 1).value or "").strip().lower()
            if key in {"output revision", "output_revision"}:
                ws.cell(r, 2).value = revision
            elif key in {"workbook path", "output_path"}:
                ws.cell(r, 2).value = workbook_path
    wb.save(path)
    wb.close()


def write_data_sheet(ws, headers, rows, table_name, price_cols, pct_cols, score_col=None, status_col=None, class_col=None, decimal_cols=None, type_col=None):
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers))
    for r_i, row in enumerate(rows, 2):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
        if status_col:
            fill = status_fill(str(row[status_col - 1]))
            if fill:
                ws.cell(r_i, status_col).fill = fill
        if class_col:
            fill = class_fill(str(row[class_col - 1]))
            if fill:
                ws.cell(r_i, class_col).fill = fill
        if type_col:
            fill = status_fill(str(row[type_col - 1]))
            if fill:
                ws.cell(r_i, type_col).fill = fill
    idx = {name: i + 1 for i, name in enumerate(headers)}
    for name in price_cols:
        if name in idx:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, idx[name]).number_format = PRICE_FMT
    for name in pct_cols:
        if name in idx:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, idx[name]).number_format = PCT_FMT
    for name in (decimal_cols or []):
        if name in idx:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, idx[name]).number_format = PRICE_FMT
    if score_col:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, score_col).number_format = "0.00"
        letter = get_column_letter(score_col)
        rng = f"{letter}2:{letter}{max(ws.max_row, 2)}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["85"], fill=GREEN_DK, font=Font(color="FFFFFF", bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["70", "84.999"], fill=GREEN_LT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["60", "69.999"], fill=YELLOW))
    if "Range Width Percent" in idx:
        letter = get_column_letter(idx["Range Width Percent"])
        rng = f"{letter}2:{letter}{max(ws.max_row, 2)}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["13"], fill=RED))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=["13"], fill=GREEN_LT))
    if rows:
        add_table(ws, table_name, 1 + len(rows), len(headers))
    else:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    autosize(ws)
    ws.sheet_view.showGridLines = False


def write_kv_sheet(ws, headers, rows, table_name, wrap_cols):
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers))
    for r_i, row in enumerate(rows, 2):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = WRAP if c_i in wrap_cols else Alignment(vertical="center")
        ws.row_dimensions[r_i].height = 36 if any(len(str(x)) > 80 for x in row) else 18
    if rows:
        add_table(ws, table_name, 1 + len(rows), len(headers))
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 92
    ws.column_dimensions["D"].width = 22
    ws.sheet_view.showGridLines = False


def build_workbook(hq, mod, all_q, diag, integrity, fps_after, created, synth, width_synth, extra_synth, meta):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "High Quality Ranges"
    price_sum = ["Range High", "Range Low", "Range EQ", "Observed Maximum High", "Observed Minimum Low", "Range Width", "Breakout Close"]
    hq_pct = ["Range Width Percent Formatted", "Inside Close Percentage", "Confirmation Progress Ratio"]
    write_data_sheet(ws1, HQ_COLS, [row_summary(r) for r in hq], "tblPrimaryRangeV1HQ", price_sum, hq_pct,
                     score_col=HQ_COLS.index("Quality Score") + 1, status_col=7, type_col=3,
                     decimal_cols=["Inside Close Ratio", "Range Width Percent"])
    if not hq:
        ws1.cell(2, 1, "No high-quality PRIMARY ranges were retained under the PRIMARY_RANGE_CONFIG_REV05 requirements.")
        ws1.cell(2, 1).fill = NOTE_FILL
    ws2 = wb.create_sheet("Range Details")
    write_data_sheet(ws2, DETAIL_COLS, [row_detail(r) for r in all_q], "tblPrimaryRangeV1Details",
                     price_sum + ["Median ATR During Range", "Breakout ATR"],
                     ["Range Width Percent Formatted", "Inside Close Percentage", "Confirmation Progress Ratio"],
                     score_col=DETAIL_COLS.index("Quality Score") + 1, status_col=7, type_col=3,
                     decimal_cols=["Inside Close Ratio", "Range Width Percent"])
    ws3 = wb.create_sheet("Boundary Touches")
    touch_rows = []
    for r in all_q:
        for t in r.touches:
            touch_rows.append([
                r.range_id, t.boundary, t.classification, t.visit_number, t.anchor_position or "", t.pair_id or "",
                excel_clock(t.pivot_time), excel_clock(turkey(t.pivot_time)), t.pivot_idx,
                excel_clock(t.pivot_confirmed_at), excel_clock(turkey(t.pivot_confirmed_at)),
                float(t.pivot_price), float(t.frozen), float(t.abs_dist),
                float(t.atr), float(t.dist_atr), float(t.high), float(t.low), float(t.close),
                t.is_distinct, t.is_anchor, t.is_additional,
                excel_clock(t.counterpart_time), t.counterpart_idx, t.pair_sep, t.passes_pair,
                t.moved_away, t.reason,
            ])
    touch_rows.sort(key=lambda x: (x[0], x[1], x[6] or datetime.min, x[8] or 0))
    write_data_sheet(ws3, TOUCH_COLS, touch_rows, "tblPrimaryRangeV1Touches",
                     ["Pivot Price", "Frozen Boundary Price", "Absolute Distance From Boundary", "ATR At Touch", "Candle High", "Candle Low", "Candle Close"],
                     [], class_col=3)
    ws4 = wb.create_sheet("EQ Crossings")
    eq_rows = []
    for r in all_q:
        for x in r.crossings:
            eq_rows.append([r.range_id, x.number, excel_clock(x.time), excel_clock(turkey(x.time)),
                            float(x.prev_close), float(x.close), float(x.eq), x.direction])
    eq_rows.sort(key=lambda x: (x[0], x[2] or datetime.min))
    write_data_sheet(ws4, EQ_COLS, eq_rows, "tblPrimaryRangeV1EQ", ["Previous Close", "Current Close", "Range EQ"], [])
    ws5 = wb.create_sheet("Moderate Ranges")
    write_data_sheet(ws5, HQ_COLS, [row_summary(r) for r in mod], "tblPrimaryRangeV1Moderate", price_sum, hq_pct,
                     score_col=HQ_COLS.index("Quality Score") + 1, status_col=7, type_col=3,
                     decimal_cols=["Inside Close Ratio", "Range Width Percent"])
    if not mod:
        ws5.cell(2, 1, "No moderate-quality PRIMARY ranges were retained under the PRIMARY_RANGE_CONFIG_REV05 requirements.")
        ws5.cell(2, 1).fill = NOTE_FILL
        ws5.cell(2, 1).font = Font(name="Calibri", italic=True, size=11)

    R, U = "CONFIG_REV05", "UNCHANGED_FROM_V1"
    params = [
        ("Dataset path", meta["dataset_root"], "Verified derived USD-M perpetual 1h Hive root", U),
        ("Symbol", meta["symbol"], "Exact Binance USD-M USDT-margined PERPETUAL symbol; not inferred from an asset nickname", U),
        ("Market", MARKET_LABEL, "Not Spot, not COIN-M, not delivery", U),
        ("Timeframe", TIMEFRAME, "Loaded directly; not resampled. This detector version rejects any other timeframe.", U),
        ("Analysis start/end UTC", f'{integrity["first"]} .. {integrity["last"]}', "Actual inclusive stored open times discovered from partitions", U),
        ("detector_version", DETECTOR_VERSION, "Reusable algorithm identity. Independent of output-run revision.", R),
        ("configuration_version", CONFIGURATION_VERSION, "Hard-filter / reporting configuration. Independent of output-run revision.", R),
        ("output_revision", meta["output_revision"], "Desktop workbook revision allocated from existing matching files", R),
        ("output_path", str(meta["workbook_path"]), "Never overwrites an existing workbook", R),
        ("Script SHA-256", meta["script_sha256"], "run_primary_range_detector.py", R),
        ("Run ID", meta["run_id"], "", R),
        ("Precision mode", "exact", "decimal128(38, 8)", U),
        ("ATR method", "Wilder with SMA seed", "SMA of first 14 TR, then (prev*13+TR)/14", U),
        ("pivot_left / pivot_right", f"{PIVOT_LEFT} / {PIVOT_RIGHT}", "Leftmost strict-left / closed-right fractal.", U),
        ("touch_tolerance", "0.25 × ATR(14)", "Candle high vs Range High, candle low vs Range Low", U),
        ("Visit grouping", "Consecutive in-tolerance confirmed pivots without meaningful retreat = one visit (earliest representative)", "Retreat = 0.5×ATR close or 1.0×ATR opposite wick", U),
        ("minimum_range_bars", MIN_RANGE_BARS, "Inclusive Duration Bars = end_idx - start_idx + 1", R),
        ("minimum_upper_anchor_touches", MIN_ANCHOR_TOUCHES, "Exactly two selected upper anchors forming a qualifying pair", R),
        ("minimum_lower_anchor_touches", MIN_ANCHOR_TOUCHES, "Exactly two selected lower anchors forming a qualifying pair", R),
        ("minimum_total_eligible_upper_visits", MIN_UPPER_TOUCHES, "Hard minimum is two distinct eligible upper visits", R),
        ("minimum_total_eligible_lower_visits", MIN_LOWER_TOUCHES, "Hard minimum is two distinct eligible lower visits", R),
        ("minimum_additional_eligible_upper_touches", MIN_ADDITIONAL_ELIGIBLE, "Descriptive only. Zero is allowed. Not a hard filter. Does not delay confirmed_at.", R),
        ("minimum_additional_eligible_lower_touches", MIN_ADDITIONAL_ELIGIBLE, "Descriptive only. Zero is allowed. Not a hard filter. Does not delay confirmed_at.", R),
        ("Additional Eligible Upper Touch Count formula", "Upper Touch Count - 2 Upper Anchors", "May be zero. Does not reject a candidate.", R),
        ("Additional Eligible Lower Touch Count formula", "Lower Touch Count - 2 Lower Anchors", "May be zero. Does not reject a candidate.", R),
        ("minimum_qualifying_touch_pair_separation", MIN_PAIR_SEP, "later_row_index - earlier_row_index >= 30. Pair members need not be adjacent.", R),
        ("Qualifying-pair formula", "exists i<j with visit_j.idx - visit_i.idx >= 30", "Intervening same-side visits do not invalidate the pair and may be additional eligible touches.", R),
        ("Anchor-pair selection", "1) earliest second-anchor confirmed_idx 2) earliest first pivot idx 3) earliest second pivot idx 4) largest separation 5) lowest row ids", "Exactly one pair per boundary. Remaining distinct eligible visits are ADDITIONAL_ELIGIBLE_TOUCH.", R),
        ("Independent boundary evaluation", "TRUE", "Upper lists never mix with lower lists.", R),
        ("minimum_eq_crossings", MIN_EQ_CROSS, "Close-to-close side change through EQ", U),
        ("minimum_inside_close_ratio", "0.97", "Unrounded Decimal. Low <= close <= High inclusive", R),
        ("maximum_efficiency_ratio", "0.30", "Zero path-length => ER=0", U),
        ("maximum_normalized_slope", "0.35", "abs(linreg_slope × bars) / width", U),
        ("minimum_normalized_width_atr", "2.0", "width / median ATR. Structural H-L, not wick extremes.", R),
        ("maximum_range_width_percent", "13.0", "((Range High - Range Low) / Range EQ) × 100. Inclusive. Unrounded.", R),
        ("required_final_structure_type", "PRIMARY", "Structure Type = PRIMARY AND Parent Range ID = none", R),
        ("non_primary_policy", "discard before every reported count", "Internal hierarchy comparison may run. Non-PRIMARY structures are then discarded from the reporting universe.", R),
        ("Hierarchy rules", "Contained interval AND High or Low differs by > 0.50 ATR. Closest-span parent. Unchanged from V1.", "Does not alter confirmed_at. A discarded non-PRIMARY range cannot replace a PRIMARY range.", U),
        ("Hierarchy classification mode", HIERARCHY_MODE, "Applied after causal detection. Not a live signal. Does not rewrite confirmed_at.", R),
        ("Processing order", "detect (causal) → hierarchy internally → discard non-PRIMARY → near-duplicate PRIMARY → merge PRIMARY → quality split → export", "confirmed_at frozen at causal confirmation; hierarchy is reporting filter only", R),
        ("Candidate safety expiration", f"{CANDIDATE_EXPIRE_BARS} completed 1h bars from range_start", "Caps SEARCHING/CANDIDATE wait. Does not split a confirmed range.", U),
        ("normal/strong breakout", "0.25 ATR two closes / 0.50 ATR one close", "Close-based; wick-only does not end the range", U),
        ("Duplicate rule", "70% overlap AND High/Low within 0.50 ATR, PRIMARY universe only", "Higher score, earlier confirm, longer, earlier start.", R),
        ("Merge rule", "Both PRIMARY AND gap<=6 AND High/Low within 0.25 ATR AND merged independently passes every hard structural rule AND remains PRIMARY", "Never merge only to manufacture 500 bars, anchors, ICR, 13% width, or PRIMARY status. Additional eligible touches are not required of a merge.", R),
        ("Containment score", "clamp((ICR-0.97)/0.03,0,1)", "Not the former 0.85 baseline", R),
        ("Touch score", "0.5*min(U/4,1)+0.5*min(L/4,1)", "Hard min is 2 distinct eligible visits/side. Additional visits may improve the score. /4 remains the score scale.", R),
        ("Duration score", "min(bars/168,1)", "Hard filter remains 500 bars", U),
        ("Quality formula", "100*(0.30*cont+0.20*touch+0.15*osc+0.15*flat+0.10*dur+0.10*bound)", "Hierarchy does not alter the numeric score", U),
        ("Detector version", DETECTOR_VERSION, "PRIMARY-only; 500/2+2 any-pair-30/ICR 0.97/width<=13; additional touches descriptive", R),
        ("Revision allocation rule", "Scan ^SYMBOL_1H_Range_Detection_rev([0-9]+).xlsx$; next = max existing + 1; rev01 if none; never fill gaps; never overwrite", "Output revision is not the algorithm version", R),
    ]
    ws6 = wb.create_sheet("Parameters")
    write_kv_sheet(ws6, ["Parameter", "Value", "Explanation", "Revision Status"], params, "tblPrimaryRangeV1Parameters", {3})

    excl = diag.get("exclusive") or {}
    src_notes = "; ".join(f"{Path(x['path']).name} sha256={x['sha256'][:16]}…" for x in integrity["source_files"][:3])
    diags = [
        ("Total candles analyzed", integrity["rows"], f'{meta["symbol"]} {TIMEFRAME} exact'),
        ("Actual first candle", integrity["first"], "UTC"),
        ("Actual last candle", integrity["last"], "UTC"),
        ("Raw PRIMARY candidate count", diag["raw_candidates"], "PRIMARY reporting universe after internal hierarchy discard of non-PRIMARY confirmed structures"),
        ("OVERLAPPING rejection note", "Individual fail counts overlap", "One PRIMARY-universe candidate may increment several counters"),
        ("PRIMARY candidates failing duration below 500", diag["fail_duration"], "0 expected: eval starts at 500th bar"),
        ("PRIMARY candidates failing total eligible upper visits below 2", diag["fail_upper_touches"], "Overlapping"),
        ("PRIMARY candidates failing total eligible lower visits below 2", diag["fail_lower_touches"], "Overlapping"),
        ("PRIMARY candidates failing one or both 2-visit requirements", diag["fail_touch_counts_either"], "Overlapping"),
        ("PRIMARY candidates with >=2 upper visits but no qualifying upper pair", diag["fail_u_ge2_no_pair"], "Overlapping"),
        ("PRIMARY candidates with >=2 lower visits but no qualifying lower pair", diag["fail_l_ge2_no_pair"], "Overlapping"),
        ("PRIMARY candidates failing upper qualifying-pair requirement", diag["fail_upper_pair"], "Overlapping"),
        ("PRIMARY candidates failing lower qualifying-pair requirement", diag["fail_lower_pair"], "Overlapping"),
        ("PRIMARY candidates failing independent-boundary validation", diag["fail_independent"], "Should be 0"),
        ("Potential upper contact events (candidate last-eval, overlapping pairs)", diag["cand_pot_u"], ""),
        ("Distinct eligible upper visits (candidate last-eval)", diag["cand_vis_u"], ""),
        ("Upper qualifying anchor touches (candidate last-eval)", diag["cand_anc_u"], ""),
        ("Additional eligible upper touches (candidate last-eval)", diag["cand_add_u"], ""),
        ("Grouped same-visit upper contacts (candidate last-eval)", diag["cand_grp_u"], ""),
        ("Ineligible upper contacts (candidate last-eval)", diag["cand_inelig_u"], ""),
        ("Potential lower contact events (candidate last-eval, overlapping pairs)", diag["cand_pot_l"], ""),
        ("Distinct eligible lower visits (candidate last-eval)", diag["cand_vis_l"], ""),
        ("Lower qualifying anchor touches (candidate last-eval)", diag["cand_anc_l"], ""),
        ("Additional eligible lower touches (candidate last-eval)", diag["cand_add_l"], ""),
        ("Grouped same-visit lower contacts (candidate last-eval)", diag["cand_grp_l"], ""),
        ("Ineligible lower contacts (candidate last-eval)", diag["cand_inelig_l"], ""),
        ("Retained-PRIMARY potential upper events", diag.get("qual_pot_u", 0), "Workbook audit"),
        ("Retained-PRIMARY distinct upper visits", diag.get("qual_vis_u", 0), ""),
        ("Retained-PRIMARY upper anchors", diag.get("qual_anc_u", 0), ""),
        ("Retained-PRIMARY additional upper", diag.get("qual_add_u", 0), ""),
        ("Retained-PRIMARY potential lower events", diag.get("qual_pot_l", 0), ""),
        ("Retained-PRIMARY distinct lower visits", diag.get("qual_vis_l", 0), ""),
        ("Retained-PRIMARY lower anchors", diag.get("qual_anc_l", 0), ""),
        ("Retained-PRIMARY additional lower", diag.get("qual_add_l", 0), ""),
        ("Retained PRIMARY ranges with zero additional eligible upper touches", diag.get("qual_zero_add_u", 0), "Descriptive; not a rejection"),
        ("Retained PRIMARY ranges with zero additional eligible lower touches", diag.get("qual_zero_add_l", 0), "Descriptive; not a rejection"),
        ("PRIMARY last-eval candidates with a qualifying upper pair and zero additional eligible upper touches", diag.get("cand_zero_add_u", 0), "Descriptive; not a hard-filter rejection"),
        ("PRIMARY last-eval candidates with a qualifying lower pair and zero additional eligible lower touches", diag.get("cand_zero_add_l", 0), "Descriptive; not a hard-filter rejection"),
        ("PRIMARY candidates failing Inside Close Ratio below 0.97", diag["fail_containment"], "Unrounded"),
        ("PRIMARY candidates failing normalized slope", diag["fail_slope"], ""),
        ("PRIMARY candidates failing efficiency ratio", diag["fail_efficiency"], ""),
        ("PRIMARY candidates failing EQ crossings", diag["fail_eq"], ""),
        ("PRIMARY candidates failing minimum normalized width below 2.0 ATR", diag["fail_width"], ""),
        ("PRIMARY candidates failing maximum Range Width Percentage above 13.0", diag["fail_width_pct"], "Unrounded structural"),
        ("Candidates equal to exactly 13.0 Range Width Percentage", diag["width_pct_exact_13"], ""),
        ("Candidates passing maximum Range Width Percentage", diag["pass_width_pct"], "last-eval snapshot"),
        ("Minimum Range Width Percentage among retained PRIMARY ranges", diag.get("qual_min_width_pct", "n/a"), ""),
        ("Maximum Range Width Percentage among retained PRIMARY ranges", diag.get("qual_max_width_pct", "n/a"), "Must be <= 13.0"),
        ("PRIMARY candidates failing causality validation", diag["fail_causality"], ""),
        ("Invalidated PRIMARY candidates", diag["invalidated"], ""),
        ("Expired PRIMARY candidates", diag["expired"], f"Safety cap {CANDIDATE_EXPIRE_BARS} 1h bars from range_start"),
        ("Exact safety-expiration parameter", f"{CANDIDATE_EXPIRE_BARS} completed 1h bars", "Disclosed; does not split a confirmed range"),
        ("EXCLUSIVE hard-rejection note", "Mutually exclusive last-eval primary reason", "Priority: duration, upper_touches, lower_touches, upper_pair, lower_pair, independent, containment, slope, efficiency, eq, width, width_pct, causality, expired, invalidated, other"),
        ("Exclusive FAIL_DURATION_BELOW_500", excl.get("duration", 0), ""),
        ("Exclusive FAIL_UPPER_TOUCHES_BELOW_2", excl.get("upper_touches", 0), ""),
        ("Exclusive FAIL_LOWER_TOUCHES_BELOW_2", excl.get("lower_touches", 0), ""),
        ("Exclusive FAIL_UPPER_PAIR", excl.get("upper_pair", 0), ""),
        ("Exclusive FAIL_LOWER_PAIR", excl.get("lower_pair", 0), ""),
        ("Exclusive FAIL_INDEPENDENT", excl.get("independent", 0), ""),
        ("Exclusive FAIL_CONTAINMENT", excl.get("containment", 0), ""),
        ("Exclusive FAIL_SLOPE", excl.get("slope", 0), ""),
        ("Exclusive FAIL_EFFICIENCY", excl.get("efficiency", 0), ""),
        ("Exclusive FAIL_EQ", excl.get("eq", 0), ""),
        ("Exclusive FAIL_WIDTH", excl.get("width", 0), ""),
        ("Exclusive FAIL_WIDTH_PCT_ABOVE_13", excl.get("width_pct", 0), ""),
        ("Exclusive FAIL_CAUSALITY", excl.get("causality", 0), ""),
        ("Exclusive EXPIRED", excl.get("expired", 0), ""),
        ("Exclusive INVALIDATED", excl.get("invalidated", 0), ""),
        ("Exclusive OTHER", excl.get("other", 0), ""),
        ("FAILED_HARD_FILTER (exclusive sum)", diag.get("failed_hard_exclusive", 0), "Sum of exclusive hard-rejection buckets"),
        ("PRIMARY near-duplicates removed", diag.get("near_dupes", 0), "PRIMARY universe only"),
        ("PRIMARY adjacent structures merged", diag.get("merged", 0), "PRIMARY+PRIMARY only"),
        ("Structurally valid PRIMARY candidates before quality filtering", diag.get("hard_pass_after_dedupe", 0), "After near-duplicate/merge on PRIMARY universe"),
        ("Low-quality hard-qualified PRIMARY candidates excluded", diag.get("low_q", 0), "Hard pass, score<60, PRIMARY"),
        ("PRIMARY candidates retained", len(all_q), "score>=60 AND PRIMARY"),
        ("RETAINED_PRIMARY_HIGH_QUALITY", len(hq), "score>=70 AND PRIMARY"),
        ("RETAINED_PRIMARY_MODERATE_QUALITY", len(mod), "60<=score<70 AND PRIMARY"),
        ("Final total retained PRIMARY count", len(all_q), ""),
        ("Open PRIMARY ranges at dataset end", sum(1 for r in all_q if r.status == "OPEN_AT_DATASET_END"), ""),
        ("Broken-up PRIMARY range count", sum(1 for r in all_q if r.status == "BROKEN_UP"), ""),
        ("Broken-down PRIMARY range count", sum(1 for r in all_q if r.status == "BROKEN_DOWN"), ""),
        ("Hierarchy classification mode", HIERARCHY_MODE, "Does not alter confirmed_at"),
        ("confirmed_at altered by hierarchy", "NO", "Validated"),
        ("Missing timestamps", integrity["missing"], "Must be 0"),
        ("Duplicate timestamps", integrity["duplicates"], "Must be 0"),
        ("Out-of-order timestamps", integrity["out_of_order"], "Must be 0"),
        ("Null OHLC values", integrity["nulls"], "Must be 0"),
        ("OHLC violations", integrity["ohlc_violations"], "Must be 0"),
        ("Source files used", len(integrity["source_files"]), src_notes),
        ("Source immutability result", fps_after.get("source", "UNCHECKED"), ""),
        ("Original workbook immutability result", fps_after.get("original_wb", "UNCHECKED"), ""),
        ("Previous revision workbook immutability result", fps_after.get("previous_wbs", "UNCHECKED"), ""),
        ("Existing code immutability result", fps_after.get("code", "UNCHECKED"), ""),
        ("Temporary-artifact cleanup result", fps_after.get("temp_cleanup", "PENDING"), "Workbook written to OS temp, validated, then atomically committed"),
        ("Synthetic pair test", "PASS" if synth["ok"] else "FAIL", str(synth)),
        ("Synthetic 13% width-boundary tests", "PASS" if width_synth["ok"] else "FAIL", str(width_synth)),
        ("Synthetic zero-additional-touch test", "PASS" if extra_synth["ok"] else "FAIL", str(extra_synth)),
        ("Warnings", diag.get("warnings") or "none", ""),
        ("Errors", "none", ""),
    ]
    ws7 = wb.create_sheet("Diagnostics")
    write_kv_sheet(ws7, ["Diagnostic", "Value", "Notes"], [(a, str(b), c) for a, b, c in diags], "tblPrimaryRangeV1Diagnostics", {3})

    readme = [
        ("Title", "Reusable Binance USD-M perpetual 1h PRIMARY range detector", "Exploratory market-structure detection. Not investment advice. Parameters were not optimized. No profitability test, signals, orders, or strategy registration."),
        ("PRIMARY-only policy", "Only PRIMARY ranges are included or counted", "Non-PRIMARY structures are classified internally, then discarded before every reported statistic. No nested counts, nested sheets, nested events, or nested examples."),
        ("Hierarchy mode", HIERARCHY_MODE, "Uses later completed structures. The PRIMARY label is a post-detection reporting filter and was not necessarily known at confirmed_at. Not a live trading signal. Does not rewrite confirmed_at."),
        ("Causality of confirmation", "confirmed_at is the first causal time all current hard structural rules are simultaneously satisfied", "Hierarchy classification does not move confirmed_at. Additional eligible touches do not delay confirmed_at."),
        ("Boundaries", "Range High/Low = median of selected pivot-cluster prices", "Observed max high / min low reported separately. Anomalous wicks do not define structural bounds."),
        ("Range EQ", "(Range High + Range Low) / 2", ""),
        ("Pivots", "3 left / 3 right", "A boundary touch based on a pivot is unavailable before pivot_confirmed_at. Second anchors cannot contribute before their causal confirmation."),
        ("Duration", "At least 500 completed 1h candles", "Inclusive count"),
        ("Anchors", "Exactly two qualifying upper anchors and two qualifying lower anchors", "Each same-side pair must have row-index separation >= 30. Pair members need not be adjacent."),
        ("Additional eligible touches", "Descriptive only. Not a hard filter.", "A range may qualify with zero additional eligible upper touches and zero additional eligible lower touches. Additional visits may improve the touch-quality score. They do not delay confirmed_at."),
        ("Qualifying pair", "Any two distinct same-side visits with row-index difference >= 30", "Need not be adjacent. Example [100,110,125,135] -> (100,135)."),
        ("Independence", "Upper and lower evaluated separately", "A High touch never counts toward Low"),
        ("ICR", "Unrounded >= 0.97", "Low <= close <= High"),
        ("Range width percent", "Unrounded ((Range High - Range Low) / Range EQ) × 100 <= 13.0", "Inclusive. Structural frozen boundaries, not wick extremes. Exact 13.0 passes."),
        ("Score", "Hard min is 2 distinct eligible visits plus a 30-bar pair per side; /4 still saturates the touch component", "Hierarchy does not change the numeric score. Low-quality (score<60) PRIMARY candidates are excluded."),
        ("Timezone", "UTC stored/computed", "Turkey columns Europe/Istanbul UTC+3"),
        ("Safety expiration", f"{CANDIDATE_EXPIRE_BARS} 1h bars from range_start", "Disclosed SEARCHING/CANDIDATE cap. UNCHANGED_FROM_V1."),
        ("Optimization / trading", "None", "No signals, orders, or profitability tests"),
        ("Isolation", "This detector lives only under backtest_system/detectors/primary_range_v1/", "It is not a trading strategy, is not auto-registered, and is not imported by unrelated modules. Future strategies must not import it unless explicitly requested."),
        ("How to run another symbol", r"python C:\Users\oranb\Desktop\backtest_system\detectors\primary_range_v1\run_primary_range_detector.py --symbol ETHUSDT", "Pass the exact perpetual symbol. The script allocates the next desktop revision automatically."),
        ("Algorithm vs output revision", f"{DETECTOR_VERSION} / {CONFIGURATION_VERSION}", "Workbook filename revNN is an output-run counter, not an algorithm change."),
        ("Code / data", "Source Parquet and previous workbooks remain read-only", "The detector recomputes from verified Parquet. It never copies an earlier Excel result."),
        ("Excel prices", "Numeric 0.00000000 plus Exact text on Range Details", ""),
        ("Creation time UTC", created, ""),
        ("Python", sys.version.split()[0], ""),
        ("PyArrow", pa.__version__, ""),
        ("Excel writer", f"openpyxl {__import__('openpyxl').__version__}", ""),
        ("Source Parquet root", meta["dataset_root"], ""),
        ("Source row count", str(integrity["rows"]), ""),
        ("Workbook path", str(meta["workbook_path"]), ""),
        ("Detector version", DETECTOR_VERSION, "PRIMARY-only; 500/2+2 any-pair-30/ICR 0.97/width<=13; additional touches descriptive"),
        ("Configuration version", CONFIGURATION_VERSION, ""),
        ("Output revision", meta["output_revision"], ""),
        ("Script SHA-256", meta["script_sha256"], ""),
        ("Run ID", meta["run_id"], ""),
    ]
    ws8 = wb.create_sheet("README")
    write_kv_sheet(ws8, ["Topic", "Statement", "Detail"], readme, "tblPrimaryRangeV1Readme", {2, 3})
    for row in ws8.iter_rows(min_row=2, max_row=ws8.max_row):
        row[0].fill = PatternFill("solid", fgColor="D6EAF8")
    order = ["High Quality Ranges", "Range Details", "Boundary Touches", "EQ Crossings", "Moderate Ranges", "Parameters", "Diagnostics", "README"]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))
    dest = Path(meta["temp_path"])
    wb.save(dest)
    wb.close()
    return dest


def validate_workbook(workbook_path, hq, mod, all_q, integrity, fps_before, prev_before, code_before, synth, width_synth, extra_synth, orig_confirmed, meta) -> list[str]:
    errs = []
    if not synth["ok"]:
        errs.append(f"synthetic pair test failed {synth}")
    if not width_synth["ok"]:
        errs.append(f"synthetic width tests failed {width_synth}")
    if not extra_synth["ok"]:
        errs.append(f"synthetic zero-additional test failed {extra_synth}")
    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        return ["workbook missing"] + errs
    wb = load_workbook(workbook_path)
    expected = ["High Quality Ranges", "Range Details", "Boundary Touches", "EQ Crossings", "Moderate Ranges", "Parameters", "Diagnostics", "README"]
    if wb.sheetnames != expected:
        errs.append(f"sheet order {wb.sheetnames}")
    hqw = wb["High Quality Ranges"]
    hq_headers = [c.value for c in hqw[1]]
    hq_data = [r for r in range(2, hqw.max_row + 1) if hqw.cell(r, 1).value and not str(hqw.cell(r, 1).value).startswith("No ")]
    if "Symbol" in hq_headers:
        si = hq_headers.index("Symbol") + 1
        for r in hq_data:
            if str(hqw.cell(r, si).value) != meta["symbol"]:
                errs.append("HQ symbol mismatch")
    if "Output Revision" in hq_headers:
        oi = hq_headers.index("Output Revision") + 1
        for r in hq_data:
            if str(hqw.cell(r, oi).value) != meta["output_revision"]:
                errs.append("HQ output revision mismatch")

    if len(hq_data) != len(hq):
        errs.append(f"HQ rows {len(hq_data)} != {len(hq)}")
    mod_ws = wb["Moderate Ranges"]
    mod_data = [r for r in range(2, mod_ws.max_row + 1) if mod_ws.cell(r, 1).value and not str(mod_ws.cell(r, 1).value).startswith("No ")]
    if len(mod_data) != len(mod):
        errs.append(f"moderate rows {len(mod_data)} != {len(mod)}")
    det = wb["Range Details"]
    headers = [c.value for c in det[1]]
    idx = {name: i + 1 for i, name in enumerate(headers) if name}

    def col(name):
        return idx.get(name)

    ids = []
    for r in range(2, det.max_row + 1):
        rid = det.cell(r, col("Range ID")).value if col("Range ID") else None
        if not rid:
            continue
        ids.append(rid)
        stype = str(det.cell(r, col("Structure Type")).value)
        parent = str(det.cell(r, col("Parent Range ID")).value)
        if stype != "PRIMARY":
            errs.append(f"{rid} structure {stype}")
        if parent != "none":
            errs.append(f"{rid} parent {parent}")
        if col("Passes PRIMARY-Only Policy") and det.cell(r, col("Passes PRIMARY-Only Policy")).value not in (True, "TRUE", "True"):
            errs.append(f"{rid} PRIMARY policy")
        H = Decimal(str(det.cell(r, col("Range High Exact")).value or det.cell(r, col("Range High")).value))
        L = Decimal(str(det.cell(r, col("Range Low Exact")).value or det.cell(r, col("Range Low")).value))
        E = Decimal(str(det.cell(r, col("Range EQ Exact")).value or det.cell(r, col("Range EQ")).value))
        W = Decimal(str(det.cell(r, col("Range Width Exact")).value or det.cell(r, col("Range Width")).value))
        WP = Decimal(str(det.cell(r, col("Range Width Percent Exact")).value or det.cell(r, col("Range Width Percent")).value))
        if abs((H + L) / Decimal(2) - E) > Decimal("0.00000001"):
            errs.append(f"{rid} EQ mismatch")
        if abs((H - L) - W) > Decimal("0.00000001"):
            errs.append(f"{rid} width mismatch")
        recalc_wp = structural_width_pct(H, L)
        if abs(recalc_wp - WP) > Decimal("0.00000001"):
            errs.append(f"{rid} width_pct mismatch")
        if WP > MAX_WIDTH_PCT:
            errs.append(f"{rid} width_pct {WP} > 13")
        obs_max = Decimal(str(det.cell(r, col("Observed Maximum High")).value))
        obs_min = Decimal(str(det.cell(r, col("Observed Minimum Low")).value))
        if abs(structural_width_pct(obs_max, obs_min) - WP) <= Decimal("0.00000001") and (obs_max != H or obs_min != L):
            errs.append(f"{rid} width_pct appears to use wick extremes")
        if float(det.cell(r, col("Quality Score")).value) < 60:
            errs.append(f"{rid} score")
        if int(det.cell(r, col("Duration Bars")).value) < 500:
            errs.append(f"{rid} duration")
        if int(det.cell(r, col("Upper Touch Count")).value) < 2:
            errs.append(f"{rid} upper count")
        if int(det.cell(r, col("Lower Touch Count")).value) < 2:
            errs.append(f"{rid} lower count")
        if int(det.cell(r, col("Upper Anchor Touch Count")).value) != 2:
            errs.append(f"{rid} upper anchors")
        if int(det.cell(r, col("Lower Anchor Touch Count")).value) != 2:
            errs.append(f"{rid} lower anchors")
        add_u = int(det.cell(r, col("Additional Eligible Upper Touch Count")).value)
        add_l = int(det.cell(r, col("Additional Eligible Lower Touch Count")).value)
        if add_u < 0 or add_l < 0:
            errs.append(f"{rid} negative additional")
        if int(det.cell(r, col("Upper Touch Count")).value) != 2 + add_u:
            errs.append(f"{rid} upper touch identity")
        if int(det.cell(r, col("Lower Touch Count")).value) != 2 + add_l:
            errs.append(f"{rid} lower touch identity")
        if det.cell(r, col("Upper Qualifying Pair Exists")).value not in (True, "TRUE"):
            errs.append(f"{rid} no upper pair")
        if det.cell(r, col("Lower Qualifying Pair Exists")).value not in (True, "TRUE"):
            errs.append(f"{rid} no lower pair")
        if int(det.cell(r, col("Upper Qualifying Pair Separation Bars")).value) < 30:
            errs.append(f"{rid} upper sep")
        if int(det.cell(r, col("Lower Qualifying Pair Separation Bars")).value) < 30:
            errs.append(f"{rid} lower sep")
        icr = Decimal(str(det.cell(r, col("Inside Close Ratio Exact")).value or det.cell(r, col("Inside Close Ratio")).value))
        if icr < MIN_INSIDE:
            errs.append(f"{rid} ICR")
        for flag in (
            "Passes Minimum 500 Bars",
            "Passes Upper Anchor Pair", "Passes Lower Anchor Pair",
            "Passes Minimum Two Total Upper Visits", "Passes Minimum Two Total Lower Visits",
            "Passes Independent Boundary Evaluation", "Passes Inside Close Ratio 0.97",
            "Passes Minimum EQ Crossings", "Passes Maximum Efficiency Ratio", "Passes Maximum Normalized Slope",
            "Passes Minimum Normalized Width", "Passes Maximum Range Width Percent 13",
            "Passes PRIMARY-Only Policy", "Passes Causality Validation",
        ):
            if col(flag) and det.cell(r, col(flag)).value not in (True, "TRUE", "True"):
                errs.append(f"{rid} {flag}")
    if len(ids) != len(set(ids)):
        errs.append("duplicate Range IDs")
    if len(ids) != len(all_q):
        errs.append(f"details {len(ids)} != {len(all_q)}")
    idset = set(ids)
    for sheet_name in ("High Quality Ranges", "Moderate Ranges", "Range Details"):
        sh = wb[sheet_name]
        hh = [c.value for c in sh[1]]
        if "Structure Type" in hh:
            ci = hh.index("Structure Type") + 1
            for r in range(2, sh.max_row + 1):
                val = sh.cell(r, ci).value
                if val and str(val) == "NESTED":
                    errs.append(f"NESTED in {sheet_name}")
    tw = wb["Boundary Touches"]
    th = [c.value for c in tw[1]]
    tidx = {name: i + 1 for i, name in enumerate(th) if name}
    vis_u = defaultdict(int)
    vis_l = defaultdict(int)
    anc_u = defaultdict(list)
    anc_l = defaultdict(list)
    add_u = defaultdict(int)
    add_l = defaultdict(int)
    for r in range(2, tw.max_row + 1):
        rid = tw.cell(r, tidx["Range ID"]).value
        if not rid:
            continue
        if rid not in idset:
            errs.append(f"orphan/non-PRIMARY {rid} on touches")
            continue
        btype = tw.cell(r, tidx["Boundary Type"]).value
        klass = tw.cell(r, tidx["Event Classification"]).value
        distinct = tw.cell(r, tidx["Is Distinct Eligible Visit"]).value in (True, "TRUE", 1)
        is_add = False
        if "Is Additional Eligible Touch" in tidx:
            is_add = tw.cell(r, tidx["Is Additional Eligible Touch"]).value in (True, "TRUE", 1)
        if distinct and btype == "UPPER":
            vis_u[rid] += 1
        if distinct and btype == "LOWER":
            vis_l[rid] += 1
        if klass == "QUALIFYING_ANCHOR_TOUCH":
            prow = int(tw.cell(r, tidx["Pivot Row Index"]).value)
            if btype == "UPPER":
                anc_u[rid].append(prow)
            elif btype == "LOWER":
                anc_l[rid].append(prow)
            else:
                errs.append(f"{rid} mixed boundary on anchor")
            if is_add:
                errs.append(f"{rid} anchor marked additional")
        if klass == "ADDITIONAL_ELIGIBLE_TOUCH":
            if btype == "UPPER":
                add_u[rid] += 1
            elif btype == "LOWER":
                add_l[rid] += 1
        if klass == "GROUPED_SAME_VISIT_CONTACT" and tw.cell(r, tidx["Is Distinct Eligible Visit"]).value in (True, "TRUE", 1):
            errs.append(f"{rid} grouped marked distinct")
        if klass == "INELIGIBLE_CONTACT" and tw.cell(r, tidx["Is Distinct Eligible Visit"]).value in (True, "TRUE", 1):
            errs.append(f"{rid} ineligible marked distinct")
    for rec in all_q:
        if vis_u.get(rec.range_id, 0) != rec.upper_touches:
            errs.append(f"{rec.range_id} upper visits mismatch")
        if vis_l.get(rec.range_id, 0) != rec.lower_touches:
            errs.append(f"{rec.range_id} lower visits mismatch")
        if add_u.get(rec.range_id, 0) != rec.add_upper:
            errs.append(f"{rec.range_id} additional upper mismatch")
        if add_l.get(rec.range_id, 0) != rec.add_lower:
            errs.append(f"{rec.range_id} additional lower mismatch")
        au = sorted(anc_u.get(rec.range_id, []))
        al = sorted(anc_l.get(rec.range_id, []))
        if len(au) != 2 or au[1] - au[0] < 30:
            errs.append(f"{rec.range_id} upper anchors {au}")
        if len(al) != 2 or al[1] - al[0] < 30:
            errs.append(f"{rec.range_id} lower anchors {al}")
        if rec.confirmed_idx < rec.start_idx + MIN_RANGE_BARS - 1:
            errs.append(f"{rec.range_id} confirmed before 500th")
        if rec.u_second_conf_idx is not None and rec.confirmed_idx < rec.u_second_conf_idx:
            errs.append(f"{rec.range_id} confirmed before upper 2nd")
        if rec.l_second_conf_idx is not None and rec.confirmed_idx < rec.l_second_conf_idx:
            errs.append(f"{rec.range_id} confirmed before lower 2nd")
        if rec.inside_ratio < MIN_INSIDE or rec.eq_cross < 2 or rec.slope > MAX_SLOPE or rec.er > MAX_ER or rec.width_atr < MIN_WIDTH_ATR or rec.width_pct > MAX_WIDTH_PCT:
            errs.append(f"{rec.range_id} hard metric fail")
        if rec.structure_type != "PRIMARY" or rec.parent_range_id != "none" or not rec.pass_primary:
            errs.append(f"{rec.range_id} not PRIMARY")
        key = (rec.start_idx, rec.confirmed_idx, rec.end_idx, rec.high, rec.low)
        orig = orig_confirmed.get(key)
        if orig is not None and rec.confirmed_at != orig:
            errs.append(f"{rec.range_id} confirmed_at altered by hierarchy")
    for r in range(2, wb["EQ Crossings"].max_row + 1):
        rid = wb["EQ Crossings"].cell(r, 1).value
        if rid and rid not in idset:
            errs.append(f"orphan/nested {rid} on EQ")
    if hq and "Quality Score" in hq_headers:
        qi = hq_headers.index("Quality Score") + 1
        for r in hq_data:
            if float(hqw.cell(r, qi).value) < 70:
                errs.append("HQ score <70")
        if "Structure Type" in hq_headers:
            si = hq_headers.index("Structure Type") + 1
            for r in hq_data:
                if str(hqw.cell(r, si).value) != "PRIMARY":
                    errs.append("HQ not PRIMARY")
    if mod:
        qi = [c.value for c in mod_ws[1]].index("Quality Score") + 1
        for r in mod_data:
            q = float(mod_ws.cell(r, qi).value)
            if q < 60 or q >= 70:
                errs.append(f"moderate score {q}")
    if hqw.freeze_panes != "A2":
        errs.append("HQ freeze")
    if not hqw.auto_filter or not hqw.auto_filter.ref:
        errs.append("HQ filter")
    if hq:
        colh = hq_headers.index("Range High") + 1
        if hqw.cell(hq_data[0], colh).number_format != PRICE_FMT:
            errs.append("price format")
    if wb.vba_archive is not None:
        errs.append("macros")
    if getattr(wb, "_external_links", None):
        errs.append("external links")
    dws = wb["Diagnostics"]
    for r in range(2, dws.max_row + 1):
        name = str(dws.cell(r, 1).value or "")
        if "NESTED" in name.upper():
            errs.append(f"non-PRIMARY count label in Diagnostics: {name}")
    fps = [fingerprint(Path(x["path"])) for x in integrity["source_files"]]
    if fps != fps_before:
        errs.append("source files changed")
    for p, before in prev_before.items():
        if Path(p).is_file() and fingerprint(Path(p)) != before:
            errs.append(f"previous workbook changed {p}")
    if not fps_equal(fingerprint_unrelated_project_code(), code_before):
        errs.append("existing code changed")
    wb.close()
    return errs


def annotate_ranges(rows, *, symbol, output_revision, script_hash, run_id, created):
    for r in rows:
        r.symbol = symbol
        r.market = MARKET_LABEL
        r.timeframe = TIMEFRAME
        r.detector_version = DETECTOR_VERSION
        r.configuration_version = CONFIGURATION_VERSION
        r.output_revision = output_revision
        r.script_sha256 = script_hash
        r.run_id = run_id
        r.creation_time_utc = created
        r.hierarchy_mode = HIERARCHY_MODE


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="run_primary_range_detector",
        description="Isolated PRIMARY range detector for Binance USD-M perpetual 1h Parquet data.",
    )
    parser.add_argument("--symbol", required=True, help="Exact perpetual symbol, e.g. BTCUSDT")
    parser.add_argument("--data-root", default=None, help="MarketData root (default C:\\MarketData or MARKET_DATA_ROOT)")
    parser.add_argument("--output-dir", default=None, help="Workbook output directory (default Windows desktop)")
    parser.add_argument("--timeframe", default=TIMEFRAME, help="Must be 1h for this detector version")
    parser.add_argument("--dry-run", action="store_true", help="Inspect sources and print next revision; write nothing")
    return parser.parse_args(argv)


def run_primary_range_detector(
    symbol: str,
    *,
    data_root: Path | None = None,
    output_dir: Path | None = None,
    timeframe: str = TIMEFRAME,
    dry_run: bool = False,
) -> dict:
    symbol = validate_symbol(symbol)
    timeframe = validate_timeframe(timeframe)
    data_root = Path(data_root) if data_root is not None else default_data_root()
    output_dir = Path(output_dir) if output_dir is not None else DEFAULT_OUTPUT_DIR
    synth = synthetic_pair_test()
    width_synth = synthetic_width_tests()
    extra_synth = synthetic_zero_additional_test()
    print("SYNTHETIC_PAIR_TEST", synth, flush=True)
    print("SYNTHETIC_WIDTH_TEST", width_synth, flush=True)
    print("SYNTHETIC_ZERO_ADDITIONAL_TEST", extra_synth, flush=True)
    inventory = inspect_source_inventory(data_root, symbol)
    rev_n, rev_label, intended = allocate_output_path(output_dir, symbol)
    thresholds = active_thresholds()
    if dry_run:
        result = {
            "dry_run": True,
            "symbol": symbol,
            "timeframe": TIMEFRAME,
            "market": MARKET_LABEL,
            "data_root": str(data_root),
            "dataset_root": inventory["dataset_root"],
            "partition_count": inventory["partition_count"],
            "coverage_start": iso_z(inventory["range_start"]),
            "coverage_end_exclusive": iso_z(inventory["range_end_exclusive"]),
            "intended_workbook": str(intended),
            "allocated_output_revision": rev_label,
            "detector_version": DETECTOR_VERSION,
            "configuration_version": CONFIGURATION_VERSION,
            "hierarchy_classification_mode": HIERARCHY_MODE,
            "thresholds": thresholds,
            "wrote_workbook": False,
        }
        print("DRY_RUN_JSON_BEGIN")
        print(json.dumps(result, indent=2, default=str))
        print("DRY_RUN_JSON_END")
        return result

    prev_before = fingerprint_matching_workbooks(output_dir, symbol)
    code_before = fingerprint_unrelated_project_code()
    data, integrity, sl, inventory = load_and_verify(symbol, data_root)
    if data is None:
        print("INTEGRITY_FAILED", integrity["problems"])
        raise PrimaryRangeDetectorError("source integrity failed: " + "; ".join(integrity["problems"]))
    times, o, h, l, c, opens = data
    fps_before = [fingerprint(Path(x["path"])) for x in integrity["source_files"]]
    atr = compute_atr(h, l, c)
    raw, diag, ph, pl = detect(times, o, h, l, c, atr)
    internal_raw = diag["raw_candidates"]
    assign_temp_ids(raw)
    classify_hierarchy(raw)
    primary_hard = [r for r in raw if r.structure_type == "PRIMARY"]
    for r in primary_hard:
        if r.confirmed_at != (r.confirmed_at_original or r.confirmed_at):
            raise RuntimeError("confirmed_at mutated by hierarchy")
    cleaned, n_dup, n_merge = dedupe_and_merge(primary_hard, atr, times, h, l, c, ph, pl)
    hard_pass = list(cleaned)
    low = [r for r in hard_pass if r.quality < Decimal("60")]
    primary_scored = [r for r in hard_pass if r.quality >= Decimal("60")]
    orig_confirmed = {
        (r.start_idx, r.confirmed_idx, r.end_idx, r.high, r.low): r.confirmed_at_original or r.confirmed_at
        for r in primary_scored
    }
    assign_primary_ids(primary_scored)
    primary_scored.sort(key=lambda r: (r.start, r.confirmed_at, r.range_id))
    created = datetime.now(tz=UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    script_hash = script_sha256()
    run_id = new_run_id()
    annotate_ranges(
        primary_scored, symbol=symbol, output_revision=rev_label,
        script_hash=script_hash, run_id=run_id, created=created,
    )
    hq = [r for r in primary_scored if r.quality >= Decimal("70")]
    mod = [r for r in primary_scored if Decimal("60") <= r.quality < Decimal("70")]
    exclusive_sum = sum(diag["exclusive"].values())
    reported_raw = exclusive_sum + len(primary_hard)
    diag["raw_candidates_internal"] = internal_raw
    diag["raw_candidates"] = reported_raw
    diag["near_dupes"] = n_dup
    diag["merged"] = n_merge
    diag["hard_pass_after_dedupe"] = len(hard_pass)
    diag["low_q"] = len(low)
    diag["failed_hard_exclusive"] = exclusive_sum
    diag["qual_pot_u"] = sum(1 for r in primary_scored for t in r.touches if t.boundary == "UPPER")
    diag["qual_pot_l"] = sum(1 for r in primary_scored for t in r.touches if t.boundary == "LOWER")
    diag["qual_vis_u"] = sum(r.upper_touches for r in primary_scored)
    diag["qual_vis_l"] = sum(r.lower_touches for r in primary_scored)
    diag["qual_anc_u"] = sum(r.u_anchor_count for r in primary_scored)
    diag["qual_anc_l"] = sum(r.l_anchor_count for r in primary_scored)
    diag["qual_add_u"] = sum(r.add_upper for r in primary_scored)
    diag["qual_add_l"] = sum(r.add_lower for r in primary_scored)
    diag["qual_zero_add_u"] = sum(1 for r in primary_scored if r.add_upper == 0)
    diag["qual_zero_add_l"] = sum(1 for r in primary_scored if r.add_lower == 0)
    if primary_scored:
        diag["qual_min_width_pct"] = format(min(r.width_pct for r in primary_scored), "f")
        diag["qual_max_width_pct"] = format(max(r.width_pct for r in primary_scored), "f")
    else:
        diag["qual_min_width_pct"] = "n/a"
        diag["qual_max_width_pct"] = "n/a"
    if len(HQ_COLS) != len(row_summary(primary_scored[0] if primary_scored else RangeRec())):
        raise RuntimeError(f"HQ_COLS mismatch {len(HQ_COLS)} {len(row_summary(RangeRec()))}")
    if primary_scored and len(DETAIL_COLS) != len(row_detail(primary_scored[0])):
        raise RuntimeError(f"DETAIL_COLS mismatch {len(DETAIL_COLS)} {len(row_detail(primary_scored[0]))}")

    rev_n, rev_label, intended = allocate_output_path(output_dir, symbol)
    annotate_ranges(
        primary_scored, symbol=symbol, output_revision=rev_label,
        script_hash=script_hash, run_id=run_id, created=created,
    )
    tmp_dir = Path(tempfile.mkdtemp(prefix="primary_range_v1_"))
    tmp_path = tmp_dir / f"{symbol}_1H_Range_Detection_{rev_label}.{run_id}.part.xlsx"
    meta = {
        "symbol": symbol,
        "dataset_root": integrity["dataset_root"],
        "workbook_path": str(intended),
        "output_revision": rev_label,
        "script_sha256": script_hash,
        "run_id": run_id,
        "temp_path": str(tmp_path),
    }
    prev_after_ok = all(fingerprint(Path(p)) == fp for p, fp in prev_before.items())
    fps_after = {
        "source": "UNCHANGED" if [fingerprint(Path(x["path"])) for x in integrity["source_files"]] == fps_before else "CHANGED",
        "previous_wbs": "UNCHANGED" if prev_after_ok else "CHANGED",
        "code": "UNCHANGED" if fps_equal(fingerprint_unrelated_project_code(), code_before) else "CHANGED",
        "temp_cleanup": "PENDING",
    }
    build_workbook(hq, mod, primary_scored, diag, integrity, fps_after, created, synth, width_synth, extra_synth, meta)
    errs = validate_workbook(
        tmp_path, hq, mod, primary_scored, integrity, fps_before, prev_before, code_before,
        synth, width_synth, extra_synth, orig_confirmed, meta,
    )
    if errs:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise PrimaryRangeDetectorError("temporary workbook validation failed: " + "; ".join(errs))

    committed = None
    last_error = None
    for _attempt in range(32):
        if intended.exists():
            rev_n, rev_label, intended = allocate_output_path(output_dir, symbol)
            meta["workbook_path"] = str(intended)
            meta["output_revision"] = rev_label
            for r in primary_scored:
                r.output_revision = rev_label
            _stamp_output_revision(tmp_path, rev_label, str(intended))
        try:
            os.rename(tmp_path, intended)
            committed = intended
            break
        except FileExistsError as exc:
            last_error = exc
            continue
        except OSError as exc:
            last_error = exc
            if intended.exists():
                continue
            shutil.rmtree(tmp_dir, ignore_errors=True)
            raise
    if committed is None:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise PrimaryRangeDetectorError(f"failed to commit workbook without overwrite: {last_error}")

    meta["workbook_path"] = str(committed)
    meta["output_revision"] = rev_label
    for r in primary_scored:
        r.output_revision = rev_label

    try:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        temp_cleanup = "REMOVED"
    except OSError:
        temp_cleanup = "PARTIAL"
    fps_after["temp_cleanup"] = temp_cleanup

    # Reopen the committed workbook and confirm it still validates.
    post_errs = validate_workbook(
        committed, hq, mod, primary_scored, integrity, fps_before, prev_before, code_before,
        synth, width_synth, extra_synth, orig_confirmed, meta,
    )
    errs = list(errs) + post_errs

    immutable_src = [fingerprint(Path(x["path"])) for x in integrity["source_files"]] == fps_before
    immutable_prev = all(Path(p).is_file() and fingerprint(Path(p)) == fp for p, fp in prev_before.items())
    immutable_code = fps_equal(fingerprint_unrelated_project_code(), code_before)
    import openpyxl as _ox
    recon = exclusive_sum + n_dup + len(low) + len(primary_scored)
    merge_adjust = n_merge
    recon_note = f"merges={n_merge}" if n_merge else "no merges"
    overwritten = str(committed.resolve()) in prev_before
    result = {
        "dry_run": False,
        "workbook": str(committed),
        "size": committed.stat().st_size if committed.exists() else 0,
        "excel_writer": f"openpyxl {_ox.__version__}",
        "detector_version": DETECTOR_VERSION,
        "configuration_version": CONFIGURATION_VERSION,
        "output_revision": rev_label,
        "symbol": symbol,
        "timeframe": TIMEFRAME,
        "dataset_root": integrity["dataset_root"],
        "rows": integrity["rows"], "first": integrity["first"], "last": integrity["last"],
        "raw_primary_candidates": reported_raw,
        "fail_duration": diag["fail_duration"],
        "fail_upper_touches": diag["fail_upper_touches"],
        "fail_lower_touches": diag["fail_lower_touches"],
        "fail_upper_pair": diag["fail_upper_pair"],
        "fail_lower_pair": diag["fail_lower_pair"],
        "fail_containment": diag["fail_containment"],
        "fail_slope": diag["fail_slope"],
        "fail_efficiency": diag["fail_efficiency"],
        "fail_eq": diag["fail_eq"],
        "fail_width": diag["fail_width"],
        "fail_width_pct": diag["fail_width_pct"],
        "fail_causality": diag["fail_causality"],
        "qual_pot_u": diag["qual_pot_u"], "qual_vis_u": diag["qual_vis_u"], "qual_anc_u": diag["qual_anc_u"], "qual_add_u": diag["qual_add_u"],
        "qual_pot_l": diag["qual_pot_l"], "qual_vis_l": diag["qual_vis_l"], "qual_anc_l": diag["qual_anc_l"], "qual_add_l": diag["qual_add_l"],
        "qual_zero_add_u": diag["qual_zero_add_u"], "qual_zero_add_l": diag["qual_zero_add_l"],
        "hard_pass_after_dedupe": len(hard_pass),
        "low_quality_excluded": len(low),
        "near_dupes": n_dup,
        "merged": n_merge,
        "high_quality_primary": len(hq),
        "moderate_primary": len(mod),
        "retained_primary": len(primary_scored),
        "qual_min_width_pct": diag.get("qual_min_width_pct"),
        "qual_max_width_pct": diag.get("qual_max_width_pct"),
        "touches": sum(len(r.touches) for r in primary_scored),
        "eq": sum(len(r.crossings) for r in primary_scored),
        "open": sum(1 for r in primary_scored if r.status == "OPEN_AT_DATASET_END"),
        "broken_up": sum(1 for r in primary_scored if r.status == "BROKEN_UP"),
        "broken_down": sum(1 for r in primary_scored if r.status == "BROKEN_DOWN"),
        "expired": diag["expired"],
        "safety_expiration": f"{CANDIDATE_EXPIRE_BARS} completed 1h bars from range_start",
        "hierarchy_mode": HIERARCHY_MODE,
        "additional_hard_filter": False,
        "non_primary_in_workbook": False,
        "overwrote_existing": overwritten,
        "exclusive": diag["exclusive"],
        "reconciliation": {
            "exclusive_hard": exclusive_sum, "near_dupes": n_dup, "low_q": len(low),
            "retained_primary": len(primary_scored), "sum": recon, "raw_primary": reported_raw,
            "note": recon_note,
        },
        "source_immutable": immutable_src,
        "previous_workbooks_immutable": immutable_prev,
        "previous_workbooks_fingerprinted": list(prev_before.keys()),
        "code_immutable": immutable_code,
        "temp_cleanup": temp_cleanup,
        "script_sha256": script_hash,
        "run_id": run_id,
        "synthetic_pair_test": synth,
        "synthetic_width_tests": width_synth,
        "synthetic_zero_additional_test": extra_synth,
        "primary_only_policy": "PASS" if all(r.structure_type == "PRIMARY" and r.parent_range_id == "none" for r in primary_scored) else "FAIL",
        "validation_errors": errs,
        "preview": [
            {
                "id": r.range_id, "structure": r.structure_type, "start": iso_z(r.start),
                "confirmed": iso_z(r.confirmed_at), "end": iso_z(r.end),
                "bars": r.duration_bars, "pre": r.pre_conf_bars, "post": r.post_conf_bars,
                "high": format(r.high, "f"), "low": format(r.low, "f"), "eq": format(r.eq, "f"),
                "width_pct": format(r.width_pct, "f"), "icr": format(r.inside_ratio, "f"),
                "upper": r.upper_touches, "u_anchors": r.u_anchor_count, "u_additional": r.add_upper,
                "lower": r.lower_touches, "l_anchors": r.l_anchor_count, "l_additional": r.add_lower,
                "u_sep": r.u_pair_sep, "l_sep": r.l_pair_sep,
                "score": str(r.quality.quantize(Decimal("0.01"))), "breakout": r.status,
            }
            for r in hq[:5]
        ],
        "warnings": [] if recon + merge_adjust == reported_raw or recon == reported_raw else [f"reconciliation sum {recon} vs raw {reported_raw} {recon_note}"],
        "integrity": {k: integrity[k] for k in ("rows", "first", "last", "nulls", "ohlc_violations", "duplicates", "out_of_order", "missing")},
    }
    print("RESULT_JSON_BEGIN")
    print(json.dumps(result, indent=2, default=str))
    print("RESULT_JSON_END")
    if errs or not immutable_src or not immutable_prev or not immutable_code or overwritten:
        raise PrimaryRangeDetectorError("post-commit validation or immutability failed")
    return result


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        run_primary_range_detector(
            args.symbol,
            data_root=Path(args.data_root) if args.data_root else None,
            output_dir=Path(args.output_dir) if args.output_dir else None,
            timeframe=args.timeframe,
            dry_run=args.dry_run,
        )
    except PrimaryRangeDetectorError as exc:
        print(f"PRIMARY_RANGE_DETECTOR_ERROR: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
