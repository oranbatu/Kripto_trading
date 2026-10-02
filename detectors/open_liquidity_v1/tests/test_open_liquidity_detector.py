"""Deterministic tests for completed-candle open liquidity."""
from __future__ import annotations

import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from data.catalog import fingerprint_file
from detectors.open_liquidity_v1.engine import (
    OPEN,
    TERMINAL,
    Bar,
    Candidate,
    _GroupBuilder,
    attach_children,
    brute_first,
    brute_suffix,
    detect_timeframe,
    first_reach,
    hierarchy_tier,
    iso_turkey,
    monthly_summary,
    origin_classification,
    representative_price,
    status_counts,
    suffix_extremes,
    wilder_atr,
    analyze,
)
from detectors.open_liquidity_v1.engine import _sparse_extreme
from detectors.open_liquidity_v1.liquidity_config import EXPECTED_CANDIDATES, EXPECTED_ROWS, PROJECT_ROOT
from detectors.open_liquidity_v1.run_open_liquidity_detector import next_revision, revision_filename, scan_revisions
from detectors.open_liquidity_v1.workbook import excel_value, round_trip_workbook, validate_ooxml_workbook, write_workbook

D = Decimal


def _series(specs, timeframe="1h", start=None):
    step = {"1h": timedelta(hours=1), "4h": timedelta(hours=4), "1d": timedelta(days=1)}[timeframe]
    origin = start or datetime(2024, 1, 1, tzinfo=timezone.utc)
    rows = []
    for index, values in enumerate(specs):
        opened = origin + step * index
        rows.append(
            Bar(
                timeframe,
                index,
                opened,
                opened + step - timedelta(milliseconds=1),
                D(str(values[0])),
                D(str(values[1])),
                D(str(values[2])),
                D(str(values[3])),
                D("10"),
                3,
            )
        )
    return rows


def _highs(candidates):
    return [item for item in candidates if item.side == "HIGH"]


def _lows(candidates):
    return [item for item in candidates if item.side == "LOW"]


def _dummy(timeframe, side, price, row=0, future=10, candidate_id="X"):
    opened = datetime(2024, 1, 1, tzinfo=timezone.utc) + timedelta(hours=row)
    return Candidate(
        candidate_id=candidate_id,
        timeframe=timeframe,
        side=side,
        row=row,
        open_time=opened,
        close_time=opened + timedelta(hours=1),
        price=D(str(price)),
        open=D("1"),
        high=D(str(price)) if side == "HIGH" else D("2"),
        low=D(str(price)) if side == "LOW" else D("1"),
        close=D("1"),
        volume=D("1"),
        trades=1,
        atr=D("1"),
        upper_wick=D("0"),
        lower_wick=D("0"),
        body=D("0"),
        wick_to_body=None,
        origin_class="BODY_AT_HIGH" if side == "HIGH" else "BODY_AT_LOW",
        status=OPEN,
        future_bars=future,
        warning="HIGH_FUTURE_OBSERVATION",
    )


class OpenLiquidityTests(unittest.TestCase):
    def test_01_later_high_below_leaves_high_open(self):
        found = _highs(detect_timeframe(_series([(100, 110, 90, 100), (100, 109, 90, 100)])))
        self.assertEqual(found[0].status, OPEN)

    def test_02_equal_high_mitigates(self):
        found = _highs(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 90, 100)])))
        self.assertEqual(found[0].status, "MITIGATED")
        self.assertTrue(found[0].exact_flag)

    def test_03_higher_high_mitigates(self):
        found = _highs(detect_timeframe(_series([(100, 110, 90, 100), (100, 111, 90, 100)])))
        self.assertEqual(found[0].status, "MITIGATED")

    def test_04_later_low_above_leaves_low_open(self):
        found = _lows(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 91, 100)])))
        self.assertEqual(found[0].status, OPEN)

    def test_05_equal_low_mitigates(self):
        found = _lows(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 90, 100)])))
        self.assertEqual(found[0].status, "MITIGATED")
        self.assertTrue(found[0].exact_flag)

    def test_06_lower_low_mitigates(self):
        found = _lows(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 89, 100)])))
        self.assertEqual(found[0].status, "MITIGATED")

    def test_07_wick_touch_without_close_beyond(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (95, 101, 94, 96)])))
        self.assertEqual(found[0].status, "MITIGATED")
        self.assertEqual(found[0].mechanism, "WICK_REACH")
        self.assertEqual(found[0].close_status, "NOT_CLOSE_MITIGATED")

    def test_08_close_beyond_sets_both_statuses(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (99, 105, 98, 102)])))
        self.assertEqual(found[0].touch_status, "MITIGATED")
        self.assertEqual(found[0].close_status, "CLOSE_MITIGATED")
        self.assertEqual(found[0].mechanism, "CLOSE_AT_OR_BEYOND")

    def test_09_origin_cannot_mitigate_itself(self):
        found = detect_timeframe(_series([(90, 100, 80, 95)]))
        self.assertTrue(all(item.mitigation_row is None for item in found))
        self.assertTrue(all(item.status == TERMINAL for item in found))

    def test_10_only_strictly_later_candles(self):
        found = detect_timeframe(_series([(90, 100, 80, 90), (90, 99, 81, 90), (90, 101, 80, 90)]))
        high = _highs(found)[0]
        self.assertGreater(high.mitigation_row, high.row)

    def test_11_first_reaching_candle_is_selected(self):
        found = _highs(detect_timeframe(_series([
            (90, 100, 80, 90),
            (90, 99, 80, 90),
            (90, 110, 80, 90),
            (90, 120, 80, 90),
        ])))
        self.assertEqual(found[0].mitigation_row, 2)

    def test_12_later_reaching_candle_does_not_replace_first(self):
        found = _highs(detect_timeframe(_series([
            (90, 100, 80, 90),
            (90, 130, 80, 90),
            (90, 140, 80, 90),
        ])))
        self.assertEqual(found[0].mitigation_row, 1)
        self.assertEqual(found[0].penetration, D("30"))

    def test_13_gap_above_high_mitigates(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (105, 110, 104, 106)])))
        self.assertEqual(found[0].mechanism, "GAP_OPEN_AT_OR_BEYOND")
        self.assertEqual(found[0].status, "MITIGATED")

    def test_14_gap_below_low_mitigates(self):
        found = _lows(detect_timeframe(_series([(100, 110, 100, 105), (90, 95, 89, 94)])))
        self.assertEqual(found[0].mechanism, "GAP_OPEN_AT_OR_BEYOND")
        self.assertEqual(found[0].status, "MITIGATED")

    def test_15_exact_decimal_comparison(self):
        rows = _series([(100, "100.00000001", 90, 100), (100, "100.00000000", 90, 100)])
        found = _highs(detect_timeframe(rows))
        self.assertEqual(found[0].status, OPEN)
        self.assertEqual(found[0].price, D("100.00000001"))

    def test_16_near_miss_does_not_mitigate(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, "99.95", 80, 90)])))
        self.assertEqual(found[0].status, OPEN)
        self.assertTrue(found[0].near_miss)
        wider = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, "99.94", 80, 90)])))
        self.assertFalse(wider[0].near_miss)

    def test_17_equal_high_closes_older_candidate(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, 100, 80, 90), (90, 99, 80, 90)])))
        self.assertEqual(found[0].status, "MITIGATED")

    def test_18_later_equal_high_creates_a_new_candidate(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, 100, 80, 90), (90, 99, 80, 90)])))
        self.assertEqual(found[1].status, OPEN)
        self.assertEqual(found[1].price, found[0].price)
        self.assertNotEqual(found[0].candidate_id, found[1].candidate_id)

    def test_19_equal_low_closes_older_candidate(self):
        found = _lows(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 90, 100), (100, 110, 91, 100)])))
        self.assertEqual(found[0].status, "MITIGATED")

    def test_20_later_equal_low_creates_a_new_candidate(self):
        found = _lows(detect_timeframe(_series([(100, 110, 90, 100), (100, 110, 90, 100), (100, 110, 91, 100)])))
        self.assertEqual(found[1].status, OPEN)
        self.assertEqual(found[1].price, D("90"))

    def test_21_equal_level_chains_reconcile(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, 100, 80, 90), (90, 99, 80, 90)])))
        self.assertEqual(found[0].chain_id, found[1].chain_id)
        self.assertEqual(found[0].chain_length, 2)
        self.assertEqual(found[0].next_equal_id, found[1].candidate_id)
        self.assertEqual(found[1].previous_equal_id, found[0].candidate_id)
        self.assertEqual(found[0].chain_survivor_id, found[1].candidate_id)
        self.assertEqual(found[0].chain_sequence, 1)
        self.assertEqual(found[1].chain_sequence, 2)

    def test_22_final_high_is_terminal(self):
        found = _highs(detect_timeframe(_series([(90, 100, 80, 90), (90, 101, 80, 90)])))
        self.assertEqual(found[-1].status, TERMINAL)

    def test_23_final_low_is_terminal(self):
        found = _lows(detect_timeframe(_series([(90, 100, 80, 90), (90, 101, 80, 90)])))
        self.assertEqual(found[-1].status, TERMINAL)

    def test_24_terminal_is_excluded_from_main_open_sheets(self):
        found = detect_timeframe(_series([(90, 100, 80, 90), (90, 99, 81, 90)]))
        main = [item for item in found if item.status == OPEN]
        self.assertTrue(all(item.status != TERMINAL for item in main))
        self.assertTrue(any(item.status == TERMINAL for item in found))

    def test_25_one_later_observation_may_be_open(self):
        found = _highs(detect_timeframe(_series([(90, 110, 80, 90), (90, 100, 80, 90)])))
        self.assertEqual(found[0].status, OPEN)
        self.assertEqual(found[0].future_bars, 1)

    def test_26_open_candidate_is_right_censored(self):
        found = _highs(detect_timeframe(_series([(90, 110, 80, 90), (90, 100, 80, 90)])))
        self.assertEqual(found[0].status, OPEN)
        self.assertGreaterEqual(found[0].future_bars, 1)

    def test_27_candidate_count_is_twice_the_candle_count(self):
        rows = _series([(90, 100, 80, 90)] * 5)
        self.assertEqual(len(detect_timeframe(rows)), 10)

    def test_28_expected_candidate_total_is_61318(self):
        self.assertEqual(EXPECTED_ROWS["1h"] * 2 + EXPECTED_ROWS["4h"] * 2 + EXPECTED_ROWS["1d"] * 2, 61318)
        self.assertEqual(EXPECTED_CANDIDATES, 61318)

    def test_29_statuses_are_mutually_exclusive(self):
        found = detect_timeframe(_series([(90, 100, 80, 90), (90, 101, 79, 90), (90, 99, 81, 90)]))
        for item in found:
            flags = [item.status == OPEN, item.status == "MITIGATED", item.status == TERMINAL]
            self.assertEqual(sum(flags), 1)

    def test_30_status_reconciliation_is_exact(self):
        result = analyze({"1h": _series([(90, 100, 80, 90), (90, 101, 79, 90), (90, 98, 82, 90)]), "4h": [], "1d": []})
        counts = status_counts(result)["1h"]
        self.assertEqual(counts[OPEN] + counts["MITIGATED"] + counts[TERMINAL], 6)

    def test_31_suffix_high_matches_brute_force(self):
        values = [D(str(item)) for item in (5, 1, 4, 4, 9, 2, 8)]
        self.assertEqual(suffix_extremes(values, True), brute_suffix(values, True))

    def test_32_suffix_low_matches_brute_force(self):
        values = [D(str(item)) for item in (5, 1, 4, 4, 9, 2, 8)]
        self.assertEqual(suffix_extremes(values, False), brute_suffix(values, False))

    def test_33_first_mitigation_lookup_matches_brute_force(self):
        values = [D(str(item)) for item in (3, 1, 4, 2, 8, 5)]
        table = _sparse_extreme(values, True)
        self.assertEqual(first_reach(table, 1, 5, D("4"), True), brute_first(values, 1, D("4"), True))
        self.assertEqual(first_reach(table, 0, 5, D("9"), True), brute_first(values, 0, D("9"), True))

    def test_34_wilder_atr_is_causal(self):
        rows = _series([(10, 12, 9, 11)] * 16)
        first = wilder_atr(rows[:15])
        second = wilder_atr(rows)
        self.assertEqual(first[13], second[13])
        self.assertIsNone(first[12])
        tr = rows[14].high - rows[14].low
        self.assertEqual(second[14], (second[13] * 13 + tr) / 14)

    def test_35_no_lookahead_before_origin_close(self):
        rows = _series([(90, 100, 80, 90), (101, 102, 99, 101)])
        found = _highs(detect_timeframe(rows))[0]
        self.assertEqual(rows[0].close_time, rows[0].open_time + timedelta(hours=1) - timedelta(milliseconds=1))
        self.assertGreater(rows[found.mitigation_row].open_time, rows[0].open_time)
        self.assertGreater(rows[found.mitigation_row].open_time, rows[0].close_time)

    def test_36_upper_wick_classification(self):
        self.assertEqual(origin_classification("HIGH", D("10"), D("12"), D("9"), D("11")), "UPPER_WICK_EXTREME")

    def test_37_body_at_high_classification(self):
        self.assertEqual(origin_classification("HIGH", D("10"), D("12"), D("9"), D("12")), "BODY_AT_HIGH")

    def test_38_lower_wick_classification(self):
        self.assertEqual(origin_classification("LOW", D("10"), D("12"), D("8"), D("11")), "LOWER_WICK_EXTREME")

    def test_39_body_at_low_classification(self):
        self.assertEqual(origin_classification("LOW", D("10"), D("12"), D("10"), D("11")), "BODY_AT_LOW")

    def test_40_wick_classification_is_not_a_hard_filter(self):
        found = detect_timeframe(_series([(10, 12, 10, 12), (11, 11, 9, 9)]))
        classes = {item.origin_class for item in found}
        self.assertIn("BODY_AT_HIGH", classes)
        self.assertIn("BODY_AT_LOW", classes)
        self.assertEqual(len(found), 4)

    def test_41_open_high_distance(self):
        result = analyze({"1h": _series([(100, 110, 90, 100), (100, 105, 95, 100)]), "4h": [], "1d": []})
        high = [item for item in result.candidates["1h"] if item.side == "HIGH" and item.status == OPEN][0]
        self.assertEqual(high.distance_price, D("10"))
        self.assertEqual(high.distance_percent, D("10"))

    def test_42_open_low_distance(self):
        result = analyze({"1h": _series([(100, 110, 90, 100), (100, 105, 95, 100)]), "4h": [], "1d": []})
        low = [item for item in result.candidates["1h"] if item.side == "LOW" and item.status == OPEN][0]
        self.assertEqual(low.distance_price, D("10"))
        self.assertEqual(low.distance_percent, D("10"))

    def test_43_nearest_high_is_ranked_first(self):
        rows = _series([(100, 130, 80, 100), (100, 110, 90, 100), (100, 105, 95, 100)])
        result = analyze({"1h": rows, "4h": [], "1d": []})
        ranked = [item for item in result.candidates["1h"] if item.side == "HIGH" and item.status == OPEN]
        ranked.sort(key=lambda item: item.side_rank)
        self.assertEqual(ranked[0].price, D("110"))
        self.assertEqual(ranked[1].price, D("130"))

    def test_44_nearest_low_is_ranked_first(self):
        rows = _series([(100, 120, 70, 100), (100, 115, 90, 100), (100, 110, 95, 100)])
        result = analyze({"1h": rows, "4h": [], "1d": []})
        ranked = [item for item in result.candidates["1h"] if item.side == "LOW" and item.status == OPEN]
        ranked.sort(key=lambda item: item.side_rank)
        self.assertEqual(ranked[0].price, D("90"))
        self.assertEqual(ranked[1].price, D("70"))

    def test_45_tier_1_all_timeframes(self):
        self.assertEqual(hierarchy_tier(["1d", "4h", "1h"]), "TIER_1_ALL_TIMEFRAMES")

    def test_46_tier_2_daily_with_intraday(self):
        self.assertEqual(hierarchy_tier(["1d", "4h"]), "TIER_2_DAILY_WITH_INTRADAY")
        self.assertEqual(hierarchy_tier(["1d", "1h"]), "TIER_2_DAILY_WITH_INTRADAY")

    def test_47_tier_3_daily_only(self):
        self.assertEqual(hierarchy_tier(["1d"]), "TIER_3_DAILY_ONLY")

    def test_48_tier_4_4h_and_1h(self):
        self.assertEqual(hierarchy_tier(["4h", "1h"]), "TIER_4_4H_AND_1H")

    def test_49_tier_5_4h_only(self):
        self.assertEqual(hierarchy_tier(["4h"]), "TIER_5_4H_ONLY")

    def test_50_tier_6_1h_only(self):
        self.assertEqual(hierarchy_tier(["1h"]), "TIER_6_1H_ONLY")

    def test_51_high_never_maps_to_low(self):
        parent = _dummy("1d", "HIGH", "100", candidate_id="H")
        child = _dummy("4h", "LOW", "100", candidate_id="L")
        attempts = []
        unmatched = attach_children([child], [_GroupBuilder(parent)], attempts)
        self.assertEqual(unmatched, [child])
        self.assertTrue(all(item.rejection_reason == "HIGH_LOW_SIDE_MISMATCH" for item in attempts))
        self.assertFalse(any(item.accepted for item in attempts))

    def test_52_one_child_maps_to_one_parent(self):
        daily = _series([(80, 100, 50, 80), (80, "99.95", 50, 80), (80, 90, 50, 80)], "1d")
        four = _series([(80, "99.97", 50, 80), (80, 90, 50, 80)], "4h")
        result = analyze({"1d": daily, "4h": four, "1h": []})
        accepted = [item for item in result.attempts if item.accepted and item.child_id.startswith("LQ-4H")]
        self.assertEqual(len(accepted), 1)
        child_ids = [item.candidate_id for item in result.candidates["4h"] if item.status == OPEN]
        self.assertEqual(len({item.group_id for item in result.candidates["4h"] if item.candidate_id in child_ids}), 1)
        crowded_daily = _series([(80, 100, 50, 80), (80, 90, 50, 80)], "1d")
        crowded_four = _series([(80, "100.08", 50, 80), (80, "100.04", 50, 80), (80, 90, 50, 80)], "4h")
        crowded = analyze({"1d": crowded_daily, "4h": crowded_four, "1h": []})
        open_four = [item for item in crowded.candidates["4h"] if item.status == OPEN and item.side == "HIGH"]
        self.assertEqual(len(open_four), 2)
        self.assertEqual(len({item.group_id for item in open_four}), 2)
        self.assertTrue(all(item.group_id for item in open_four))

    def test_53_group_span_prevents_chain_merging(self):
        daily = _series([(80, 100, 50, 80), (80, 90, 50, 80)], "1d")
        four = _series([(80, "100.09", 50, 80), (80, 90, 50, 80)], "4h")
        hour = _series([(80, "100.19", 50, 80), (80, 90, 50, 80)], "1h")
        result = analyze({"1d": daily, "4h": four, "1h": hour})
        daily_group = next(group for group in result.groups if "1d" in group.members and group.side == "HIGH")
        self.assertNotIn("1h", daily_group.members)
        self.assertIn("4h", daily_group.members)
        anchor = _dummy("1d", "HIGH", "100", candidate_id="D")
        wide = _dummy("4h", "HIGH", "100.30", candidate_id="W")
        child = _dummy("1h", "HIGH", "100.04", candidate_id="H")
        builder = _GroupBuilder(anchor)
        builder.add(wide)
        attempts = []
        unmatched = attach_children([child], [builder], attempts)
        self.assertEqual(unmatched, [child])
        self.assertEqual(attempts[0].rejection_reason, "GROUP_SPAN_ABOVE_0_20")

    def test_54_higher_timeframe_sets_representative_price(self):
        members = {
            "1d": _dummy("1d", "HIGH", "100", candidate_id="D"),
            "1h": _dummy("1h", "HIGH", "100.05", candidate_id="H"),
        }
        self.assertEqual(representative_price(members), D("100"))
        four = _dummy("4h", "HIGH", "80", candidate_id="F")
        hour = _dummy("1h", "HIGH", "90", candidate_id="Z")
        self.assertEqual(representative_price({"4h": four, "1h": hour}), D("80"))

    def test_55_utc_to_turkey_uses_europe_istanbul(self):
        moment = datetime(2024, 6, 15, 21, tzinfo=timezone.utc)
        rendered = iso_turkey(moment)
        self.assertTrue(rendered.startswith("2024-06-16T00:00:00.000+03:00"))
        self.assertEqual(moment.astimezone(ZoneInfo("Europe/Istanbul")).isoformat(timespec="milliseconds"), rendered)

    def test_56_revision_starts_at_rev00(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertEqual(scan_revisions(Path(folder)), [])
            self.assertEqual(next_revision([]), 0)
            self.assertEqual(revision_filename(0), "BTCUSDT_MTF_Open_Liquidity_rev00.xlsx")

    def test_57_revision_increments_from_maximum(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "BTCUSDT_MTF_Open_Liquidity_rev02.xlsx").write_bytes(b"x")
            (root / "BTCUSDT_MTF_Open_Liquidity_rev00.xlsx").write_bytes(b"x")
            self.assertEqual(next_revision(scan_revisions(root)), 3)

    def test_58_revision_gaps_are_not_filled(self):
        self.assertEqual(next_revision([0, 4]), 5)
        self.assertNotEqual(revision_filename(next_revision([0, 4])), "BTCUSDT_MTF_Open_Liquidity_rev01.xlsx")

    def test_59_existing_workbook_is_not_reused(self):
        existing = [0, 2]
        chosen = revision_filename(next_revision(existing))
        self.assertNotIn(chosen, {revision_filename(item) for item in existing})

    def test_60_filename_race_rescans(self):
        planned = next_revision([0, 2])
        appeared = [0, 2, planned]
        self.assertEqual(next_revision(appeared), planned + 1)

    def test_61_source_parquet_remains_unchanged(self):
        root = Path(r"C:\MarketData\derived\binance\futures\um\perpetual\1d\symbol=BTCUSDT")
        files = list(root.rglob("*.parquet"))
        self.assertTrue(files)
        before = fingerprint_file(files[0])
        after = fingerprint_file(files[0])
        self.assertEqual(before, after)

    def test_62_range_detector_is_not_imported(self):
        source = _package_source()
        self.assertNotIn("import detectors.primary_range_v1", source)
        self.assertNotIn("from detectors.primary_range_v1", source)

    def test_63_swing_detector_is_not_imported(self):
        source = _package_source()
        self.assertNotIn("import detectors.hierarchical_swing_v4", source)
        self.assertNotIn("from detectors.hierarchical_swing_v4", source)

    def test_64_price_touch_detector_is_not_imported(self):
        source = _package_source()
        self.assertNotIn("import detectors.price_touch_hierarchy_v1", source)
        self.assertNotIn("from detectors.price_touch_hierarchy_v1", source)

    def test_65_strategy_code_is_not_imported(self):
        source = _package_source()
        self.assertNotIn("import strategy", source)
        self.assertNotIn("strategies", source)

    def test_66_no_separate_project_is_created(self):
        package = PROJECT_ROOT / "detectors" / "open_liquidity_v1"
        self.assertTrue(package.is_dir())
        self.assertEqual(package.parent.parent, PROJECT_ROOT)
        self.assertFalse((Path(r"C:\Users\oranb\Desktop") / "open_liquidity_v1").exists())

    def test_67_no_detector_python_outside_the_project(self):
        desktop = Path(r"C:\Users\oranb\Desktop")
        stray = list(desktop.glob("*open_liquidity*.py"))
        self.assertEqual(stray, [])
        script = PROJECT_ROOT / "detectors" / "open_liquidity_v1" / "run_open_liquidity_detector.py"
        self.assertTrue(str(script).startswith(str(PROJECT_ROOT)))

    def test_68_two_runs_match(self):
        bars = {
            "1h": _series([(90, 120, 80, 100), (100, 110, 90, 100), (100, 105, 95, 101)]),
            "4h": _series([(90, 120, 80, 100), (100, 111, 91, 100)], "4h"),
            "1d": _series([(90, 130, 70, 100), (100, 110, 90, 100)], "1d"),
        }
        first = analyze(bars)
        second = analyze(bars)
        for timeframe in bars:
            left = [(item.candidate_id, item.status, item.mitigation_row, item.price) for item in first.candidates[timeframe]]
            right = [(item.candidate_id, item.status, item.mitigation_row, item.price) for item in second.candidates[timeframe]]
            self.assertEqual(left, right)
        self.assertEqual([group.group_id for group in first.groups], [group.group_id for group in second.groups])

    def test_69_workbook_status_counts_reconcile(self):
        result = analyze({"1h": _series([(90, 120, 80, 100), (100, 110, 90, 100)]), "4h": [], "1d": []})
        counts = status_counts(result)["1h"]
        context = {
            "counts": {"1h": {**counts, "rows": 2}, "4h": {OPEN: 0, "MITIGATED": 0, TERMINAL: 0, "HIGH_OPEN": 0, "LOW_OPEN": 0, "rows": 0}, "1d": {OPEN: 0, "MITIGATED": 0, TERMINAL: 0, "HIGH_OPEN": 0, "LOW_OPEN": 0, "rows": 0}},
            "as_of": "2024-01-01T01:59:59.999Z",
            "revision_name": "BTCUSDT_MTF_Open_Liquidity_rev00.xlsx",
            "diagnostic_rows": [["Section", "Metric", "Value"], ["reconciliation", "status_reconciliation", "PASS"]],
        }
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "book.xlsx"
            write_workbook(path, result, context, "readme")
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames[0], "Executive Summary")
                self.assertEqual(len(workbook.sheetnames), 13)
                statuses = [row[6] for row in workbook["Candidate Audit"].iter_rows(min_row=2, values_only=True)]
                self.assertEqual(statuses.count(OPEN) + statuses.count("MITIGATED") + statuses.count(TERMINAL), len(statuses))
            finally:
                workbook.close()

    def test_70_monthly_origin_counts_reconcile(self):
        result = analyze({"1h": _series([(90, 120, 80, 100), (100, 110, 90, 100), (100, 105, 95, 100)]), "4h": [], "1d": []})
        rows = [item for item in monthly_summary(result) if item["timeframe"] == "1h"]
        created = sum(int(item["high_created"]) + int(item["low_created"]) for item in rows)
        opened = sum(int(item["high_open"]) + int(item["low_open"]) for item in rows)
        mitigated = sum(int(item["high_mitigated"]) + int(item["low_mitigated"]) for item in rows)
        terminal = sum(int(item["terminal"]) for item in rows)
        self.assertEqual(created, 6)
        self.assertEqual(opened + mitigated + terminal, created)


class WorkbookCompatibilityTests(unittest.TestCase):
    def _sample_workbook(self, folder: Path):
        result = analyze({"1h": _series([(90, 120, 80, 100), (100, 110, 90, 100)]), "4h": [], "1d": []})
        counts = status_counts(result)["1h"]
        context = {
            "counts": {
                "1h": {**counts, "rows": 2},
                "4h": {OPEN: 0, "MITIGATED": 0, TERMINAL: 0, "HIGH_OPEN": 0, "LOW_OPEN": 0, "rows": 0},
                "1d": {OPEN: 0, "MITIGATED": 0, TERMINAL: 0, "HIGH_OPEN": 0, "LOW_OPEN": 0, "rows": 0},
            },
            "as_of": "2024-01-01T01:59:59.999Z",
            "revision_name": "BTCUSDT_MTF_Open_Liquidity_rev02.xlsx",
            "diagnostic_rows": [["Section", "Metric", "Value"], ["reconciliation", "status_reconciliation", "PASS"]],
        }
        path = folder / "BTCUSDT_MTF_Open_Liquidity_rev02.xlsx"
        write_workbook(path, result, context, "readme line\nsecond line")
        return path, result

    def test_71_generated_file_is_a_zip_xlsx(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            self.assertTrue(path.name.endswith(".xlsx"))
            self.assertEqual(path.read_bytes()[:4], b"PK\x03\x04")
            import zipfile
            self.assertTrue(zipfile.is_zipfile(path))

    def test_72_required_ooxml_members_exist(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            import zipfile
            with zipfile.ZipFile(path) as archive:
                for member in (
                    "[Content_Types].xml",
                    "_rels/.rels",
                    "docProps/app.xml",
                    "docProps/core.xml",
                    "xl/workbook.xml",
                    "xl/_rels/workbook.xml.rels",
                    "xl/styles.xml",
                    "xl/theme/theme1.xml",
                    "xl/worksheets/sheet1.xml",
                ):
                    self.assertIn(member, archive.namelist())
                self.assertIsNone(archive.testzip())

    def test_73_every_xml_entry_parses(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            validate_ooxml_workbook(path)

    def test_74_workbook_reopens_in_normal_mode(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
            try:
                self.assertEqual(len(workbook.sheetnames), 13)
            finally:
                workbook.close()

    def test_75_workbook_reopens_in_read_only_mode(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                rows = list(workbook["README"].iter_rows(values_only=True))
                self.assertGreaterEqual(len(rows), 2)
            finally:
                workbook.close()

    def test_76_thirteen_sheets_stay_in_order(self):
        from detectors.open_liquidity_v1.liquidity_config import WORKBOOK_SHEETS
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(tuple(workbook.sheetnames), WORKBOOK_SHEETS)
            finally:
                workbook.close()

    def test_77_illegal_xml_characters_are_removed(self):
        self.assertEqual(excel_value("A\x00B\x01C"), "A B C")

    def test_78_oversized_strings_are_limited(self):
        self.assertEqual(len(excel_value("x" * 40000)), 32767)

    def test_79_nan_and_infinity_are_blank(self):
        self.assertIsNone(excel_value(float("nan")))
        self.assertIsNone(excel_value(float("inf")))
        self.assertIsNone(excel_value(Decimal("NaN")))

    def test_80_decimal_values_are_numeric(self):
        self.assertEqual(excel_value(Decimal("79570.90000000")), 79570.9)

    def test_81_aware_datetimes_are_text(self):
        moment = datetime(2024, 1, 1, tzinfo=timezone.utc)
        self.assertEqual(excel_value(moment), "2024-01-01T00:00:00.000Z")

    def test_82_table_names_are_unique_and_references_valid(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            validate_ooxml_workbook(path)

    def test_83_empty_result_sheets_remain_valid(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook["4H Open Liquidity"]["A2"].value, "NO_RESULTS")
                self.assertEqual(len(workbook["4H Open Liquidity"].tables), 0)
                self.assertTrue(workbook["4H Open Liquidity"].auto_filter.ref)
                audit = workbook["Candidate Audit"]
                events = workbook["Mitigation Events"]
                self.assertEqual(len(audit.tables), 0)
                self.assertEqual(len(events.tables), 0)
                self.assertTrue(audit.auto_filter.ref)
                self.assertTrue(events.auto_filter.ref)
                self.assertEqual(audit.freeze_panes, "A2")
                self.assertEqual(events.freeze_panes, "A2")
                for title in ("Executive Summary", "Diagnostics", "README"):
                    self.assertEqual(len(workbook[title].tables), 0)
            finally:
                workbook.close()

    def test_84_worksheets_stay_inside_excel_limits(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            workbook = load_workbook(path, read_only=True)
            try:
                for worksheet in workbook.worksheets:
                    self.assertLessEqual(worksheet.max_row, 1_048_576)
                    self.assertLessEqual(worksheet.max_column, 16_384)
                    self.assertLessEqual(len(worksheet.title), 31)
            finally:
                workbook.close()

    def test_85_no_macros_or_external_links(self):
        import zipfile
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            with zipfile.ZipFile(path) as archive:
                self.assertFalse(any("vba" in name.lower() for name in archive.namelist()))
                self.assertFalse(any("externalLink" in name for name in archive.namelist()))

    def test_86_round_trip_preserves_summary(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            copy = Path(folder) / "roundtrip.xlsx"
            round_trip_workbook(path, copy)
            copy.unlink()

    def test_87_revision_does_not_overwrite_failed_workbooks(self):
        self.assertEqual(next_revision([0, 1]), 2)
        chosen = revision_filename(next_revision([0, 1]))
        self.assertNotIn(chosen, {"BTCUSDT_MTF_Open_Liquidity_rev00.xlsx", "BTCUSDT_MTF_Open_Liquidity_rev01.xlsx"})

    def test_88_temporary_workbook_is_closed_before_move(self):
        source = (PROJECT_ROOT / "detectors" / "open_liquidity_v1" / "run_open_liquidity_detector.py").read_text(encoding="utf-8")
        temporary = source.index(".tmp.xlsx")
        replace = source.index("os.replace(temporary, final_path)")
        close = (PROJECT_ROOT / "detectors" / "open_liquidity_v1" / "workbook.py").read_text(encoding="utf-8").index("workbook.close()")
        self.assertLess(temporary, replace)
        self.assertGreater(close, 0)

    def test_89_final_workbook_reopens_after_atomic_move(self):
        with tempfile.TemporaryDirectory() as folder:
            path, _result = self._sample_workbook(Path(folder))
            final = Path(folder) / "moved.xlsx"
            path.replace(final)
            validate_ooxml_workbook(final)

    def test_90_analytical_counts_remain_identical(self):
        with tempfile.TemporaryDirectory() as folder:
            path, result = self._sample_workbook(Path(folder))
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                statuses = [row[6] for row in workbook["Candidate Audit"].iter_rows(min_row=2, values_only=True)]
            finally:
                workbook.close()
            expected = [item.status for rows in result.candidates.values() for item in rows]
            self.assertEqual(statuses.count(OPEN) + statuses.count("MITIGATED") + statuses.count(TERMINAL), len(expected))

    def test_91_detector_logic_module_is_unchanged_by_workbook_import(self):
        engine = PROJECT_ROOT / "detectors" / "open_liquidity_v1" / "engine.py"
        text = engine.read_text(encoding="utf-8")
        self.assertNotIn("openpyxl", text)
        self.assertIn("def detect_timeframe", text)
        self.assertIn("def attach_children", text)


def _package_source() -> str:
    root = PROJECT_ROOT / "detectors" / "open_liquidity_v1"
    parts = []
    for path in root.rglob("*.py"):
        if "tests" in path.parts or "__pycache__" in path.parts:
            continue
        parts.append(path.read_text(encoding="utf-8"))
    return "\n".join(parts)


if __name__ == "__main__":
    unittest.main()
