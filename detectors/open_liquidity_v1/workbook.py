"""Excel workbook for the open-liquidity detector."""
from __future__ import annotations

import math
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.open_liquidity_v1.engine import (
    MITIGATED,
    OPEN,
    TERMINAL,
    AnalysisResult,
    Candidate,
    iso_turkey,
    iso_utc,
    monthly_summary,
)
from detectors.open_liquidity_v1.engine import HOURS_PER_BAR as BAR_HOURS
from detectors.open_liquidity_v1.liquidity_config import DETECTOR_VERSION, TIER_ORDER, WORKBOOK_SHEETS

PRICE = "0.00000000"
PERCENT_POINTS = '0.00000000"%"'
EXCEL_MAX_TEXT = 32767
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
OOXML_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _solid(rgb: str) -> PatternFill:
    """Solid fill with both colors set.

    Excel rejects a differential style whose solid pattern has a foreground
    color and no background color. That defect makes the whole workbook fail
    to open.
    """
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


HEADER_FILL = _solid("1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
HIGH_FILL = _solid("F4CCCC")
LOW_FILL = _solid("D9EAD3")
TIER1_FILL = _solid("FCE5CD")
OLD_FILL = _solid("FFF2CC")
LOW_OBS_FILL = _solid("D0E2F3")
EQUAL_FILL = _solid("EAD1DC")
TERMINAL_FILL = _solid("D9D9D9")
NEAREST_FILL = _solid("C9DAF8")
WRAP = Alignment(wrap_text=True, vertical="center")
HEADER_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
CF_ROW_LIMIT = 5000
PLAIN_RANGE_SHEETS = frozenset({
    "Executive Summary",
    "Candidate Audit",
    "Mitigation Events",
    "Diagnostics",
    "README",
})
TABLE_NAMES = {
    "Hierarchical Open": "tblHierarchicalOpen",
    "1D Open Liquidity": "tbl1DOpenLiquidity",
    "4H Open Liquidity": "tbl4HOpenLiquidity",
    "1H Open Liquidity": "tbl1HOpenLiquidity",
    "Nearest Above Below": "tblNearestAboveBelow",
    "Cross-TF Mapping": "tblCrossTFMapping",
    "Monthly Origins": "tblMonthlyOrigins",
    "Parameters": "tblParameters",
}
_EMPTY_TYPED_CELL = re.compile(br'<c\b[^>]*\bt="(?:n|b|s|str|inlineStr|d|e)"[^>]*/>')


def write_workbook(path: Path, result: AnalysisResult, context: dict[str, Any], readme: str) -> None:
    workbook = Workbook()
    sheets = {
        "Executive Summary": _executive(result, context),
        "Hierarchical Open": _hierarchy(result, context),
        "1D Open Liquidity": _open_sheet(result, "1d", context),
        "4H Open Liquidity": _open_sheet(result, "4h", context),
        "1H Open Liquidity": _open_sheet(result, "1h", context),
        "Nearest Above Below": _nearest(result, context),
        "Candidate Audit": _audit(result, context),
        "Mitigation Events": _mitigations(result),
        "Cross-TF Mapping": _mapping(result),
        "Monthly Origins": _monthly(result),
        "Parameters": _parameters(),
        "Diagnostics": context["diagnostic_rows"],
        "README": _readme_rows(readme),
    }
    first = workbook.active
    first.title = WORKBOOK_SHEETS[0]
    for name in WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(name)
    for name in WORKBOOK_SHEETS:
        _write_sheet(workbook[name], sheets[name], TABLE_NAMES.get(name))
    _highlight(workbook, result)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def validate_saved_workbook(path: Path, expected_candidates: int) -> None:
    validate_ooxml_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        audit = workbook["Candidate Audit"]
        rows = audit.max_row - 1
        if rows != expected_candidates:
            raise RuntimeError(f"audit rows {rows} != {expected_candidates}")
    finally:
        workbook.close()


def validate_ooxml_workbook(path: Path) -> None:
    """Require a genuine non-macro Office Open XML workbook."""
    payload = path.read_bytes()
    if payload[:4] != b"PK\x03\x04":
        raise RuntimeError("workbook does not begin with the ZIP signature")
    if path.suffix.lower() != ".xlsx":
        raise RuntimeError("workbook extension is not .xlsx")
    if not zipfile.is_zipfile(path):
        raise RuntimeError("zipfile.is_zipfile returned false")
    required = (
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/worksheets/sheet1.xml",
    )
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("ZipFile.testzip found a corrupt member")
        names = set(archive.namelist())
        for member in required:
            if member not in names:
                raise RuntimeError(f"missing OOXML member {member}")
        if any("vba" in name.lower() for name in names):
            raise RuntimeError("workbook contains a VBA part")
        if any("externalLink" in name for name in names):
            raise RuntimeError("workbook contains an external link")
        for name in sorted(names):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                payload = archive.read(name)
                if _EMPTY_TYPED_CELL.search(payload):
                    raise RuntimeError(f"{name} contains a typed cell with no value")
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise RuntimeError(f"XML parse failed for {name}: {exc}") from exc
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                    _validate_sheet_xml(name, root)
        _validate_style_xml(ET.fromstring(archive.read("xl/styles.xml")))
        _validate_workbook_relationships(archive)
    _validate_openpyxl_workbook(path)


def _validate_sheet_xml(name: str, root: ET.Element) -> None:
    dimension = root.find("m:dimension", OOXML_NS)
    auto_filter = root.find("m:autoFilter", OOXML_NS)
    if auto_filter is not None and dimension is not None:
        filter_ref = auto_filter.attrib.get("ref", "")
        dimension_ref = dimension.attrib.get("ref", "")
        if filter_ref and filter_ref != dimension_ref:
            raise RuntimeError(f"{name} autoFilter {filter_ref} does not match dimension {dimension_ref}")


def _validate_style_xml(root: ET.Element) -> None:
    for fill in root.findall("m:dxfs/m:dxf/m:fill/m:patternFill", OOXML_NS):
        if fill.attrib.get("patternType") == "solid":
            if fill.find("m:fgColor", OOXML_NS) is None or fill.find("m:bgColor", OOXML_NS) is None:
                raise RuntimeError("solid conditional-format fill is missing fgColor or bgColor")


def _validate_workbook_relationships(archive: zipfile.ZipFile) -> None:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.attrib.get("Id"): rel.attrib.get("Target", "") for rel in relationships}
    for sheet in root.findall("m:sheets/m:sheet", OOXML_NS):
        title = sheet.attrib.get("name", "")
        if not title or len(title) > 31:
            raise RuntimeError(f"invalid sheet name {title!r}")
        relation = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = targets.get(relation, "").lstrip("/")
        member = target if target.startswith("xl/") else "xl/" + target
        if member not in archive.namelist():
            raise RuntimeError(f"sheet {title} references missing part {member}")


def _validate_openpyxl_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError(f"sheet order {workbook.sheetnames}")
        if workbook.vba_archive is not None:
            raise RuntimeError("workbook contains macros")
        if getattr(workbook, "_external_links", []):
            raise RuntimeError("workbook contains external links")
        names = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > EXCEL_MAX_ROWS or worksheet.max_column > EXCEL_MAX_COLUMNS:
                raise RuntimeError(f"{worksheet.title} exceeds an Excel worksheet limit")
            if worksheet.title in PLAIN_RANGE_SHEETS and worksheet.tables:
                raise RuntimeError(f"{worksheet.title} must stay a plain filtered range")
            if worksheet.title in {"Candidate Audit", "Mitigation Events"} and not worksheet.auto_filter.ref:
                raise RuntimeError(f"{worksheet.title} is missing its AutoFilter")
            for table in worksheet.tables.values():
                names.append(table.name)
                _validate_table(worksheet, table)
        if len(names) != len(set(names)):
            raise RuntimeError(f"duplicate table names {names}")
        for defined_name in workbook.defined_names.values():
            attr = getattr(defined_name, "attr_text", "") or ""
            if "#REF!" in attr or "#NAME?" in attr:
                raise RuntimeError(f"invalid defined name {defined_name.name}")
    finally:
        workbook.close()
    streamed = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if tuple(streamed.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError("read-only sheet order mismatch")
        for worksheet in streamed.worksheets:
            for _row in worksheet.iter_rows():
                pass
    finally:
        streamed.close()


def _validate_table(worksheet, table) -> None:
    name = table.name
    if not name or not (name[0].isalpha() or name[0] == "_") or not all(ch.isalnum() or ch == "_" for ch in name):
        raise RuntimeError(f"invalid table display name {name}")
    ref = table.ref
    if ":" not in ref:
        raise RuntimeError(f"table {name} reference is not rectangular")
    start, end = ref.split(":")
    min_col, min_row, max_col, max_row = _range_bounds(start, end)
    if min_row < 1 or max_row > worksheet.max_row or min_col < 1 or max_col > worksheet.max_column:
        raise RuntimeError(f"table {name} reference {ref} is outside {worksheet.title}")
    if max_row <= min_row:
        raise RuntimeError(f"table {name} has no data row")
    headers = [worksheet.cell(min_row, col).value for col in range(min_col, max_col + 1)]
    if any(value in (None, "") for value in headers) or len(headers) != len(set(headers)):
        raise RuntimeError(f"table {name} headers are blank or duplicated")


def _range_bounds(start: str, end: str) -> tuple[int, int, int, int]:
    from openpyxl.utils import coordinate_to_tuple

    start_row, start_col = coordinate_to_tuple(start)
    end_row, end_col = coordinate_to_tuple(end)
    return start_col, start_row, end_col, end_row


def round_trip_workbook(source: Path, copy_path: Path) -> None:
    workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    try:
        original_sheets = tuple(workbook.sheetnames)
        original_summary = [tuple(row) for row in workbook["Executive Summary"].iter_rows(values_only=True)]
        original_dimensions = {worksheet.title: worksheet.dimensions for worksheet in workbook.worksheets}
        original_headers = {worksheet.title: tuple(cell.value for cell in worksheet[1]) for worksheet in workbook.worksheets}
        workbook.save(copy_path)
    finally:
        workbook.close()
    reopened = load_workbook(copy_path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(reopened.sheetnames) != original_sheets:
            raise RuntimeError("round trip changed the sheet order")
        summary = [tuple(row) for row in reopened["Executive Summary"].iter_rows(values_only=True)]
        if summary != original_summary:
            raise RuntimeError("round trip changed the executive summary")
        for worksheet in reopened.worksheets:
            if worksheet.dimensions != original_dimensions[worksheet.title]:
                raise RuntimeError(f"round trip changed dimensions for {worksheet.title}")
            headers = tuple(cell.value for cell in worksheet[1])
            if headers != original_headers[worksheet.title]:
                raise RuntimeError(f"round trip changed headers for {worksheet.title}")
    finally:
        reopened.close()


def _table_token(name: str) -> str:
    return "".join(ch for ch in name if ch.isalnum())


def _write_sheet(ws, rows: Sequence[Sequence[Any]], table_name: str | None) -> None:
    prepared = _prepare_rows(rows)
    for row in prepared:
        ws.append(row)
    if not prepared:
        return
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = HEADER_BORDER
    width = len(prepared[0])
    for index, header in enumerate(prepared[0], start=1):
        letter = get_column_letter(index)
        label = header if isinstance(header, str) else ""
        ws.column_dimensions[letter].width = min(42, max(12, len(label) + 2))
    ws.row_dimensions[1].height = 30
    if ws.max_row > EXCEL_MAX_ROWS or ws.max_column > EXCEL_MAX_COLUMNS:
        raise RuntimeError(f"{ws.title} exceeds an Excel worksheet limit")
    if len(ws.title) > 31:
        raise RuntimeError(f"sheet name exceeds 31 characters: {ws.title}")
    if len(prepared) >= 2:
        ref = f"A1:{get_column_letter(width)}{len(prepared)}"
        ws.auto_filter.ref = ref
        has_real_rows = prepared[1][0] != "NO_RESULTS"
        if table_name and has_real_rows:
            table = Table(displayName=_table_display_name(table_name), ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    _apply_number_formats(ws)


def _readme_rows(readme: str) -> list[list[Any]]:
    lines = _clean_text(readme).splitlines() or ["(empty)"]
    return [["README Line"], *[[line] for line in lines]]


def _prepare_rows(rows: Sequence[Sequence[Any]]) -> list[list[Any]]:
    prepared = [[excel_value(value) for value in row] for row in rows]
    if not prepared:
        return prepared
    width = max(len(row) for row in prepared)
    header = list(prepared[0]) + [None] * (width - len(prepared[0]))
    header = _unique_headers(header)
    body = []
    for row in prepared[1:]:
        padded = list(row) + [None] * (width - len(row))
        body.append(padded[:width])
    if not body:
        body.append(["NO_RESULTS"] + [None] * (width - 1))
    return [header, *body]


def _unique_headers(header: Sequence[Any]) -> list[Any]:
    seen: dict[str, int] = {}
    unique = []
    for index, value in enumerate(header, start=1):
        text = value if isinstance(value, str) and value.strip() else f"Column{index}"
        count = seen.get(text, 0) + 1
        seen[text] = count
        unique.append(text if count == 1 else f"{text}_{count}")
    return unique


def _table_display_name(name: str) -> str:
    token = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in name)
    if not token or not (token[0].isalpha() or token[0] == "_"):
        token = "T_" + token
    return token[:40]


def excel_value(value: Any) -> Any:
    """Convert a Python value into a scalar Excel can store."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        return float(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, datetime):
        if value.tzinfo is not None and value.tzinfo.utcoffset(value) is not None:
            return value.astimezone(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        text = _clean_text(value)
        return text[:EXCEL_MAX_TEXT] or None
    if isinstance(value, (dict, list, tuple, set)):
        text = _clean_text(str(value))
        return text[:EXCEL_MAX_TEXT] or None
    text = _clean_text(str(value))
    return text[:EXCEL_MAX_TEXT] or None


def _clean_text(text: str) -> str:
    kept = []
    for char in text:
        code = ord(char)
        if char in "\t\n\r" or code >= 32:
            kept.append(char)
        else:
            kept.append(" ")
    return "".join(kept)


def _header_column(ws, header: str) -> str:
    for cell in ws[1]:
        if cell.value == header:
            return get_column_letter(cell.column)
    return ""


def _apply_number_formats(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    price_cols: list[int] = []
    percent_cols: list[int] = []
    for index, name in enumerate(headers, start=1):
        if not isinstance(name, str):
            continue
        lower = name.lower()
        if any(token in lower for token in ("utc", "turkey", "time", " id", "flag", "status", "rank", "bars", "count", "warning", "version", "reason", "section", "side", "month", "classification", "mechanism", "chain", "sequence", "trades", "tier", "censored", "inclusion")):
            continue
        if "percent" in lower or "percentage" in lower or lower.endswith("rate"):
            percent_cols.append(index)
        elif any(token in lower for token in ("price", "atr", "wick size", "volume", "penetration", "span", "distance", "origin open", "origin high", "origin low", "origin close", "mitigation open", "mitigation high", "mitigation low", "mitigation close")):
            price_cols.append(index)
    if not price_cols and not percent_cols:
        return
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for index in price_cols:
            cell = row[index - 1]
            if isinstance(cell.value, (int, float)):
                cell.number_format = PRICE
        for index in percent_cols:
            cell = row[index - 1]
            if isinstance(cell.value, (int, float)):
                cell.number_format = PERCENT_POINTS


def _highlight(workbook, result: AnalysisResult) -> None:
    _side_rules(workbook["Hierarchical Open"], "D")
    for name in ("1D Open Liquidity", "4H Open Liquidity", "1H Open Liquidity"):
        _side_rules(workbook[name], "D")
    _side_rules(workbook["Nearest Above Below"], "C")
    _formula(workbook["Hierarchical Open"], "E", '"TIER_1_ALL_TIMEFRAMES"', TIER1_FILL)
    _formula(workbook["Mitigation Events"], "R", '"TRUE"', EQUAL_FILL)
    for name in ("1D Open Liquidity", "4H Open Liquidity", "1H Open Liquidity", "Hierarchical Open"):
        ws = workbook[name]
        warning = _header_column(ws, "Observation Warning")
        if warning:
            _formula(ws, warning, '"VERY_LOW_FUTURE_OBSERVATION"', LOW_OBS_FILL)
        age = _header_column(ws, "Candidate Age Days")
        if age and ws.max_row >= 2:
            span = f"${age}2:${age}{ws.max_row}"
            ws.conditional_formatting.add(span, FormulaRule(formula=[f"{age}2>=365"], fill=OLD_FILL))
        rank = _header_column(ws, "Side Rank")
        if rank:
            _formula(ws, rank, "1", NEAREST_FILL)
    _formula(workbook["Candidate Audit"], "G", f'"{TERMINAL}"', TERMINAL_FILL)


def _side_rules(ws, column: str) -> None:
    if ws.max_row < 2:
        return
    span = f"${column}2:${column}{ws.max_row}"
    ws.conditional_formatting.add(span, FormulaRule(formula=[f'{column}2="HIGH"'], fill=HIGH_FILL))
    ws.conditional_formatting.add(span, FormulaRule(formula=[f'{column}2="LOW"'], fill=LOW_FILL))


def _formula(ws, column: str, literal: str, fill: PatternFill) -> None:
    if ws.max_row < 2 or not column:
        return
    if ws.max_row > CF_ROW_LIMIT:
        expected = literal[1:-1] if len(literal) >= 2 and literal.startswith('"') and literal.endswith('"') else literal
        if isinstance(expected, str) and expected.isdigit():
            expected = int(expected)
        index = column_index_from_string(column)
        for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index, max_row=ws.max_row):
            if cell.value == expected:
                cell.fill = fill
        return
    span = f"${column}2:${column}{ws.max_row}"
    ws.conditional_formatting.add(span, FormulaRule(formula=[f"{column}2={literal}"], fill=fill))


def _num(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)


def _bool(value: bool) -> str:
    return "TRUE" if value else "FALSE"


def _dt(value: datetime | None, turkey: bool = False) -> str:
    if value is None:
        return ""
    return iso_turkey(value) if turkey else iso_utc(value)


def _age_days(result: AnalysisResult, candidate: Candidate) -> Decimal:
    end = result.final_close_time[candidate.timeframe]
    return Decimal(str((end - candidate.open_time).total_seconds())) / Decimal("86400")


def _hours(candidate: Candidate) -> Decimal:
    return Decimal(candidate.future_bars) * BAR_HOURS[candidate.timeframe]


def _executive(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    rows = [["Section", "Metric", "Value"]]
    counts = context["counts"]
    rows.extend(
        [
            ["Market", "Symbol", "BTCUSDT"],
            ["Market", "TradingView", "BINANCE:BTCUSDT.P"],
            ["Market", "Market", "Binance USD-M PERPETUAL"],
            ["Market", "Dataset start", "2024-01-01T00:00:00.000Z"],
            ["Market", "Final as-of", context["as_of"]],
            ["Market", "Timeframes", "1h, 4h, 1d"],
            ["Market", "Detector version", DETECTOR_VERSION],
            ["Market", "Workbook revision", context["revision_name"]],
            ["Definition", "Equality closes liquidity", "TRUE"],
            ["Definition", "Wick touch closes liquidity", "TRUE"],
            ["Definition", "Close beyond required", "FALSE"],
            ["Definition", "Main results require a later candle", "TRUE"],
            ["Definition", "Open results are right-censored at dataset end", "TRUE"],
            ["Definition", "Trading signal", "FALSE. This workbook is market-structure description only."],
        ]
    )
    for timeframe in ("1h", "4h", "1d"):
        item = counts[timeframe]
        rows.append(["Source", f"{timeframe} rows", item["rows"]])
        rows.append(["Open High", timeframe, item["HIGH_OPEN"]])
        rows.append(["Open Low", timeframe, item["LOW_OPEN"]])
        rows.append(["Mitigated", timeframe, item[MITIGATED]])
        rows.append(["Terminal", timeframe, item[TERMINAL]])
    rows.append(["Hierarchy", "Groups", len(result.groups)])
    for tier in TIER_ORDER:
        rows.append(["Hierarchy Tier", tier, sum(1 for group in result.groups if group.tier == tier)])
    for label, group in (("Nearest open High", _nth(result, "HIGH", 1)), ("Nearest open Low", _nth(result, "LOW", 1))):
        rows.append(["Nearest", label, _group_text(group)])
    for timeframe in ("1h", "4h", "1d"):
        for side in ("HIGH", "LOW"):
            oldest = _edge(result, timeframe, side, oldest=True)
            newest = _edge(result, timeframe, side, oldest=False)
            rows.append(["Oldest", f"{timeframe} {side}", _candidate_text(oldest)])
            rows.append(["Newest", f"{timeframe} {side}", _candidate_text(newest)])
    return rows


def _nth(result: AnalysisResult, side: str, rank: int):
    sided = [group for group in result.groups if group.side == side and group.side_rank == rank]
    return sided[0] if sided else None


def _edge(result: AnalysisResult, timeframe: str, side: str, oldest: bool) -> Candidate | None:
    pool = [item for item in result.candidates[timeframe] if item.status == OPEN and item.side == side]
    if not pool:
        return None
    return min(pool, key=lambda item: item.open_time) if oldest else max(pool, key=lambda item: item.open_time)


def _group_text(group) -> str:
    if group is None:
        return "NONE"
    return f"{group.group_id} {group.representative} {group.tier} {group.distance_percent}"


def _candidate_text(candidate: Candidate | None) -> str:
    if candidate is None:
        return "NONE"
    return f"{candidate.candidate_id} {candidate.price} {iso_utc(candidate.open_time)}"


def _hierarchy(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    header = [
        "Overall Rank", "Side Rank", "Hierarchy Group ID", "Side", "Hierarchy Tier", "Representative Price",
        "Distance from Final 1h Close", "Distance Percent", "Timeframes Present", "Timeframe Count",
        "1D Candidate ID", "4H Candidate ID", "1H Candidate ID",
        "1D Origin UTC", "1D Origin Turkey", "4H Origin UTC", "4H Origin Turkey", "1H Origin UTC", "1H Origin Turkey",
        "1D Future Observation Bars", "4H Future Observation Bars", "1H Future Observation Bars",
        "Oldest Origin", "Newest Origin", "Maximum Group Span Percent", "Right Censored", "Censoring Reason",
        "As-Of Timestamp", "Observation Warning", "Detector Version",
    ]
    rows = [header]
    for group in result.groups:
        members = group.members
        rows.append([
            group.rank, group.side_rank, group.group_id, group.side, group.tier, _num(group.representative),
            _num(group.distance_price), _num(group.distance_percent), "+".join(tf for tf in ("1d", "4h", "1h") if tf in members),
            len(members),
            _id(members, "1d"), _id(members, "4h"), _id(members, "1h"),
            _origin(members, "1d", False), _origin(members, "1d", True),
            _origin(members, "4h", False), _origin(members, "4h", True),
            _origin(members, "1h", False), _origin(members, "1h", True),
            _future(members, "1d"), _future(members, "4h"), _future(members, "1h"),
            iso_utc(group.oldest), iso_utc(group.newest), _num(group.max_span),
            "TRUE", "DATASET_END", context["as_of"], group.warning, DETECTOR_VERSION,
        ])
    return rows


def _id(members, timeframe: str) -> str:
    return members[timeframe].candidate_id if timeframe in members else ""


def _origin(members, timeframe: str, turkey: bool) -> str:
    if timeframe not in members:
        return ""
    return _dt(members[timeframe].open_time, turkey)


def _future(members, timeframe: str):
    return members[timeframe].future_bars if timeframe in members else ""


def _open_sheet(result: AnalysisResult, timeframe: str, context: dict[str, Any]) -> list[list[Any]]:
    header = [
        "Timeframe Rank", "Side Rank", "Candidate ID", "Side", "Liquidity Price",
        "Origin Candle Open UTC", "Origin Candle Close UTC", "Origin Candle Open Turkey", "Origin Candle Close Turkey",
        "Candidate Available UTC", "Candidate Available Turkey",
        "Origin Open", "Origin High", "Origin Low", "Origin Close", "Origin Volume", "Origin Trades", "Origin ATR",
        "Wick Size", "Wick Percentage", "Wick ATR", "Wick Classification",
        "Equal-Level Chain ID", "Equal-Level Sequence",
        "Future Observation Bars", "Future Observation Hours", "Future Observation Days",
        "Candidate Age Bars", "Candidate Age Days", "Final Reference Close",
        "Distance Price", "Distance Percent", "Distance ATR",
        "Touch-Based Status", "Close-Based Status", "Right Censored", "Censoring Reason",
        "Observation Warning", "Hierarchy Group ID", "Hierarchy Tier", "Detector Version",
    ]
    rows = [header]
    chosen = [item for item in result.candidates[timeframe] if item.status == OPEN]
    if timeframe not in result.final_close:
        return rows
    reference = result.final_close[timeframe]
    chosen.sort(key=lambda item: (0 if item.side == "HIGH" else 1, item.side_rank))
    for candidate in chosen:
        wick = candidate.upper_wick if candidate.side == "HIGH" else candidate.lower_wick
        wick_percent = wick / candidate.price * Decimal("100") if candidate.price else Decimal("0")
        wick_atr = (wick / candidate.atr) if candidate.atr not in (None, Decimal("0")) else None
        hours = _hours(candidate)
        rows.append([
            candidate.timeframe_rank, candidate.side_rank, candidate.candidate_id, candidate.side, _num(candidate.price),
            iso_utc(candidate.open_time), iso_utc(candidate.close_time), iso_turkey(candidate.open_time), iso_turkey(candidate.close_time),
            iso_utc(candidate.close_time), iso_turkey(candidate.close_time),
            _num(candidate.open), _num(candidate.high), _num(candidate.low), _num(candidate.close),
            _num(candidate.volume), candidate.trades, _num(candidate.atr),
            _num(wick), _num(wick_percent), _num(wick_atr), candidate.origin_class,
            candidate.chain_id, candidate.chain_sequence,
            candidate.future_bars, _num(hours), _num(hours / Decimal("24")),
            candidate.future_bars, _num(_age_days(result, candidate)), _num(reference),
            _num(candidate.distance_price), _num(candidate.distance_percent), _num(candidate.distance_atr),
            candidate.touch_status, candidate.close_status, "TRUE", "DATASET_END",
            candidate.warning, candidate.group_id, candidate.tier, DETECTOR_VERSION,
        ])
    return rows


def _nearest(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    header = [
        "Section", "Rank", "Side", "Price", "Distance", "Distance Percent", "Timeframes Present",
        "Origin Time", "Age Days", "Future Observation Bars", "Wick Classification", "Hierarchy Tier",
        "Right Censored",
    ]
    rows = [header]
    sections = [
        ("Nearest hierarchical open Highs", [group for group in result.groups if group.side == "HIGH"]),
        ("Nearest hierarchical open Lows", [group for group in result.groups if group.side == "LOW"]),
    ]
    for title, groups in sections:
        for group in groups[:25]:
            member = _representative_member(group.members)
            rows.append([
                title, group.side_rank, group.side, _num(group.representative), _num(group.distance_price),
                _num(group.distance_percent), "+".join(tf for tf in ("1d", "4h", "1h") if tf in group.members),
                iso_utc(member.open_time), _num(_age_days(result, member)), member.future_bars,
                member.origin_class, group.tier, "TRUE",
            ])
    for timeframe, label in (("1d", "1d"), ("4h", "4h"), ("1h", "1h")):
        for side, side_label in (("HIGH", "Highs"), ("LOW", "Lows")):
            title = f"Nearest {label} open {side_label}"
            pool = [item for item in result.candidates[timeframe] if item.status == OPEN and item.side == side]
            pool.sort(key=lambda item: item.side_rank)
            for candidate in pool[:25]:
                rows.append([
                    title, candidate.side_rank, candidate.side, _num(candidate.price), _num(candidate.distance_price),
                    _num(candidate.distance_percent), timeframe, iso_utc(candidate.open_time),
                    _num(_age_days(result, candidate)), candidate.future_bars, candidate.origin_class,
                    candidate.tier, "TRUE",
                ])
    return rows


def _representative_member(members: dict[str, Candidate]) -> Candidate:
    for timeframe in ("1d", "4h", "1h"):
        if timeframe in members:
            return members[timeframe]
    raise RuntimeError("empty hierarchy group")


def _audit(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    header = [
        "Candidate ID", "Timeframe", "Side", "Price", "Origin Time", "Availability Time",
        "Primary Status", "Touch-Based Status", "Close-Based Status", "First Mitigation ID",
        "First Mitigation Time", "Bars To Mitigation", "Future Observation Bars", "Terminal Flag",
        "Equal-Level Chain ID", "Previous Equal Candidate ID", "Next Equal Candidate ID",
        "Equal-Level Sequence", "Chain Length", "Final Surviving Candidate",
        "Main-Sheet Inclusion", "Exclusion Reason",
    ]
    rows = [header]
    for timeframe in ("1d", "4h", "1h"):
        bars = result.bars[timeframe]
        for candidate in result.candidates[timeframe]:
            included = candidate.status == OPEN
            if included:
                reason = ""
            elif candidate.status == TERMINAL:
                reason = "TERMINAL_NO_LATER_CANDLE"
            else:
                reason = "MITIGATED_BY_LATER_COMPLETED_CANDLE"
            mitigation_id = ""
            mitigation_time = ""
            bars_to = ""
            if candidate.mitigation_row is not None:
                later = bars[candidate.mitigation_row]
                mitigation_id = f"{timeframe}:{candidate.mitigation_row + 1}"
                mitigation_time = iso_utc(later.open_time)
                bars_to = candidate.mitigation_row - candidate.row
            rows.append([
                candidate.candidate_id, timeframe, candidate.side, _num(candidate.price),
                iso_utc(candidate.open_time), iso_utc(candidate.close_time),
                candidate.status, candidate.touch_status, candidate.close_status,
                mitigation_id, mitigation_time, bars_to, candidate.future_bars,
                _bool(candidate.status == TERMINAL),
                candidate.chain_id, candidate.previous_equal_id, candidate.next_equal_id,
                candidate.chain_sequence, candidate.chain_length, candidate.chain_survivor_id,
                _bool(included), reason,
            ])
    return rows


def _mitigations(result: AnalysisResult) -> list[list[Any]]:
    header = [
        "Candidate ID", "Timeframe", "Side", "Liquidity Price", "Origin UTC", "Origin Turkey",
        "Candidate Availability UTC", "Candidate Availability Turkey",
        "First Mitigation Candle Open UTC", "First Mitigation Candle Close UTC",
        "First Mitigation Candle Open Turkey", "First Mitigation Candle Close Turkey",
        "Mitigation Open", "Mitigation High", "Mitigation Low", "Mitigation Close",
        "First Mitigation Mechanism", "Exact Equality Flag", "Gap Flag", "Wick Reach Flag",
        "Body Reach Flag", "Close Beyond Flag", "Penetration Price", "Penetration Percentage",
        "Bars To Mitigation", "Hours To Mitigation", "Days To Mitigation",
    ]
    rows = [header]
    for timeframe in ("1d", "4h", "1h"):
        bars = result.bars[timeframe]
        for candidate in result.candidates[timeframe]:
            if candidate.status != "MITIGATED" or candidate.mitigation_row is None:
                continue
            later = bars[candidate.mitigation_row]
            bars_to = candidate.mitigation_row - candidate.row
            hours = Decimal(bars_to) * BAR_HOURS[timeframe]
            rows.append([
                candidate.candidate_id, timeframe, candidate.side, _num(candidate.price),
                iso_utc(candidate.open_time), iso_turkey(candidate.open_time),
                iso_utc(candidate.close_time), iso_turkey(candidate.close_time),
                iso_utc(later.open_time), iso_utc(later.close_time),
                iso_turkey(later.open_time), iso_turkey(later.close_time),
                _num(later.open), _num(later.high), _num(later.low), _num(later.close),
                candidate.mechanism, _bool(candidate.exact_flag), _bool(candidate.gap_flag),
                _bool(candidate.wick_flag), _bool(candidate.body_flag), _bool(candidate.close_flag),
                _num(candidate.penetration), _num(candidate.penetration_percent),
                bars_to, _num(hours), _num(hours / Decimal("24")),
            ])
    return rows


def _mapping(result: AnalysisResult) -> list[list[Any]]:
    header = [
        "Mapping Attempt ID", "Side", "Parent Timeframe", "Parent Candidate ID", "Parent Price",
        "Child Timeframe", "Child Candidate ID", "Child Price", "Relative Distance Percent",
        "Proposed Group Span", "Tolerance Pass", "Group Span Pass", "Accepted", "Selected Parent",
        "Rejection Reason", "Tie-Break Outcome",
    ]
    rows = [header]
    for attempt in result.attempts:
        rows.append([
            attempt.attempt_id, attempt.side, attempt.parent_timeframe, attempt.parent_id, _num(attempt.parent_price),
            attempt.child_timeframe, attempt.child_id, _num(attempt.child_price), _num(attempt.distance_percent),
            _num(attempt.proposed_span), _bool(attempt.tolerance_pass), _bool(attempt.span_pass),
            _bool(attempt.accepted), attempt.selected_parent, attempt.rejection_reason, attempt.tie_break,
        ])
    return rows


def _monthly(result: AnalysisResult) -> list[list[Any]]:
    header = [
        "Month", "Timeframe", "Source Candle Count", "High Candidates Created", "Low Candidates Created",
        "High Candidates Still Open", "Low Candidates Still Open", "High Candidates Mitigated",
        "Low Candidates Mitigated", "Terminal Candidates", "Median Bars To Mitigation", "Maximum Bars To Mitigation",
        "Open Rate", "Mitigation Rate", "Oldest Surviving High", "Oldest Surviving Low",
    ]
    rows = [header]
    for item in monthly_summary(result):
        rows.append([
            item["month"], item["timeframe"], item["candles"], item["high_created"], item["low_created"],
            item["high_open"], item["low_open"], item["high_mitigated"], item["low_mitigated"], item["terminal"],
            _num(item["median_bars"]) if item["median_bars"] is not None else "",
            item["max_bars"] if item["max_bars"] is not None else "",
            _num(item["open_rate"]), _num(item["mitigation_rate"]),
            iso_utc(item["oldest_high"]) if item["oldest_high"] else "",
            iso_utc(item["oldest_low"]) if item["oldest_low"] else "",
        ])
    return rows


def _parameters() -> list[list[Any]]:
    rows = [["Parameter", "Value"]]
    entries = [
        ("symbol", "BTCUSDT"),
        ("tradingview_symbol", "BINANCE:BTCUSDT.P"),
        ("market", "Binance USD-M PERPETUAL"),
        ("timeframes", "1h, 4h, 1d"),
        ("completed_candles_only", "TRUE"),
        ("high_candidate_source", "EVERY_COMPLETED_CANDLE_HIGH"),
        ("low_candidate_source", "EVERY_COMPLETED_CANDLE_LOW"),
        ("high_primary_mitigation", "LATER_HIGH_GREATER_THAN_OR_EQUAL_TO_LEVEL"),
        ("low_primary_mitigation", "LATER_LOW_LESS_THAN_OR_EQUAL_TO_LEVEL"),
        ("equality_counts_as_mitigation", "TRUE"),
        ("wick_touch_counts_as_mitigation", "TRUE"),
        ("close_beyond_required", "FALSE"),
        ("origin_candle_may_self_mitigate", "FALSE"),
        ("primary_touch_tolerance", "0"),
        ("near_miss_tolerance_percent", "0.05"),
        ("wick_required_for_candidate", "FALSE"),
        ("terminal_without_later_candle_in_main_results", "FALSE"),
        ("minimum_later_observation_bars_for_main_result", "1"),
        ("cross_timeframe_match_tolerance_percent", "0.10"),
        ("maximum_group_price_span_percent", "0.20"),
        ("timeframe_priority", "1d, 4h, 1h"),
        ("main_rank_priority", "hierarchy_tier_then_distance"),
        ("computational_timezone", "UTC"),
        ("display_timezone", "Europe/Istanbul"),
        ("timestamp_excel_storage", "ISO_8601_TEXT"),
        ("excel_format", "XLSX_OFFICE_OPEN_XML"),
        ("excel_writer", "OPENPYXL_3_1_5"),
        ("excel_macro_enabled", "FALSE"),
        ("excel_external_links", "FALSE"),
        ("large_audit_sheet_mode", "PLAIN_RANGE_WITH_AUTOFILTER"),
        ("candidate_audit_formal_excel_table", "FALSE"),
        ("mitigation_events_formal_excel_table", "FALSE"),
        ("implementation_mode", "PERMANENT_SCRIPT_INSIDE_EXISTING_PROJECT"),
        ("project_root", r"C:\Users\oranb\Desktop\backtest_system"),
        ("detector_package", r"C:\Users\oranb\Desktop\backtest_system\detectors\open_liquidity_v1"),
        ("detector_version", DETECTOR_VERSION),
        ("open_high_formula", "suffix_future_high < origin_high"),
        ("open_low_formula", "suffix_future_low > origin_low"),
        ("relative_distance_formula", "abs(a-b)/((a+b)/2)*100"),
        ("high_distance_formula", "liquidity_price - final_close"),
        ("low_distance_formula", "final_close - liquidity_price"),
        ("wilder_atr_formula", "(previous_ATR * 13 + current_TR) / 14 after SMA of the first 14 true ranges"),
        ("tie_break_rank", "distance, future observation bars, candidate age, wick ATR, more recent origin, candidate ID"),
        ("tie_break_hierarchy", "smaller distance, higher parent timeframe, greater future observation, older origin, candidate ID"),
        ("tie_break_group_rank", "tier, distance from final 1h close, future-observation hours, oldest origin, representative price"),
        ("representative_price", "1d price if present, else 4h price, else 1h price"),
        ("observation_warning", "future_bars/candle_count: <1% VERY_LOW, <5% LOW, <20% MODERATE, else HIGH"),
        ("mechanism_priority", "GAP_OPEN_AT_OR_BEYOND, EXACT_EXTREME_TOUCH, CLOSE_AT_OR_BEYOND, BODY_REACH, WICK_REACH"),
        ("not_a_trading_signal", "TRUE"),
    ]
    rows.extend(entries)
    return rows
