"""Simplified workbook for the 4h Special Swing detector."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from detectors.swing_open_close_4h_special_v1.config import (
    MIN_BOUNDARY,
    RESULT_COLUMNS,
    SHEETS,
    STANDARD_CLASS,
    SWING_HIGH,
    USER_PARAMETERS,
    ZERO_CLASS,
)
from detectors.swing_open_close_4h_special_v1.engine import Analysis, turkey_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
HIGH_FILL = PatternFill("solid", fgColor="F4CCCC")
HIGH_ALT = PatternFill("solid", fgColor="F8E0E0")
LOW_FILL = PatternFill("solid", fgColor="D9EAD3")
LOW_ALT = PatternFill("solid", fgColor="E7F3E4")
ZERO_CLASS_FILL = PatternFill("solid", fgColor="F6B26B")
STANDARD_CLASS_FILL = PatternFill("solid", fgColor="9FC5E8")
WRAP = Alignment(wrap_text=True, vertical="center")
EIGHT = Decimal("0.00000001")


def shown(value: Decimal) -> Decimal:
    return value.quantize(EIGHT, rounding=ROUND_HALF_UP)


def _row(item, bars) -> list:
    opened = bars[item.open_row]
    closed = bars[item.close_row]
    extreme = bars[item.extreme_row]
    return [
        item.swing_id,
        item.direction,
        item.formation_class,
        item.interior,
        item.total,
        turkey_text(opened.open_time),
        shown(opened.open),
        shown(opened.high),
        shown(opened.low),
        turkey_text(closed.open_time),
        turkey_text(closed.close_time),
        shown(item.close_price),
        shown(closed.high),
        shown(closed.low),
        turkey_text(extreme.open_time),
        "HIGH" if item.direction == SWING_HIGH else "LOW",
        shown(item.extreme_price),
        shown(item.structure_high),
        shown(item.structure_low),
        shown(item.open_width_percent),
        shown(item.close_width_percent),
    ]


def _write(book: Workbook, title: str, rows: list[list]) -> None:
    sheet = book.create_sheet(title)
    sheet.append(list(RESULT_COLUMNS))
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, row in enumerate(rows, start=2):
        direction = row[1]
        alternate = index % 2 == 0
        fill = (HIGH_ALT if alternate else HIGH_FILL) if direction == SWING_HIGH else (LOW_ALT if alternate else LOW_FILL)
        for cell in sheet[index]:
            cell.fill = fill
        formation = row[2]
        sheet.cell(index, 3).fill = ZERO_CLASS_FILL if formation == ZERO_CLASS else STANDARD_CLASS_FILL
        for column in (7, 8, 9, 12, 13, 14, 17, 18, 19):
            sheet.cell(index, column).number_format = "0.00000000"
        for column in (20, 21):
            sheet.cell(index, column).number_format = "0.00000000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:U{max(1, sheet.max_row)}"
    sheet.row_dimensions[1].height = 32
    for column in range(1, len(RESULT_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 28


def write_workbook(path: Path, result: Analysis) -> None:
    book = Workbook()
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=False)
    book.remove(book.active)
    bars = result.bars
    ordered = sorted(result.displayed, key=lambda item: (item.open_row, item.close_row, item.swing_id))
    rows = [_row(item, bars) for item in ordered]
    _write(book, "Special Swings", rows)
    _write(book, "Swing Highs", [row for row in rows if row[1] == SWING_HIGH])
    _write(book, "Swing Lows", [row for row in rows if row[1] != SWING_HIGH])
    parameters = book.create_sheet("Parameters")
    parameters.append(["Parameter", "Value"])
    for name, value in USER_PARAMETERS:
        parameters.append([name, value])
    for cell in parameters[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    parameters.freeze_panes = "A2"
    parameters.auto_filter.ref = f"A1:B{parameters.max_row}"
    parameters.column_dimensions["A"].width = 48
    parameters.column_dimensions["B"].width = 78
    if book.sheetnames != list(SHEETS):
        raise RuntimeError(f"sheet order {book.sheetnames}")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def validate_workbook(path: Path, result: Analysis) -> dict:
    normal = load_workbook(path, read_only=False, data_only=False)
    readonly = load_workbook(path, read_only=True, data_only=False)
    data_only = load_workbook(path, read_only=True, data_only=True)
    try:
        if normal.sheetnames != list(SHEETS) or readonly.sheetnames != list(SHEETS):
            raise RuntimeError(f"sheets {normal.sheetnames}")
        if normal.vba_archive is not None or list(normal._external_links):
            raise RuntimeError("workbook contains macros or external links")
        headers = [cell.value for cell in next(normal["Special Swings"].iter_rows(max_row=1))]
        if tuple(headers) != RESULT_COLUMNS:
            raise RuntimeError(f"unexpected columns {headers}")
        for name in ("Swing Highs", "Swing Lows"):
            sheet_headers = [cell.value for cell in next(normal[name].iter_rows(max_row=1))]
            if tuple(sheet_headers) != RESULT_COLUMNS:
                raise RuntimeError(f"{name} columns differ")
        if any(header and "UTC" in str(header).upper() for header in headers):
            raise RuntimeError("UTC column present")
        values = {row[0]: row[1] for row in normal["Parameters"].iter_rows(min_row=2, values_only=True)}
        if values != dict(USER_PARAMETERS):
            raise RuntimeError("parameters sheet does not match the user-facing list")
        if values["Maximum Boundary Percent"] != "None — No Upper Limit":
            raise RuntimeError("maximum boundary is still limited")
        if values["Zero-Interior Confirmation"] != "Full Return to Swing Open Open":
            raise RuntimeError("zero-interior confirmation changed")
        if values["Normal Standard Confirmation"] != "Minimum One-Third Swing Open Body Penetration":
            raise RuntimeError("normal Standard confirmation changed")
        collected: dict[str, list] = {}
        minimum = Decimal(MIN_BOUNDARY)
        for sheet_name in ("Special Swings", "Swing Highs", "Swing Lows"):
            collected[sheet_name] = []
            for row in normal[sheet_name].iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                collected[sheet_name].append(row)
                formation = row[2]
                interior = int(row[3])
                total = int(row[4])
                if formation == ZERO_CLASS:
                    if interior != 0 or total != 2:
                        raise RuntimeError("zero-interior row has the wrong candle count")
                elif formation == STANDARD_CLASS:
                    if not 1 <= interior <= 5 or not 3 <= total <= 7:
                        raise RuntimeError(f"normal Standard interior {interior} is outside 1-5")
                else:
                    raise RuntimeError(f"unexpected formation class {formation}")
                open_boundary = Decimal(str(row[19]))
                close_boundary = Decimal(str(row[20]))
                if open_boundary < minimum or close_boundary < minimum:
                    raise RuntimeError("workbook boundary is below 1.30")
                if "UTC" in "".join(str(value) for value in row):
                    raise RuntimeError("UTC text appeared in a result row")
                if row[1] not in {"SWING_HIGH", "SWING_LOW"}:
                    raise RuntimeError(f"unexpected swing type {row[1]}")
            sheet_ids = [row[0] for row in collected[sheet_name]]
            if len(sheet_ids) != len(set(sheet_ids)):
                raise RuntimeError(f"duplicate Swing ID on {sheet_name}")
        high_ids = {row[0] for row in collected["Swing Highs"]}
        low_ids = {row[0] for row in collected["Swing Lows"]}
        all_ids = {row[0] for row in collected["Special Swings"]}
        if high_ids & low_ids or high_ids | low_ids != all_ids:
            raise RuntimeError("direction sheets do not partition Special Swings")
        if any(row[1] != "SWING_HIGH" for row in collected["Swing Highs"]):
            raise RuntimeError("Swing Highs contains a non-high")
        if any(row[1] != "SWING_LOW" for row in collected["Swing Lows"]):
            raise RuntimeError("Swing Lows contains a non-low")
        if data_only["Parameters"]["A1"].value not in (None, "Parameter"):
            raise RuntimeError("data_only workbook did not expose the header")
        if normal["Special Swings"].max_row - 1 != len(result.displayed):
            raise RuntimeError("workbook row count does not match primary results")
    finally:
        normal.close()
        readonly.close()
        data_only.close()
    return {"workbook_validation": "PASS", "sheets": 4}
