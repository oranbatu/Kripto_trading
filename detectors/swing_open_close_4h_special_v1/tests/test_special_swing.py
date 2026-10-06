"""Tests for the 4h Special Swing detector."""
from __future__ import annotations

import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from detectors.swing_open_close_4h_special_v1.config import (
    ALTERNATIVE_ENABLED,
    COMPACT_ENABLED,
    DATASET,
    DERIVED,
    LONGER,
    MAX_INTERIOR,
    NORMAL_MIN_INTERIOR,
    STANDARD_CLASS,
    ZERO_CLASS,
    ZERO_EXCEPTION_ENABLED,
    NESTED,
    PRIMARY,
    PROTECTED_ROOT,
    RESULT_COLUMNS,
    SHEETS,
    STANDARD_ENABLED,
    SWING_HIGH,
    SWING_LOW,
    TMP_DIR,
    USER_PARAMETERS,
)
from detectors.swing_open_close_4h_special_v1.engine import (
    Bar,
    analyze,
    body_threshold,
    boundary_reason,
    duration_reason,
    zero_adjacency_reason,
    signature,
    swing_open_body,
    turkey_text,
    wick_only_reason,
    zero_direction_reason,
    zero_return_reason,
)
from detectors.swing_open_close_4h_special_v1.run_swing_open_close_4h_special_detector import (
    allocate_revision,
    fingerprint_tree,
)
from detectors.swing_open_close_4h_special_v1.workbook import validate_workbook, write_workbook

D = Decimal
START = datetime(2026, 1, 1, tzinfo=timezone.utc)
PROTECTED_BEFORE = fingerprint_tree(PROTECTED_ROOT)
SOURCE_BEFORE = fingerprint_tree(Path(DATASET))


def bars_from(specs: list[tuple[str, str, str, str]]) -> list[Bar]:
    rows = []
    for index, (opened, high, low, close) in enumerate(specs):
        moment = START + timedelta(hours=4 * index)
        rows.append(Bar(
            row=index,
            open_time=moment,
            close_time=moment + timedelta(hours=4, milliseconds=-1),
            open=D(opened),
            high=D(high),
            low=D(low),
            close=D(close),
            volume=D("10"),
            quote_volume=D("100"),
            trades=2,
        ))
    return rows


def flat_above(price: str = "101") -> tuple[str, str, str, str]:
    return (price, price, price, price)


def swing_at(open_row: int, result):
    return [item for item in result.raw if item.open_row == open_row and item.direction == SWING_HIGH]


def search_at(open_row: int, result, direction=SWING_HIGH):
    return next(item for item in result.searches if item.open_row == open_row and item.direction == direction)


class SpecialSwingTests(unittest.TestCase):
    def test_only_standard_method_is_enabled(self):
        self.assertTrue(STANDARD_ENABLED)
        self.assertTrue(ZERO_EXCEPTION_ENABLED)
        self.assertFalse(COMPACT_ENABLED)
        self.assertFalse(ALTERNATIVE_ENABLED)
        result = analyze(self._valid(3))
        self.assertEqual(result.counters["alternative"], 0)
        self.assertEqual(result.counters["compact"], 0)
        self.assertEqual(result.counters["non_standard"], 0)
        self.assertTrue(result.raw)
        self.assertTrue(all(item.direction in {SWING_HIGH, SWING_LOW} for item in result.raw))

    def test_duration_boundaries_and_search_horizon(self):
        self.assertEqual(NORMAL_MIN_INTERIOR, 1)
        self.assertEqual(MAX_INTERIOR, 5)
        self.assertIsNone(duration_reason(1))
        self.assertIsNone(duration_reason(5))
        self.assertIsNone(duration_reason(0))
        self.assertEqual(duration_reason(6), "STANDARD_DURATION_ABOVE_5")
        one = analyze(self._valid(1))
        five = analyze(self._valid(5))
        self.assertEqual(swing_at(0, one)[0].interior, 1)
        self.assertEqual(swing_at(0, one)[0].formation_class, STANDARD_CLASS)
        self.assertFalse(swing_at(0, one)[0].zero_exception)
        self.assertEqual(swing_at(0, one)[0].extreme_source, "STRICTLY_INTERIOR_CANDLE_ONLY")
        self.assertTrue(swing_at(0, one)[0].open_row < swing_at(0, one)[0].extreme_row < swing_at(0, one)[0].close_row)
        self.assertEqual(swing_at(0, five)[0].interior, 5)
        self.assertEqual(swing_at(0, five)[0].total, 7)
        late = analyze(self._return_at(6))
        self.assertFalse(swing_at(0, late))
        self.assertLessEqual(search_at(0, late).max_interior, 5)
        self.assertEqual(search_at(0, late).terminal, "SEARCH_HORIZON_EXHAUSTED_AT_5_INTERIOR_CANDLES")
        self.assertEqual(late.counters["maximum_boundary_rejection_count"], 0)
        self.assertTrue(all(item.total <= 7 for item in late.raw))

    def test_first_binding_return_is_final_and_dataset_censor_works(self):
        first = analyze(self._valid(1) + bars_from([flat_above("90")]))
        self.assertEqual(swing_at(0, first)[0].close_row, 2)
        failed = analyze(bars_from([
            ("103", "104", "99", "100"),
            ("101.5", "102.5", "100", "101"),
            ("100", "106", "99", "100"),
        ]))
        self.assertEqual(search_at(0, failed).terminal, "ZERO_INTERIOR_SWING_HIGH_OPEN_NOT_BULLISH")
        self.assertFalse(swing_at(0, failed))
        self.assertFalse(any(item.open_row == 0 and item.formation_class == ZERO_CLASS for item in failed.raw))
        censored = analyze(bars_from([
            ("100", "100.4", "99", "100.2"),
            flat_above(),
            flat_above(),
            flat_above(),
        ]))
        self.assertEqual(search_at(0, censored).terminal, "SEARCH_CENSORED_BY_DATASET_END")
        self.assertEqual(signature(first), signature(analyze(self._valid(1) + bars_from([flat_above("90")]))))

    def test_boundary_interval_is_inclusive_and_exact(self):
        self.assertIsNone(boundary_reason(D("1.30"), D("1.30")))
        self.assertIsNone(boundary_reason(D("3.50"), D("3.50")))
        self.assertIsNone(boundary_reason(D("1.80"), D("2.20")))
        low = analyze(self._band("101.29999999"))
        high = analyze(self._band("103.50000001"))
        self.assertEqual(search_at(0, low).terminal, "BOTH_BOUNDARIES_BELOW_MINIMUM")
        self.assertGreater(swing_at(0, high)[0].open_width_percent, D("3.50"))
        self.assertFalse(swing_at(0, low))
        self.assertEqual(swing_at(0, analyze(self._band("105")))[0].open_width_percent, D("5"))
        self.assertEqual(swing_at(0, analyze(self._band("101.3")))[0].open_width_percent, D("1.3"))
        self.assertEqual(swing_at(0, analyze(self._band("103.5")))[0].open_width_percent, D("3.5"))
        open_low = analyze(self._high_widths(D("101.2"), D("101.2") / D("1.02")))
        self.assertEqual(search_at(0, open_low).terminal, "OPEN_BOUNDARY_BELOW_1_30")
        close_high = analyze(self._high_widths(D("102"), D("102") / D("1.036")))
        self.assertGreater(swing_at(0, close_high)[0].close_width_percent, D("3.50"))
        self.assertEqual(boundary_reason(D("2.00"), D("1.20")), "CLOSE_BOUNDARY_BELOW_1_30")
        self.assertIsNone(boundary_reason(D("3.60"), D("2.00")))
        self.assertEqual(boundary_reason(D("1.20"), D("1.20")), "BOTH_BOUNDARIES_BELOW_MINIMUM")
        self.assertIsNone(boundary_reason(D("3.60"), D("3.60")))
        self.assertIsNone(boundary_reason(D("5"), D("10")))
        self.assertIsNone(boundary_reason(D("1.30"), D("3.50")))
        self.assertIsNotNone(boundary_reason(D("1.29999999"), D("2.00")))
        self.assertIsNone(boundary_reason(D("2.00"), D("3.50000001")))
        displayed_as_minimum = D("1.299999995")
        from detectors.swing_open_close_4h_special_v1.workbook import shown
        self.assertEqual(shown(displayed_as_minimum), D("1.30000000"))
        visual = analyze(self._band("101.299999995"))
        self.assertEqual(search_at(0, visual).terminal, "BOTH_BOUNDARIES_BELOW_MINIMUM")
        self.assertFalse(swing_at(0, visual))
        doji = analyze(bars_from([
            ("100", "102", "99", "100"),
            ("100.2", "103", "99", "99.5"),
        ]))
        self.assertEqual(search_at(0, doji).terminal, "ZERO_INTERIOR_DOJI_NOT_ALLOWED")
        self.assertFalse(swing_at(0, doji))

    def test_zero_interior_two_candle_class(self):
        low = analyze(bars_from([
            ("100", "100.5", "98", "99"),
            ("98.5", "100.2", "99", "100"),
        ]))
        low_swing = next(item for item in low.raw if item.direction == SWING_LOW and item.open_row == 0)
        self.assertEqual(low_swing.formation_class, ZERO_CLASS)
        self.assertEqual(low_swing.extreme_price, D("98"))
        self.assertEqual(low_swing.extreme_row, 0)
        self.assertEqual(low_swing.total, 2)
        high = analyze(bars_from([
            ("100", "100.4", "99", "100.2"),
            ("100.1", "102", "99.5", "100"),
        ]))
        high_swing = swing_at(0, high)[0]
        self.assertEqual(high_swing.formation_class, ZERO_CLASS)
        self.assertEqual(high_swing.extreme_price, D("102"))
        self.assertEqual(high_swing.extreme_row, 1)
        self.assertLess(high.bars[high_swing.open_row].close_time, high.bars[high_swing.close_row].close_time)
        tied = analyze(bars_from([
            ("100", "102", "99", "100.3"),
            ("100.2", "102", "99.5", "100"),
        ]))
        tied_swing = swing_at(0, tied)[0]
        self.assertTrue(tied_swing.zero_plateau)
        self.assertEqual(tied_swing.zero_match_count, 2)
        self.assertEqual(tied_swing.extreme_row, 1)
        exact_low = analyze(self._zero_high("101.3"))
        exact_high = analyze(self._zero_high("103.5"))
        self.assertEqual(swing_at(0, exact_low)[0].open_width_percent, D("1.3"))
        self.assertEqual(swing_at(0, exact_high)[0].open_width_percent, D("3.5"))
        self.assertEqual(search_at(0, analyze(self._zero_high("101.2"))).terminal, "BOTH_BOUNDARIES_BELOW_MINIMUM")
        above_zero = analyze(self._zero_high("105"))
        self.assertEqual(swing_at(0, above_zero)[0].open_width_percent, D("5"))
        self.assertEqual(above_zero.counters["maximum_boundary_rejection_count"], 0)
        bearish_close = analyze(bars_from([
            ("100", "100.5", "98", "99"),
            ("100.4", "100.6", "99", "100"),
        ]))
        self.assertEqual(search_at(0, bearish_close, SWING_LOW).terminal, "ZERO_INTERIOR_SWING_LOW_CLOSE_NOT_BULLISH")
        bullish_open = analyze(bars_from([
            ("100", "100.5", "98", "100.2"),
            ("100.1", "100.4", "99", "100.2"),
        ]))
        self.assertEqual(search_at(0, bullish_open, SWING_LOW).terminal, "ZERO_INTERIOR_SWING_LOW_OPEN_NOT_BEARISH")
        self.assertEqual(zero_return_reason(SWING_HIGH, D("100"), D("100.1")), "ZERO_INTERIOR_REFERENCE_RETURN_NOT_ACHIEVED")
        self.assertIsNone(zero_return_reason(SWING_HIGH, D("100"), D("100")))
        self.assertEqual(
            zero_direction_reason(SWING_HIGH, D("100"), D("100.2"), D("100.1"), D("100.2")),
            "ZERO_INTERIOR_SWING_HIGH_CLOSE_NOT_BEARISH",
        )
        continued = analyze(bars_from([
            ("100", "100.4", "99", "100.4"),
            ("100.2", "102", "99.8", "100.5"),
            ("100.1", "100.4", "99.5", "100"),
        ]))
        self.assertFalse(any(item.open_row == 0 and item.formation_class == ZERO_CLASS for item in continued.raw))
        self.assertTrue(any(item.open_row == 0 and item.formation_class == STANDARD_CLASS for item in continued.raw))
        family = analyze(bars_from([
            ("99.00", "99.40", "98.50", "99.20"),
            ("100.00", "100.40", "99.50", "100.30"),
            ("100.20", "102.00", "99.60", "100.00"),
            ("99.40", "99.80", "98.90", "99.50"),
            ("99.20", "99.60", "98.80", "99.00"),
        ]))
        members = [item for item in family.raw if item.extreme_price == D("102") and item.direction == SWING_HIGH]
        self.assertTrue(any(item.formation_class == ZERO_CLASS and item.status == PRIMARY for item in members))
        self.assertTrue(any(item.formation_class == STANDARD_CLASS and item.status == DERIVED and item.total > 2 for item in members))
        self.assertEqual(signature(high), signature(analyze(bars_from([
            ("100", "100.4", "99", "100.2"),
            ("100.1", "102", "99.5", "100"),
        ]))))

    def test_normal_standard_one_third_body_penetration(self):
        bullish = swing_open_body(D("100"), D("103"))
        bearish = swing_open_body(D("103"), D("100"))
        self.assertEqual(bullish, (D("100"), D("103"), D("3")))
        self.assertEqual(bearish, bullish)
        high_threshold = body_threshold(SWING_HIGH, *bullish)
        low_threshold = body_threshold(SWING_LOW, *bullish)
        self.assertEqual(high_threshold, D("102"))
        self.assertEqual(low_threshold, D("101"))
        wick_threshold = body_threshold(SWING_HIGH, D("90"), D("110"), D("20"))
        self.assertNotEqual(wick_threshold, high_threshold)
        exact = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.4", "103.4", "101", "102.8"),
            ("102.2", "102.6", "101.5", "102"),
            ("100", "104", "99", "100"),
        ]))
        exact_swing = swing_at(0, exact)[0]
        self.assertEqual(exact_swing.formation_class, STANDARD_CLASS)
        self.assertEqual(exact_swing.close_row, 2)
        self.assertFalse(exact_swing.full_open_return)
        self.assertEqual(exact_swing.penetration_fraction, D(1) / D(3))
        self.assertEqual(exact_swing.body_threshold, D("102"))
        deeper = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.4", "103.4", "101", "102.8"),
            ("102.2", "102.6", "101.5", "101.5"),
        ]))
        self.assertGreater(swing_at(0, deeper)[0].penetration_fraction, D(1) / D(3))
        self.assertFalse(swing_at(0, deeper)[0].full_open_return)
        shy = D("102") + D("1E-18")
        missed = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.4", "103.4", "101", "102.8"),
            ("102.2", "102.6", "101.5", format(shy, "f")),
        ]))
        self.assertFalse(swing_at(0, missed))
        wick = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.4", "103.4", "101", "102.8"),
            ("102.4", "102.7", "101", "102.5"),
            ("102.1", "102.4", "101.2", "102"),
        ]))
        self.assertEqual(wick_only_reason(SWING_HIGH, D("102.7"), D("101"), D("102.5"), D("102")), "NORMAL_STANDARD_WICK_ONLY_BODY_PENETRATION")
        self.assertGreaterEqual(search_at(0, wick).wick_only, 1)
        self.assertEqual(swing_at(0, wick)[0].close_row, 3)
        self.assertEqual(wick.counters["wick_only_confirmations"], 0)
        doji = analyze(bars_from([
            ("100", "101", "99", "100"),
            ("100.2", "100.4", "99.5", "100.3"),
        ]))
        self.assertEqual(search_at(0, doji).terminal, "NORMAL_STANDARD_SWING_OPEN_BODY_ZERO")
        self.assertFalse(swing_at(0, doji))
        failed = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("101", "101.2", "100.5", "102.5"),
            ("101.5", "101.8", "101", "102"),
            ("100", "103.2", "99", "101"),
        ]))
        self.assertFalse(swing_at(0, failed))
        self.assertEqual(search_at(0, failed).terminal, "BOTH_BOUNDARIES_BELOW_MINIMUM")
        low = analyze(bars_from([
            ("100", "101", "96", "97"),
            ("98.5", "99", "96.5", "97.9"),
            ("98.2", "99", "97", "98"),
        ]))
        low_swing = next(item for item in low.raw if item.direction == SWING_LOW and item.open_row == 0)
        self.assertEqual(low_swing.formation_class, STANDARD_CLASS)
        self.assertFalse(low_swing.full_open_return)
        self.assertEqual(low_swing.penetration_fraction, D(1) / D(3))
        partial_only = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102", "102.5", "101", "102"),
        ]))
        self.assertFalse(any(item.open_row == 0 and item.formation_class == ZERO_CLASS for item in partial_only.raw))
        self.assertEqual(body_threshold(SWING_HIGH, *bullish), body_threshold(SWING_HIGH, *bullish))
        self.assertEqual(zero_adjacency_reason(1, 2, 0), "ZERO_INTERIOR_ENDPOINTS_NOT_ADJACENT")
        self.assertIsNone(zero_adjacency_reason(0, 1, 0))
        zero_high = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.2", "103.4", "101.5", "102"),
        ]))
        self.assertFalse(swing_at(0, zero_high))
        self.assertEqual(search_at(0, zero_high).terminal, "SEARCH_CENSORED_BY_DATASET_END")
        self.assertEqual(zero_return_reason(SWING_HIGH, D("100"), D("102")), "ZERO_INTERIOR_REFERENCE_RETURN_NOT_ACHIEVED")
        zero_low = analyze(bars_from([
            ("100", "100.4", "96.8", "97"),
            ("97.5", "98.2", "96.6", "98"),
        ]))
        self.assertFalse(any(item.direction == SWING_LOW and item.open_row == 0 for item in zero_low.raw))
        self.assertEqual(zero_return_reason(SWING_LOW, D("100"), D("98")), "ZERO_INTERIOR_REFERENCE_RETURN_NOT_ACHIEVED")
        shy_zero = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.2", "103.4", "101.5", format(D("102") + D("1E-18"), "f")),
        ]))
        self.assertFalse(swing_at(0, shy_zero))
        self.assertEqual(search_at(0, shy_zero).terminal, "SEARCH_CENSORED_BY_DATASET_END")
        zero_wick = analyze(bars_from([
            ("100", "101", "99", "103"),
            ("102.4", "102.8", "101", "102.5"),
        ]))
        self.assertFalse(swing_at(0, zero_wick))
        self.assertEqual(search_at(0, zero_wick).terminal, "SEARCH_CENSORED_BY_DATASET_END")

    def test_primary_family_nesting_and_sequential_structures(self):
        family = analyze(bars_from([
            ("100", "100.5", "99", "100.5"),
            ("100", "101", "99.5", "100.8"),
            ("100.2", "101", "99.8", "100.6"),
            ("100.4", "102", "100", "101"),
            ("100.8", "101.2", "100", "100.7"),
            ("100.5", "101", "100", "100.4"),
            ("100.3", "100.6", "99.5", "100"),
        ]))
        members = [item for item in family.raw if item.extreme_price == D("102") and item.direction == SWING_HIGH]
        self.assertGreaterEqual(len(members), 2)
        winner = min(members, key=lambda item: item.total)
        self.assertEqual(winner.status, PRIMARY)
        self.assertTrue(any(item.status == DERIVED and item.total > winner.total for item in members))
        nested = analyze(bars_from([
            ("99.50", "100.20", "99.00", "99.90"),
            ("100.00", "102.20", "99.40", "100.20"),
            ("99.70", "101.60", "99.30", "100.15"),
            ("100.10", "100.50", "99.60", "100.10"),
            ("100.10", "101.90", "99.40", "100.05"),
            ("99.60", "100.20", "99.20", "99.50"),
        ]))
        statuses = {item.open_row: item.status for item in nested.raw if item.direction == SWING_HIGH}
        self.assertEqual(statuses[1], PRIMARY)
        self.assertEqual(statuses[2], NESTED)
        self.assertEqual(statuses[0], LONGER)
        sequential = analyze(
            self._valid(1) + bars_from([flat_above("101"), flat_above("101")]) + self._valid(1, start_price="110")
        )
        displayed_opens = [item.open_row for item in sequential.displayed if item.direction == SWING_HIGH]
        self.assertIn(0, displayed_opens)
        self.assertTrue(any(row >= 5 for row in displayed_opens))
        self.assertFalse(any(item.interior == 0 for item in sequential.raw))
        mixed = analyze(self._valid(3) + self._low_block())
        directions = {item.direction for item in mixed.displayed}
        self.assertIn(SWING_HIGH, directions)
        self.assertIn(SWING_LOW, directions)
        tied = analyze(bars_from([
            ("100", "100.4", "99", "100.4"),
            ("99.70", "101", "99.2", "100.5"),
            ("101.5", "101.6", "99.8", "100.6"),
            ("100.3", "102", "100", "101"),
            ("100.4", "101.1", "100", "100.5"),
            ("100.2", "101", "99.5", "99.95"),
            ("99.9", "100.2", "99.2", "99.60"),
        ]))
        pair = [item for item in tied.raw if item.extreme_price == D("102") and item.direction == SWING_HIGH and item.open_row in {0, 1}]
        self.assertEqual(min(pair, key=lambda item: item.completion_error).status, PRIMARY)

    def test_workbook_timestamps_and_immutability(self):
        moment = datetime(2026, 1, 5, 4, tzinfo=timezone.utc)
        close = datetime(2026, 9, 15, 23, 59, 59, 999000, tzinfo=timezone.utc)
        self.assertEqual(moment.tzinfo, timezone.utc)
        self.assertEqual(turkey_text(moment), "2026-01-05 07:00:00+03:00")
        self.assertEqual(turkey_text(close), "2026-09-16 02:59:59.999+03:00")
        result = analyze(self._valid(1))
        item = next(row for row in result.displayed if row.open_row == 0)
        self.assertEqual(turkey_text(result.bars[item.open_row].open_time), "2026-01-01 03:00:00+03:00")
        self.assertEqual(turkey_text(result.bars[item.close_row].open_time), "2026-01-01 11:00:00+03:00")
        self.assertEqual(turkey_text(result.bars[item.close_row].close_time), "2026-01-01 14:59:59.999+03:00")
        self.assertEqual(turkey_text(result.bars[item.extreme_row].open_time), "2026-01-01 07:00:00+03:00")
        directory = TMP_DIR / "revision-test"
        directory.mkdir(parents=True, exist_ok=True)
        existing = directory / "BTCUSDT_4H_Special_Swings_rev00.xlsx"
        existing.write_bytes(b"keep")
        (directory / "BTCUSDT_4H_Swing_Structure_rev09.xlsx").write_bytes(b"ignore")
        revision, path = allocate_revision(directory)
        self.assertEqual((revision, path.name), (1, "BTCUSDT_4H_Special_Swings_rev01.xlsx"))
        self.assertEqual(existing.read_bytes(), b"keep")
        workbook = directory / "sample.xlsx"
        write_workbook(workbook, result)
        report = validate_workbook(workbook, result)
        self.assertEqual(report["workbook_validation"], "PASS")
        from openpyxl import load_workbook
        book = load_workbook(workbook)
        self.assertEqual(tuple(book.sheetnames), SHEETS)
        headers = [cell.value for cell in next(book["Special Swings"].iter_rows(max_row=1))]
        self.assertEqual(tuple(headers), RESULT_COLUMNS)
        self.assertFalse(any("UTC" in str(header).upper() for header in headers))
        self.assertIsNone(book.vba_archive)
        self.assertFalse(list(book._external_links))
        parameters = [tuple(row) for row in book["Parameters"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(parameters, list(USER_PARAMETERS))
        self.assertIn(("Maximum Boundary Percent", "None — No Upper Limit"), parameters)
        self.assertIn(("Minimum Boundary Percent", "1.30%"), parameters)
        self.assertIn(("Boundary Above 3.50% Allowed", "Yes"), parameters)
        self.assertIn(("Zero-Interior Confirmation", "Full Return to Swing Open Open"), parameters)
        self.assertIn(("Zero-Interior One-Third Body Rule", "Disabled"), parameters)
        self.assertIn(("Normal Standard Confirmation", "Minimum One-Third Swing Open Body Penetration"), parameters)
        values = [row[0] for row in book["Special Swings"].iter_rows(min_row=2, values_only=True) if row[0]]
        self.assertEqual(len(values), len(set(values)))
        self.assertEqual(book["Special Swings"].max_row - 1, len(result.displayed))
        book.close()
        self.assertEqual(PROTECTED_BEFORE, fingerprint_tree(PROTECTED_ROOT))
        self.assertEqual(SOURCE_BEFORE, fingerprint_tree(Path(DATASET)))
        self.assertEqual(signature(result), signature(analyze(self._valid(1))))

    def _zero_high(self, extreme: str) -> list[Bar]:
        return bars_from([
            ("100", "100.4", "99", "100.2"),
            ("100.1", extreme, "99.5", "100"),
        ])

    def _valid(self, interior: int, start_price: str = "100") -> list[Bar]:
        reference = D(start_price)
        extreme = reference * D("1.02")
        specs = [(str(reference), str(reference + D("0.2")), str(reference - D("1")), str(reference + D("0.4")))]
        for offset in range(1, interior):
            specs.append((str(reference + D("0.3")), str(reference + D("0.6")), str(reference), str(reference + D("0.5"))))
        specs.append((str(reference + D("0.3")), str(extreme), str(reference), str(reference + D("0.5"))))
        specs.append((str(reference), str(reference + D("0.2")), str(reference - D("1")), str(reference)))
        return bars_from(specs)

    def _return_at(self, interior: int) -> list[Bar]:
        specs = [("100", "100.4", "99", "100.4")]
        for _ in range(interior):
            specs.append(flat_above())
        specs.append(("101", "101.2", "99", "100"))
        specs.append(flat_above())
        return bars_from(specs)

    def _band(self, extreme: str) -> list[Bar]:
        return bars_from([
            ("100", "100.4", "99", "100.4"),
            ("100.2", extreme, "100", "100.5"),
            ("100.3", "100.6", "100", "100.4"),
            ("100.2", "100.5", "100", "100.3"),
            ("100", "100.2", "99", "100"),
        ])

    def _high_widths(self, extreme: Decimal, close: Decimal) -> list[Bar]:
        return bars_from([
            ("100", "100.4", "99", "100.4"),
            ("100.2", format(extreme, "f"), "99.5", "100.6"),
            ("100.3", "100.5", "100", "100.4"),
            ("100.2", "100.4", "100", "100.3"),
            ("100", "100.2", "90", format(close, "f")),
        ])

    def _low_widths(self, extreme: str, close: str) -> list[Bar]:
        return bars_from([
            ("100", "100.4", "99.8", "99.9"),
            ("99.8", "100", extreme, "99.7"),
            ("99.7", "99.9", "99.6", "99.8"),
            ("99.8", "100", "99.7", "99.9"),
            ("99.9", "101", "99.6", close),
        ])

    def _low_block(self) -> list[Bar]:
        moment_shift = len(self._valid(3))
        rows = []
        specs = [
            ("100", "100.5", "99", "99.8"),
            ("99.7", "100", "98", "99.2"),
            ("99.4", "99.8", "98.4", "99.1"),
            ("99.2", "99.6", "98.2", "99"),
            ("99", "100.2", "98.8", "100"),
        ]
        for index, (opened, high, low, close) in enumerate(specs):
            moment = START + timedelta(hours=4 * (moment_shift + index))
            rows.append(Bar(
                row=moment_shift + index,
                open_time=moment,
                close_time=moment + timedelta(hours=4, milliseconds=-1),
                open=D(opened),
                high=D(high),
                low=D(low),
                close=D(close),
                volume=D("1"),
                quote_volume=D("1"),
                trades=1,
            ))
        return rows


if __name__ == "__main__":
    unittest.main()
