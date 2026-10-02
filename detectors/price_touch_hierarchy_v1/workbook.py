"""Excel workbook for the price-touch hierarchy detector."""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.price_touch_hierarchy_v1.engine import (
    AnalysisResult,
    HierarchyGroup,
    LevelStats,
    directional_balance,
    gap_mean,
    gap_median,
    iso_turkey,
    iso_utc,
    level_persistence,
    rejection_reasons,
)
from detectors.price_touch_hierarchy_v1.touch_config import (
    CONFIG,
    CROSS_TF_TOLERANCE,
    DATASET_END,
    DETECTOR_VERSION,
    FINAL_HALF_WIDTH,
    GRID_ANCHOR,
    GRID_GROWTH,
    GRID_STEP_PERCENT,
    MAX_GROUP_SPAN,
    MIN_DISTINCT_DATES,
    MIN_EPISODES,
    MIN_TOUCHES,
    PEAK_RADIUS,
    SAME_TF_SEPARATION,
    SCORE_WEIGHTS,
    TIMEFRAME_WEIGHT,
    TOP_TOUCH_GROUPS,
    WORKBOOK_SHEETS,
)

D0 = Decimal("0")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FILL_TOP = PatternFill("solid", fgColor="C6EFCE")
FILL_REJECT = PatternFill("solid", fgColor="FCE4D6")
FILL_NEAR = PatternFill("solid", fgColor="FFF2CC")
FILL_TIER = {
    "TIER_1_ALL_TIMEFRAMES": PatternFill("solid", fgColor="548235"),
    "TIER_2_DAILY_WITH_INTRADAY": PatternFill("solid", fgColor="5B9BD5"),
    "TIER_3_DAILY_ONLY": PatternFill("solid", fgColor="8064A2"),
    "TIER_4_4H_AND_1H": PatternFill("solid", fgColor="ED7D31"),
    "TIER_5_4H_ONLY": PatternFill("solid", fgColor="FFC000"),
    "TIER_6_1H_ONLY": PatternFill("solid", fgColor="D9D9D9"),
}
TAB_COLOR = {
    "1D Levels": "8064A2",
    "4H Levels": "ED7D31",
    "1H Levels": "5B9BD5",
    "Hierarchical Ranking": "548235",
}

PRICE_FORMAT = "0.00000000"
PERCENT_FORMAT = "0.00%"
SCORE_FORMAT = "0.000000"
INT_FORMAT = "#,##0"


def _num(value: Decimal | int | float | None) -> float | int | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def _pct_points(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value / Decimal(100))


def _write_table(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]], formats: dict[int, str], table_name: str) -> None:
    ws.append(list(headers))
    for row in rows:
        if len(row) != len(headers):
            raise RuntimeError(f"{ws.title} row width {len(row)} != {len(headers)}")
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(headers))}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for index, header in enumerate(headers, start=1):
        width = min(42, max(12, len(header) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width
        fmt = formats.get(index)
        if not fmt:
            continue
        for cell in ws.iter_cols(min_col=index, max_col=index, min_row=2, max_row=ws.max_row):
            for item in cell:
                if isinstance(item.value, (int, float)):
                    item.number_format = fmt
    ws.sheet_properties.tabColor = TAB_COLOR.get(ws.title, "1F4E79")


def _formats_for(headers: Sequence[str]) -> dict[int, str]:
    formats: dict[int, str] = {}
    for index, header in enumerate(headers, start=1):
        name = header.lower()
        if any(token in name for token in ("flag", " pass", "label", "reason", " id", "tier", "version", "timeframe", "classification")):
            continue
        if "percentile" in name or name.endswith(" rate") or "balance" in name or "share" in name or "percent" in name:
            formats[index] = PERCENT_FORMAT
        elif "score" in name:
            formats[index] = SCORE_FORMAT
        elif any(token in name for token in ("price", "zone lower", "zone upper", "width usdt")) or name in {"open", "high", "low", "close"}:
            formats[index] = PRICE_FORMAT
        elif any(token in name for token in ("touch", "episode", "count", "rank", "candle", "bin", "shortfall")):
            formats[index] = INT_FORMAT
    return formats


def _level_row(level: LevelStats, candle_count: int) -> list[Any]:
    width = level.zone_upper - level.zone_lower
    width_pct = width / level.price if level.price else D0
    touch_rate = Decimal(level.touches) / Decimal(candle_count) if candle_count else D0
    candles_per_episode = Decimal(level.touches) / Decimal(level.episodes) if level.episodes else D0
    return [
        level.rank,
        level.level_id,
        _num(level.price),
        _num(level.zone_lower),
        _num(level.zone_upper),
        _num(width),
        _num(width_pct),
        MIN_TOUCHES[level.timeframe],
        True,
        level.raw_atomic_touches,
        level.touches,
        _num(level.percentiles.get("touch", D0)),
        _num(touch_rate),
        level.episodes,
        MIN_EPISODES[level.timeframe],
        True,
        level.distinct_dates,
        MIN_DISTINCT_DATES[level.timeframe],
        True,
        _num(candles_per_episode),
        level.longest_episode,
        level.distinct_weeks,
        level.distinct_months,
        level.distinct_years,
        iso_utc(level.first_touch) if level.first_touch else "",
        iso_turkey(level.first_touch) if level.first_touch else "",
        iso_utc(level.last_touch) if level.last_touch else "",
        iso_turkey(level.last_touch) if level.last_touch else "",
        iso_utc(level.first_eligible) if level.first_eligible else "",
        iso_turkey(level.first_eligible) if level.first_eligible else "",
        level.close_inside,
        level.body,
        level.wick,
        level.cross_through,
        level.bullish,
        level.bearish,
        level.doji,
        _num(directional_balance(level.bullish, level.bearish)),
        level_persistence(level),
        _num(level.smoothed_density),
        PEAK_RADIUS,
        level.suppressed_neighbor_count,
        level.group_id,
        DETECTOR_VERSION,
    ]


LEVEL_HEADERS = [
    "Timeframe Rank", "Level ID", "Representative Price", "Zone Lower", "Zone Upper",
    "Zone Width USDT", "Zone Width Percent", "Applicable Minimum-Touch Threshold",
    "Touch Threshold Passed", "Raw Atomic-Bin Touch Count",
    "Final Recounted Completed-Candle Touch Count", "Touch-Count Percentile",
    "Candle Touch Rate", "Independent Episodes", "Minimum Episodes Required",
    "Episode Threshold Passed", "Distinct UTC Dates", "Minimum Distinct Dates Required",
    "Distinct-Date Threshold Passed", "Candles per Episode", "Longest Episode",
    "Distinct Weeks", "Distinct Months", "Distinct Years",
    "First Touch UTC", "First Touch Turkey", "Last Touch UTC", "Last Touch Turkey",
    "First Eligible UTC", "First Eligible Turkey", "Close-Inside Count",
    "Body-Intersection Count", "Wick-Only Count", "Full Cross-Through Count",
    "Bullish Count", "Bearish Count", "Doji Count", "Directional Balance",
    "Temporal Persistence Label", "Smoothed Density", "Local Peak Radius",
    "Suppressed Neighbor Count", "Parent Hierarchy Group ID", "Detector Version",
]


def _group_touch_sum(group: HierarchyGroup, timeframe: str, attr: str) -> int | str:
    level = group.members.get(timeframe)
    if level is None:
        return ""
    return getattr(level, attr)


def _group_row(group: HierarchyGroup) -> list[Any]:
    width = group.consolidated_upper - group.consolidated_lower
    width_pct = width / group.representative_price if group.representative_price else D0
    canonical = group.members[group.canonical_timeframe]
    dates = canonical.distinct_dates
    return [
        group.rank,
        group.group_id,
        group.tier,
        _num(group.score),
        _num(group.representative_price),
        _num(group.consolidated_lower),
        _num(group.consolidated_upper),
        _num(width),
        _num(width_pct),
        "+".join(group.timeframes),
        len(group.timeframes),
        group.members["1d"].level_id if "1d" in group.members else "",
        group.members["4h"].level_id if "4h" in group.members else "",
        group.members["1h"].level_id if "1h" in group.members else "",
        group.members["1d"].rank if "1d" in group.members else "",
        group.members["4h"].rank if "4h" in group.members else "",
        group.members["1h"].rank if "1h" in group.members else "",
        _group_touch_sum(group, "1d", "touches"),
        _group_touch_sum(group, "4h", "touches"),
        _group_touch_sum(group, "1h", "touches"),
        _group_touch_sum(group, "1d", "episodes"),
        _group_touch_sum(group, "4h", "episodes"),
        _group_touch_sum(group, "1h", "episodes"),
        group.canonical_visits,
        dates,
        canonical.distinct_weeks,
        canonical.distinct_months,
        canonical.distinct_years,
        iso_utc(group.first_touch) if group.first_touch else "",
        iso_turkey(group.first_touch) if group.first_touch else "",
        iso_utc(group.last_touch) if group.last_touch else "",
        iso_turkey(group.last_touch) if group.last_touch else "",
        iso_utc(group.first_eligible) if group.first_eligible else "",
        iso_turkey(group.first_eligible) if group.first_eligible else "",
        canonical.close_inside,
        canonical.body,
        canonical.wick,
        canonical.cross_through,
        canonical.bullish,
        canonical.bearish,
        canonical.doji,
        _num(directional_balance(canonical.bullish, canonical.bearish)),
        group.persistence,
        _pct_points(group.max_child_distance),
        _pct_points(group.max_span),
        "EX_POST_HISTORICAL_STRUCTURE",
        DETECTOR_VERSION,
    ]


RANK_HEADERS = [
    "Overall Hierarchical Rank", "Hierarchy Group ID", "Hierarchy Tier", "Hierarchy Score",
    "Representative Price", "Consolidated Zone Lower", "Consolidated Zone Upper",
    "Zone Width USDT", "Zone Width Percent", "Timeframes Present", "Timeframe Count",
    "1D Level ID", "4H Level ID", "1H Level ID", "1D Rank", "4H Rank", "1H Rank",
    "1D Completed-Candle Touches", "4H Completed-Candle Touches", "1H Completed-Candle Touches",
    "1D Independent Episodes", "4H Independent Episodes", "1H Independent Episodes",
    "Canonical Independent Market Visits", "Distinct Dates", "Distinct Weeks", "Distinct Months",
    "Distinct Years", "First Touch UTC", "First Touch Turkey", "Last Touch UTC", "Last Touch Turkey",
    "First Eligible UTC", "First Eligible Turkey", "Close-Inside Touches", "Body-Intersection Touches",
    "Wick-Only Touches", "Cross-Through Touches", "Bullish Touches", "Bearish Touches", "Doji Touches",
    "Directional Balance", "Temporal Persistence Label", "Maximum Child-to-Parent Distance Percent",
    "Maximum Group Span Percent", "Historical Classification", "Detector Version",
]


def _summary_rows(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    def add(item: str, value: Any) -> list[Any]:
        return [item, value]

    rows = [
        add("Symbol", "BTCUSDT"),
        add("TradingView Symbol", "BINANCE:BTCUSDT.P"),
        add("Market", "Binance USD-M USDT-margined PERPETUAL"),
        add("Analysis Start UTC", "2024-01-01T00:00:00.000Z"),
        add("Analysis End UTC", iso_utc(DATASET_END)),
        add("Timeframes", "1h, 4h, 1d"),
        add("1h Source Rows", len(result.bars.get("1h", []))),
        add("4h Source Rows", len(result.bars.get("4h", []))),
        add("1d Source Rows", len(result.bars.get("1d", []))),
        add("1h meaningful-level threshold", "40 completed-candle touches"),
        add("4h meaningful-level threshold", "20 completed-candle touches"),
        add("1d meaningful-level threshold", "10 completed-candle touches"),
        add("Qualified 1h Levels", len(result.levels.get("1h", []))),
        add("Qualified 4h Levels", len(result.levels.get("4h", []))),
        add("Qualified 1d Levels", len(result.levels.get("1d", []))),
        add("Rejected 1h Candidates", sum(1 for level in result.rejected if level.timeframe == "1h")),
        add("Rejected 4h Candidates", sum(1 for level in result.rejected if level.timeframe == "4h")),
        add("Rejected 1d Candidates", sum(1 for level in result.rejected if level.timeframe == "1d")),
        add("Hierarchy Groups", len(result.groups)),
        add("Detector Version", DETECTOR_VERSION),
        add("Workbook Revision", context.get("revision_name", "")),
        add("One Candle Touch Rule", "One completed candle contributes at most one touch to the same zone."),
        add("Intrabar Rule", "Intrabar repeated crossings are not counted. The intrabar path is unknown."),
        add("Cross-Timeframe Count Rule", "Raw 1h, 4h, and 1d touch counts are never added together."),
        add("Result Class", "Historical price-interaction structure. Not a trade signal."),
    ]
    for tier in (
        "TIER_1_ALL_TIMEFRAMES",
        "TIER_2_DAILY_WITH_INTRADAY",
        "TIER_3_DAILY_ONLY",
        "TIER_4_4H_AND_1H",
        "TIER_5_4H_ONLY",
        "TIER_6_1H_ONLY",
    ):
        rows.append(add(tier, sum(1 for group in result.groups if group.tier == tier)))
    for timeframe, label in (("1d", "Most-Touched 1d"), ("4h", "Most-Touched 4h"), ("1h", "Most-Touched 1h")):
        levels = result.levels.get(timeframe, [])
        if not levels:
            rows.append(add(label, "NONE"))
            continue
        best = min(levels, key=lambda level: level.rank)
        rows.append(add(label, f"{best.level_id} price={best.price} touches={best.touches}"))
    if context.get("most_persistent"):
        rows.append(add("Most Persistent Level", context["most_persistent"]))
    if context.get("highest_canonical"):
        rows.append(add("Highest Canonical Independent Visits", context["highest_canonical"]))
    for group in result.groups[:20]:
        rows.append(add(
            f"Top {group.rank:02d}",
            f"{group.tier} {group.representative_price} score={group.score} visits={group.canonical_visits} { '+'.join(group.timeframes)}",
        ))
    return rows


def _mapping_rows(result: AnalysisResult) -> list[list[Any]]:
    rows = []
    for attempt in result.attempts:
        rows.append([
            attempt.attempt_id,
            attempt.parent_timeframe,
            attempt.parent_level_id,
            _num(attempt.parent_price),
            attempt.child_timeframe,
            attempt.child_level_id,
            _num(attempt.child_price),
            _pct_points(attempt.distance_percent),
            attempt.zone_overlap,
            _pct_points(attempt.proposed_span_percent),
            attempt.tolerance_pass,
            attempt.span_pass,
            attempt.accepted,
            attempt.rejection_reason,
            attempt.selected_parent,
            attempt.tie_break,
        ])
    return rows


MAPPING_HEADERS = [
    "Mapping Attempt ID", "Parent Timeframe", "Parent Level ID", "Parent Price",
    "Child Timeframe", "Child Level ID", "Child Price", "Relative Distance Percent",
    "Zone Overlap", "Proposed Group Span Percent", "Match Tolerance Pass", "Group Span Pass",
    "Accepted", "Rejection Reason", "Selected Parent", "Tie-Break Outcome",
]


def _suppressed_rows(result: AnalysisResult) -> list[list[Any]]:
    winners = {}
    for timeframe_levels in result.levels.values():
        for level in timeframe_levels:
            winners[level.candidate_id] = level
    for level in result.suppressed:
        winners.setdefault(level.candidate_id, level)
    rows = []
    for level in result.suppressed:
        winner = winners.get(level.suppressed_by)
        rows.append([
            level.candidate_id,
            level.timeframe,
            _num(level.price),
            level.suppressed_by,
            _num(winner.price) if winner else "",
            _pct_points(level.suppression_distance) if level.suppression_distance is not None else "",
            level.touches,
            level.episodes,
            level.distinct_months,
            "SUPPRESSED_AS_NEAR_DUPLICATE",
            level.would_be_rank or "",
        ])
    return rows


SUPPRESSED_HEADERS = [
    "Suppressed Level ID", "Timeframe", "Suppressed Price", "Winning Level ID", "Winning Price",
    "Relative Distance Percent", "Original Completed-Candle Touches", "Original Independent Episodes",
    "Original Distinct Months", "Suppression Reason", "Would-Have-Been Timeframe Rank",
]


def _rejected_rows(result: AnalysisResult) -> list[list[Any]]:
    rows = []
    for level in result.rejected:
        primary, reasons = rejection_reasons(level)
        shortfall = max(0, MIN_TOUCHES[level.timeframe] - level.touches)
        rows.append([
            level.candidate_id,
            level.timeframe,
            _num(level.price),
            _num(level.zone_lower),
            _num(level.zone_upper),
            level.touches,
            MIN_TOUCHES[level.timeframe],
            shortfall,
            level.episodes,
            MIN_EPISODES[level.timeframe],
            level.distinct_dates,
            MIN_DISTINCT_DATES[level.timeframe],
            level.touches >= MIN_TOUCHES[level.timeframe],
            level.episodes >= MIN_EPISODES[level.timeframe],
            level.distinct_dates >= MIN_DISTINCT_DATES[level.timeframe],
            primary,
            "; ".join(reasons),
            level.would_be_rank or "",
            iso_utc(level.first_touch) if level.first_touch else "",
            iso_utc(level.last_touch) if level.last_touch else "",
        ])
    return rows


REJECTED_HEADERS = [
    "Candidate ID", "Timeframe", "Representative Price", "Zone Lower", "Zone Upper",
    "Completed-Candle Touches", "Required Completed-Candle Touches", "Touch Shortfall",
    "Independent Episodes", "Required Independent Episodes", "Distinct UTC Dates",
    "Required Distinct UTC Dates", "Touch Threshold Pass", "Episode Threshold Pass",
    "Distinct-Date Threshold Pass", "Primary Rejection Reason", "All Rejection Reasons",
    "Would-Have-Been Rank", "First Touch", "Last Touch",
]


def _touch_rows(result: AnalysisResult) -> list[list[Any]]:
    rows = []
    for group in result.groups[:TOP_TOUCH_GROUPS]:
        for timeframe in group.timeframes:
            level = group.members[timeframe]
            for event in level.events:
                direction = "DOJI" if event.doji else "BULLISH" if event.bullish else "BEARISH"
                rows.append([
                    group.group_id,
                    group.rank,
                    timeframe,
                    level.level_id,
                    iso_utc(event.open_time),
                    iso_utc(event.close_time),
                    iso_turkey(event.open_time),
                    iso_turkey(event.close_time),
                    _num(event.open),
                    _num(event.high),
                    _num(event.low),
                    _num(event.close),
                    _num(event.zone_lower),
                    _num(event.zone_upper),
                    event.touch_type,
                    event.open_inside,
                    event.close_inside,
                    event.high_inside,
                    event.low_inside,
                    event.cross_through,
                    direction,
                    event.episode_id,
                    event.episode_start,
                    event.episode_end,
                    event.contribution,
                ])
    return rows


TOUCH_HEADERS = [
    "Hierarchy Group ID", "Overall Rank", "Timeframe", "Source Level ID",
    "Candle Open UTC", "Candle Close UTC", "Candle Open Turkey", "Candle Close Turkey",
    "Open", "High", "Low", "Close", "Zone Lower", "Zone Upper", "Primary Touch Type",
    "Open Inside Flag", "Close Inside Flag", "High Inside Flag", "Low Inside Flag",
    "Full Cross-Through Flag", "Bullish/Bearish/Doji", "Episode ID", "Episode Start Flag",
    "Episode End Flag", "Touch Contribution",
]


def _month_span(level: LevelStats) -> list[str]:
    if not level.months:
        return []
    start = min(level.months)
    end = max(level.months)
    months = []
    year, month = start
    while (year, month) <= end:
        months.append(f"{year:04d}-{month:02d}")
        month += 1
        if month == 13:
            month = 1
            year += 1
    return months


def _monthly_rows(result: AnalysisResult) -> list[list[Any]]:
    rows = []
    for group in result.groups:
        for timeframe in group.timeframes:
            level = group.members[timeframe]
            by_month: dict[str, list] = {}
            for event in level.events:
                key = event.open_time.astimezone(timezone.utc).strftime("%Y-%m")
                by_month.setdefault(key, []).append(event)
            cumulative_touches = 0
            cumulative_episodes = 0
            seen_dates: set = set()
            episodes_seen = 0
            already_eligible = False
            for month in _month_span(level):
                events = by_month.get(month, [])
                starts = sum(1 for event in events if event.episode_start)
                cumulative_touches += len(events)
                cumulative_episodes += starts
                for event in events:
                    seen_dates.add(event.open_time.astimezone(timezone.utc).date())
                    if event.episode_start:
                        episodes_seen += 1
                eligible = (
                    cumulative_touches >= MIN_TOUCHES[timeframe]
                    and episodes_seen >= MIN_EPISODES[timeframe]
                    and len(seen_dates) >= MIN_DISTINCT_DATES[timeframe]
                )
                reached = eligible and not already_eligible
                if eligible:
                    already_eligible = True
                rows.append([
                    month,
                    timeframe,
                    group.group_id,
                    len(events),
                    starts,
                    sum(1 for event in events if event.touch_type == "CLOSE_INSIDE_ZONE"),
                    sum(1 for event in events if event.touch_type == "BODY_INTERSECTION"),
                    sum(1 for event in events if event.touch_type == "WICK_ONLY_INTERSECTION"),
                    sum(1 for event in events if event.cross_through),
                    sum(1 for event in events if event.bullish),
                    sum(1 for event in events if event.bearish),
                    len({event.open_time.astimezone(timezone.utc).date() for event in events}),
                    iso_utc(events[0].open_time) if events else "",
                    iso_utc(events[-1].open_time) if events else "",
                    cumulative_touches,
                    cumulative_episodes,
                    eligible,
                    MIN_TOUCHES[timeframe],
                    reached,
                    True,
                ])
    return rows


MONTH_HEADERS = [
    "Year-Month", "Timeframe", "Hierarchy Group ID", "Completed-Candle Touches",
    "Independent Episodes", "Close-Inside Touches", "Body-Intersection Touches",
    "Wick-Only Touches", "Cross-Through Touches", "Bullish Touches", "Bearish Touches",
    "Distinct Touched Dates", "First Monthly Touch", "Last Monthly Touch",
    "Cumulative Touches Through Month End", "Cumulative Episodes Through Month End",
    "Eligible By Month End", "Applicable Touch Threshold", "Threshold Reached During Month",
    "Historical Rank Known Only Ex Post",
]


def _parameter_rows() -> list[list[Any]]:
    rows = [
        ["symbol", "BTCUSDT"],
        ["tradingview_symbol", "BINANCE:BTCUSDT.P"],
        ["market", "Binance USD-M PERPETUAL"],
        ["timeframes", "1h, 4h, 1d"],
        ["analysis_start_utc", "2024-01-01T00:00:00Z"],
        ["analysis_end_utc", "2026-09-15T23:59:59.999Z"],
        ["completed_candles_only", "TRUE"],
        ["maximum_touch_per_candle_per_zone", 1],
        ["intrabar_path_inference", "FORBIDDEN"],
        ["price_grid_type", "LOGARITHMIC"],
        ["price_grid_anchor_usdt", "1.00000000"],
        ["price_grid_step_percent", "0.10"],
        ["price_grid_growth_factor", "1.001"],
        ["density_smoothing_kernel", "[0.25, 0.50, 0.25]"],
        ["local_peak_radius_bins", PEAK_RADIUS],
        ["final_zone_half_width_percent", "0.10"],
        ["minimum_same_timeframe_level_separation_percent", "0.50"],
        ["cross_timeframe_match_tolerance_percent", "0.30"],
        ["maximum_group_center_span_percent", "0.60"],
        ["1h_minimum_completed_candle_touches", 40],
        ["1h_minimum_independent_episodes", 3],
        ["1h_minimum_distinct_dates", 3],
        ["4h_minimum_completed_candle_touches", 20],
        ["4h_minimum_independent_episodes", 3],
        ["4h_minimum_distinct_dates", 3],
        ["1d_minimum_completed_candle_touches", 10],
        ["1d_minimum_independent_episodes", 3],
        ["1d_minimum_distinct_dates", 3],
        ["timeframe_weight_1d", 4],
        ["timeframe_weight_4h", 2],
        ["timeframe_weight_1h", 1],
        ["touch_score_weight", "0.60"],
        ["episode_score_weight", "0.20"],
        ["persistence_score_weight", "0.15"],
        ["engagement_score_weight", "0.05"],
        ["top_touch_event_groups", TOP_TOUCH_GROUPS],
        ["implementation_mode", CONFIG.implementation_mode],
        ["project_root", str(CONFIG.project_root)],
        ["detector_package", str(CONFIG.package_dir)],
        ["detector_version", DETECTOR_VERSION],
        ["formula_grid_center", "grid_center(k) = 1.00000000 * 1.001 ** k"],
        ["formula_atomic_bounds", "lower = sqrt(P(k-1)*P(k)); upper = sqrt(P(k)*P(k+1))"],
        ["formula_touch", "high >= zone_lower AND low <= zone_upper, counted once per candle"],
        ["formula_final_zone", "lower = P * 0.999; upper = P * 1.001"],
        ["formula_relative_distance", "abs(a-b) / ((a+b)/2) * 100"],
        ["formula_score", "100 * (0.60*touch + 0.20*episode + 0.15*persistence + 0.05*engagement)"],
        ["formula_smoothing", "0.25*raw[k-1] + 0.50*raw[k] + 0.25*raw[k+1]; missing neighbor is 0"],
        ["tie_break_peak", "higher raw touches, episodes, months, close-inside, earlier first touch, lower grid index"],
        ["tie_break_nms", "higher final touches, episodes, months, close+body, then lower price"],
        ["tie_break_rank", "touches, episodes, months, years, close-inside, body, span, then lower price"],
        ["tie_break_parent", "higher timeframe, smaller distance, higher touch percentile, higher episode percentile, lower level id"],
        ["persistence_rules", "RECENT_OR_EMERGING if months<=2 or span<60d; FREQUENT_BUT_CONCENTRATED if most-active month share>=0.45; FREQUENT_AND_PERSISTENT if months>=6 and share<=0.35; else INFREQUENT_BUT_PERSISTENT"],
        ["close_time_1h", "open + 1 hour - 1 millisecond"],
        ["close_time_4h", "open + 4 hours - 1 millisecond"],
        ["close_time_1d", "open + 1 UTC day - 1 millisecond"],
        ["grid_anchor_constant", str(GRID_ANCHOR)],
        ["grid_growth_constant", str(GRID_GROWTH)],
        ["grid_step_constant", str(GRID_STEP_PERCENT)],
        ["final_half_width_constant", str(FINAL_HALF_WIDTH)],
        ["same_tf_separation_constant", str(SAME_TF_SEPARATION)],
        ["cross_tf_tolerance_constant", str(CROSS_TF_TOLERANCE)],
        ["max_group_span_constant", str(MAX_GROUP_SPAN)],
        ["score_weights_constant", str(SCORE_WEIGHTS)],
        ["timeframe_weights_constant", str(TIMEFRAME_WEIGHT)],
    ]
    return rows


def _paint_rank(ws) -> None:
    tier_col = 3
    for row in range(2, ws.max_row + 1):
        tier = ws.cell(row, tier_col).value
        fill = FILL_TIER.get(str(tier))
        if fill is not None:
            ws.cell(row, tier_col).fill = fill
        rank = ws.cell(row, 1).value
        if isinstance(rank, int) and rank <= 20:
            ws.cell(row, 1).fill = FILL_TOP
    if ws.max_row >= 2:
        ws.conditional_formatting.add(
            f"A2:A{ws.max_row}",
            CellIsRule(operator="lessThanOrEqual", formula=["20"], fill=FILL_TOP),
        )


def _paint_rejections(ws) -> None:
    if ws.max_row < 2:
        return
    ws.conditional_formatting.add(
        f"M2:M{ws.max_row}",
        CellIsRule(operator="equal", formula=["FALSE"], fill=FILL_REJECT),
    )


def write_workbook(path: Path, result: AnalysisResult, context: dict[str, Any], readme_text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    summary = wb.create_sheet("Executive Summary")
    _write_table(summary, ["Item", "Value"], _summary_rows(result, context), {}, "T_ExecutiveSummary")

    ranking = wb.create_sheet("Hierarchical Ranking")
    rank_rows = [_group_row(group) for group in result.groups]
    _write_table(ranking, RANK_HEADERS, rank_rows, _formats_for(RANK_HEADERS), "T_HierarchicalRanking")
    _paint_rank(ranking)

    for title, timeframe, table_name in (
        ("1D Levels", "1d", "T_Levels1D"),
        ("4H Levels", "4h", "T_Levels4H"),
        ("1H Levels", "1h", "T_Levels1H"),
    ):
        sheet = wb.create_sheet(title)
        candle_count = len(result.bars.get(timeframe, []))
        level_rows = [
            _level_row(level, candle_count)
            for level in sorted(result.levels.get(timeframe, []), key=lambda item: item.rank)
        ]
        _write_table(sheet, LEVEL_HEADERS, level_rows, _formats_for(LEVEL_HEADERS), table_name)

    mapping = wb.create_sheet("Cross-TF Mapping")
    _write_table(mapping, MAPPING_HEADERS, _mapping_rows(result), _formats_for(MAPPING_HEADERS), "T_CrossTFMapping")
    _paint_rejections(mapping)

    touches = wb.create_sheet("Top Level Touches")
    _write_table(touches, TOUCH_HEADERS, _touch_rows(result), _formats_for(TOUCH_HEADERS), "T_TopLevelTouches")

    monthly = wb.create_sheet("Monthly Persistence")
    _write_table(monthly, MONTH_HEADERS, _monthly_rows(result), _formats_for(MONTH_HEADERS), "T_MonthlyPersistence")

    suppressed = wb.create_sheet("Suppressed Levels")
    _write_table(suppressed, SUPPRESSED_HEADERS, _suppressed_rows(result), _formats_for(SUPPRESSED_HEADERS), "T_SuppressedLevels")

    rejected = wb.create_sheet("Rejected Candidates")
    rejected_rows = _rejected_rows(result)
    _write_table(rejected, REJECTED_HEADERS, rejected_rows, _formats_for(REJECTED_HEADERS), "T_RejectedCandidates")
    if rejected.max_row >= 2:
        shortfall_col = REJECTED_HEADERS.index("Touch Shortfall") + 1
        letter = get_column_letter(shortfall_col)
        rejected.conditional_formatting.add(
            f"{letter}2:{letter}{rejected.max_row}",
            CellIsRule(operator="equal", formula=["1"], fill=FILL_NEAR),
        )

    parameters = wb.create_sheet("Parameters")
    _write_table(parameters, ["Parameter", "Value"], _parameter_rows(), {}, "T_Parameters")

    diagnostics = wb.create_sheet("Diagnostics")
    _write_table(diagnostics, ["Section", "Key", "Value"], context["diagnostic_rows"], {}, "T_Diagnostics")

    readme = wb.create_sheet("README")
    readme_rows = [[line] for line in readme_text.splitlines()] or [[""]]
    _write_table(readme, ["README"], readme_rows, {}, "T_Readme")

    if wb.sheetnames != list(WORKBOOK_SHEETS):
        raise RuntimeError(f"sheet order mismatch: {wb.sheetnames}")
    wb.save(path)


def validate_saved_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if tuple(wb.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError(f"workbook sheets {wb.sheetnames}")
        if wb.vba_archive is not None:
            raise RuntimeError("workbook contains macros")
        diagnostics = wb["Diagnostics"]
        values = {}
        for row in diagnostics.iter_rows(min_row=2, values_only=True):
            if row and row[1] is not None:
                values[str(row[1])] = row[2]
        if str(values.get("duplicate_touch_keys")) != "0":
            raise RuntimeError("duplicate touch keys are not zero")
        if str(values.get("1h_minimum_completed_candle_touches")) != "40":
            raise RuntimeError("1h threshold was not 40")
        if str(values.get("4h_minimum_completed_candle_touches")) != "20":
            raise RuntimeError("4h threshold was not 20")
        if str(values.get("1d_minimum_completed_candle_touches")) != "10":
            raise RuntimeError("1d threshold was not 10")
        links = getattr(wb, "_external_links", []) or []
        for external in links:
            raise RuntimeError(f"external link present: {external}")
        return values
    finally:
        wb.close()
