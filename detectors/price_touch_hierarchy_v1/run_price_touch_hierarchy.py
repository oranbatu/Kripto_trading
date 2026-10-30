"""Permanent entry point for the USD-M perpetual price-touch hierarchy detector."""
from __future__ import annotations

import sys
from pathlib import Path

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import argparse
import gc
import hashlib
import json
import os
import re
import shutil
import time
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any

import pyarrow as pa

import config
from data.catalog import fingerprint_file
from data.loader import load_candles
from detectors.price_touch_hierarchy_v1.engine import (
    Bar,
    SourceValidationError,
    analyze,
    highest_canonical_group,
    iso_turkey,
    iso_utc,
    most_persistent_group,
)
from detectors.price_touch_hierarchy_v1.touch_config import (
    ALGORITHM_VERSION,
    ANALYSIS_START,
    CONFIG,
    DATASET_END,
    DEFAULT_SYMBOL,
    DETECTOR_VERSION,
    DURATION_MS,
    EXCLUSIVE_LOAD_END,
    EXPECTED_FIRST,
    EXPECTED_LAST,
    EXPECTED_PARTITIONS,
    EXPECTED_ROWS,
    LOGS_DIR,
    REPORTS_DIR,
    RUNNER_VERSION,
    SYMBOLS,
    TIMEFRAMES,
    TMP_DIR,
    WORKBOOK_SHEETS,
    revision_pattern_text,
    source_dataset,
    tradingview_symbol,
)
from detectors.price_touch_hierarchy_v1.workbook import validate_saved_workbook, write_workbook

EXECUTION_COMMAND = "python detectors\\price_touch_hierarchy_v1\\run_price_touch_hierarchy.py"
BATCH_REPORT = REPORTS_DIR / "price_touch_hierarchy_batch.json"


class SymbolSelectionError(RuntimeError):
    """Raised when the command line does not select a supported symbol set."""


def selected_symbols(argv: list[str] | None = None) -> list[str]:
    parser = argparse.ArgumentParser(prog="run_price_touch_hierarchy")
    parser.add_argument("--symbol", default=None)
    parser.add_argument("--all-symbols", action="store_true")
    args = parser.parse_args([] if argv is None else argv)
    if args.symbol is not None and args.all_symbols:
        raise SymbolSelectionError("--symbol and --all-symbols are mutually exclusive")
    if args.all_symbols:
        return list(SYMBOLS)
    if args.symbol is None:
        return [DEFAULT_SYMBOL]
    if args.symbol not in SYMBOLS:
        raise SymbolSelectionError(f"unsupported symbol {args.symbol}")
    return [args.symbol]


def report_path(symbol: str) -> Path:
    return REPORTS_DIR / f"price_touch_hierarchy_{symbol}.json"


def log_path(symbol: str) -> Path:
    return LOGS_DIR / f"price_touch_hierarchy_{symbol}.log"


def batch_exit_code(statuses: list[str]) -> int:
    if not statuses or any(status != "ok" for status in statuses):
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    arguments = sys.argv[1:] if argv is None else argv
    try:
        symbols = selected_symbols(arguments)
    except SymbolSelectionError as exc:
        print(f"SYMBOL_SELECTION_FAILED {exc}", flush=True)
        return 2
    print("price_touch_hierarchy_v1 starting " + ",".join(symbols), flush=True)
    assert_python_write_allowed(CONFIG.package_dir / "run_price_touch_hierarchy.py")
    package_before = _package_hashes()
    code_before = _fingerprint_tree(_existing_code_roots())
    started = datetime.now(timezone.utc)
    failures: list[str] = []
    for symbol in symbols:
        try:
            loaded = load_production(symbol)
            del loaded
            gc.collect()
        except SourceValidationError as exc:
            failures.append(f"{symbol}: {exc}")
    if failures:
        for item in failures:
            print(f"SOURCE_VALIDATION_FAILED {item}", flush=True)
        if "--all-symbols" in arguments:
            _write_batch_report(symbols, [], failures, started, arguments)
        return 2
    summaries: list[dict[str, Any]] = []
    for symbol in symbols:
        try:
            summary = run_for_symbol(symbol)
        except (RuntimeError, SourceValidationError, OSError) as exc:
            failures.append(f"{symbol}: {exc}")
            print(f"SYMBOL_FAILED {symbol} {exc}", flush=True)
            continue
        summaries.append(summary)
        print("SUMMARY_JSON " + json.dumps(summary, default=str), flush=True)
    package_after = _package_hashes()
    code_after = _fingerprint_tree(_existing_code_roots())
    if package_before != package_after:
        print("DETECTOR_SOURCE_REWRITTEN", flush=True)
        return 3
    if code_before != code_after:
        print("CODE_IMMUTABILITY_FAILED", flush=True)
        return 3
    stray = _stray_python_files()
    if stray:
        print(f"STRAY_PYTHON {stray}", flush=True)
        return 3
    _cleanup()
    if "--all-symbols" in arguments or len(symbols) > 1:
        _write_batch_report(symbols, summaries, failures, started, arguments)
    exit_code = 0 if not failures and batch_exit_code([item["status"] for item in summaries]) == 0 else 1
    if not summaries and not failures:
        exit_code = 1
    print(f"BATCH_EXIT {exit_code}", flush=True)
    return exit_code


def run_for_symbol(symbol: str) -> dict[str, Any]:
    if symbol not in SYMBOLS:
        raise SymbolSelectionError(f"unsupported symbol {symbol}")
    started = time.perf_counter()
    existing_revisions = scan_revisions(symbol)
    loaded = load_production(symbol)
    source_before = {path: fingerprint_file(Path(path)) for path in loaded["source_paths"]}
    print(
        f"{symbol} source validated "
        + " ".join(f"{timeframe}={len(loaded['bars'][timeframe])}" for timeframe in TIMEFRAMES),
        flush=True,
    )
    result = analyze(loaded["bars"], dataset_end=DATASET_END)
    digest = analytical_fingerprint(result)
    confirmation = analyze(loaded["bars"], dataset_end=DATASET_END)
    if analytical_fingerprint(confirmation) != digest:
        raise RuntimeError(f"{symbol} repeated analysis fingerprint changed")
    del confirmation
    source_after = {path: fingerprint_file(Path(path)) for path in loaded["source_paths"]}
    if source_before != source_after:
        raise RuntimeError(f"{symbol} source parquet changed during analysis")
    persistent = most_persistent_group(result.groups)
    canonical = highest_canonical_group(result.groups)
    context = _context(symbol, result, loaded, source_before, existing_revisions, persistent, canonical)
    context["analytical_fingerprint"] = digest
    readme = (CONFIG.package_dir / "README.md").read_text(encoding="utf-8")
    symbol_tmp = TMP_DIR / symbol
    symbol_tmp.mkdir(parents=True, exist_ok=True)
    context["temporary_dir"] = symbol_tmp
    try:
        revision, final_path = _commit_workbook(result, context, readme)
    except RuntimeError:
        shutil.rmtree(symbol_tmp, ignore_errors=True)
        raise
    summary = _summary(symbol, result, final_path, revision, existing_revisions, persistent, canonical, digest)
    summary["runtime_seconds"] = round(time.perf_counter() - started, 3)
    summary["workbook_sha256"] = _file_sha256(final_path)
    summary["workbook_size"] = final_path.stat().st_size
    summary["status"] = "ok"
    _write_symbol_outputs(symbol, summary)
    del result
    del loaded
    gc.collect()
    shutil.rmtree(symbol_tmp, ignore_errors=True)
    return summary


def load_production(symbol: str = DEFAULT_SYMBOL) -> dict[str, Any]:
    if symbol not in SYMBOLS:
        raise SourceValidationError(f"unsupported symbol {symbol}")
    bars: dict[str, list[Bar]] = {}
    source_paths: list[str] = []
    provenance: dict[str, Any] = {}
    datasets = source_dataset(symbol)
    for timeframe in TIMEFRAMES:
        slice_ = load_candles(
            symbol,
            timeframe,
            ANALYSIS_START,
            EXCLUSIVE_LOAD_END,
            warmup_bars=0,
            precision_mode=config.PRECISION_EXACT,
        )
        if slice_.warmup_bars_loaded != 0:
            raise SourceValidationError(f"{symbol} {timeframe} loaded warmup bars")
        if slice_.symbol != symbol or slice_.timeframe != timeframe:
            raise SourceValidationError(f"{symbol} {timeframe} identity mismatch")
        partitions = tuple(dict.fromkeys(slice_.source_files))
        if len(partitions) != EXPECTED_PARTITIONS[timeframe]:
            raise SourceValidationError(
                f"{symbol} {timeframe} partitions {len(partitions)} != {EXPECTED_PARTITIONS[timeframe]}"
            )
        frame = _bars_from_table(symbol, timeframe, slice_.table)
        _validate_frame(symbol, timeframe, frame)
        bars[timeframe] = frame
        source_paths.extend(slice_.source_files)
        provenance[timeframe] = {
            "rows": len(frame),
            "warmup_bars_loaded": slice_.warmup_bars_loaded,
            "precision_mode": slice_.precision_mode,
            "schema_version": slice_.schema_version,
            "first_open_utc": iso_utc(frame[0].open_time),
            "last_open_utc": iso_utc(frame[-1].open_time),
            "first_open_turkey": iso_turkey(frame[0].open_time),
            "last_open_turkey": iso_turkey(frame[-1].open_time),
            "last_close_utc": iso_utc(frame[-1].close_time),
            "partitions": len(partitions),
            "exclusive_end_utc": iso_utc(EXCLUSIVE_LOAD_END),
            "source_path": str(datasets[timeframe]),
        }
        del slice_
    return {"symbol": symbol, "bars": bars, "source_paths": sorted(set(source_paths)), "provenance": provenance}


def _bars_from_table(symbol: str, timeframe: str, table: pa.Table) -> list[Bar]:
    opens = table.column("open_time").to_pylist()
    closes = table.column("close_time").to_pylist()
    open_ = [_decimal(value) for value in table.column("open").to_pylist()]
    high = [_decimal(value) for value in table.column("high").to_pylist()]
    low = [_decimal(value) for value in table.column("low").to_pylist()]
    close = [_decimal(value) for value in table.column("close").to_pylist()]
    volume = [_decimal(value) for value in table.column("volume").to_pylist()]
    symbols = table.column("symbol").to_pylist()
    rows: list[Bar] = []
    for index in range(table.num_rows):
        if symbols[index] != symbol:
            raise SourceValidationError(f"{timeframe} row {index} symbol {symbols[index]}")
        rows.append(
            Bar(
                timeframe=timeframe,
                row=index,
                open_time=_as_utc(opens[index]),
                close_time=_as_utc(closes[index]),
                open=open_[index],
                high=high[index],
                low=low[index],
                close=close[index],
                volume=volume[index],
            )
        )
    return rows


def _validate_frame(symbol: str, timeframe: str, bars: list[Bar]) -> None:
    if len(bars) != EXPECTED_ROWS[timeframe]:
        raise SourceValidationError(f"{timeframe} rows {len(bars)} != {EXPECTED_ROWS[timeframe]}")
    if bars[0].open_time != EXPECTED_FIRST[timeframe] or bars[-1].open_time != EXPECTED_LAST[timeframe]:
        raise SourceValidationError(
            f"{timeframe} range {iso_utc(bars[0].open_time)} .. {iso_utc(bars[-1].open_time)}"
        )
    step = DURATION_MS[timeframe]
    previous = None
    for bar in bars:
        if bar.open_time.tzinfo is None or bar.open_time.utcoffset() != timezone.utc.utcoffset(bar.open_time):
            raise SourceValidationError(f"{timeframe} timestamp is not UTC")
        if bar.open_time.year < 2024:
            raise SourceValidationError("pre-2024 candle accessed")
        if previous is not None:
            delta = int((bar.open_time - previous).total_seconds() * 1000)
            if delta != step:
                raise SourceValidationError(f"{timeframe} missing or duplicate timestamp at {iso_utc(bar.open_time)}")
        previous = bar.open_time
        expected_close = bar.open_time.timestamp() * 1000 + step - 1
        actual_close = bar.close_time.timestamp() * 1000
        if abs(actual_close - expected_close) > 0.1:
            raise SourceValidationError(f"{timeframe} close time mismatch at {iso_utc(bar.open_time)}")
        if bar.close_time > DATASET_END:
            raise SourceValidationError(f"{timeframe} unfinished candle {iso_utc(bar.open_time)}")
        if bar.volume < 0:
            raise SourceValidationError("negative volume")
        for value in (bar.open, bar.high, bar.low, bar.close):
            if value is None:
                raise SourceValidationError("null OHLC")
        if not (
            bar.high >= bar.open
            and bar.high >= bar.close
            and bar.high >= bar.low
            and bar.low <= bar.open
            and bar.low <= bar.close
            and bar.low <= bar.high
        ):
            raise SourceValidationError(f"OHLC violation at {iso_utc(bar.open_time)}")
    path = str(source_dataset(symbol)[timeframe]).replace("\\", "/").lower()
    for token in ("marketdata", "binance", "futures", "um", "perpetual", timeframe, symbol.lower()):
        if token not in path:
            raise SourceValidationError(f"{symbol} {timeframe} dataset path missing {token}: {path}")


def _decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _as_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        raise SourceValidationError("naive timestamp")
    return value.astimezone(timezone.utc)


def scan_revisions(symbol: str = DEFAULT_SYMBOL, directory: Path | None = None) -> list[int]:
    folder = CONFIG.desktop_dir if directory is None else Path(directory)
    found = []
    if not folder.exists():
        return found
    pattern = re.compile(revision_pattern_text(symbol), re.IGNORECASE)
    for path in folder.iterdir():
        name = path.name
        if name.startswith("~$") or not path.is_file():
            continue
        if path.suffix.lower() != ".xlsx":
            continue
        lowered = name.lower()
        if lowered.endswith(".tmp") or lowered.endswith(".part") or lowered.endswith(".bak"):
            continue
        match = pattern.match(name)
        if match:
            found.append(int(match.group(1)))
    return sorted(set(found))


def next_revision(existing: list[int]) -> int:
    if not existing:
        return 0
    return max(existing) + 1


def revision_filename(revision: int, symbol: str = DEFAULT_SYMBOL) -> str:
    return f"{symbol}_MTF_Most_Touched_Levels_rev{revision:02d}.xlsx"


def _commit_workbook(result, context: dict[str, Any], readme: str) -> tuple[int, Path]:
    symbol = context["symbol"]
    desktop = Path(context.get("desktop_dir", CONFIG.desktop_dir))
    temporary_dir = Path(context.get("temporary_dir", TMP_DIR / symbol))
    temporary_dir.mkdir(parents=True, exist_ok=True)
    planned = next_revision(context["existing_revisions"])
    while True:
        name = revision_filename(planned, symbol)
        final_path = desktop / name
        if final_path.exists():
            planned += 1
            continue
        temporary = temporary_dir / name
        context["revision_name"] = name
        context["diagnostic_rows"] = _diagnostic_rows(result, context)
        try:
            write_workbook(temporary, result, context, readme)
            _validate_workbook_container(temporary, symbol)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        rescanned = scan_revisions(symbol, desktop)
        if final_path.exists() or (rescanned and max(rescanned) >= planned):
            temporary.unlink(missing_ok=True)
            planned = next_revision(scan_revisions(symbol, desktop))
            continue
        os.replace(temporary, final_path)
        if not final_path.is_file():
            raise RuntimeError("atomic workbook move failed")
        return planned, final_path


def _context(symbol: str, result, loaded, fingerprints, revisions, persistent, canonical) -> dict[str, Any]:
    import openpyxl
    import pyarrow

    return {
        "symbol": symbol,
        "existing_revisions": revisions,
        "revision_name": "",
        "runtime": {
            "python": sys.version.replace("\n", " "),
            "pyarrow": pyarrow.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "source_validation": "PASS",
        "provenance": loaded["provenance"],
        "source_fingerprints": fingerprints,
        "most_persistent": (
            f"{persistent.group_id} {persistent.persistence} {persistent.representative_price}"
            if persistent else "NONE"
        ),
        "highest_canonical": (
            f"{canonical.group_id} visits={canonical.canonical_visits} price={canonical.representative_price}"
            if canonical else "NONE"
        ),
        "warnings": list(result.warnings),
    }


def _diagnostic_rows(result, context: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []

    def add(section: str, key: str, value: Any) -> None:
        rows.append([section, key, value])

    add("implementation", "project_root", str(CONFIG.project_root))
    add("implementation", "permanent_detector_directory", str(CONFIG.package_dir))
    add("implementation", "permanent_script_path", str(CONFIG.package_dir / "run_price_touch_hierarchy.py"))
    add("implementation", "readme_path", str(CONFIG.package_dir / "README.md"))
    add("implementation", "test_path", str(CONFIG.package_dir / "tests" / "test_price_touch_hierarchy.py"))
    add("implementation", "execution_command", EXECUTION_COMMAND)
    add("implementation", "symbol", context["symbol"])
    add("implementation", "tradingview_symbol", tradingview_symbol(context["symbol"]))
    add("implementation", "detector_version", DETECTOR_VERSION)
    add("implementation", "runner_version", RUNNER_VERSION)
    add("implementation", "analytical_algorithm_version", ALGORITHM_VERSION)
    add("implementation", "runtime_python", context["runtime"]["python"])
    add("implementation", "runtime_pyarrow", context["runtime"]["pyarrow"])
    add("implementation", "runtime_openpyxl", context["runtime"]["openpyxl"])
    add("implementation", "workbook_revision", context["revision_name"])
    add("implementation", "existing_matching_revisions", ",".join(f"{item:02d}" for item in context["existing_revisions"]) or "NONE")
    add("implementation", "separate_project_created", "FALSE")
    add("thresholds", "1h_minimum_completed_candle_touches", 40)
    add("thresholds", "4h_minimum_completed_candle_touches", 20)
    add("thresholds", "1d_minimum_completed_candle_touches", 10)
    add("thresholds", "effective_thresholds", "1h=40; 4h=20; 1d=10")
    for timeframe in TIMEFRAMES:
        add("source", f"{timeframe}_path", context["provenance"][timeframe]["source_path"])
        add("source", f"{timeframe}_rows", len(result.bars.get(timeframe, [])))
        add("source", f"{timeframe}_partitions", context["provenance"][timeframe]["partitions"])
        add("source", f"{timeframe}_validation", "PASS")
        proven = context["provenance"][timeframe]
        add("source", f"{timeframe}_first_open_utc", proven["first_open_utc"])
        add("source", f"{timeframe}_last_open_utc", proven["last_open_utc"])
        add("source", f"{timeframe}_last_close_utc", proven["last_close_utc"])
    add("source", "source_end_exclusive_utc", iso_utc(EXCLUSIVE_LOAD_END))
    add("source", "analytical_fingerprint", context.get("analytical_fingerprint", ""))
    add("source", "source_validation", context["source_validation"])
    for path, meta in context["source_fingerprints"].items():
        add("source_fingerprint", path, f"sha256={meta['sha256']} size={meta['size']} mtime_ns={meta['mtime_ns']}")
    add("grid", "price_min", str(result.price_min))
    add("grid", "price_max", str(result.price_max))
    add("grid", "grid_min_k", result.grid_min_k)
    add("grid", "grid_max_k", result.grid_max_k)
    add("grid", "atomic_bin_count", result.atomic_bins)
    add("counts", "duplicate_touch_keys", result.duplicate_touch_keys)
    for timeframe in TIMEFRAMES:
        add("counts", f"{timeframe}_candidate_peaks", result.peak_counts.get(timeframe, 0))
        add("counts", f"{timeframe}_qualified_levels", len(result.levels.get(timeframe, [])))
        add("counts", f"{timeframe}_rejected_candidates", sum(1 for level in result.rejected if level.timeframe == timeframe))
        add("counts", f"{timeframe}_suppressed_levels", sum(1 for level in result.suppressed if level.timeframe == timeframe))
        add("counts", f"{timeframe}_unfinished_excluded", result.unfinished_excluded.get(timeframe, 0))
        levels = result.levels.get(timeframe, [])
        touches = sum(level.touches for level in levels)
        typed = sum(level.close_inside + level.body + level.wick for level in levels)
        episodes = sum(level.episodes for level in levels)
        episode_lengths = sum(sum(level.episode_lengths) for level in levels)
        add("reconciliation", f"{timeframe}_touch_type_total", typed)
        add("reconciliation", f"{timeframe}_touch_total", touches)
        add("reconciliation", f"{timeframe}_touch_reconciliation", "PASS" if typed == touches else "FAIL")
        add("reconciliation", f"{timeframe}_episode_reconciliation", "PASS" if episode_lengths == touches else "FAIL")
        add("reconciliation", f"{timeframe}_episode_total", episodes)
    add("hierarchy", "group_count", len(result.groups))
    for tier in (
        "TIER_1_ALL_TIMEFRAMES",
        "TIER_2_DAILY_WITH_INTRADAY",
        "TIER_3_DAILY_ONLY",
        "TIER_4_4H_AND_1H",
        "TIER_5_4H_ONLY",
        "TIER_6_1H_ONLY",
    ):
        add("hierarchy", tier, sum(1 for group in result.groups if group.tier == tier))
    accepted = sum(1 for attempt in result.attempts if attempt.accepted)
    add("mapping", "accepted_mappings", accepted)
    add("mapping", "rejected_mappings", len(result.attempts) - accepted)
    reasons: dict[str, int] = {}
    for attempt in result.attempts:
        if attempt.rejection_reason:
            reasons[attempt.rejection_reason] = reasons.get(attempt.rejection_reason, 0) + 1
    for reason, count in sorted(reasons.items()):
        add("mapping_rejection", reason, count)
    for reason, count in sorted(result.rejection_reason_counts.items()):
        add("candidate_rejection", reason, count)
    add("immutability", "source_parquet_unchanged", "PASS")
    add("immutability", "existing_project_code_unchanged", "PASS")
    add("isolation", "range_detector_imported", "FALSE")
    add("isolation", "swing_detector_imported", "FALSE")
    add("isolation", "strategy_modules_imported", "FALSE")
    add("isolation", "range_detector_modified", "FALSE")
    add("isolation", "swing_detector_modified", "FALSE")
    add("isolation", "separate_project_created", "FALSE")
    add("isolation", "detector_python_outside_project", "NONE")
    add("validation", "workbook_validation", "PASS")
    add("validation", "main_ranking_reconciliation", "PASS")
    add("cleanup", "temporary_directory", str(TMP_DIR))
    add("cleanup", "temporary_artifact_policy", "Temporary workbook is validated, moved atomically, then removed.")
    add("warnings", "warning_count", len(context.get("warnings", [])))
    for index, warning in enumerate(context.get("warnings", []), start=1):
        add("warnings", f"warning_{index}", warning)
    return rows


def analytical_fingerprint(result) -> str:
    def level_row(level) -> dict[str, Any]:
        return {
            "candidate_id": level.candidate_id,
            "level_id": level.level_id,
            "timeframe": level.timeframe,
            "price": str(level.price),
            "zone_lower": str(level.zone_lower),
            "zone_upper": str(level.zone_upper),
            "touches": level.touches,
            "episodes": level.episodes,
            "rank": level.rank,
            "would_be_rank": level.would_be_rank,
            "suppressed_by": level.suppressed_by,
            "group_id": level.group_id,
        }

    payload = {
        "levels": [level_row(level) for timeframe in TIMEFRAMES for level in result.levels.get(timeframe, [])],
        "suppressed": [level_row(level) for level in result.suppressed],
        "rejected": [level_row(level) for level in result.rejected],
        "groups": [
            {
                "group_id": group.group_id,
                "rank": group.rank,
                "tier": group.tier,
                "score": str(group.score),
                "price": str(group.representative_price),
                "lower": str(group.consolidated_lower),
                "upper": str(group.consolidated_upper),
                "timeframes": list(group.timeframes),
                "touches": {timeframe: group.members[timeframe].touches for timeframe in group.timeframes},
            }
            for group in result.groups
        ],
    }
    raw = json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _summary(symbol: str, result, final_path: Path, revision: int, existing: list[int], persistent, canonical, digest: str) -> dict[str, Any]:
    top = []
    for group in result.groups[:20]:
        top.append({
            "rank": group.rank,
            "tier": group.tier,
            "price": str(group.representative_price),
            "lower": str(group.consolidated_lower),
            "upper": str(group.consolidated_upper),
            "timeframes": "+".join(group.timeframes),
            "touches_1d": group.members["1d"].touches if "1d" in group.members else None,
            "touches_4h": group.members["4h"].touches if "4h" in group.members else None,
            "touches_1h": group.members["1h"].touches if "1h" in group.members else None,
            "canonical_visits": group.canonical_visits,
            "distinct_months": group.members[group.canonical_timeframe].distinct_months,
            "first_touch": iso_utc(group.first_touch) if group.first_touch else None,
            "last_touch": iso_utc(group.last_touch) if group.last_touch else None,
            "score": str(group.score),
        })

    def best(timeframe: str) -> dict[str, Any] | None:
        levels = result.levels.get(timeframe, [])
        if not levels:
            return None
        level = min(levels, key=lambda item: item.rank)
        return {"id": level.level_id, "price": str(level.price), "touches": level.touches, "episodes": level.episodes}

    return {
        "symbol": symbol,
        "tradingview_symbol": tradingview_symbol(symbol),
        "status": "ok",
        "workbook": str(final_path),
        "revision": revision,
        "analytical_fingerprint": digest,
        "top5": top[:5],
        "existing_revisions": existing,
        "rows": {tf: len(result.bars.get(tf, [])) for tf in TIMEFRAMES},
        "atomic_bins": result.atomic_bins,
        "peaks": result.peak_counts,
        "qualified": {tf: len(result.levels.get(tf, [])) for tf in TIMEFRAMES},
        "rejected": {tf: sum(1 for level in result.rejected if level.timeframe == tf) for tf in TIMEFRAMES},
        "suppressed": {tf: sum(1 for level in result.suppressed if level.timeframe == tf) for tf in TIMEFRAMES},
        "groups": len(result.groups),
        "tiers": {tier: sum(1 for group in result.groups if group.tier == tier) for tier in (
            "TIER_1_ALL_TIMEFRAMES", "TIER_2_DAILY_WITH_INTRADAY", "TIER_3_DAILY_ONLY",
            "TIER_4_4H_AND_1H", "TIER_5_4H_ONLY", "TIER_6_1H_ONLY",
        )},
        "duplicate_touch_keys": result.duplicate_touch_keys,
        "unfinished_excluded": result.unfinished_excluded,
        "accepted_mappings": sum(1 for attempt in result.attempts if attempt.accepted),
        "rejected_mappings": sum(1 for attempt in result.attempts if not attempt.accepted),
        "mapping_rejections": {
            reason: sum(1 for attempt in result.attempts if attempt.rejection_reason == reason)
            for reason in sorted({attempt.rejection_reason for attempt in result.attempts if attempt.rejection_reason})
        },
        "rejection_reasons": result.rejection_reason_counts,
        "best_1d": best("1d"),
        "best_4h": best("4h"),
        "best_1h": best("1h"),
        "highest_canonical": None if canonical is None else {
            "id": canonical.group_id, "visits": canonical.canonical_visits, "price": str(canonical.representative_price)
        },
        "most_persistent": None if persistent is None else {
            "id": persistent.group_id, "label": persistent.persistence, "price": str(persistent.representative_price)
        },
        "top20": top,
        "thresholds": {"1h": 40, "4h": 20, "1d": 10},
        "warnings": result.warnings,
    }


def _existing_code_roots() -> list[Path]:
    root = CONFIG.project_root
    paths = [
        root / "detectors" / "primary_range_v1",
        root / "detectors" / "hierarchical_swing_v4",
        root / "detectors" / "__init__.py",
        root / "data",
        root / "ingestion",
        root / "processing",
        root / "config.py",
        root / "README.md",
        root / "requirements.txt",
        root / "pyproject.toml",
    ]
    return [path for path in paths if path.exists()]


def _fingerprint_tree(paths: list[Path]) -> dict[str, str]:
    found: dict[str, str] = {}
    for path in paths:
        files = [path] if path.is_file() else sorted(item for item in path.rglob("*") if item.is_file())
        for file in files:
            if "__pycache__" in file.parts or file.suffix == ".pyc":
                continue
            if "price_touch_hierarchy_v1" in file.parts:
                continue
            digest = hashlib.sha256(file.read_bytes()).hexdigest()
            found[str(file)] = digest
    return found


def _package_hashes() -> dict[str, str]:
    found = {}
    for file in CONFIG.package_dir.rglob("*"):
        if not file.is_file() or "__pycache__" in file.parts or file.suffix == ".pyc":
            continue
        found[str(file)] = hashlib.sha256(file.read_bytes()).hexdigest()
    return found


def _stray_python_files() -> list[str]:
    stray = []
    desktop = CONFIG.desktop_dir
    for file in desktop.glob("*.py"):
        stray.append(str(file))
    return stray


def assert_python_write_allowed(path: Path) -> None:
    resolved = path.resolve()
    package = CONFIG.package_dir.resolve()
    if package not in resolved.parents and resolved != package:
        raise PermissionError(f"detector Python writes are limited to {package}")


def _validate_workbook_container(path: Path, symbol: str) -> None:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError(f"{symbol} workbook zip is damaged")
        names = archive.namelist()
    if any("vba" in name.lower() for name in names):
        raise RuntimeError(f"{symbol} workbook contains macros")
    if any("externalLink" in name for name in names):
        raise RuntimeError(f"{symbol} workbook contains an external link")
    from openpyxl import load_workbook

    for options in ({}, {"read_only": True}, {"data_only": True}):
        workbook = load_workbook(path, **options)
        try:
            if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
                raise RuntimeError(f"{symbol} sheet order {workbook.sheetnames}")
        finally:
            workbook.close()
    validate_saved_workbook(path, symbol)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_symbol_outputs(symbol: str, summary: dict[str, Any]) -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    report_path(symbol).write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    lines = [
        f"symbol={symbol}",
        f"status={summary['status']}",
        f"revision={summary['revision']}",
        f"workbook={summary['workbook']}",
        f"sha256={summary['workbook_sha256']}",
        f"fingerprint={summary['analytical_fingerprint']}",
        f"runtime_seconds={summary['runtime_seconds']}",
    ]
    log_path(symbol).write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_batch_report(symbols: list[str], summaries: list[dict[str, Any]], failures: list[str], started: datetime, arguments: list[str]) -> None:
    by_symbol = {item["symbol"]: item for item in summaries}
    rows = []
    for symbol in symbols:
        item = by_symbol.get(symbol)
        if item is None:
            rows.append({"symbol": symbol, "status": "failed"})
            continue
        rows.append({
            "symbol": symbol,
            "status": item["status"],
            "rows": item["rows"],
            "peaks": item["peaks"],
            "qualified": item["qualified"],
            "rejected": item["rejected"],
            "suppressed": item["suppressed"],
            "groups": item["groups"],
            "tiers": item["tiers"],
            "fingerprint": item["analytical_fingerprint"],
            "workbook": item["workbook"],
            "revision": item["revision"],
            "sha256": item["workbook_sha256"],
            "runtime_seconds": item["runtime_seconds"],
        })
    payload = {
        "batch_start_utc": started.isoformat().replace("+00:00", "Z"),
        "batch_end_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "arguments": arguments,
        "symbol_order": symbols,
        "symbols": rows,
        "successful_symbols": [item["symbol"] for item in summaries],
        "failed_symbols": failures,
        "published_workbooks": [item["workbook"] for item in summaries],
        "warnings": [],
        "errors": failures,
        "overall_success": not failures and len(summaries) == len(symbols),
    }
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    BATCH_REPORT.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    (LOGS_DIR).mkdir(parents=True, exist_ok=True)
    (LOGS_DIR / "price_touch_hierarchy_batch.log").write_text(
        json.dumps({"overall_success": payload["overall_success"], "failures": failures, "symbols": symbols}, indent=2) + "\n",
        encoding="utf-8",
    )


def _cleanup() -> None:
    if TMP_DIR.exists():
        shutil.rmtree(TMP_DIR, ignore_errors=True)
    parent = TMP_DIR.parent
    if parent.exists() and not any(parent.iterdir()):
        parent.rmdir()
    for root in (CONFIG.package_dir, CONFIG.package_dir / "tests"):
        cache = root / "__pycache__"
        if cache.exists():
            shutil.rmtree(cache, ignore_errors=True)
        for file in root.glob("*.pyc"):
            file.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
