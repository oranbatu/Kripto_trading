"""Excel audit workbook for hierarchical swing detection. No macros and no external links."""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.hierarchical_swing_v4.engine import DetectionResult, iso_turkey, iso_utc
from detectors.hierarchical_swing_v4.swing_config import CONFIG, WORKBOOK_SHEETS

PRICE_FMT = "0.00000000"
PCT_FMT = "0.00%"
INT_FMT = "#,##0"
HEADER_FILL = PatternFill("solid", fgColor="1F2933")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
TAB_COLORS = {
    "Swing Summary": "334E68",
    "Major Swing Points": "1F7A4D",
    "Major Swing Highs": "1F7A4D",
    "Major Swing Lows": "1F7A4D",
    "Major Swing Legs": "1F7A4D",
    "Internal Swings": "C56A1A",
    "Raw Pivot Candidates": "6B4C9A",
    "Raw Formations": "2F6FED",
    "Formation Candles": "2F6FED",
    "State Transitions": "7A4E2D",
    "Monthly Diagnostics": "3D5A40",
    "Forward Evaluation": "1F7A4D",
    "Parameters": "24527A",
    "Diagnostics": "243B53",
    "README": "243B53",
}
FILL_BOOTSTRAP = PatternFill("solid", fgColor="FFE08A")
FILL_REPLACEMENT = PatternFill("solid", fgColor="F6C945")
FILL_UNRESOLVED = PatternFill("solid", fgColor="F4B6C2")
FILL_LONG = PatternFill("solid", fgColor="F8D48A")
FILL_ERROR = PatternFill("solid", fgColor="F5A3A3")


def _num(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return value
    return value


def _pct(value: Any) -> Any:
    if value is None or value == "":
        return None
    return float(Decimal(value) / Decimal(100))


def _txt(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    text = str(value)
    if len(text) > 32000:
        return text[:31900] + "...[truncated]"
    return text


def _ohlc_parts(ohlc: Any) -> tuple[Any, Any, Any, Any]:
    if not ohlc:
        return (None, None, None, None)
    return tuple(_num(part) for part in ohlc)


def _write_table(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]], formats: dict[str, str]) -> None:
    if rows and len(rows[0]) != len(headers):
        raise RuntimeError(f"{ws.title} has {len(rows[0])} values for {len(headers)} headers")
    for row in rows:
        if len(row) != len(headers):
            raise RuntimeError(f"{ws.title} row width {len(row)} != {len(headers)}")
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        width = min(42, max(12, len(header) + 2))
        ws.column_dimensions[letter].width = width
        number_format = formats.get(header)
        if not number_format:
            continue
        for cell in ws.iter_rows(min_row=2, min_col=index, max_col=index, max_row=ws.max_row):
            cell[0].number_format = number_format
    if ws.max_row >= 2 and ws.max_column >= 1:
        name = "T_" + "".join(ch for ch in ws.title if ch.isalnum())
        table = Table(displayName=name[:40], ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = TAB_COLORS.get(ws.title, "334E68")
    ws.oddHeader.left.text = ws.title
    ws.oddFooter.right.text = CONFIG.detector_version


def _paint(ws, row_index: int, fill: PatternFill | None) -> None:
    if fill is None:
        return
    for cell in ws[row_index]:
        cell.fill = fill


def _major_row(swing) -> list[Any]:
    start_ohlc = _ohlc_parts(getattr(swing, "leg_start_ohlc", None))
    end_ohlc = _ohlc_parts(getattr(swing, "leg_end_ohlc", None))
    previous = swing.previous
    same_side = getattr(swing, "structure_label", "")
    return [
        swing.major_swing_id,
        swing.sequence,
        "MAJOR_SWING",
        "TRUE",
        swing.type,
        swing.structure_label,
        swing.formation_id,
        swing.pivot_id,
        swing.internal_swing_id,
        iso_utc(swing.extreme_time),
        iso_turkey(swing.extreme_time),
        iso_utc(swing.formation_confirmed_at),
        iso_turkey(swing.formation_confirmed_at),
        iso_utc(swing.internal_confirmed_at),
        iso_turkey(swing.internal_confirmed_at),
        iso_utc(swing.activated_at),
        iso_turkey(swing.activated_at),
        iso_utc(swing.reversal_confirmed_at),
        iso_turkey(swing.reversal_confirmed_at),
        iso_utc(swing.confirmed_at),
        iso_turkey(swing.confirmed_at),
        swing.age,
        _txt(swing.replacement_history),
        _num(swing.price),
        _num(swing.atr),
        getattr(previous, "major_swing_id", getattr(previous, "internal_swing_id", "")),
        same_side,
        _num(swing.displacement_price),
        _num(swing.displacement_atr),
        _pct(swing.displacement_percent),
        "TRUE" if swing.displacement_atr_passed else "FALSE",
        "TRUE" if swing.displacement_percent_passed else "FALSE",
        _num(swing.reversal_price),
        _num(swing.reversal_atr),
        _pct(swing.reversal_percent),
        "TRUE" if swing.reversal_atr_passed else "FALSE",
        "TRUE" if swing.reversal_percent_passed else "FALSE",
        swing.spacing_bars,
        "TRUE" if swing.spacing_passed else "FALSE",
        _num(swing.prominence_atr),
        "TRUE" if swing.prominence_passed else "FALSE",
        _num(getattr(swing, "absolute_amplitude", None)),
        _pct(getattr(swing, "amplitude_percent_from_start", None)),
        _pct(getattr(swing, "peak_to_trough_percent", None)),
        _pct(getattr(swing, "trough_to_peak_percent", None)),
        _pct(getattr(swing, "symmetric_width_percent", None)),
        _num(getattr(swing, "amplitude_atr_median", None)),
        getattr(swing, "duration_bars", None),
        _num(getattr(swing, "duration_hours", None)),
        _num(getattr(swing, "duration_days", None)),
        getattr(swing, "confirmation_to_confirmation_bars", None),
        _num(getattr(swing, "path_length", None)),
        _num(getattr(swing, "net_close_change", None)),
        _num(getattr(swing, "leg_efficiency", None)),
        getattr(swing, "bullish_candles", None),
        getattr(swing, "bearish_candles", None),
        getattr(swing, "unchanged_candles", None),
        _num(getattr(swing, "internal_pullback", None)),
        _num(getattr(swing, "realized_volatility", None)),
        getattr(swing, "raw_formations_inside_leg", None),
        getattr(swing, "raw_pivots_inside_leg", None),
        getattr(swing, "internal_swings_inside_leg", None),
        _num(getattr(swing, "extreme_body", None)),
        _num(getattr(swing, "extreme_upper_wick", None)),
        _num(getattr(swing, "extreme_lower_wick", None)),
        _num(getattr(swing, "close_location", None)),
        _num(getattr(swing, "base_volume", None)),
        _num(getattr(swing, "quote_volume", None)),
        getattr(swing, "trade_count", None),
        _num(getattr(swing, "taker_buy_ratio", None)),
        _num(getattr(swing, "pivot_volume_to_trailing_50_mean", None)),
        _num(getattr(swing, "total_leg_volume", None)),
        _num(getattr(swing, "vwap", None)),
        _num(getattr(swing, "significance_score", None)),
        getattr(swing, "significance_label", ""),
        "TRUE" if getattr(swing, "forward_data_censored", False) else "FALSE",
        getattr(swing, "forward_censoring_reason", ""),
        iso_utc(getattr(swing, "leg_start_open_time", None)),
        iso_utc(getattr(swing, "leg_start_close_time", None)),
        iso_turkey(getattr(swing, "leg_start_open_time", None)),
        iso_turkey(getattr(swing, "leg_start_close_time", None)),
        iso_utc(getattr(swing, "leg_end_open_time", None)),
        iso_utc(getattr(swing, "leg_end_close_time", None)),
        iso_turkey(getattr(swing, "leg_end_open_time", None)),
        iso_turkey(getattr(swing, "leg_end_close_time", None)),
        *start_ohlc,
        *end_ohlc,
        getattr(swing, "leg_start_class", ""),
        getattr(swing, "metric_class_geometry", "CAUSAL_AT_MAJOR_CONFIRMATION"),
        "EX_POST_FORWARD_METRIC",
        CONFIG.detector_version,
    ]


MAJOR_HEADERS = [
    "Major Swing ID", "Sequence", "Hierarchy Level", "Inside Requested Analysis Window",
    "Type", "Structure Label", "Formation ID", "Raw Pivot ID", "Internal Swing ID",
    "Extreme UTC", "Extreme Turkey", "Formation Confirmation UTC", "Formation Confirmation Turkey",
    "Internal Confirmation UTC", "Internal Confirmation Turkey",
    "Major Candidate Activation UTC", "Major Candidate Activation Turkey",
    "Major Reversal UTC", "Major Reversal Turkey", "Major Confirmation UTC", "Major Confirmation Turkey",
    "Candidate Age", "Replacement History", "Price", "ATR",
    "Previous Opposite Major Or Seed", "Structure Label Repeated",
    "Displacement Price", "Displacement ATR", "Displacement Percent",
    "Displacement ATR Pass", "Displacement Percent Pass",
    "Reversal Price", "Reversal ATR", "Reversal Percent",
    "Reversal ATR Pass", "Reversal Percent Pass",
    "Spacing Bars", "Spacing Pass", "Prominence ATR", "Prominence Pass",
    "Absolute Amplitude", "Amplitude Percent From Start", "Peak To Trough Percent",
    "Trough To Peak Percent", "Symmetric Width Percent", "Amplitude ATR Median",
    "Duration Bars", "Duration Hours", "Duration Days", "Confirmation To Confirmation Bars",
    "Path Length", "Net Close Change", "Leg Efficiency",
    "Bullish Candles", "Bearish Candles", "Unchanged Candles", "Internal Pullback",
    "Realized Volatility", "Raw Formations Inside Leg", "Raw Pivots Inside Leg",
    "Internal Swings Inside Leg", "Extreme Body", "Extreme Upper Wick", "Extreme Lower Wick",
    "Close Location", "Base Volume", "Quote Volume", "Trade Count", "Taker Buy Ratio",
    "Trailing Volume Ratio", "Total Leg Volume", "VWAP",
    "Significance Score", "Significance Label", "Forward Data Censored", "Forward Censoring Reason",
    "Start Candle Open Time UTC", "Start Candle Close Time UTC",
    "Start Candle Open Time Turkey", "Start Candle Close Time Turkey",
    "End Candle Open Time UTC", "End Candle Close Time UTC",
    "End Candle Open Time Turkey", "End Candle Close Time Turkey",
    "Start Open", "Start High", "Start Low", "Start Close",
    "End Open", "End High", "End Low", "End Close",
    "Leg Start Class", "Geometry Metric Class", "Forward Metric Class", "Detector Version",
]

PRICE_HEADERS = {
    "Price", "ATR", "Displacement Price", "Reversal Price", "Absolute Amplitude",
    "Path Length", "Net Close Change", "Internal Pullback", "Extreme Body",
    "Extreme Upper Wick", "Extreme Lower Wick", "Base Volume", "Quote Volume",
    "Total Leg Volume", "VWAP", "Start Open", "Start High", "Start Low", "Start Close",
    "End Open", "End High", "End Low", "End Close", "Open", "High", "Low", "Close",
    "Volume", "Quote Asset Volume", "Reference Price", "Extreme Price", "Outbound Price",
    "Return Price", "Touch Price", "Candidate Price", "Last Close", "Missing Reversal Price",
    "Candidate ATR", "Required Reversal", "Left Prominence", "Right Prominence",
    "Two Sided Prominence", "Swing Open", "Swing High", "Swing Low", "Swing Close",
    "Extreme Open", "Extreme High", "Extreme Low", "Extreme Close",
    "Close Open", "Close High", "Close Low", "Close Close",
}
PCT_HEADERS = {
    "Displacement Percent", "Reversal Percent", "Amplitude Percent From Start",
    "Peak To Trough Percent", "Trough To Peak Percent", "Symmetric Width Percent",
    "Outbound Percent", "Return Percent", "Return Percent Reference Basis",
    "Percentage Prominence", "Missing Reversal Percent", "Taker Buy Ratio",
    "Close Location", "Leg Efficiency",
}


def _formats(headers: Sequence[str]) -> dict[str, str]:
    formats = {}
    for header in headers:
        if header in PRICE_HEADERS or header.endswith(" Price") or header.endswith(" ATR"):
            formats[header] = PRICE_FMT
        elif header in PCT_HEADERS or header.endswith(" Percent"):
            formats[header] = PCT_FMT
        elif header in {"Sequence", "Candidate Age", "Spacing Bars", "Duration Bars", "Trade Count"}:
            formats[header] = INT_FMT
    return formats


def _internal_rows(result: DetectionResult) -> tuple[list[str], list[list[Any]], list[PatternFill | None]]:
    headers = [
        "Record Class", "Internal Swing ID", "Sequence", "Type", "Structure Label",
        "Formation ID", "Raw Pivot ID", "Extreme UTC", "Extreme Turkey",
        "Raw Pivot Confirmation UTC", "Raw Pivot Confirmation Turkey",
        "Candidate Activation UTC", "Candidate Activation Turkey",
        "Reversal Confirmation UTC", "Reversal Confirmation Turkey",
        "Internal Confirmation UTC", "Internal Confirmation Turkey",
        "Candidate Age", "Replacement History", "Price", "ATR",
        "Displacement Price", "Displacement ATR", "Displacement Pass",
        "Spacing Bars", "Spacing Pass", "Reversal Price", "Reversal ATR", "Reversal Pass",
        "Incoming Leg Start Class",
        "Start Candle Open Time UTC", "Start Candle Close Time UTC",
        "Start Candle Open Time Turkey", "Start Candle Close Time Turkey",
        "End Candle Open Time UTC", "End Candle Close Time UTC",
        "End Candle Open Time Turkey", "End Candle Close Time Turkey",
        "Start Open", "Start High", "Start Low", "Start Close",
        "End Open", "End High", "End Low", "End Close",
        "Left Prominence", "Right Prominence", "Two Sided Prominence", "Prominence ATR",
        "Absolute Amplitude", "Duration Bars", "Leg Efficiency", "Path Length",
        "Became Major Swing", "Major Swing ID", "Major Role", "Exact Non Promotion Reason",
        "Geometry Metric Class", "Detector Version",
    ]
    rows: list[list[Any]] = []
    fills: list[PatternFill | None] = []
    seed = result.internal_seed
    if seed is not None:
        values = {header: "" for header in headers}
        values.update(
            {
                "Record Class": "INTERNAL_BOOTSTRAP_SEED",
                "Type": seed.type,
                "Structure Label": "BOOTSTRAP_SEED",
                "Formation ID": seed.formation_id,
                "Raw Pivot ID": seed.pivot_id,
                "Extreme UTC": iso_utc(seed.open_time),
                "Extreme Turkey": iso_turkey(seed.open_time),
                "Raw Pivot Confirmation UTC": iso_utc(seed.confirm_time),
                "Raw Pivot Confirmation Turkey": iso_turkey(seed.confirm_time),
                "Price": _num(seed.price),
                "ATR": _num(seed.atr),
                "Incoming Leg Start Class": "CONTEXT_ONLY",
                "Became Major Swing": "FALSE",
                "Major Role": "INTERNAL_BOOTSTRAP_SEED",
                "Exact Non Promotion Reason": "INTERNAL_BOOTSTRAP_SEED",
                "Geometry Metric Class": "CONTEXT_ONLY",
                "Detector Version": CONFIG.detector_version,
            }
        )
        rows.append([values[header] for header in headers])
        fills.append(FILL_BOOTSTRAP)
    for swing in result.internal_swings:
        prom = swing.pivot.prominence
        start_ohlc = _ohlc_parts(getattr(swing, "leg_start_ohlc", None))
        end_ohlc = _ohlc_parts(getattr(swing, "leg_end_ohlc", None))
        rows.append([
            swing.record_class, swing.internal_swing_id, swing.sequence, swing.type, swing.structure_label,
            swing.formation_id, swing.pivot_id, iso_utc(swing.extreme_time), iso_turkey(swing.extreme_time),
            iso_utc(swing.raw_pivot_confirmed_at), iso_turkey(swing.raw_pivot_confirmed_at),
            iso_utc(swing.activated_at), iso_turkey(swing.activated_at),
            iso_utc(swing.reversal_confirmed_at), iso_turkey(swing.reversal_confirmed_at),
            iso_utc(swing.confirmed_at), iso_turkey(swing.confirmed_at),
            swing.age, _txt(swing.replacement_history), _num(swing.price), _num(swing.atr),
            _num(swing.displacement_price), _num(swing.displacement_atr), "TRUE" if swing.displacement_passed else "FALSE",
            swing.spacing_bars, "TRUE" if swing.spacing_passed else "FALSE",
            _num(swing.reversal_price), _num(swing.reversal_atr), "TRUE" if swing.reversal_passed else "FALSE",
            getattr(swing, "leg_start_class", ""),
            iso_utc(getattr(swing, "leg_start_open_time", None)),
            iso_utc(getattr(swing, "leg_start_close_time", None)),
            iso_turkey(getattr(swing, "leg_start_open_time", None)),
            iso_turkey(getattr(swing, "leg_start_close_time", None)),
            iso_utc(getattr(swing, "leg_end_open_time", None)),
            iso_utc(getattr(swing, "leg_end_close_time", None)),
            iso_turkey(getattr(swing, "leg_end_open_time", None)),
            iso_turkey(getattr(swing, "leg_end_close_time", None)),
            *start_ohlc, *end_ohlc,
            _num(prom["left_prominence"]), _num(prom["right_prominence"]),
            _num(prom["two_sided_prominence"]), _num(swing.pivot.prominence_atr_at_confirmation),
            _num(getattr(swing, "absolute_amplitude", None)),
            getattr(swing, "duration_bars", None),
            _num(getattr(swing, "leg_efficiency", None)),
            _num(getattr(swing, "path_length", None)),
            "TRUE" if swing.became_major_swing else "FALSE",
            swing.major_swing_id,
            swing.major_role,
            swing.major_non_promotion_reason,
            getattr(swing, "metric_class_geometry", "CAUSAL_AT_INTERNAL_CONFIRMATION"),
            CONFIG.detector_version,
        ])
        fill = None
        if swing.major_role == "MAJOR_BOOTSTRAP_SEED":
            fill = FILL_BOOTSTRAP
        elif swing.replacement_history:
            fill = FILL_REPLACEMENT
        elif swing.major_role == "UNRESOLVED_CANDIDATE":
            fill = FILL_UNRESOLVED
        elif swing.max_age > 168:
            fill = FILL_LONG
        fills.append(fill)
    return headers, rows, fills


def _pivot_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    for pivot in result.pivots:
        prom = pivot.prominence
        rows.append([
            pivot.pivot_id, pivot.type, pivot.formation_id, pivot.extreme_row, pivot.confirm_row,
            iso_utc(pivot.open_time), iso_turkey(pivot.open_time),
            iso_utc(pivot.confirm_time), iso_turkey(pivot.confirm_time),
            _num(pivot.price), _num(pivot.atr), _num(pivot.atr_at_extreme), _num(pivot.atr_at_confirmation),
            _pct(pivot.outbound_percent), _pct(pivot.return_percent),
            pivot.left_bars, pivot.right_bars, "TRUE" if pivot.symmetric else "FALSE",
            _num(prom["left_prominence"]), _num(prom["right_prominence"]),
            _num(prom["two_sided_prominence"]), _num(prom["max_two_sided_prominence"]),
            _pct(prom["percentage_prominence"]), _num(pivot.prominence_atr_at_confirmation),
            "TRUE" if prom["left_censored"] else "FALSE",
            "TRUE" if prom["right_censored"] else "FALSE",
            pivot.internal_role, pivot.internal_disposition, pivot.candidate_id,
            "FORMATION_DERIVED_ASYMMETRIC", CONFIG.detector_version,
        ])
    return rows


def _formation_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    bars = result.bars
    for formation in result.formations:
        def bar_at(row: int | None):
            if row is None or row < 0 or row >= len(bars):
                return None
            return bars[row]

        opener = bar_at(formation.get("open_row") if formation.get("open_row") is not None and formation.get("open_row") >= 0 else None)
        extreme = bar_at(formation.get("extreme_row"))
        close = bar_at(formation.get("close_row") if formation.get("close_row") is not None and 0 <= formation.get("close_row", -1) < len(bars) else None)
        rows.append([
            formation["formation_id"], formation["formation_type"],
            formation.get("open_row"), formation.get("extreme_row"), formation.get("close_row"),
            iso_utc(opener.open_time) if opener else "",
            iso_turkey(opener.open_time) if opener else "",
            iso_utc(extreme.open_time) if extreme else "",
            iso_turkey(extreme.open_time) if extreme else "",
            iso_utc(close.open_time) if close else "",
            iso_turkey(close.open_time) if close else "",
            iso_utc(close.close_time) if close else "",
            iso_turkey(close.close_time) if close else "",
            formation.get("total_bars"), formation.get("interior_bars"),
            formation.get("left_bars"), formation.get("right_bars"),
            formation.get("asymmetry_difference"),
            _num(formation.get("asymmetry_ratio")),
            "TRUE" if formation.get("symmetric") else "FALSE",
            formation.get("extreme_position"), formation.get("allowed_positions"),
            "TRUE" if formation.get("central_pass") else "FALSE",
            _num(opener.open) if opener else None, _num(opener.high) if opener else None,
            _num(opener.low) if opener else None, _num(opener.close) if opener else None,
            _num(extreme.open) if extreme else None, _num(extreme.high) if extreme else None,
            _num(extreme.low) if extreme else None, _num(extreme.close) if extreme else None,
            _num(close.open) if close else None, _num(close.high) if close else None,
            _num(close.low) if close else None, _num(close.close) if close else None,
            _num(formation.get("reference_price")), _num(formation.get("extreme_price")),
            _num(formation.get("outbound_price")), _pct(formation.get("outbound_percent")),
            _num(formation.get("return_price")), _pct(formation.get("return_percent")),
            _pct(formation.get("return_percent_reference_basis")),
            _num(formation.get("touch_price")),
            "TRUE" if formation.get("revisit_pass") else "FALSE",
            formation.get("plateau_start_row"), formation.get("plateau_end_row"), formation.get("plateau_length"),
            formation.get("qualification_status"), formation.get("primary_disposition"),
            formation.get("specific_reason"), formation.get("representative_formation_id"),
            formation.get("duplicate_count"), _txt(",".join(formation.get("duplicate_ids") or [])),
            formation.get("dedup_reason"), "FORMATION_DERIVED_ASYMMETRIC", "NONE", "NONE",
            CONFIG.detector_version,
        ])
    return rows


FORMATION_HEADERS = [
    "Formation ID", "Formation Type", "Swing Open Row", "Extreme Row", "Swing Close Row",
    "Swing Open UTC", "Swing Open Turkey", "Extreme UTC", "Extreme Turkey",
    "Swing Close UTC", "Swing Close Turkey", "Formation Confirmation UTC", "Formation Confirmation Turkey",
    "Total Formation Bars", "Interior Candle Count", "Formation Left Bars", "Formation Right Bars",
    "Asymmetry Difference", "Asymmetry Ratio", "Symmetric",
    "Extreme Position", "Allowed Central Positions", "Central Position Pass",
    "Swing Open", "Swing High", "Swing Low", "Swing Close",
    "Extreme Open", "Extreme High", "Extreme Low", "Extreme Close",
    "Close Open", "Close High", "Close Low", "Close Close",
    "Reference Price", "Extreme Price", "Outbound Price", "Outbound Percent",
    "Return Price", "Return Percent", "Return Percent Reference Basis", "Touch Price",
    "Close Candle Reference Touch", "Plateau Start Row", "Plateau End Row", "Plateau Length",
    "Qualification Status", "Primary Disposition", "Specific Reason",
    "Representative Formation ID", "Duplicate Count", "Duplicate Formation IDs", "Deduplication Reason",
    "Pivot Span Mode", "Fixed Pivot Left Bars", "Fixed Pivot Right Bars", "Detector Version",
]


def _candle_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    qualified = [
        formation for formation in result.formations
        if formation["primary_disposition"] == "QUALIFIED_RAW_FORMATION" and formation.get("emits_pivot")
    ]
    for formation in qualified:
        start = int(formation["open_row"])
        end = int(formation["close_row"])
        extreme_row = int(formation["extreme_row"])
        plateau_start = formation.get("plateau_start_row")
        plateau_end = formation.get("plateau_end_row")
        reference = formation.get("reference_price")
        extreme_price = formation.get("extreme_price")
        for row in range(start, end + 1):
            bar = result.bars[row]
            position = row - start + 1
            interior_position = position - 1 if start < row < end else None
            if row == start:
                role = "Swing Open Candle"
            elif row == end:
                role = "Swing Close Candle"
            elif row == extreme_row:
                role = "Extreme Candle"
            else:
                role = f"Interior Candle {interior_position}"
            plateau_member = (
                plateau_start is not None and plateau_end is not None and int(plateau_start) <= row <= int(plateau_end)
            )
            rows.append([
                formation["formation_id"], role, position, interior_position,
                iso_utc(bar.open_time), iso_utc(bar.close_time),
                iso_turkey(bar.open_time), iso_turkey(bar.close_time),
                _num(bar.open), _num(bar.high), _num(bar.low), _num(bar.close),
                _num(bar.volume), bar.trades,
                "TRUE" if row == start else "FALSE",
                "TRUE" if row == extreme_row else "FALSE",
                "TRUE" if row == end else "FALSE",
                "TRUE" if plateau_member else "FALSE",
                "TRUE" if row == extreme_row else "FALSE",
                _num(bar.close - reference) if reference is not None else None,
                _num(bar.close - extreme_price) if extreme_price is not None else None,
                "TRUE",
            ])
    return rows


def _transition_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    for event in result.transitions:
        rows.append([
            event["engine"], iso_utc(event["event_time"]), iso_turkey(event["event_time"]), event["row"],
            event["previous_state"], event["new_state"], event["trigger"],
            event["formation_id"], event["raw_pivot_id"], event["internal_swing_id"], event["major_swing_id"],
            event["candidate_id"], event["candidate_type"], _num(event["candidate_price"]) if event["candidate_price"] != "" else None,
            _txt(event["displacement_status"]), _txt(event["spacing_status"]),
            _txt(event["prominence_status"]), _txt(event["reversal_status"]),
            "TRUE" if event["replacement_event"] else "FALSE",
            "TRUE" if event["confirmation_event"] else "FALSE",
            event["reason"],
        ])
    return rows


def _monthly_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    for row in result.monthly:
        rows.append([
            row["utc_month"], row["source_candles"], row["formation_count"], row["qualified_formation_count"],
            row["rejected_formation_count"], row["rejection_breakdown"],
            row["raw_pivot_highs"], row["raw_pivot_lows"],
            row["new_internal_high_candidates"], row["new_internal_low_candidates"],
            row["confirmed_internal_highs"], row["confirmed_internal_lows"],
            row["replaced_internal_candidates"], row["rejected_internal_candidates"],
            row["active_internal_state"],
            row["new_major_high_candidates"], row["new_major_low_candidates"],
            row["confirmed_major_highs"], row["confirmed_major_lows"],
            row["replaced_major_candidates"], row["major_non_promotion_count"],
            row["active_major_state"], row["maximum_candidate_age"],
            row["ignored_valid_reversal_count"], row["state_transition_errors"],
        ])
    return rows


def _forward_rows(result: DetectionResult) -> list[list[Any]]:
    rows = []
    for swing in result.major_swings:
        rows.append([
            swing.major_swing_id, swing.type, _num(swing.price),
            swing.next_major_swing_id, swing.bars_to_next_pivot, swing.bars_to_next_confirmation,
            _num(swing.max_favorable_excursion), _num(swing.max_adverse_excursion),
            swing.retest_count, iso_utc(swing.first_close_beyond_time), _num(swing.first_close_beyond_price),
            "TRUE" if swing.forward_data_censored else "FALSE",
            swing.forward_censoring_reason,
            "EX_POST_FORWARD_METRIC",
            "KNOWN_AFTER_NEXT_SWING" if swing.next_major_swing_id else "EX_POST_FORWARD_METRIC",
        ])
    return rows


def _summary_rows(result: DetectionResult, context: dict[str, Any]) -> list[list[Any]]:
    unique = result.unique
    lines = [
        ("Symbol", CONFIG.symbol),
        ("Timeframe", CONFIG.timeframe),
        ("TradingView Symbol", CONFIG.tradingview_symbol),
        ("Detector Version", CONFIG.detector_version),
        ("Workbook Revision", context.get("revision_name", "")),
        ("Analysis Start UTC", iso_utc(CONFIG.analysis_start_utc)),
        ("Analysis End UTC", iso_utc(CONFIG.analysis_end_utc)),
        ("Source Rows", len(result.bars)),
        ("Qualified Formations", unique["unique_qualified_formations"]),
        ("Raw Pivots", unique["unique_raw_pivots"]),
        ("Confirmed Internal Swings", unique["unique_confirmed_internal_swings"]),
        ("Confirmed Major Swings", unique["unique_confirmed_major_swings"]),
        ("Internal Bootstrap Seed", result.internal_seed.pivot_id if result.internal_seed else "NONE"),
        ("Major Bootstrap Seed", result.major_seed.internal_swing_id if result.major_seed else "NONE"),
        ("Ignored Valid Reversals", result.internal_engine.ignored_valid + (result.major_engine.ignored_valid if result.major_engine else 0)),
        ("State Transition Errors", result.internal_engine.state_errors + (result.major_engine.state_errors if result.major_engine else 0)),
        ("Engine Independence", "PASS" if result.independence_pass else "FAIL"),
        ("Pivot Span Mode", CONFIG.pivot_span_mode),
        ("Candidate Timeout", CONFIG.candidate_timeout),
    ]
    for swing in result.major_swings[:12]:
        lines.append((
            f"Major Preview {swing.major_swing_id}",
            f"{swing.type} {swing.structure_label} price={swing.price} extreme={iso_utc(swing.extreme_time)} score={swing.significance_score}",
        ))
    if not result.major_swings:
        lines.append(("Major Preview", "No confirmed major swing. See Internal Swings for exact non-promotion reasons."))
    return [[key, value] for key, value in lines]


def _parameter_rows() -> list[list[Any]]:
    items = [
        ("source_symbol", CONFIG.symbol),
        ("source_timeframe", CONFIG.timeframe),
        ("analysis_start_utc", iso_utc(CONFIG.analysis_start_utc)),
        ("analysis_end_utc", iso_utc(CONFIG.analysis_end_utc)),
        ("expected_rows", CONFIG.expected_rows),
        ("minimum_total_formation_candles", CONFIG.minimum_total_formation_candles),
        ("maximum_total_formation_candles", CONFIG.maximum_total_formation_candles),
        ("minimum_interior_candles", CONFIG.minimum_interior_candles),
        ("maximum_interior_candles", CONFIG.maximum_interior_candles),
        ("pivot_span_mode", CONFIG.pivot_span_mode),
        ("fixed_pivot_left_bars", CONFIG.fixed_pivot_left_bars),
        ("fixed_pivot_right_bars", CONFIG.fixed_pivot_right_bars),
        ("extreme_position_rule", CONFIG.extreme_position_rule),
        ("formation_reference", CONFIG.formation_reference),
        ("minimum_outbound_percent", str(CONFIG.minimum_outbound_percent)),
        ("minimum_return_percent", str(CONFIG.minimum_return_percent)),
        ("formation_confirmation", CONFIG.formation_confirmation),
        ("internal_bootstrap_enabled", "TRUE"),
        ("internal_minimum_displacement_atr", str(CONFIG.internal_minimum_displacement_atr)),
        ("internal_reversal_atr", str(CONFIG.internal_reversal_atr)),
        ("internal_minimum_spacing_bars", CONFIG.internal_minimum_spacing_bars),
        ("major_bootstrap_enabled", "TRUE"),
        ("major_minimum_displacement_atr", str(CONFIG.major_minimum_displacement_atr)),
        ("major_minimum_displacement_percent", str(CONFIG.major_minimum_displacement_percent)),
        ("major_reversal_atr", str(CONFIG.major_reversal_atr)),
        ("major_reversal_percent", str(CONFIG.major_reversal_percent)),
        ("major_minimum_spacing_bars", CONFIG.major_minimum_spacing_bars),
        ("major_minimum_prominence_atr", str(CONFIG.major_minimum_prominence_atr)),
        ("reversal_evaluation_frequency", CONFIG.reversal_evaluation_frequency),
        ("candidate_timeout", CONFIG.candidate_timeout),
        ("detector_version", CONFIG.detector_version),
        ("implementation_mode", CONFIG.implementation_mode),
        ("authorized_project_root", str(CONFIG.project_root)),
        ("authorized_detector_package", str(CONFIG.package_dir)),
        ("return_touch_price_definition", "swing_open_close when the swing-close range contains it"),
        ("return_gate", "reference-basis return percent >= 1.00, which equals outbound when the touch price is the swing-open close"),
        ("extreme_basis_return", "recorded as Return Percent = distance / extreme price; not a second tighter gate"),
        ("central_three_odd", "center=(N+1)/2; positions center-1, center, center+1"),
        ("central_three_even", "lower_center=N/2; positions lower_center-1, lower_center, lower_center+1"),
        ("atr_formula", "TR=max(high-low, abs(high-prev_close), abs(low-prev_close)); seed=SMA(first 14 TR); then (prev*13+TR)/14"),
        ("prominence_formula", "Walk outward until a strictly more extreme print. High prominence is distance down to the minimum low. Low prominence is distance up to the maximum high. Two-sided prominence is the minimum of the two sides. Bars after the evaluation bar are excluded."),
        ("processing_order", "On each completed bar: validate, update ATR, release formations whose swing close just completed, deduplicate, release raw pivots, update the internal candidate, evaluate the internal reversal on the close, confirm immediately when valid, feed new internal swings to the major engine, evaluate the major reversal on the same close."),
        ("bootstrap", "INTERNAL_BOOTSTRAP_SEARCH and MAJOR_BOOTSTRAP_SEARCH. Seeds initialize direction only and are excluded from normal swing counts."),
        ("timezone", "Computation is UTC. Display conversion uses ZoneInfo('Europe/Istanbul') and does not add three hours manually."),
        ("close_time", "open_time + 1 hour - 1 millisecond, computed in UTC"),
    ]
    return [[key, value] for key, value in items]


def write_workbook(path: Path, result: DetectionResult, context: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    summary = wb.create_sheet("Swing Summary")
    _write_table(summary, ["Item", "Value"], _summary_rows(result, context), {})

    major_rows = [_major_row(swing) for swing in result.major_swings]
    major_formats = _formats(MAJOR_HEADERS)
    points = wb.create_sheet("Major Swing Points")
    _write_table(points, MAJOR_HEADERS, major_rows, major_formats)
    highs = wb.create_sheet("Major Swing Highs")
    _write_table(highs, MAJOR_HEADERS, [row for row, swing in zip(major_rows, result.major_swings) if swing.type == "HIGH"], major_formats)
    lows = wb.create_sheet("Major Swing Lows")
    _write_table(lows, MAJOR_HEADERS, [row for row, swing in zip(major_rows, result.major_swings) if swing.type == "LOW"], major_formats)
    legs = wb.create_sheet("Major Swing Legs")
    _write_table(legs, MAJOR_HEADERS, major_rows, major_formats)
    for sheet, swings in (
        (points, result.major_swings),
        (highs, [swing for swing in result.major_swings if swing.type == "HIGH"]),
        (lows, [swing for swing in result.major_swings if swing.type == "LOW"]),
        (legs, result.major_swings),
    ):
        for offset, swing in enumerate(swings, start=2):
            fill = None
            if swing.max_age > 2160:
                fill = FILL_LONG
            elif swing.replacement_history:
                fill = FILL_REPLACEMENT
            _paint(sheet, offset, fill)

    internal_headers, internal_rows, internal_fills = _internal_rows(result)
    internal_sheet = wb.create_sheet("Internal Swings")
    _write_table(internal_sheet, internal_headers, internal_rows, _formats(internal_headers))
    for offset, fill in enumerate(internal_fills, start=2):
        _paint(internal_sheet, offset, fill)

    pivot_headers = [
        "Raw Pivot ID", "Type", "Formation ID", "Extreme Row", "Confirmation Row",
        "Extreme UTC", "Extreme Turkey", "Raw Pivot Confirmation UTC", "Raw Pivot Confirmation Turkey",
        "Price", "ATR At Candidate", "ATR At Extreme", "ATR At Confirmation",
        "Outbound Percent", "Return Percent", "Formation Left Bars", "Formation Right Bars", "Symmetric",
        "Left Prominence", "Right Prominence", "Two Sided Prominence", "Max Two Sided Prominence",
        "Percentage Prominence", "Prominence ATR", "Left Prominence Censored", "Right Prominence Censored",
        "Internal Role", "Internal Disposition", "Candidate ID", "Pivot Span Mode", "Detector Version",
    ]
    pivots = wb.create_sheet("Raw Pivot Candidates")
    pivot_rows = _pivot_rows(result)
    _write_table(pivots, pivot_headers, pivot_rows, _formats(pivot_headers))
    for offset, pivot in enumerate(result.pivots, start=2):
        fill = FILL_BOOTSTRAP if pivot.internal_role == "INTERNAL_BOOTSTRAP_SEED" else None
        if pivot.internal_role == "REPLACED_CANDIDATE":
            fill = FILL_REPLACEMENT
        _paint(pivots, offset, fill)

    formations = wb.create_sheet("Raw Formations")
    _write_table(formations, FORMATION_HEADERS, _formation_rows(result), _formats(FORMATION_HEADERS))

    candle_headers = [
        "Formation ID", "Candle Role", "Formation Position", "Interior Position",
        "Open Time UTC", "Close Time UTC", "Open Time Turkey", "Close Time Turkey",
        "Open", "High", "Low", "Close", "Volume", "Trades",
        "Is Swing Open", "Is Extreme", "Is Swing Close", "Is Plateau Member",
        "Is Representative Extreme", "Distance From Reference", "Distance From Extreme",
        "Known At Formation Confirmation",
    ]
    candles = wb.create_sheet("Formation Candles")
    _write_table(candles, candle_headers, _candle_rows(result), _formats(candle_headers))

    transition_headers = [
        "Engine", "Event UTC", "Event Turkey", "Row", "Previous State", "New State", "Trigger",
        "Formation ID", "Raw Pivot ID", "Internal Swing ID", "Major Swing ID",
        "Candidate ID", "Candidate Type", "Candidate Price",
        "Displacement Status", "Spacing Status", "Prominence Status", "Reversal Status",
        "Replacement Event", "Confirmation Event", "Reason",
    ]
    transitions = wb.create_sheet("State Transitions")
    transition_rows = _transition_rows(result)
    _write_table(transitions, transition_headers, transition_rows, _formats(transition_headers))
    for offset, event in enumerate(result.transitions, start=2):
        reason = str(event["reason"])
        fill = None
        if "BOOTSTRAP" in event["trigger"] or "BOOTSTRAP" in reason:
            fill = FILL_BOOTSTRAP
        elif event["replacement_event"]:
            fill = FILL_REPLACEMENT
        elif "UNRESOLVED" in reason or "HARD_" in reason:
            fill = FILL_ERROR if "HARD_" in reason else FILL_UNRESOLVED
        _paint(transitions, offset, fill)

    monthly_headers = [
        "UTC Month", "Source Candle Count", "Formation Count", "Qualified Formation Count",
        "Rejected Formation Count", "Rejected Formation Breakdown",
        "Raw Pivot High Count", "Raw Pivot Low Count",
        "New Internal High Candidates", "New Internal Low Candidates",
        "Confirmed Internal Swing Highs", "Confirmed Internal Swing Lows",
        "Replaced Internal Candidates", "Rejected Internal Candidates", "Active Internal State At Month End",
        "New Major High Candidates", "New Major Low Candidates",
        "Confirmed Major Swing Highs", "Confirmed Major Swing Lows",
        "Replaced Major Candidates", "Major Non Promotion Counts", "Active Major State At Month End",
        "Maximum Candidate Age", "Ignored Valid Reversal Count", "State Transition Errors",
    ]
    monthly = wb.create_sheet("Monthly Diagnostics")
    _write_table(monthly, monthly_headers, _monthly_rows(result), {})

    forward_headers = [
        "Major Swing ID", "Type", "Price", "Next Major Swing", "Bars To Next Pivot",
        "Bars To Next Confirmation", "Maximum Favorable Excursion", "Maximum Adverse Excursion",
        "Retests", "First Close Beyond UTC", "First Close Beyond Price",
        "Forward Data Censored", "Forward Censoring Reason", "Excursion Metric Class", "Next Swing Metric Class",
    ]
    forward = wb.create_sheet("Forward Evaluation")
    _write_table(forward, forward_headers, _forward_rows(result), _formats(forward_headers))

    parameters = wb.create_sheet("Parameters")
    _write_table(parameters, ["Parameter", "Value"], _parameter_rows(), {})

    diagnostics = wb.create_sheet("Diagnostics")
    _write_table(diagnostics, ["Section", "Key", "Value"], context["diagnostic_rows"], {})

    readme = wb.create_sheet("README")
    readme_lines = context.get("readme_lines") or ["README missing"]
    _write_table(readme, ["Line"], [[line] for line in readme_lines], {})

    if tuple(wb.sheetnames) != WORKBOOK_SHEETS:
        raise RuntimeError(f"sheet order mismatch: {wb.sheetnames}")
    wb.properties.creator = "hierarchical_swing_v4"
    wb.properties.title = context.get("revision_name", CONFIG.detector_version)
    wb.properties.subject = "Exploratory hierarchical swing detection. Not a trading strategy."
    wb.properties.created = datetime.now(timezone.utc).replace(tzinfo=None)
    wb.properties.modified = wb.properties.created
    if path.exists():
        raise FileExistsError(f"refusing to overwrite {path}")
    wb.save(path)
    wb.close()


def _diagnostic_map(ws) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 3 and row[1]:
            mapping[str(row[1])] = "" if row[2] is None else str(row[2])
    return mapping


def validate_saved_workbook(path: Path, result: DetectionResult, context: dict[str, Any]) -> list[str]:
    problems: list[str] = []
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    try:
        if tuple(wb.sheetnames) != WORKBOOK_SHEETS:
            problems.append(f"sheet order {wb.sheetnames}")
        if wb.vba_archive is not None:
            problems.append("workbook contains macros")
        if getattr(wb, "_external_links", None):
            problems.append("workbook contains external links")
        diagnostics = _diagnostic_map(wb["Diagnostics"])
        if diagnostics.get("source_rows") != str(len(result.bars)):
            problems.append("diagnostics source_rows mismatch")
        if diagnostics.get("workbook_revision") != context.get("revision_name"):
            problems.append("diagnostics revision mismatch")
        if diagnostics.get("per_bar_counter_label") != "PER_BAR_DIAGNOSTIC_ONLY":
            problems.append("per-bar counters are not labeled")
        points = wb["Major Swing Points"]
        major_count = max(0, points.max_row - 1)
        if major_count != len(result.major_swings):
            problems.append(f"major sheet rows {major_count} != {len(result.major_swings)}")
        if result.major_swings:
            headers = [cell.value for cell in points[1]]
            hierarchy = headers.index("Hierarchy Level")
            window = headers.index("Inside Requested Analysis Window")
            for row in points.iter_rows(min_row=2, max_row=points.max_row, values_only=True):
                if row[hierarchy] != "MAJOR_SWING" or row[window] != "TRUE":
                    problems.append("major row missing hierarchy flags")
                    break
        internal = wb["Internal Swings"]
        headers = [cell.value for cell in internal[1]]
        class_index = headers.index("Record Class")
        normal = 0
        seeds = 0
        for row in internal.iter_rows(min_row=2, max_row=internal.max_row, values_only=True):
            if row[class_index] == "NORMAL_CONFIRMED":
                normal += 1
            elif row[class_index] == "INTERNAL_BOOTSTRAP_SEED":
                seeds += 1
        if normal != len(result.internal_swings):
            problems.append("internal normal count mismatch")
        if (seeds == 1) != (result.internal_seed is not None):
            problems.append("bootstrap seed row mismatch")
        formations = wb["Raw Formations"]
        headers = [cell.value for cell in formations[1]]
        disposition_index = headers.index("Primary Disposition")
        total_index = headers.index("Total Formation Bars")
        interior_index = headers.index("Interior Candle Count")
        central_index = headers.index("Central Position Pass")
        left_index = headers.index("Formation Left Bars")
        right_index = headers.index("Formation Right Bars")
        qualified = 0
        for row in formations.iter_rows(min_row=2, max_row=formations.max_row, values_only=True):
            if row[disposition_index] != "QUALIFIED_RAW_FORMATION":
                continue
            qualified += 1
            total = row[total_index]
            interior = row[interior_index]
            if total is None or not (5 <= int(total) <= 10):
                problems.append("qualified formation length outside 5-10")
                break
            if interior is None or not (3 <= int(interior) <= 8):
                problems.append("qualified interior outside 3-8")
                break
            if row[central_index] != "TRUE":
                problems.append("qualified formation failed central-three rule")
                break
            if row[left_index] in (None, "") or row[right_index] in (None, ""):
                problems.append("qualified formation missing asymmetric spans")
                break
        if qualified != result.unique["unique_qualified_formations"]:
            problems.append("qualified formation count mismatch")
        monthly = wb["Monthly Diagnostics"]
        source_total = 0
        for row in monthly.iter_rows(min_row=2, max_row=monthly.max_row, values_only=True):
            source_total += int(row[1] or 0)
        if source_total != len(result.bars):
            problems.append("monthly source total mismatch")
        if diagnostics.get("ignored_valid_reversal_count") != "0":
            problems.append("workbook reports ignored reversals")
        if diagnostics.get("state_transition_error_count") != "0":
            problems.append("workbook reports state errors")
    finally:
        wb.close()
    return problems
