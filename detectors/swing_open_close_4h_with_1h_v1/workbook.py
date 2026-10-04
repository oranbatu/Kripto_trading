"""Conservative Excel workbook for the 2026 4h Swing Open/Close detector."""
from __future__ import annotations

import math
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.swing_open_close_4h_with_1h_v1.engine import (
    ALTERNATIVE_METHOD,
    COMPACT_SWING,
    REFERENCE_METHOD,
    STANDARD_SWING,
    SWING_HIGH,
    _returned,
    AnalysisResult,
    Swing,
    iso_turkey,
    iso_utc,
)
from detectors.swing_open_close_4h_with_1h_v1.swing_config import (
    DETECTOR_VERSION,
    PLAIN_RANGE_SHEETS,
    TABLE_NAMES,
    WORKBOOK_SHEETS,
    parameter_registry,
)

PRICE = "0.00000000"
PERCENT_POINTS = '0.00000000"%"'
EXCEL_MAX_TEXT = 32767
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
OOXML_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_EMPTY_TYPED_CELL = re.compile(br'<c\b[^>]*\bt="(?:n|b|s|str|inlineStr|d|e)"[^>]*/>')
LARGE_SHEETS = frozenset({"Formation Candles", "Candidate Audit", "Derived Same Extreme"})
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


HEADER_FILL = _solid("1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
HEADER_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
HIGH_FILL = _solid("F4CCCC")
LOW_FILL = _solid("D9EAD3")
COMPACT_FILL = _solid("C9DAF8")
STANDARD_FILL = _solid("FCE5CD")
EXACT_FILL = _solid("FFF2CC")
CROSS_FILL = _solid("D0E2F3")
FAIL_FILL = _solid("E6B8B7")
ALTERNATIVE_FILL = _solid("EAD1DC")
WRAP = Alignment(wrap_text=True, vertical="center")


def write_workbook(path: Path, result: AnalysisResult, context: dict[str, Any], readme: str) -> None:
    workbook = Workbook()
    sheets = {
        "Executive Summary": _executive(result, context),
        "Primary Swings": _swings(result.primaries),
        "Primary Swing Highs": _swings([item for item in result.primaries if item.direction == SWING_HIGH]),
        "Primary Swing Lows": _swings([item for item in result.primaries if item.direction != SWING_HIGH]),
        "Compact Primary": _swings([item for item in result.primaries if item.formation_class == COMPACT_SWING]),
        "Standard Primary": _swings([item for item in result.primaries if item.formation_class == STANDARD_SWING]),
        "Alternative Primary": _swings([
            item for item in result.primaries
            if item.detection_method == ALTERNATIVE_METHOD and item.found_through_alternative
        ]),
        "Derived Same Extreme": _derived(result),
        "Formation Candles": _formation(result),
        "Candidate Audit": _audit(result),
        "Overlap Analysis": _overlaps(result.primaries),
        "Forward Evaluation": _forward(result.primaries),
        "Parameters": _parameters(),
        "Diagnostics": context["diagnostic_rows"],
        "README": _readme_rows(readme),
    }
    workbook.active.title = WORKBOOK_SHEETS[0]
    for name in WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(name)
    for name in WORKBOOK_SHEETS:
        _write_sheet(workbook[name], sheets[name], None if name in PLAIN_RANGE_SHEETS else TABLE_NAMES.get(name))
    _highlight(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def validate_saved_workbook(path: Path, expected_attempts: int) -> None:
    validate_ooxml_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = workbook["Candidate Audit"].max_row - 1
        if rows != expected_attempts:
            raise RuntimeError(f"candidate audit rows {rows} != {expected_attempts}")
    finally:
        workbook.close()


def validate_ooxml_workbook(path: Path) -> None:
    payload = path.read_bytes()
    if payload[:4] != b"PK\x03\x04":
        raise RuntimeError("workbook does not begin with the ZIP signature")
    if path.suffix.lower() != ".xlsx":
        raise RuntimeError("workbook extension is not .xlsx")
    if not zipfile.is_zipfile(path):
        raise RuntimeError("zipfile.is_zipfile returned false")
    required = (
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/worksheets/sheet1.xml",
    )
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("ZipFile.testzip found a corrupt member")
        names = set(archive.namelist())
        for member in required:
            if member not in names:
                raise RuntimeError(f"missing OOXML member {member}")
        if any("vba" in name.lower() for name in names):
            raise RuntimeError("workbook contains a VBA part")
        if any("externalLink" in name for name in names):
            raise RuntimeError("workbook contains an external link")
        for name in sorted(names):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                if _EMPTY_TYPED_CELL.search(archive.read(name)):
                    raise RuntimeError(f"{name} contains a typed cell with no value")
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise RuntimeError(f"XML parse failed for {name}: {exc}") from exc
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name:
                    _validate_sheet_xml(name, root)
        _validate_relationships(archive)
    _validate_openpyxl_workbook(path)


def _validate_parameter_sheet(worksheet) -> None:
    from detectors.swing_open_close_4h_with_1h_v1.swing_config import assert_parameter_registry

    assert_parameter_registry()
    expected = parameter_registry()
    headers = [cell.value for cell in worksheet[1]]
    required = [
        "Category", "Parameter Name", "Effective Value", "Data Type", "Unit",
        "Hard Rule or Descriptive", "Formula or Allowed Values", "Description",
        "Source of Parameter", "Used In",
    ]
    if headers[:10] != required:
        raise RuntimeError(f"Parameters headers {headers[:10]}")
    rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    if rows and rows[0][0] == "NO_RESULTS":
        raise RuntimeError("Parameters sheet is empty")
    listed = [row[1] for row in rows]
    if listed != [item.name for item in expected]:
        raise RuntimeError("Parameters sheet does not match the code parameter registry")
    values = [row[2] for row in rows]
    if values != [item.value for item in expected]:
        raise RuntimeError("Parameters sheet values do not match the code parameter registry")


def _validate_sheet_xml(name: str, root: ET.Element) -> None:
    dimension = root.find("m:dimension", OOXML_NS)
    auto_filter = root.find("m:autoFilter", OOXML_NS)
    if auto_filter is not None and dimension is not None:
        if auto_filter.attrib.get("ref", "") != dimension.attrib.get("ref", ""):
            raise RuntimeError(f"{name} autoFilter does not match the used range")


def _validate_relationships(archive: zipfile.ZipFile) -> None:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib.get("Id"): item.attrib.get("Target", "") for item in relationships}
    for sheet in root.findall("m:sheets/m:sheet", OOXML_NS):
        title = sheet.attrib.get("name", "")
        if not title or len(title) > 31:
            raise RuntimeError(f"invalid sheet name {title!r}")
        relation = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = targets.get(relation, "").lstrip("/")
        member = target if target.startswith("xl/") else "xl/" + target
        if member not in archive.namelist():
            raise RuntimeError(f"sheet {title} references missing part {member}")


def _validate_openpyxl_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
            raise RuntimeError(f"sheet order {workbook.sheetnames}")
        if workbook.vba_archive is not None:
            raise RuntimeError("workbook contains macros")
        if getattr(workbook, "_external_links", []):
            raise RuntimeError("workbook contains external links")
        names: list[str] = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > EXCEL_MAX_ROWS or worksheet.max_column > EXCEL_MAX_COLUMNS:
                raise RuntimeError(f"{worksheet.title} exceeds an Excel worksheet limit")
            if worksheet.title in LARGE_SHEETS or worksheet.title in PLAIN_RANGE_SHEETS:
                if worksheet.tables:
                    raise RuntimeError(f"{worksheet.title} must stay a plain filtered range")
            expected = TABLE_NAMES.get(worksheet.title)
            if expected is not None:
                if worksheet["A2"].value == "NO_RESULTS":
                    if worksheet.tables:
                        raise RuntimeError(f"{worksheet.title} created a table for NO_RESULTS")
                elif not worksheet.tables:
                    raise RuntimeError(f"{worksheet.title} is missing its Excel table")
            for table in worksheet.tables.values():
                names.append(table.name)
                _validate_table(worksheet, table)
        if len(names) != len(set(names)):
            raise RuntimeError(f"duplicate table names {names}")
        _validate_parameter_sheet(workbook["Parameters"])
    finally:
        workbook.close()
    streamed = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for worksheet in streamed.worksheets:
            for _row in worksheet.iter_rows():
                pass
    finally:
        streamed.close()


def _validate_table(worksheet, table) -> None:
    from openpyxl.utils import coordinate_to_tuple

    if not table.name or not (table.name[0].isalpha() or table.name[0] == "_"):
        raise RuntimeError(f"invalid table name {table.name}")
    start, end = table.ref.split(":")
    start_row, start_col = coordinate_to_tuple(start)
    end_row, end_col = coordinate_to_tuple(end)
    if end_row <= start_row or end_row > worksheet.max_row or end_col > worksheet.max_column:
        raise RuntimeError(f"table {table.name} reference {table.ref} is invalid")
    headers = [worksheet.cell(start_row, col).value for col in range(start_col, end_col + 1)]
    if any(value in (None, "") for value in headers) or len(headers) != len(set(headers)):
        raise RuntimeError(f"table {table.name} headers are missing or duplicated")


def round_trip_workbook(source: Path, copy_path: Path) -> None:
    workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    try:
        workbook.save(copy_path)
    finally:
        workbook.close()
    validate_ooxml_workbook(copy_path)


def _write_sheet(ws, rows: Sequence[Sequence[Any]], table_name: str | None) -> None:
    prepared = _prepare_rows(rows)
    for row in prepared:
        ws.append(row)
    if not prepared:
        return
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = HEADER_BORDER
    for index, header in enumerate(prepared[0], start=1):
        label = header if isinstance(header, str) else ""
        ws.column_dimensions[get_column_letter(index)].width = min(42, max(12, len(label) + 2))
    ws.row_dimensions[1].height = 30
    if len(prepared) >= 2:
        ref = f"A1:{get_column_letter(len(prepared[0]))}{len(prepared)}"
        ws.auto_filter.ref = ref
        if table_name and prepared[1][0] != "NO_RESULTS":
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    if ws.max_row <= 20000:
        _apply_number_formats(ws)


def _prepare_rows(rows: Sequence[Sequence[Any]]) -> list[list[Any]]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    prepared = []
    for row in rows:
        padded = list(row) + [None] * (width - len(row))
        prepared.append([excel_value(value) for value in padded])
    if len(prepared) == 1:
        prepared.append(["NO_RESULTS"] + [None] * (width - 1))
    return prepared


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        return float(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, datetime):
        return iso_utc(value)
    if isinstance(value, str):
        cleaned = _clean_text(value)
        return cleaned[:EXCEL_MAX_TEXT] or None
    cleaned = _clean_text(str(value))
    return cleaned[:EXCEL_MAX_TEXT] or None


def _clean_text(text: str) -> str:
    return _ILLEGAL_XML.sub("", text)


def _apply_number_formats(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    for column, name in enumerate(headers, start=1):
        if not isinstance(name, str):
            continue
        if "Exact" in name:
            continue
        if any(token in name for token in ("UTC", "Turkey", "Row", "Count", "Label", "Class", "Type", "Status", "Flag", "Version", "Bars", "Hours", "Days", "Score", "Trades", "Position", "Sequence", "View", "Reason", "Role", "ID", "Id")):
            continue
        if "Percent" in name:
            number_format = PERCENT_POINTS
        elif any(token in name for token in ("Price", "Open", "High", "Low", "Close", "Reference", "Extreme", "Volume", "Overshoot", "Movement", "Excursion", "Body", "Wick")):
            number_format = PRICE
        else:
            continue
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2, max_row=ws.max_row):
            for item in cell:
                if isinstance(item.value, (int, float)) and not isinstance(item.value, bool):
                    item.number_format = number_format


def _highlight(workbook: Workbook) -> None:
    for worksheet in workbook.worksheets:
        if worksheet.title in LARGE_SHEETS or worksheet.max_row > 2500:
            continue
        headers = {cell.value: cell.column for cell in worksheet[1] if isinstance(cell.value, str)}
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            kind = row[headers["Swing Type"] - 1].value if "Swing Type" in headers else None
            formation = row[headers["Formation Class"] - 1].value if "Formation Class" in headers else None
            completion = row[headers["Completion Type"] - 1].value if "Completion Type" in headers else None
            status = row[headers["Status"] - 1].value if "Status" in headers else None
            fill = None
            if kind == SWING_HIGH:
                fill = HIGH_FILL
            elif kind == "SWING_LOW":
                fill = LOW_FILL
            if formation == COMPACT_SWING:
                fill = COMPACT_FILL
            elif formation == STANDARD_SWING:
                fill = STANDARD_FILL
            if completion == "EXACT_REFERENCE_CLOSE":
                fill = EXACT_FILL
            elif completion == "CROSSED_REFERENCE_CLOSE":
                fill = CROSS_FILL
            if isinstance(status, str) and status.startswith("REJECTED_WIDTH"):
                fill = FAIL_FILL
            method = row[headers["Detection Method"] - 1].value if "Detection Method" in headers else None
            found = row[headers["Found Through Alternative Method"] - 1].value if "Found Through Alternative Method" in headers else None
            if method == ALTERNATIVE_METHOD or found is True:
                fill = ALTERNATIVE_FILL
            if fill is not None:
                row[0].fill = fill


def _exact(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return format(value, "f")


def _time(value: datetime | None, turkey: bool = False) -> str | None:
    if value is None:
        return None
    return iso_turkey(value) if turkey else iso_utc(value)


def _swings(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[name for name, _getter in SWING_COLUMNS]]
    for swing in swings:
        rows.append([getter(swing) for _name, getter in SWING_COLUMNS])
    return rows


def _formation(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Primary Swing ID", "Raw Swing ID", "Family ID", "Primary Status", "Detection Method", "Found Through Alternative Method",
        "Swing Type", "Formation Class", "Candle Role", "Formation Position",
        "Interior Position", "Open UTC", "Open Turkey", "Close UTC", "Close Turkey",
        "Open", "High", "Low", "Close", "ATR", "ATR Available", "Reference Price", "Extreme Price",
        "Reached Reference", "Volume", "Quote Volume", "Trades",
        "Taker Buy Base", "Taker Buy Quote", "Is Swing Open", "Is Extreme", "Is Swing Close",
        "Is Plateau Member", "Is Representative Extreme", "Distance From Reference",
        "Distance From Extreme", "Known At UTC", "Timing Class",
    ]]
    for swing in result.swings:
        for offset, row in enumerate(range(swing.open_row, swing.close_row + 1)):
            candle = result.bars[row]
            if row == swing.open_row:
                role = "SWING_OPEN"
                interior_position = None
            elif row == swing.close_row:
                role = "SWING_CLOSE"
                interior_position = None
            elif row == swing.extreme_row:
                role = "EXTREME"
                interior_position = row - swing.open_row
            else:
                role = "INTERIOR"
                interior_position = row - swing.open_row
            rows.append([
                swing.primary_id, swing.swing_id, swing.family_id, swing.role, swing.detection_method, swing.found_through_alternative,
                swing.direction, swing.formation_class, role, offset + 1, interior_position,
                iso_utc(candle.open_time), iso_turkey(candle.open_time), iso_utc(candle.close_time), iso_turkey(candle.close_time),
                candle.open, candle.high, candle.low, candle.close, candle.atr, candle.atr_available,
                swing.reference, swing.extreme_price, _returned(swing.direction, candle.close, swing.reference),
                candle.volume, candle.quote_volume, candle.trades,
                candle.taker_base, candle.taker_quote,
                row == swing.open_row,
                row in swing.equal_extreme_rows,
                row == swing.close_row,
                swing.plateau_start <= row <= swing.plateau_end and row in swing.equal_extreme_rows,
                row == swing.extreme_row,
                candle.close - swing.reference,
                candle.close - swing.extreme_price,
                iso_utc(candle.close_time),
                "KNOWN_AT_SWING_CONFIRMATION",
            ])
    return rows


def _audit(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Attempt ID", "Direction", "Detection Method", "Swing Open Row", "Swing Open UTC", "Swing Open Turkey",
        "Reference Price", "Maximum Search Interior Candles", "Actual Search Interior Candles",
        "Search Horizon End Row", "Search Horizon Exhausted", "Search Censored by Dataset End",
        "Maximum Observable Interior Candles", "Missing Search Candles", "Search Terminal Reason",
        "Evaluations Within Horizon", "Evaluations Beyond Horizon",
        "Width Achieved Row", "Width Achieved UTC", "First Return Row",
        "Extreme Row", "Extreme Price", "Interior Count", "Duration Class", "Applicable Rule ID",
        "Applicable Minimum Width Percentage", "Standard Boundary Pass", "Open To Extreme Pass", "Close To Extreme Pass", "Both Widths Pass",
        "Open To Extreme Percent Exact", "Close To Extreme Percent Exact", "Status",
        "Raw Swing ID", "Family ID", "Primary Swing ID", "Consolidation Status",
        "Rejection Or Derivation Reason", "Terminal Flag",
    ]]
    for attempt in result.attempts:
        rejected = None if attempt.status.startswith("CONFIRMED") else attempt.status
        reason = attempt.derivation_reason or rejected
        rows.append([
            attempt.attempt_id, attempt.direction, attempt.detection_method, attempt.open_row, iso_utc(attempt.open_time), iso_turkey(attempt.open_time),
            attempt.reference, attempt.maximum_search_interior, attempt.actual_search_interior,
            attempt.search_horizon_end_row, attempt.search_horizon_exhausted, attempt.search_censored,
            attempt.maximum_observable_interior, attempt.missing_search_candles, attempt.search_terminal_reason,
            attempt.evaluations_within_horizon, attempt.evaluations_beyond_horizon,
            attempt.width_achieved_row, _time(attempt.width_achieved_time), attempt.close_row,
            attempt.extreme_row, attempt.extreme_price, attempt.interior_count, attempt.formation_class, attempt.applicable_rule_id,
            attempt.applicable_minimum_width, "2.00", attempt.open_width_pass, attempt.close_width_pass, attempt.both_widths_passed,
            _exact(attempt.open_to_extreme_percent), _exact(attempt.close_to_extreme_percent), attempt.status,
            attempt.swing_id, attempt.family_id, attempt.primary_swing_id, attempt.consolidation_status,
            reason, attempt.terminal,
        ])
    return rows


def _overlaps(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[
        "Primary Swing ID", "Raw Swing ID", "Extreme Family ID", "Swing Type", "Overlap Classification", "Overlap Count", "Same Type Overlap",
        "Opposite Type Overlap", "Contained Swing Count", "Parent Swing IDs", "Child Swing IDs",
        "Shared Candle Count", "Shared Candle Percent", "Alternating View",
    ]]
    for swing in swings:
        rows.append([
            swing.primary_id, swing.swing_id, swing.family_id, swing.direction, swing.overlap_class, swing.overlap_count, swing.same_type_overlap,
            swing.opposite_type_overlap, swing.contained_count, swing.parent_ids, swing.child_ids,
            swing.shared_candles, swing.shared_percent, swing.alternating_view,
        ])
    return rows


def _forward(swings: Sequence[Swing]) -> list[list[Any]]:
    rows = [[
        "Swing ID", "Next Same Type", "Next Opposite Type", "Bars To Next", "Favorable Excursion",
        "Adverse Excursion", "First Break Row", "First Close Beyond Row", "Retest Extreme Row",
        "Retest Reference Row", "Forward Data Censored", "Forward Censoring Reason", "Timing Class",
    ]]
    for swing in swings:
        rows.append([
            swing.primary_id, swing.next_same_id, swing.next_opposite_id, swing.bars_to_next,
            swing.favorable_excursion, swing.adverse_excursion, swing.first_break_row, swing.first_close_beyond_row,
            swing.retest_extreme_row, swing.retest_reference_row, swing.forward_censored, swing.forward_censor_reason,
            "EX_POST_FORWARD_METRIC",
        ])
    return rows


def _derived(result: AnalysisResult) -> list[list[Any]]:
    rows = [[
        "Derived Raw Swing ID", "Extreme Family ID", "Primary Swing ID", "Swing Type", "Extreme Event ID",
        "Exact Extreme Price", "Raw Formation Class", "Primary Formation Class",
        "Raw Open Row", "Raw Extreme Row", "Raw Close Row",
        "Primary Open Row", "Primary Extreme Row", "Primary Close Row",
        "Raw Open UTC", "Raw Open Turkey", "Raw Extreme UTC", "Raw Extreme Turkey", "Raw Close UTC", "Raw Close Turkey",
        "Raw Total Formation Candles", "Primary Total Formation Candles", "Extra Candle Count",
        "Raw Interior Count", "Primary Interior Count",
        "Raw Applicable Minimum Width", "Primary Applicable Minimum Width",
        "Raw Open To Extreme Percent Exact", "Raw Close To Extreme Percent Exact",
        "Raw Completion Error Percent Exact", "Derivation Reason", "Tie-Break Details",
        "Excluded From Principal Results",
    ]]
    primaries = {item.primary_id: item for item in result.primaries}
    for swing in result.derived:
        primary = primaries[swing.primary_id]
        rows.append([
            swing.swing_id, swing.family_id, swing.primary_id, swing.direction, swing.extreme_event_id,
            swing.extreme_price, swing.formation_class, primary.formation_class,
            swing.open_row, swing.extreme_row, swing.close_row,
            primary.open_row, primary.extreme_row, primary.close_row,
            iso_utc(swing.open_time), iso_turkey(swing.open_time), iso_utc(swing.extreme_open_time), iso_turkey(swing.extreme_open_time),
            iso_utc(swing.close_open_time), iso_turkey(swing.close_open_time),
            swing.total_candles, primary.total_candles, swing.extra_candles,
            swing.interior_count, primary.interior_count,
            swing.required_width, primary.required_width,
            _exact(swing.open_to_extreme_percent), _exact(swing.close_to_extreme_percent),
            _exact(abs(swing.completion_difference_percent)), swing.derivation_reason, swing.tie_break_reason,
            True,
        ])
    return rows


def _executive(result: AnalysisResult, context: dict[str, Any]) -> list[list[Any]]:
    from collections import Counter

    statuses = Counter(item.status for item in result.attempts)
    rows = [["Metric", "Value"]]
    proven = context.get("provenance", {})
    raw_count = len(result.swings)
    primary_count = len(result.primaries)
    derived_count = len(result.derived)
    family_count = len(result.families)
    sizes = [family.primary.family_size for family in result.families]
    extra = [item.extra_candles for item in result.derived]
    reduction = Decimal("0") if raw_count == 0 else Decimal(derived_count) / Decimal(raw_count) * Decimal("100")
    rows.extend([
        ["Headline Policy", "Headline Swing results include PRIMARY_SWING only."],
        ["Detector Version", DETECTOR_VERSION],
        ["Selected Rows", proven.get("rows", len(result.bars))],
        ["First Open UTC", proven.get("first_open_utc", "")],
        ["Last Open UTC", proven.get("last_open_utc", "")],
        ["Last Close UTC", proven.get("last_close_utc", "")],
        ["Pre-2026 Rows Used", 0],
        ["Directional Attempts", len(result.attempts)],
        ["Raw Confirmed Swings", raw_count],
        ["Primary Swings", primary_count],
        ["Derived Same-Extreme Swings", derived_count],
        ["Extreme Families", family_count],
        ["Primary Swing Highs", sum(1 for item in result.primaries if item.direction == SWING_HIGH)],
        ["Primary Swing Lows", sum(1 for item in result.primaries if item.direction != SWING_HIGH)],
        ["Compact Primary", sum(1 for item in result.primaries if item.formation_class == COMPACT_SWING)],
        ["Standard Primary", sum(1 for item in result.primaries if item.formation_class == STANDARD_SWING)],
        ["Alternative Primary", sum(1 for item in result.primaries if item.detection_method == ALTERNATIVE_METHOD)],
        ["Reference Return Raw", sum(1 for item in result.swings if item.detection_method == REFERENCE_METHOD)],
        ["Alternative Raw", sum(1 for item in result.swings if item.detection_method == ALTERNATIVE_METHOD)],
        ["Cross Method Families", sum(1 for item in result.families if item.cross_method)],
        ["Duration Gap Rejections", statuses.get("REJECTED_DURATION_GAP_7_INTERIOR_BARS", 0)],
        ["Width Rejections 3.50", statuses.get("REJECTED_WIDTH_BELOW_3_50", 0)],
        ["Standard Boundary Pass", "2.00"],
        ["Raw To Primary Reduction Percent", reduction],
        ["Average Family Size", Decimal("0") if family_count == 0 else Decimal(raw_count) / Decimal(family_count)],
        ["Maximum Family Size", max(sizes) if sizes else 0],
        ["Families With One Member", sum(1 for size in sizes if size == 1)],
        ["Families With Multiple Members", sum(1 for size in sizes if size > 1)],
        ["Average Extra Candles Removed", Decimal("0") if not extra else Decimal(sum(extra)) / Decimal(len(extra))],
        ["Maximum Extra Candles Removed", max(extra) if extra else 0],
        ["Exact Returns", sum(1 for item in result.primaries if item.completion_type == "EXACT_REFERENCE_CLOSE")],
        ["Crossed Returns", sum(1 for item in result.primaries if item.completion_type == "CROSSED_REFERENCE_CLOSE")],
        ["Plateau Formations", sum(1 for item in result.primaries if item.plateau_length > 1)],
        ["Primary Designation", "POST_DETECTION_SAME_EXTREME_CONSOLIDATION"],
        ["Raw Swing Confirmation", "CAUSAL_AT_SWING_CLOSE"],
    ])
    for status, count in sorted(statuses.items()):
        rows.append([status, count])
    for month, count in context.get("monthly_counts", {}).items():
        rows.append([f"Month {month}", count])
    return rows


def _parameters() -> list[list[Any]]:
    header = [
        "Category", "Parameter Name", "Effective Value", "Data Type", "Unit",
        "Hard Rule or Descriptive", "Formula or Allowed Values", "Description",
        "Source of Parameter", "Used In",
    ]
    rows = [header]
    for item in parameter_registry():
        rows.append([
            item.category, item.name, item.value, item.data_type, item.unit, item.rule,
            item.formula, item.description, item.source, item.used_in,
        ])
    return rows


def _readme_rows(readme: str) -> list[list[Any]]:
    lines = [line for line in readme.splitlines() if line.strip()]
    return [["README Line"], *([line] for line in lines)]


SWING_COLUMNS: list[tuple[str, Any]] = [
    ("Primary Swing ID", lambda swing: swing.primary_id),
    ("Raw Swing ID", lambda swing: swing.swing_id),
    ("Extreme Family ID", lambda swing: swing.family_id),
    ("Family Size", lambda swing: swing.family_size),
    ("Derived Member Count", lambda swing: swing.derived_member_count),
    ("Primary Selection Rank", lambda swing: swing.primary_rank),
    ("Primary Selection Reason", lambda swing: swing.selection_reason),
    ("Tie-Break Reason", lambda swing: swing.tie_break_reason),
    ("Primary Designation Known At UTC", lambda swing: _time(swing.primary_known_at)),
    ("Primary Designation Known At Turkey", lambda swing: _time(swing.primary_known_at, turkey=True)),
    ("Swing ID", lambda swing: swing.primary_id),
    ("Sequence", lambda swing: swing.primary_sequence),
    ("Detection Method", lambda swing: swing.detection_method),
    ("Found Through Alternative Method", lambda swing: swing.found_through_alternative),
    ("Reference Return Required", lambda swing: swing.reference_return_required),
    ("Reference Return Achieved", lambda swing: swing.reference_return_achieved),
    ("Maximum Search Interior Candles", lambda swing: swing.maximum_search_interior),
    ("Actual Search Interior Candles", lambda swing: swing.actual_search_interior),
    ("Search Horizon End Row", lambda swing: swing.search_horizon_end_row),
    ("Search Horizon Exhausted", lambda swing: swing.search_horizon_exhausted),
    ("Search Censored by Dataset End", lambda swing: swing.search_censored),
    ("Maximum Observable Interior Candles", lambda swing: swing.maximum_observable_interior),
    ("Search Terminal Reason", lambda swing: swing.search_terminal_reason),
    ("Swing Type", lambda swing: swing.direction),
    ("Structure Label", lambda swing: swing.structure_label),
    ("Formation Class", lambda swing: swing.formation_class),
    ("Swing Open Row", lambda swing: swing.open_row),
    ("Extreme Row", lambda swing: swing.extreme_row),
    ("Swing Close Row", lambda swing: swing.close_row),
    ("Swing Open UTC", lambda swing: iso_utc(swing.open_time)),
    ("Swing Open Turkey", lambda swing: iso_turkey(swing.open_time)),
    ("Swing Open Close Time UTC", lambda swing: iso_utc(swing.open_close_time)),
    ("Swing Open Close Time Turkey", lambda swing: iso_turkey(swing.open_close_time)),
    ("Extreme UTC", lambda swing: iso_utc(swing.extreme_open_time)),
    ("Extreme Turkey", lambda swing: iso_turkey(swing.extreme_open_time)),
    ("Extreme Close Time UTC", lambda swing: iso_utc(swing.extreme_close_time)),
    ("Extreme Close Time Turkey", lambda swing: iso_turkey(swing.extreme_close_time)),
    ("Swing Close UTC", lambda swing: iso_utc(swing.close_open_time)),
    ("Swing Close Turkey", lambda swing: iso_turkey(swing.close_open_time)),
    ("Swing Close Close Time UTC", lambda swing: iso_utc(swing.confirmed_at)),
    ("Swing Close Close Time Turkey", lambda swing: iso_turkey(swing.confirmed_at)),
    ("Confirmed At UTC", lambda swing: iso_utc(swing.confirmed_at)),
    ("Confirmed At Turkey", lambda swing: iso_turkey(swing.confirmed_at)),
    ("Width Achieved Row", lambda swing: swing.width_achieved_row),
    ("Width Achieved At UTC", lambda swing: _time(swing.width_achieved_time)),
    ("Width Achieved At Turkey", lambda swing: _time(swing.width_achieved_time, turkey=True)),
    ("Swing Reference Price", lambda swing: swing.reference),
    ("Extreme Price", lambda swing: swing.extreme_price),
    ("Swing Close Price", lambda swing: swing.close_price),
    ("Swing Open High", lambda swing: swing.open_candle_high),
    ("Swing Open Low", lambda swing: swing.open_candle_low),
    ("Swing Open Candle Close", lambda swing: swing.open_candle_close),
    ("Extreme Candle Open", lambda swing: swing.extreme_candle_open),
    ("Extreme Candle High", lambda swing: swing.extreme_candle_high),
    ("Extreme Candle Low", lambda swing: swing.extreme_candle_low),
    ("Extreme Candle Close", lambda swing: swing.extreme_candle_close),
    ("Swing Close Candle Open", lambda swing: swing.close_candle_open),
    ("Swing Close Candle High", lambda swing: swing.close_candle_high),
    ("Swing Close Candle Low", lambda swing: swing.close_candle_low),
    ("Open To Extreme Price", lambda swing: swing.open_to_extreme_price),
    ("Open To Extreme Percent", lambda swing: swing.open_to_extreme_percent),
    ("Open To Extreme Percent Exact", lambda swing: _exact(swing.open_to_extreme_percent)),
    ("Close To Extreme Price", lambda swing: swing.close_to_extreme_price),
    ("Close To Extreme Percent", lambda swing: swing.close_to_extreme_percent),
    ("Close To Extreme Percent Exact", lambda swing: _exact(swing.close_to_extreme_percent)),
    ("Applicable Minimum Width Percent", lambda swing: swing.required_width),
    ("Applicable Rule ID", lambda swing: swing.applicable_rule_id),
    ("Open To Extreme Pass", lambda swing: swing.open_width_pass),
    ("Close To Extreme Pass", lambda swing: swing.close_width_pass),
    ("Both Width Rules Passed", lambda swing: swing.both_widths_passed),
    ("Completion Type", lambda swing: swing.completion_type),
    ("Completion Difference Price", lambda swing: swing.completion_difference_price),
    ("Completion Difference Percent", lambda swing: swing.completion_difference_percent),
    ("Completion Overshoot Price", lambda swing: swing.overshoot_price),
    ("Completion Overshoot Percent", lambda swing: swing.overshoot_percent),
    ("Interior Candle Count", lambda swing: swing.interior_count),
    ("Total Formation Candles", lambda swing: swing.total_candles),
    ("Formation Left Bars", lambda swing: swing.left_bars),
    ("Formation Right Bars", lambda swing: swing.right_bars),
    ("Left Duration Hours", lambda swing: swing.left_duration_hours),
    ("Right Duration Hours", lambda swing: swing.right_duration_hours),
    ("Extreme Position", lambda swing: swing.extreme_position),
    ("Interior Extreme Position", lambda swing: swing.interior_position),
    ("Plateau Start Row", lambda swing: swing.plateau_start),
    ("Plateau End Row", lambda swing: swing.plateau_end),
    ("Plateau Length", lambda swing: swing.plateau_length),
    ("Equal Extreme Count", lambda swing: swing.equal_extreme_count),
    ("Duration Hours", lambda swing: swing.duration_hours),
    ("Coverage Hours", lambda swing: swing.coverage_hours),
    ("Duration Days", lambda swing: swing.duration_days),
    ("Bullish Candles", lambda swing: swing.bullish_candles),
    ("Bearish Candles", lambda swing: swing.bearish_candles),
    ("Doji Candles", lambda swing: swing.doji_candles),
    ("Up Closes", lambda swing: swing.up_closes),
    ("Down Closes", lambda swing: swing.down_closes),
    ("Path Length", lambda swing: swing.path_length),
    ("Net Close Movement", lambda swing: swing.net_close_movement),
    ("Close Volatility", lambda swing: swing.close_volatility),
    ("Direction Changes", lambda swing: swing.direction_changes),
    ("Open Body", lambda swing: swing.open_body),
    ("Extreme Body", lambda swing: swing.extreme_body),
    ("Close Body", lambda swing: swing.close_body),
    ("Base Volume", lambda swing: swing.base_volume),
    ("Quote Volume", lambda swing: swing.quote_volume),
    ("Trades", lambda swing: swing.trades),
    ("Average Volume", lambda swing: swing.average_volume),
    ("Median Volume", lambda swing: swing.median_volume),
    ("Overlap Classification", lambda swing: swing.overlap_class),
    ("Contained Swing Count", lambda swing: swing.contained_count),
    ("Quality Score", lambda swing: swing.quality_score),
    ("Quality Label", lambda swing: swing.quality_label),
    ("Next Same Type", lambda swing: swing.next_same_id),
    ("Next Opposite Type", lambda swing: swing.next_opposite_id),
    ("Favorable Excursion", lambda swing: swing.favorable_excursion),
    ("Adverse Excursion", lambda swing: swing.adverse_excursion),
    ("Forward Data Censored", lambda swing: swing.forward_censored),
    ("Forward Censoring Reason", lambda swing: swing.forward_censor_reason),
    ("Inside 2026 Window", lambda swing: swing.inside_2026_window),
    ("Detector Version", lambda swing: DETECTOR_VERSION),
]
