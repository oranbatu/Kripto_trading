"""Workbook writer for the BTCUSDT 30-minute special swing detector."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.swing_open_close_30m_special_v1.config import (
    FORMATION_CLASS,
    MIN_BOUNDARY,
    RESULT_COLUMNS,
    SHEETS,
    SWING_HIGH,
    USER_PARAMETERS,
)
from detectors.swing_open_close_30m_special_v1.engine import Analysis, candle_direction, turkey_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
HIGH_FILL = PatternFill("solid", fgColor="F4CCCC")
HIGH_ALT = PatternFill("solid", fgColor="F8E0E0")
LOW_FILL = PatternFill("solid", fgColor="D9EAD3")
LOW_ALT = PatternFill("solid", fgColor="E7F3E4")
REFERENCE_FILL = PatternFill("solid", fgColor="D6EAF8")
PEAK_FILL = PatternFill("solid", fgColor="FDEDEC")
DIP_FILL = PatternFill("solid", fgColor="E8F8F5")
WRAP = Alignment(wrap_text=True, vertical="center")
EIGHT = Decimal("0.00000001")
PRICE_COLUMNS = (5, 6, 7, 8, 10, 12, 14, 16)
REFERENCE_COLUMNS = range(4, 9)
PEAK_COLUMNS = range(9, 13)
DIP_COLUMNS = range(13, 17)
TABLE_NAMES = {
    "Special Swings": "SpecialSwings",
    "Swing Highs": "SwingHighs",
    "Swing Lows": "SwingLows",
    "Parameters": "DetectorParameters",
}
FORBIDDEN_HEADERS = (
    "SWING ID",
    "UTC",
    "BOUNDARY",
    "EXTREMUM",
    "COMPATIBILITY",
    "VALIDATION",
    "DIAGNOSTIC",
    "RAW ID",
    "FAMILY",
)


def shown(value: Decimal) -> Decimal:
    return value.quantize(EIGHT, rounding=ROUND_HALF_UP)


def _optional(value) -> Decimal | None:
    return None if value is None else shown(value)


def _time_price(bars, row, price) -> tuple:
    if row is None or price is None:
        return None, None
    return turkey_text(bars[row].open_time), shown(price)


def _row(item, bars) -> list:
    opened = bars[item.open_row]
    closed = bars[item.close_row]
    reference = None if item.body_reference_row is None else bars[item.body_reference_row]
    peak_1_time, peak_1_price = _time_price(bars, item.peak_1_row, item.peak_1_price)
    peak_2_time, peak_2_price = _time_price(bars, item.peak_2_row, item.peak_2_price)
    dip_1_time, dip_1_price = _time_price(bars, item.dip_1_row, item.dip_1_price)
    dip_2_time, dip_2_price = _time_price(bars, item.dip_2_row, item.dip_2_price)
    return [
        item.direction,
        turkey_text(opened.open_time),
        turkey_text(closed.open_time),
        None if reference is None else turkey_text(reference.open_time),
        _optional(None if reference is None else reference.open),
        _optional(None if reference is None else reference.high),
        _optional(None if reference is None else reference.low),
        _optional(None if reference is None else reference.close),
        peak_1_time,
        peak_1_price,
        peak_2_time,
        peak_2_price,
        dip_1_time,
        dip_1_price,
        dip_2_time,
        dip_2_price,
    ]


def _paint(sheet, rows: list[list]) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, row in enumerate(rows, start=2):
        fill = (HIGH_ALT if index % 2 == 0 else HIGH_FILL) if row[0] == SWING_HIGH else (LOW_ALT if index % 2 == 0 else LOW_FILL)
        for cell in sheet[index]:
            cell.fill = fill
        for column in REFERENCE_COLUMNS:
            sheet.cell(index, column).fill = REFERENCE_FILL
        for column in PEAK_COLUMNS:
            sheet.cell(index, column).fill = PEAK_FILL
        for column in DIP_COLUMNS:
            sheet.cell(index, column).fill = DIP_FILL
        for column in PRICE_COLUMNS:
            sheet.cell(index, column).number_format = "0.00000000"
    sheet.freeze_panes = "A2"
    last_row = max(1, sheet.max_row)
    sheet.auto_filter.ref = f"A1:P{last_row}"
    sheet.row_dimensions[1].height = 32
    for column in range(1, 17):
        sheet.column_dimensions[get_column_letter(column)].width = 28
        sheet.column_dimensions[get_column_letter(column)].hidden = False
    table = Table(displayName=TABLE_NAMES[sheet.title], ref=f"A1:P{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def write_workbook(path: Path, result: Analysis) -> None:
    if len(RESULT_COLUMNS) != 16:
        raise RuntimeError("result contract is not 16 columns")
    book = Workbook()
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=False)
    book.remove(book.active)
    ordered = sorted(result.displayed, key=lambda item: (item.open_row, item.close_row, item.swing_id))
    rows = [_row(item, result.bars) for item in ordered]
    for title, selected in (
        ("Special Swings", rows),
        ("Swing Highs", [row for row in rows if row[0] == SWING_HIGH]),
        ("Swing Lows", [row for row in rows if row[0] != SWING_HIGH]),
    ):
        sheet = book.create_sheet(title)
        sheet.append(list(RESULT_COLUMNS))
        for row in selected:
            if len(row) != 16:
                raise RuntimeError("result row is not 16 columns")
            sheet.append(row)
        _paint(sheet, selected)
    parameters = book.create_sheet("Parameters")
    parameters.append(["Parameter", "Value"])
    for name, value in USER_PARAMETERS:
        parameters.append([name, value])
    for cell in parameters[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    parameters.freeze_panes = "A2"
    parameters.auto_filter.ref = f"A1:B{parameters.max_row}"
    parameters.column_dimensions["A"].width = 52
    parameters.column_dimensions["B"].width = 64
    parameter_table = Table(displayName="DetectorParameters", ref=f"A1:B{parameters.max_row}")
    parameter_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    parameters.add_table(parameter_table)
    if book.sheetnames != list(SHEETS):
        raise RuntimeError(f"sheet order {book.sheetnames}")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def _cell_decimal(value) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return Decimal(format(value, ".8f"))


def _blank(value) -> bool:
    return value in (None, "")


def _validate_result_row(row, item, bars) -> None:
    if len([value for value in row if not _blank(value) or True]) < 16:
        raise RuntimeError("result row is shorter than 16 columns")
    if any(not _blank(value) for value in row[16:]):
        raise RuntimeError("result row contains a value outside columns A:P")
    opened = bars[item.open_row]
    closed = bars[item.close_row]
    if item.formation_class != FORMATION_CLASS or item.interior < 6 or item.interior > 9:
        raise RuntimeError("confirmed interior is outside 6-9")
    if item.open_width_percent < MIN_BOUNDARY or item.close_width_percent < MIN_BOUNDARY:
        raise RuntimeError("confirmed boundary is below 1.00")
    if row[0] != item.direction:
        raise RuntimeError("swing type does not match the confirmed structure")
    if row[1] != turkey_text(opened.open_time) or row[2] != turkey_text(closed.open_time):
        raise RuntimeError("Turkey open timestamps do not match the source candles")
    if "UTC" in str(row[1]) or "UTC" in str(row[2]):
        raise RuntimeError("UTC timestamp leaked into a result sheet")
    if item.direction == SWING_HIGH:
        if candle_direction(opened.open, opened.close) != "BULLISH" or candle_direction(closed.open, closed.close) != "BEARISH":
            raise RuntimeError("swing high direction is wrong")
    elif candle_direction(opened.open, opened.close) != "BEARISH" or candle_direction(closed.open, closed.close) != "BULLISH":
        raise RuntimeError("swing low direction is wrong")
    if item.body_reference_row is None:
        if any(not _blank(row[index]) for index in range(3, 16)):
            raise RuntimeError("missing reference fabricated a result value")
        return
    candle = bars[item.body_reference_row]
    if row[3] != turkey_text(candle.open_time):
        raise RuntimeError("reference Turkey time does not match the source candle")
    compared = ((4, candle.open), (5, candle.high), (6, candle.low), (7, candle.close))
    for index, value in compared:
        if _cell_decimal(row[index]) != shown(value):
            raise RuntimeError("reference OHLC does not match the source candle")
    if item.body_reference_price != candle.close:
        raise RuntimeError("reference price is not the reference candle Close")
    if item.direction == SWING_HIGH:
        if any(not _blank(row[index]) for index in (12, 13, 14, 15)):
            raise RuntimeError("swing high contains a dip value")
        _validate_extreme(row, 8, bars, item.peak_1_row, item.peak_1_price, "high")
        _validate_extreme(row, 10, bars, item.peak_2_row, item.peak_2_price, "high")
        return
    if any(not _blank(row[index]) for index in (8, 9, 10, 11)):
        raise RuntimeError("swing low contains a peak value")
    _validate_extreme(row, 12, bars, item.dip_1_row, item.dip_1_price, "low")
    _validate_extreme(row, 14, bars, item.dip_2_row, item.dip_2_price, "low")


def _validate_extreme(row, start: int, bars, owner_row, price, field: str) -> None:
    candle = bars[owner_row]
    expected_price = candle.high if field == "high" else candle.low
    if row[start] != turkey_text(candle.open_time):
        raise RuntimeError("segment owner Turkey time does not match the source candle")
    if _cell_decimal(row[start + 1]) != shown(expected_price) or _cell_decimal(row[start + 1]) != shown(price):
        raise RuntimeError("segment price does not match the owner candle")


def _contract(sheet) -> None:
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=16))]
    if tuple(headers) != RESULT_COLUMNS or len(headers) != 16:
        raise RuntimeError(f"{sheet.title} headers differ")
    if sheet.max_column != 16:
        raise RuntimeError(f"{sheet.title} has {sheet.max_column} columns")
    joined = " ".join(str(header).upper() for header in headers)
    if any(token in joined for token in FORBIDDEN_HEADERS):
        raise RuntimeError(f"{sheet.title} contains a forbidden header")
    for row in sheet.iter_rows(min_col=17, max_col=max(sheet.max_column, 17), values_only=True):
        if any(not _blank(value) for value in row):
            raise RuntimeError(f"{sheet.title} contains a value outside columns A:P")
    if list(sheet.merged_cells.ranges):
        raise RuntimeError(f"{sheet.title} contains merged cells")
    if len(sheet.tables) != 1:
        raise RuntimeError(f"{sheet.title} table count is {len(sheet.tables)}")
    table = next(iter(sheet.tables.values()))
    if table.ref != f"A1:P{sheet.max_row}":
        raise RuntimeError(f"{sheet.title} table range is {table.ref}")
    for key, dimension in sheet.column_dimensions.items():
        index = column_index_from_string(getattr(dimension, "index", None) or key)
        if dimension.hidden or index > 16:
            raise RuntimeError(f"{sheet.title} has a hidden or extra column")
    for dimension in sheet.row_dimensions.values():
        if dimension.hidden:
            raise RuntimeError(f"{sheet.title} has a hidden row")


def validate_workbook(path: Path, result: Analysis) -> dict:
    normal = load_workbook(path, read_only=False, data_only=False)
    readonly = load_workbook(path, read_only=True, data_only=False)
    data_only = load_workbook(path, read_only=True, data_only=True)
    try:
        if normal.sheetnames != list(SHEETS) or readonly.sheetnames != list(SHEETS) or data_only.sheetnames != list(SHEETS):
            raise RuntimeError(f"sheets {normal.sheetnames}")
        if any(sheet.sheet_state != "visible" for sheet in normal.worksheets):
            raise RuntimeError("workbook contains a hidden worksheet")
        if normal.vba_archive is not None or list(normal._external_links):
            raise RuntimeError("workbook contains macros or external links")
        for book in (normal, readonly, data_only):
            for name in ("Special Swings", "Swing Highs", "Swing Lows"):
                header = next(book[name].iter_rows(min_row=1, max_row=1, max_col=16, values_only=True))
                if tuple(header) != RESULT_COLUMNS:
                    raise RuntimeError(f"{name} columns differ")
        for name in ("Special Swings", "Swing Highs", "Swing Lows"):
            _contract(normal[name])
        parameter_header = next(normal["Parameters"].iter_rows(min_row=1, max_row=1, max_col=2, values_only=True))
        if parameter_header != ("Parameter", "Value") or normal["Parameters"].max_column != 2:
            raise RuntimeError("parameters sheet is not limited to Parameter and Value")
        if len(normal["Parameters"].tables) != 1:
            raise RuntimeError("parameters sheet table count is wrong")
        parameter_table = next(iter(normal["Parameters"].tables.values()))
        if not parameter_table.ref.startswith("A1:B"):
            raise RuntimeError("parameters table is not limited to columns A:B")
        values = {row[0]: row[1] for row in normal["Parameters"].iter_rows(min_row=2, max_col=2, values_only=True)}
        if values != dict(USER_PARAMETERS):
            raise RuntimeError("parameters sheet does not match the configured list")
        blob = " ".join(str(item) for item in values.values())
        if "0.80%" in blob or values["Maximum Interior Candles"] in {"10", "20"} or values["Minimum Interior Candles"] == "7":
            raise RuntimeError("obsolete boundary or horizon is present")
        if "7–10" in blob or "7_TO_10" in blob or "BELOW_7" in blob or "ABOVE_10" in blob:
            raise RuntimeError("obsolete 7-10 configuration is present")
        if "Lowest Close Among Bearish Interior Candles" in blob or "Highest Close Among Bullish Interior Candles" in blob:
            raise RuntimeError("obsolete single-candle reference rule is present")
        if values["Swing High Reference-To-Close Rule"] != "Reference Close >= Swing Close Close":
            raise RuntimeError("swing high reference-to-close rule is missing")
        if values["Swing Low Reference-To-Close Rule"] != "Reference Close <= Swing Close Close":
            raise RuntimeError("swing low reference-to-close rule is missing")
        if values["Failed Compatibility Binds"] != "No" or values["Search Continues After Compatibility Failure"] != "Yes":
            raise RuntimeError("compatibility continuation rule is wrong")
        if values["Search Extension Beyond Nine Interiors"] != "No" or values["Second-Best Reference May Replace Failed Min/Max"] != "No":
            raise RuntimeError("compatibility search limit is wrong")
        if "MINIMAL_16_COLUMN_EXCEL_V12" not in values["Configuration Version"]:
            raise RuntimeError("configuration version is not V12")
        if values["Minimum Boundary Percent"] != "1.00%" or "0.90%" in blob:
            raise RuntimeError("minimum boundary is not 1.00%")
        if values["Result Worksheet Column Count"] != "16" or values["Hidden Result Columns"] != "None":
            raise RuntimeError("result column contract is wrong")
        if values["Extra Result Columns"] != "None":
            raise RuntimeError("extra result columns are configured")
        if values["Peak/Dip Fields Are Hard Filters"] != "No" or values["Peak/Dip Used For Boundary Calculation"] != "No":
            raise RuntimeError("peak and dip fields are configured as filters")
        if values["Peak/Dip Used For Primary Selection"] != "No":
            raise RuntimeError("peak and dip fields are configured as primary inputs")
        if values["Pre-Close Interior Candle Eligible As Reference"] != "No" or values["Reference Row Maximum"] != "close_row - 2":
            raise RuntimeError("pre-close reference exclusion is missing")
        if values["Reference OHLC Exported"] != "Open, High, Low, Close":
            raise RuntimeError("reference OHLC export is missing")
        if values["First Eligible Swing Close Offset"] != "open_row + 7" or values["Final Eligible Swing Close Offset"] != "open_row + 10":
            raise RuntimeError("close offsets are not +7 and +10")
        ordered = sorted(result.displayed, key=lambda item: (item.open_row, item.close_row, item.swing_id))
        grouped = {
            "Special Swings": ordered,
            "Swing Highs": [item for item in ordered if item.direction == SWING_HIGH],
            "Swing Lows": [item for item in ordered if item.direction != SWING_HIGH],
        }
        counts = {}
        for sheet_name, items in grouped.items():
            rows = [row for row in normal[sheet_name].iter_rows(min_row=2, max_col=16, values_only=True) if row[0]]
            counts[sheet_name] = len(rows)
            if len(rows) != len(items):
                raise RuntimeError(f"{sheet_name} row count {len(rows)} != {len(items)}")
            for row, item in zip(rows, items):
                _validate_result_row(row, item, result.bars)
        return {
            "sheets": list(normal.sheetnames),
            "rows": counts,
            "readable": True,
            "macros": False,
            "external_links": False,
            "result_sheet_column_count": 16,
            "result_sheet_last_column": "P",
            "result_sheet_wrong_header_count": 0,
            "result_sheet_wrong_header_order_count": 0,
            "result_sheet_extra_column_count": 0,
            "result_sheet_hidden_column_count": 0,
            "result_sheet_technical_column_count": 0,
            "result_sheet_utc_column_count": 0,
            "swing_high_nonblank_dip_field_count": 0,
            "swing_low_nonblank_peak_field_count": 0,
            "reference_ohlc_mismatch_count": 0,
            "peak_price_mismatch_count": 0,
            "dip_price_mismatch_count": 0,
            "parameters_sheet_column_count": 2,
            "hidden_result_column_count": 0,
            "extra_result_column_count": 0,
        }
    finally:
        normal.close()
        readonly.close()
        data_only.close()
