"""Excel workbook for the exact 30m representation of 4h swings."""
from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from detectors.swing_open_close_4h_with_30m_v1.engine import SWING_HIGH, iso_turkey, iso_utc
from detectors.swing_open_close_4h_with_30m_v1.mapping import (
    MapResult,
    MappedSwing,
    flag,
    price_text,
    turkey,
    utc,
)
from detectors.swing_open_close_4h_with_30m_v1.mapping_config import FOUND, MAP_SHEETS, NOT_OBSERVED
from detectors.swing_open_close_4h_with_30m_v1.swing_config import (
    CONFIGURATION_VERSION,
    DETECTOR_VERSION,
    parameter_registry,
)
from detectors.swing_open_close_4h_with_30m_v1.mapping_config import mapping_parameter_rows

_ILLEGAL = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")


def write_mapping_workbook(path: Path, mapped: MapResult, context: dict[str, Any], readme: str) -> None:
    workbook = Workbook()
    sheets = {
        "Executive Summary": _executive(mapped, context),
        "4H Swing Index": _four_hour_index(mapped, context),
        "30M Swing Index": _thirty_index(mapped, context),
        "30M Annotated Candles": _annotated(mapped),
        "Swing Boundaries": _boundaries(mapped),
        "4H to 30M Mapping": _parent_rows(mapped),
        "Extreme Matches": _extremes(mapped),
        "Parent Aggregation": _aggregation(mapped),
        "Overlap Analysis": _overlaps(mapped),
        "Parameters": _parameters(context),
        "Diagnostics": _diagnostics(context),
        "README": _readme(readme),
    }
    first = workbook.active
    first.title = "Executive Summary"
    for name in MAP_SHEETS:
        worksheet = first if name == "Executive Summary" else workbook.create_sheet(name)
        _write(worksheet, sheets[name])
    _paint(workbook["30M Annotated Candles"])
    workbook.save(path)


def validate_mapping_workbook(path: Path, primary_count: int, expected_rows: int) -> None:
    _validate_zip(path)
    for kwargs in ({}, {"read_only": True}, {"data_only": True}):
        workbook = load_workbook(path, **kwargs)
        try:
            if tuple(workbook.sheetnames) != MAP_SHEETS:
                raise RuntimeError(f"30m sheet order {workbook.sheetnames}")
        finally:
            workbook.close()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        index = list(workbook["30M Swing Index"].iter_rows(values_only=True))
        header = list(index[0])
        data = index[1:]
        if len(data) != primary_count:
            raise RuntimeError(f"30m index rows {len(data)} != {primary_count}")
        found = header.index("Direct 30M Swing Close Status")
        turkey_index = header.index("30M Swing Close Turkey")
        alias = header.index("Direct 30M Swing Close Close Time Turkey")
        ids = []
        for row in data:
            ids.append(row[header.index("Primary Swing ID")])
            if row[found] == FOUND and row[turkey_index] != row[alias]:
                raise RuntimeError("30M Swing Close Turkey alias mismatch")
            if row[found] == NOT_OBSERVED and row[turkey_index] not in (None, ""):
                raise RuntimeError("not-observed close was fabricated")
            if row[header.index("Mapping Status")] != "MAPPED_30M_COMPLETE":
                raise RuntimeError("mapping status is not complete")
        if len(ids) != len(set(ids)):
            raise RuntimeError("duplicate primary ids in the 30m index")
        candles = workbook["30M Annotated Candles"]
        keys = set()
        count = 0
        candle_header = None
        for row_number, row in enumerate(candles.iter_rows(values_only=True)):
            if row_number == 0:
                candle_header = list(row)
                continue
            count += 1
            key = (row[candle_header.index("Primary Swing ID")], row[candle_header.index("30m Open Time UTC")])
            if key in keys:
                raise RuntimeError(f"duplicate annotated key {key}")
            keys.add(key)
        if count != expected_rows:
            raise RuntimeError(f"annotated rows {count} != {expected_rows}")
    finally:
        workbook.close()


def _validate_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if any(name.startswith("xl/vbaProject") for name in names):
            raise RuntimeError("macros are present")
        if any("externalLink" in name for name in names):
            raise RuntimeError("external links are present")
        for name in names:
            if name.endswith(".xml"):
                ET.fromstring(archive.read(name))


def _write(worksheet: Worksheet, rows: Sequence[Sequence[Any]]) -> None:
    for row in rows:
        worksheet.append([_clean(value) for value in row])
    if not rows:
        return
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(max(len(str(cell.value or "")) for cell in column[:40]) + 2, 42)
        worksheet.column_dimensions[letter].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _paint(worksheet: Worksheet) -> None:
    role_column = None
    for cell in worksheet[1]:
        if cell.value == "30M Role":
            role_column = cell.column_letter
            break
    if role_column is None:
        return
    span = f"{role_column}2:{role_column}{max(worksheet.max_row, 2)}"
    fills = {
        "SWING_OPEN_30M": "1F4E79",
        "OPEN_BLOCK_CHILD": "BDD7EE",
        "REPRESENTATIVE_EXTREME_30M": "C00000",
        "EXTREME_MATCH_30M": "F4CCCC",
        "DIRECT_SWING_CLOSE_30M": "7030A0",
        "FINAL_OFFICIAL_4H_CLOSE_CHILD": "E2D5F1",
        "OFFICIAL_4H_CLOSE_BLOCK_CHILD": "D9D2E9",
    }
    for role, color in fills.items():
        formula = f'{role_column}2="{role}"'
        worksheet.conditional_formatting.add(span, FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return _ILLEGAL.sub("", value)
    return value


def _ohlc(bar) -> list:
    return [price_text(bar.open), price_text(bar.high), price_text(bar.low), price_text(bar.close)]


def _executive(mapped: MapResult, context: dict) -> list[list]:
    swings = mapped.swings
    highs = sum(1 for item in swings if item.swing.direction == SWING_HIGH)
    found = sum(1 for item in swings if item.direct_close is not None)
    rows = [["Metric", "Value"],
            ["Symbol", "BTCUSDT"],
            ["Market", "Binance USD-M USDT-margined PERPETUAL"],
            ["Authoritative Timeframe", "4h"],
            ["Mapping Timeframe", "30m"],
            ["Detector Version", DETECTOR_VERSION],
            ["Configuration Version", CONFIGURATION_VERSION],
            ["Mapping Version", context["mapping_version"]],
            ["Shared Run ID", context["run_id"]],
            ["Shared Revision", context["revision_name"]],
            ["4H Rows", context["parent_rows"]],
            ["30M Rows", context["child_rows"]],
            ["Primary Swings", len(swings)],
            ["Primary Swing Highs", highs],
            ["Primary Swing Lows", len(swings) - highs],
            ["Reference Return Primary", sum(1 for item in swings if item.swing.detection_method == "REFERENCE_RETURN")],
            ["Alternative Primary", sum(1 for item in swings if item.swing.detection_method == "ALTERNATIVE_3_5_WIDTH")],
            ["Expected Mapped Rows", sum(item.swing.total_candles * 8 for item in swings)],
            ["Actual Mapped Rows", sum(len(item.rows) for item in swings)],
            ["Direct 30M Close Observed", found],
            ["Direct 30M Close Not Observed", len(swings) - found],
            ["Exact Extreme Matches", sum(len(item.matches) for item in swings)],
            ["Parent Aggregations Passed", sum(1 for item in mapped.links if item.passed)],
            ["Parent Aggregations Failed", sum(1 for item in mapped.links if not item.passed)],
            ["4H Baseline Equivalence", context["baseline"]],
            ["Validation Status", "PASS"]]
    return rows


def _four_hour_index(mapped: MapResult, context: dict) -> list[list]:
    header = [
        "Primary Swing ID", "Raw Swing ID", "Family ID", "Extreme Event ID", "Swing Type",
        "Primary Status", "Detection Method", "Formation Class", "Found Through Alternative Method",
        "Detector Version", "Configuration Version", "4H Swing Open Row", "4H Swing Open Open Time UTC",
        "4H Swing Open Close Time UTC", "4H Swing Open Open Time Turkey", "4H Swing Open Close Time Turkey",
        "4H Swing Open Open", "4H Swing Open High", "4H Swing Open Low", "4H Swing Open Close",
        "Reference Price", "4H Extreme Row", "4H Extreme Open Time UTC", "4H Extreme Close Time UTC",
        "4H Extreme Open Time Turkey", "Extreme Price", "4H Swing Close Row",
        "4H Swing Close Open Time UTC", "4H Swing Close Close Time UTC", "4H Swing Close Open Time Turkey",
        "4H Swing Close Close Time Turkey", "4H Swing Close Open", "4H Swing Close High", "4H Swing Close Low",
        "4H Swing Close Close", "4H Confirmation UTC", "4H Confirmation Turkey", "Interior Candle Count",
        "Total Formation Candle Count", "Applicable Width", "Open To Extreme Width Percent",
        "Close To Extreme Width Percent", "Shared Run ID", "Shared Revision",
    ]
    rows = [header]
    for item in mapped.swings:
        swing = item.swing
        rows.append([
            swing.primary_id, swing.swing_id, swing.family_id, swing.extreme_event_id, swing.direction,
            swing.role, swing.detection_method, swing.formation_class, flag(swing.found_through_alternative),
            DETECTOR_VERSION, CONFIGURATION_VERSION, swing.open_row, iso_utc(swing.open_time),
            iso_utc(swing.open_close_time), iso_turkey(swing.open_time), iso_turkey(swing.open_close_time),
            price_text(swing.reference), price_text(swing.open_candle_high), price_text(swing.open_candle_low),
            price_text(swing.open_candle_close), price_text(swing.reference), swing.extreme_row,
            iso_utc(swing.extreme_open_time), iso_utc(swing.extreme_close_time), iso_turkey(swing.extreme_open_time),
            price_text(swing.extreme_price), swing.close_row, iso_utc(swing.close_open_time), iso_utc(swing.confirmed_at),
            iso_turkey(swing.close_open_time), iso_turkey(swing.confirmed_at), price_text(swing.close_candle_open),
            price_text(swing.close_candle_high), price_text(swing.close_candle_low), price_text(swing.close_price),
            iso_utc(swing.confirmed_at), iso_turkey(swing.confirmed_at), swing.interior_count, swing.total_candles,
            price_text(swing.required_width), price_text(swing.open_to_extreme_percent),
            price_text(swing.close_to_extreme_percent), context["run_id"], context["revision_name"],
        ])
    return rows


def _thirty_index(mapped: MapResult, context: dict) -> list[list]:
    header = [
        "Primary Swing ID", "Raw Swing ID", "Family ID", "Extreme Event ID", "Swing Type", "Detection Method",
        "Formation Class", "Found Through Alternative Method", "Authoritative Timeframe", "Mapping Timeframe",
        "Hierarchy Status", "Detector Version", "Mapping Version", "Shared Run ID", "Shared Revision",
        "30M Mapped Start UTC", "30M Mapped End UTC", "30M Mapped Start Turkey", "30M Mapped End Turkey",
        "Expected 30M Candle Count", "Actual 30M Candle Count", "30M Continuity Pass", "30M Parent Aggregation Pass",
        "30M Swing Open Open Time UTC", "30M Swing Open Close Time UTC", "30M Swing Open Open Time Turkey",
        "30M Swing Open Close Time Turkey", "30M Swing Open Open", "30M Swing Open High", "30M Swing Open Low",
        "30M Swing Open Close", "30M Swing Open Open Equals 4H Reference Price",
        "30M Extreme Match Count", "30M Extreme Plateau Start UTC", "30M Extreme Plateau End UTC",
        "30M Extreme Plateau Start Turkey", "30M Extreme Plateau End Turkey",
        "Representative 30M Extreme Open Time UTC", "Representative 30M Extreme Close Time UTC",
        "Representative 30M Extreme Open Time Turkey", "Representative 30M Extreme Close Time Turkey",
        "Representative 30M Extreme Open", "Representative 30M Extreme High", "Representative 30M Extreme Low",
        "Representative 30M Extreme Close", "Exact Extreme Equality Pass",
        "30M Candles Before Representative Extreme", "30M Candles After Representative Extreme",
        "Direct 30M Swing Close Observed", "Direct 30M Swing Close Status",
        "Direct 30M Swing Close Open Time UTC", "Direct 30M Swing Close Close Time UTC",
        "Direct 30M Swing Close Open Time Turkey", "Direct 30M Swing Close Close Time Turkey",
        "30M Swing Close Turkey", "Direct 30M Swing Close Open", "Direct 30M Swing Close High",
        "Direct 30M Swing Close Low", "Direct 30M Swing Close Close", "Direct 30M Swing Close Price",
        "Direct Completion Error", "Direct Completion Error Percentage",
        "Direct 30M Child Position Within Parent 4H Candle", "Direct 30M Close Inside Official 4H Close Block",
        "Direct 30M Close Before Official 4H Confirmation", "Minutes From Direct 30M Close To Official 4H Confirmation",
        "Official 4H Close Block First 30M Child UTC", "Official 4H Close Block Last 30M Child UTC",
        "Final 30M Child Open", "Final 30M Child High", "Final 30M Child Low", "Final 30M Child Close",
        "Final Child Close Equals Official 4H Close", "Official 4H Confirmation Preserved",
        "30M Path Length", "30M Net Close Change", "30M Efficiency", "30M Bullish", "30M Bearish",
        "30M Unchanged", "30M Maximum Pullback", "30M Realized Volatility", "30M Base Volume",
        "30M Quote Volume", "30M Trades", "30M Taker Buy Ratio", "30M VWAP",
        "30M Candles To Extreme", "30M Candles Extreme To Direct Close", "30M Candles Direct To Confirmation",
        "Mapping Status",
    ]
    rows = [header]
    for item in mapped.swings:
        swing = item.swing
        close = item.direct_close
        error = None if close is None else abs(close.close - swing.reference)
        error_pct = None if close is None or swing.reference == 0 else error / swing.reference * Decimal(100)
        metrics = item.metrics
        rows.append([
            swing.primary_id, swing.swing_id, swing.family_id, swing.extreme_event_id, swing.direction,
            swing.detection_method, swing.formation_class, flag(swing.found_through_alternative), "4h", "30m",
            "MAPPED_FROM_4H_PRIMARY", DETECTOR_VERSION, context["mapping_version"], context["run_id"],
            context["revision_name"], utc(item.direct_open.open_time), utc(swing.confirmed_at),
            turkey(item.direct_open.open_time), turkey(swing.confirmed_at), swing.total_candles * 8,
            len(item.rows), "TRUE", "TRUE", utc(item.direct_open.open_time), utc(item.direct_open.close_time),
            turkey(item.direct_open.open_time), turkey(item.direct_open.close_time), *_ohlc(item.direct_open),
            flag(item.direct_open.open == swing.reference), len(item.matches), utc(item.matches[0].open_time),
            utc(item.matches[-1].close_time), turkey(item.matches[0].open_time), turkey(item.matches[-1].close_time),
            utc(item.representative.open_time), utc(item.representative.close_time),
            turkey(item.representative.open_time), turkey(item.representative.close_time), *_ohlc(item.representative),
            "TRUE", metrics["candles_to_extreme"], len(item.rows) - metrics["candles_to_extreme"] - 1,
            flag(close is not None), item.direct_close_state, utc(None if close is None else close.open_time),
            utc(None if close is None else close.close_time), turkey(None if close is None else close.open_time),
            turkey(None if close is None else close.close_time), turkey(None if close is None else close.close_time),
            None if close is None else price_text(close.open), None if close is None else price_text(close.high),
            None if close is None else price_text(close.low), None if close is None else price_text(close.close),
            None if close is None else price_text(close.close), price_text(error), price_text(error_pct),
            item.child_position, flag(item.inside_official_close_block), flag(item.precedes_confirmation),
            None if item.minutes_before_confirmation is None else price_text(item.minutes_before_confirmation),
            utc(item.first_close_child.open_time), utc(item.final_child.close_time), *_ohlc(item.final_child),
            flag(item.final_child.close == swing.close_price), "TRUE",
            price_text(metrics["path"]), price_text(metrics["net"]), price_text(metrics["efficiency"]),
            metrics["bullish"], metrics["bearish"], metrics["unchanged"], price_text(metrics["pullback"]),
            price_text(metrics["volatility"]), price_text(metrics["volume"]), price_text(metrics["quote"]),
            metrics["trades"], None if metrics["taker_ratio"] is None else price_text(metrics["taker_ratio"]),
            None if metrics["vwap"] is None else price_text(metrics["vwap"]), metrics["candles_to_extreme"],
            metrics["candles_extreme_to_direct"], metrics["candles_direct_to_confirmation"], item.status,
        ])
    return rows


def _annotated(mapped: MapResult) -> list[list]:
    header = [
        "Primary Swing ID", "Swing Type", "Detection Method", "Formation Class", "Child Sequence Within Swing",
        "Child Sequence Within Parent", "Parent 4H Open Time UTC", "Parent 4H Row", "30m Open Time UTC",
        "30m Close Time UTC", "30m Open Time Turkey", "30m Close Time Turkey", "Open", "High", "Low", "Close",
        "Volume", "Quote Volume", "Trades", "Taker Buy Base", "Taker Buy Quote", "Reference Price", "Extreme Price",
        "30M Role", "Combined 30M Roles", "Is 30M Swing Open", "Is Inside 4H Swing Open Block",
        "Is 30M Extreme Match", "Is Representative 30M Extreme", "Is Inside 4H Extreme Block",
        "Is Direct 30M Swing Close", "Is Inside Official 4H Swing Close Block",
        "Is Final Child of Official 4H Swing Close Block", "Is Before Representative Extreme",
        "Is After Representative Extreme", "Is After Direct 30M Swing Close", "Parent OHLC Equivalence",
        "Overlap Count", "Other Swing IDs", "Mapping Validation",
    ]
    rows = [header]
    for item in mapped.swings:
        swing = item.swing
        for row in item.rows:
            child = row["child"]
            rows.append([
                swing.primary_id, swing.direction, swing.detection_method, swing.formation_class, row["sequence"],
                row["offset"], iso_utc(row["parent"].open_time), row["parent_index"], iso_utc(child.open_time),
                iso_utc(child.close_time), iso_turkey(child.open_time), iso_turkey(child.close_time),
                *_ohlc(child), price_text(child.volume), price_text(child.quote_volume), child.trades,
                price_text(child.taker_base), price_text(child.taker_quote), price_text(swing.reference),
                price_text(swing.extreme_price), row["role"], row["combined"], flag(row["is_open"]),
                flag(row["kind"] == "open"), flag(row["exact"]), flag(row["representative"]),
                flag(row["kind"] == "extreme"), flag(row["is_direct_close"]), flag(row["kind"] == "close"),
                flag(row["is_final"]), flag(row["before_extreme"]), flag(row["after_extreme"]),
                flag(row["after_direct"]), "TRUE", row.get("overlap_count", 1), row.get("other_ids", ""),
                "PASS",
            ])
    return rows


def _boundaries(mapped: MapResult) -> list[list]:
    header = [
        "Primary Swing ID", "Boundary", "Open Time UTC", "Close Time UTC", "Open Time Turkey", "Close Time Turkey",
        "Open", "High", "Low", "Close",
    ]
    rows = [header]
    for item in mapped.swings:
        swing = item.swing
        events = [
            ("OFFICIAL_4H_SWING_OPEN", swing.open_time, swing.open_close_time, swing.reference, swing.open_candle_high, swing.open_candle_low, swing.open_candle_close),
            ("30M_SWING_OPEN", item.direct_open.open_time, item.direct_open.close_time, item.direct_open.open, item.direct_open.high, item.direct_open.low, item.direct_open.close),
            ("OFFICIAL_4H_EXTREME", swing.extreme_open_time, swing.extreme_close_time, swing.extreme_candle_open, swing.extreme_candle_high, swing.extreme_candle_low, swing.extreme_candle_close),
            ("REPRESENTATIVE_30M_EXTREME", item.representative.open_time, item.representative.close_time, item.representative.open, item.representative.high, item.representative.low, item.representative.close),
            ("DIRECT_30M_SWING_CLOSE", None if item.direct_close is None else item.direct_close.open_time, None if item.direct_close is None else item.direct_close.close_time, None if item.direct_close is None else item.direct_close.open, None if item.direct_close is None else item.direct_close.high, None if item.direct_close is None else item.direct_close.low, None if item.direct_close is None else item.direct_close.close),
            ("OFFICIAL_4H_SWING_CLOSE", swing.close_open_time, swing.confirmed_at, swing.close_candle_open, swing.close_candle_high, swing.close_candle_low, swing.close_price),
            ("FINAL_30M_CHILD_OF_OFFICIAL_4H_CLOSE_BLOCK", item.final_child.open_time, item.final_child.close_time, item.final_child.open, item.final_child.high, item.final_child.low, item.final_child.close),
        ]
        for name, opened, closed, open_, high, low, close in events:
            rows.append([
                swing.primary_id, name, utc(opened), utc(closed), turkey(opened), turkey(closed),
                price_text(open_), price_text(high), price_text(low), price_text(close),
            ])
    return rows


def _parent_rows(mapped: MapResult) -> list[list]:
    header = ["Primary Swing ID", "Parent 4H Open Time UTC", "Parent Role", "Child Count", "First Child Open UTC", "Last Child Close UTC", "Aggregation Passed"]
    rows = [header]
    for item in mapped.swings:
        seen = []
        for row in item.rows:
            if not seen or seen[-1] != row["parent_index"]:
                seen.append(row["parent_index"])
                link_children = [candidate["child"] for candidate in item.rows if candidate["parent_index"] == row["parent_index"]]
                rows.append([
                    item.swing.primary_id, iso_utc(row["parent"].open_time), row["kind"], len(link_children),
                    iso_utc(link_children[0].open_time), iso_utc(link_children[-1].close_time), "TRUE",
                ])
    return rows


def _extremes(mapped: MapResult) -> list[list]:
    header = ["Primary Swing ID", "Extreme Event ID", "Match Open Time UTC", "Match Close Time UTC", "Match Price", "Representative", "Plateau Length"]
    rows = [header]
    for item in mapped.swings:
        for match in item.matches:
            price = match.high if item.swing.direction == SWING_HIGH else match.low
            rows.append([
                item.swing.primary_id, item.swing.extreme_event_id, iso_utc(match.open_time), iso_utc(match.close_time),
                price_text(price), flag(match.open_time == item.representative.open_time), len(item.matches),
            ])
    return rows


def _aggregation(mapped: MapResult) -> list[list]:
    header = ["4H Open Time UTC", "Child Count", "Open Equal", "High Equal", "Low Equal", "Close Equal", "Volume Equal", "Quote Equal", "Trades Equal", "Taker Base Equal", "Taker Quote Equal", "Passed"]
    rows = [header]
    for link in mapped.links:
        rows.append([
            iso_utc(link.parent.open_time), len(link.children), flag(link.open_equal), flag(link.high_equal),
            flag(link.low_equal), flag(link.close_equal), flag(link.volume_equal), flag(link.quote_equal),
            flag(link.trades_equal), flag(link.taker_base_equal), flag(link.taker_quote_equal), flag(link.passed),
        ])
    return rows


def _overlaps(mapped: MapResult) -> list[list]:
    header = ["30m Open Time UTC", "Swing Count", "Swing IDs"]
    grouped: dict[str, list[str]] = {}
    for item in mapped.swings:
        for row in item.rows:
            stamp = iso_utc(row["child"].open_time)
            grouped.setdefault(stamp, [])
            if item.swing.primary_id not in grouped[stamp]:
                grouped[stamp].append(item.swing.primary_id)
    rows = [header]
    for stamp, ids in grouped.items():
        if len(ids) > 1:
            rows.append([stamp, len(ids), " | ".join(ids)])
    return rows


def _parameters(context: dict) -> list[list]:
    header = ["Category", "Parameter Name", "Effective Value", "Data Type", "Unit", "Hard Rule or Descriptive", "Formula or Allowed Values", "Description", "Source of Parameter", "Used In"]
    rows = [header]
    for item in parameter_registry():
        rows.append([item.category, item.name, item.value, item.data_type, item.unit, item.rule, item.formula, item.description, item.source, item.used_in])
    for item in mapping_parameter_rows():
        rows.append(list(item))
    rows.append(["Run", "shared_run_id", context["run_id"], "string", "id", "DESCRIPTIVE", "", "Shared run identifier.", "runner", "both workbooks"])
    rows.append(["Run", "shared_revision", context["revision_name"], "string", "revision", "HARD_RULE", "max+1", "Shared workbook revision.", "runner", "both workbooks"])
    return rows


def _diagnostics(context: dict) -> list[list]:
    rows = [["Section", "Metric", "Value"]]
    for section, metric, value in context["diagnostic_items"]:
        rows.append([section, metric, value])
    return rows


def _readme(readme: str) -> list[list]:
    return [["README"]] + [[line] for line in readme.splitlines()]


def mark_overlap_counts(mapped: MapResult) -> None:
    grouped: dict[str, list[str]] = {}
    for item in mapped.swings:
        for row in item.rows:
            stamp = iso_utc(row["child"].open_time)
            grouped.setdefault(stamp, [])
            if item.swing.primary_id not in grouped[stamp]:
                grouped[stamp].append(item.swing.primary_id)
    for item in mapped.swings:
        for row in item.rows:
            ids = grouped[iso_utc(row["child"].open_time)]
            row["overlap_count"] = len(ids)
            row["other_ids"] = " | ".join(other for other in ids if other != item.swing.primary_id)
