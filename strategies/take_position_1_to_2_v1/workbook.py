"""Excel report for the 1:2 R take-position backtest."""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.swing_open_close_30m_special_v1.engine import turkey_text
from strategies.take_position_1_to_2_v1.config import (
    CONFIGURATION_VERSION,
    PARAMETERS,
    PERFORMANCE_LABEL,
    RISK_TO_REWARD,
    SHEETS,
    STRATEGY_VERSION,
    TRADE_COLUMNS,
)
from strategies.take_position_1_to_2_v1.models import Trade

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
LOSS_FILL = PatternFill("solid", fgColor="FFCDD2")
OPEN_FILL = PatternFill("solid", fgColor="D6EAF8")
INVALID_FILL = PatternFill("solid", fgColor="F2F2F2")
SHORT_FILL = PatternFill("solid", fgColor="F8E0E0")
LONG_FILL = PatternFill("solid", fgColor="E7F3E4")
NOTE_FONT = Font(bold=True, color="7B241C", name="Calibri", size=14)
WRAP = Alignment(wrap_text=True, vertical="center")
EIGHT = Decimal("0.00000001")
PRICE_INDEXES = {9, 11, 12, 13, 14, 26, 35, 36, 41}
PERCENT_INDEXES = {15, 16}
R_INDEX = 27
OUTCOME_INDEX = 21
HEADLINE = (
    "Resolved Trades",
    "Wins",
    "Losses",
    "Win Rate",
    "Theoretical Gross Break-Even Win Rate",
    "Win Rate Minus Break-Even",
    "Total Gross R",
    "Average Gross R",
    "Average Hours To Exit",
    "Average Hours To Take Profit",
    "Average Hours To Stop Loss",
    "Open/Censored Trades",
)


def shown(value: Decimal | None) -> Decimal | None:
    if value is None:
        return None
    return value.quantize(EIGHT, rounding=ROUND_HALF_UP)


def _iso(moment: datetime | None) -> str | None:
    if moment is None:
        return None
    value = moment.astimezone(timezone.utc)
    if value.microsecond:
        return value.strftime("%Y-%m-%dT%H:%M:%S.") + f"{value.microsecond // 1000:03d}Z"
    return value.strftime("%Y-%m-%dT%H:%M:%SZ")


def _hours(trade: Trade) -> Decimal | None:
    if trade.duration_minutes is None:
        return None
    return Decimal(trade.duration_minutes) / Decimal(60)


def trade_row(trade: Trade, detector_version: str) -> list:
    hours = _hours(trade)
    valid = trade.outcome != "INVALID"
    return [
        trade.trade_id,
        trade.swing_id,
        trade.swing_type,
        trade.direction,
        _iso(trade.signal_time),
        turkey_text(trade.signal_time),
        _iso(trade.entry_time),
        None if trade.entry_time is None else turkey_text(trade.entry_time),
        shown(trade.entry_price),
        "CONFIRMATION_CANDLE_CLOSE" if trade.entry_price is not None else None,
        shown(trade.stop_price),
        shown(trade.target_price),
        shown(trade.risk_price),
        shown(trade.reward_price),
        shown(trade.risk_percent),
        shown(trade.reward_percent),
        RISK_TO_REWARD if valid else None,
        shown(Decimal(2)) if valid else None,
        shown(trade.peak_2_high),
        shown(trade.dip_2_low),
        trade.outcome,
        trade.outcome_group,
        trade.exit_reason,
        _iso(trade.exit_time),
        None if trade.exit_time is None else turkey_text(trade.exit_time),
        shown(trade.exit_price),
        shown(trade.realized_r),
        trade.duration_minutes,
        shown(hours) if hours is not None else None,
        shown(hours / Decimal(24)) if hours is not None else None,
        shown(hours / Decimal("0.5")) if hours is not None else None,
        "Yes" if trade.intrabar_ambiguity else "No",
        "Yes" if trade.stop_gap else "No",
        "Yes" if trade.take_profit_gap else "No",
        shown(trade.mfe_price),
        shown(trade.mae_price),
        shown(trade.mfe_r),
        shown(trade.mae_r),
        "Yes" if trade.censored else "No",
        trade.censoring_reason or None,
        shown(trade.last_price),
        "Yes" if trade.overlapping else "No",
        trade.concurrent_at_entry,
        detector_version,
        STRATEGY_VERSION,
        CONFIGURATION_VERSION,
    ]


def _paint_trades(sheet, rows: list[list]) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, row in enumerate(rows, start=2):
        outcome = row[OUTCOME_INDEX - 1]
        if outcome == "WIN":
            fill = WIN_FILL
        elif outcome in {"LOSS", "LOSS_AMBIGUOUS"}:
            fill = LOSS_FILL
        elif outcome == "INVALID":
            fill = INVALID_FILL
        else:
            fill = OPEN_FILL
        if row[3] == "SHORT" and outcome not in {"WIN", "LOSS", "LOSS_AMBIGUOUS"}:
            fill = SHORT_FILL
        elif row[3] == "LONG" and outcome not in {"WIN", "LOSS", "LOSS_AMBIGUOUS", "INVALID"}:
            fill = LONG_FILL
        for cell in sheet[index]:
            cell.fill = fill
        for column in PRICE_INDEXES:
            sheet.cell(index, column).number_format = "0.00000000"
        for column in PERCENT_INDEXES:
            sheet.cell(index, column).number_format = "0.00000000"
        for column in (18, 27, 29, 30, 31, 37, 38):
            sheet.cell(index, column).number_format = "0.00000000"
    sheet.freeze_panes = "A2"
    last_row = max(1, sheet.max_row)
    last_column = get_column_letter(len(TRADE_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    sheet.row_dimensions[1].height = 32
    for column in range(1, len(TRADE_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    table = Table(displayName=sheet.title.replace(" ", ""), ref=f"A1:{last_column}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)
    if last_row > 1:
        r_column = get_column_letter(R_INDEX)
        sheet.conditional_formatting.add(
            f"{r_column}2:{r_column}{last_row}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=WIN_FILL),
        )
        sheet.conditional_formatting.add(
            f"{r_column}2:{r_column}{last_row}",
            CellIsRule(operator="lessThan", formula=["0"], fill=LOSS_FILL),
        )
        outcome_column = get_column_letter(OUTCOME_INDEX)
        sheet.conditional_formatting.add(
            f"{outcome_column}2:{outcome_column}{last_row}",
            FormulaRule(formula=[f'{outcome_column}2="WIN"'], fill=WIN_FILL),
        )
        sheet.conditional_formatting.add(
            f"{outcome_column}2:{outcome_column}{last_row}",
            FormulaRule(formula=[f'OR({outcome_column}2="LOSS",{outcome_column}2="LOSS_AMBIGUOUS")'], fill=LOSS_FILL),
        )


def _metric_table(sheet, anchor_row: int, title: str, summary: dict, table_name: str) -> int:
    sheet.cell(anchor_row, 1, title).font = Font(bold=True, size=14, name="Calibri")
    for column, header in enumerate(("Metric", "Value"), start=1):
        cell = sheet.cell(anchor_row + 1, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    labels = (
        ("Eligible Signals", "eligible_signal_count"),
        ("Valid Trades", "valid_trade_count"),
        ("Invalid Trades", "invalid_trade_count"),
        ("Entered Trades", "entered_trade_count"),
        ("Resolved Trades", "resolved_trade_count"),
        ("Wins", "win_count"),
        ("Losses", "loss_count"),
        ("Ambiguous Losses", "ambiguous_loss_count"),
        ("Win Rate", "win_rate_percent"),
        ("Loss Rate", "loss_rate_percent"),
        ("Theoretical Gross Break-Even Win Rate", "break_even_win_rate_percent"),
        ("Win Rate Minus Break-Even", "win_rate_minus_break_even_percentage_points"),
        ("Total Gross R", "total_gross_r"),
        ("Average Gross R", "average_gross_r"),
        ("Median Gross R", "median_gross_r"),
        ("Average Hours To Exit", "average_hours_to_exit"),
        ("Median Hours To Exit", "median_hours_to_exit"),
        ("Minimum Hours To Exit", "minimum_hours_to_exit"),
        ("Maximum Hours To Exit", "maximum_hours_to_exit"),
        ("Average Hours To Take Profit", "average_hours_to_take_profit"),
        ("Median Hours To Take Profit", "median_hours_to_take_profit"),
        ("Average Hours To Stop Loss", "average_hours_to_stop_loss"),
        ("Median Hours To Stop Loss", "median_hours_to_stop_loss"),
        ("Fastest Take Profit Hours", "fastest_take_profit_hours"),
        ("Slowest Take Profit Hours", "slowest_take_profit_hours"),
        ("Fastest Stop Loss Hours", "fastest_stop_loss_hours"),
        ("Slowest Stop Loss Hours", "slowest_stop_loss_hours"),
        ("Average Risk Percent", "average_risk_percent"),
        ("Median Risk Percent", "median_risk_percent"),
        ("Minimum Risk Percent", "minimum_risk_percent"),
        ("Maximum Risk Percent", "maximum_risk_percent"),
        ("Average Reward Percent", "average_reward_percent"),
        ("Same 1m Collisions", "same_1m_collision_count"),
        ("Stop Gaps", "stop_gap_count"),
        ("Take Profit Gaps", "take_profit_gap_count"),
        ("Overlapping Trades", "overlap_count"),
        ("Open/Censored Trades", "open_censored_count"),
        ("Unentered/Censored Trades", "unentered_censored_count"),
    )
    for offset, (label, key) in enumerate(labels, start=2):
        value = summary[key]
        sheet.cell(anchor_row + offset, 1, label)
        cell = sheet.cell(anchor_row + offset, 2, None if value is None else (format(value, "f") if isinstance(value, Decimal) else value))
        if label in HEADLINE:
            sheet.cell(anchor_row + offset, 1).font = Font(bold=True, name="Calibri")
            cell.font = Font(bold=True, name="Calibri")
    header_row = anchor_row + 1
    last_row = anchor_row + 1 + len(labels)
    table = Table(displayName=table_name, ref=f"A{header_row}:B{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)
    return last_row + 3


def _key_table(sheet, rows: list[tuple[str, object]], display_name: str) -> None:
    sheet.append(["Item", "Value"])
    for key, value in rows:
        sheet.append([key, "" if value is None else value])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 56
    sheet.column_dimensions["B"].width = 88
    last = max(1, sheet.max_row)
    table = Table(displayName=display_name, ref=f"A1:B{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)


def write_workbook(path: Path, trades: list[Trade], summaries: dict[str, dict], diagnostics: list[tuple[str, object]], detector_version: str) -> None:
    book = Workbook()
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=False)
    book.remove(book.active)
    summary = book.create_sheet("Performance Summary")
    summary["A1"] = "Risk : Reward = 1 : 2"
    summary["A1"].font = NOTE_FONT
    summary["A2"] = "Take Profit = +2R. Stop Loss = -1R."
    summary["A2"].font = NOTE_FONT
    summary["A3"] = f"{PERFORMANCE_LABEL}. Historical simulation only. Not investment advice."
    row_number = 5
    for title in ("ALL TRADES", "SWING HIGH SHORTS", "SWING LOW LONGS"):
        row_number = _metric_table(summary, row_number, title, summaries[title], "Summary" + title.split()[-1].title())
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 28
    groups = {
        "All Trades": trades,
        "Short Trades": [trade for trade in trades if trade.direction == "SHORT"],
        "Long Trades": [trade for trade in trades if trade.direction == "LONG"],
        "Open and Censored": [trade for trade in trades if trade.outcome in {"OPEN_CENSORED", "UNENTERED_CENSORED"}],
        "Invalid Signals": [trade for trade in trades if trade.outcome == "INVALID"],
    }
    for title, selected in groups.items():
        sheet = book.create_sheet(title)
        sheet.append(list(TRADE_COLUMNS))
        rows = [trade_row(trade, detector_version) for trade in selected]
        for row in rows:
            if len(row) != len(TRADE_COLUMNS):
                raise RuntimeError("trade row width changed")
            sheet.append(row)
        _paint_trades(sheet, rows)
        if title == "Open and Censored":
            sheet.cell(1, len(TRADE_COLUMNS) + 1, "Distance To Stop")
            sheet.cell(1, len(TRADE_COLUMNS) + 2, "Distance To Take Profit")
            sheet.cell(1, len(TRADE_COLUMNS) + 3, "Last Available Time UTC")
            for index, trade in enumerate(selected, start=2):
                sheet.cell(index, len(TRADE_COLUMNS) + 1, shown(trade.distance_to_stop))
                sheet.cell(index, len(TRADE_COLUMNS) + 2, shown(trade.distance_to_target))
                sheet.cell(index, len(TRADE_COLUMNS) + 3, _iso(trade.last_available_time))
    parameters = book.create_sheet("Parameters")
    _key_table(parameters, list(PARAMETERS), "StrategyParameters")
    diagnostics_sheet = book.create_sheet("Diagnostics")
    _key_table(diagnostics_sheet, diagnostics, "BacktestDiagnostics")
    readme = book.create_sheet("README")
    _key_table(readme, [
        ("Purpose", "Historical 1:2 R simulation of confirmed Primary 30m special swings."),
        ("Orders", "No live, broker, or exchange orders are created."),
        ("Signal source", "Existing detector engine. The detector workbook is not parsed."),
        ("Entry", "Swing Close Close, after that 30m candle is complete."),
        ("Short stop", "Peak 2 High. Target is exactly two risk units below entry."),
        ("Long stop", "Dip 2 Low. Target is exactly two risk units above entry."),
        ("Payoff", "A normal win is +2R. A normal loss is -1R."),
        ("Exit scan", "Canonical 1m candles strictly after confirmation."),
        ("Same-minute collision", "Stop Loss First. Counted as a loss of -1R."),
        ("Unresolved trades", "Remain censored. They are excluded from win rate."),
        ("Costs", PERFORMANCE_LABEL),
        ("Advice", "This workbook is not investment advice."),
    ], "BacktestReadme")
    if book.sheetnames != list(SHEETS):
        raise RuntimeError(f"sheet order {book.sheetnames}")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def validate_workbook(path: Path, trades: list[Trade], summaries: dict[str, dict]) -> dict:
    normal = load_workbook(path, read_only=False, data_only=False)
    readonly = load_workbook(path, read_only=True, data_only=False)
    data_only = load_workbook(path, read_only=True, data_only=True)
    try:
        if normal.sheetnames != list(SHEETS) or readonly.sheetnames != list(SHEETS) or data_only.sheetnames != list(SHEETS):
            raise RuntimeError(f"sheets {normal.sheetnames}")
        if normal.vba_archive is not None or list(normal._external_links):
            raise RuntimeError("workbook contains macros or external links")
        for name in ("All Trades", "Short Trades", "Long Trades", "Open and Censored", "Invalid Signals"):
            headers = [cell.value for cell in next(normal[name].iter_rows(min_row=1, max_row=1, max_col=len(TRADE_COLUMNS)))]
            if tuple(headers) != TRADE_COLUMNS:
                raise RuntimeError(f"{name} headers differ")
            if not normal[name].tables:
                raise RuntimeError(f"{name} has no table")
        expected = {
            "All Trades": trades,
            "Short Trades": [trade for trade in trades if trade.direction == "SHORT"],
            "Long Trades": [trade for trade in trades if trade.direction == "LONG"],
            "Open and Censored": [trade for trade in trades if trade.outcome in {"OPEN_CENSORED", "UNENTERED_CENSORED"}],
            "Invalid Signals": [trade for trade in trades if trade.outcome == "INVALID"],
        }
        counts = {}
        for name, selected in expected.items():
            rows = [row for row in normal[name].iter_rows(min_row=2, max_col=len(TRADE_COLUMNS), values_only=True) if row[0]]
            counts[name] = len(rows)
            if len(rows) != len(selected):
                raise RuntimeError(f"{name} count {len(rows)} != {len(selected)}")
            for row, trade in zip(rows, selected):
                if row[3] != trade.direction:
                    raise RuntimeError("trade direction does not match the detector swing")
                if name == "Short Trades" and row[3] != "SHORT":
                    raise RuntimeError("short sheet contains a long")
                if name == "Long Trades" and row[3] != "LONG":
                    raise RuntimeError("long sheet contains a short")
                if trade.outcome == "INVALID":
                    continue
                if trade.direction == "SHORT" and (row[19] not in (None, "") or row[18] in (None, "")):
                    raise RuntimeError("short row peak/dip population is wrong")
                if trade.direction == "LONG" and (row[18] not in (None, "") or row[19] in (None, "")):
                    raise RuntimeError("long row peak/dip population is wrong")
                if row[16] != "1:2" or Decimal(str(row[17])) != Decimal("2.00000000"):
                    raise RuntimeError("trade is not labeled 1:2 with a 2R target")
        parameter_values = {row[0]: row[1] for row in normal["Parameters"].iter_rows(min_row=2, max_col=2, values_only=True)}
        if parameter_values["Risk To Reward"] != "1:2" or parameter_values["Normal Take Profit Outcome"] != "+2R":
            raise RuntimeError("parameters are not 1:2")
        if parameter_values["Theoretical Gross Break-Even Win Rate"] != "33.3333%":
            raise RuntimeError("break-even label is wrong")
        if normal["Performance Summary"]["A1"].value != "Risk : Reward = 1 : 2":
            raise RuntimeError("1:2 banner is missing")
        if summaries["ALL TRADES"]["resolved_trade_count"] != counts["All Trades"] - counts["Open and Censored"] - counts["Invalid Signals"]:
            raise RuntimeError("resolved count does not reconcile")
        return {"sheets": list(normal.sheetnames), "rows": counts, "readable": True, "macros": False, "external_links": False}
    finally:
        normal.close()
        readonly.close()
        data_only.close()
