"""Excel report for the 1:1.5 R take-position backtest."""
from __future__ import annotations

import re
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from detectors.swing_open_close_30m_special_v1.engine import turkey_text
from strategies.take_position_1_to_1_5_v1.config import (
    CONFIGURATION_VERSION,
    DATA_RANGE,
    PARAMETERS,
    PERFORMANCE_LABEL,
    REVISION_PATTERN,
    RISK_REWARD_RATIO,
    SAME_BAR_POLICY,
    SHEETS,
    STRATEGY_NAME,
)
from strategies.take_position_1_to_1_5_v1.models import Trade

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
LOSS_FILL = PatternFill("solid", fgColor="FFCDD2")
OPEN_FILL = PatternFill("solid", fgColor="D6EAF8")
AMBIGUOUS_FILL = PatternFill("solid", fgColor="FDEBD0")
GAP_FILL = PatternFill("solid", fgColor="F5CBA7")
NOTE_FONT = Font(bold=True, color="7B241C", name="Calibri", size=14)
WRAP = Alignment(wrap_text=True, vertical="center")
EIGHT = Decimal("0.00000001")
TRADE_COLUMNS = (
    "Trade ID", "Strategy Name", "Configuration Version", "Symbol", "Market",
    "Signal Timeframe", "Exit Resolution Timeframe", "Swing Signal ID", "Swing Type",
    "Trade Direction", "Primary Status", "Nested Status", "Derived Status",
    "Swing Open Time UTC", "Swing Open Time Turkey", "Swing Close Open Time UTC",
    "Swing Close Open Time Turkey", "Signal Time UTC", "Signal Time Turkey",
    "Entry Time UTC", "Entry Time Turkey", "Entry Price", "Stop Reference Type",
    "Stop Reference Time UTC", "Stop Reference Time Turkey", "Stop Reference Price",
    "Peak 2 High", "Dip 2 Low", "Stop Loss Price", "Initial Risk Price",
    "Initial Risk Percent", "Reward Multiple", "Risk Multiple", "Requested Risk To Reward",
    "Take Profit Distance", "Take Profit Price", "Verified Reward To Risk",
    "First Eligible 1m UTC", "First Eligible 1m Turkey", "Exit Time UTC", "Exit Time Turkey",
    "Exit Price", "Exit Reason", "Outcome", "Resolved", "Forward Censored",
    "Forward Censoring Reason", "Intended R", "Realized R", "Gross Percent Return",
    "Minutes To Exit", "Hours To Exit", "1m Bars To Exit", "MFE Price", "MFE R",
    "MAE Price", "MAE R", "Same Minute Both Levels Hit", "Ambiguity Policy", "Gap Type",
    "Gap Fill", "Slippage Beyond Stop", "Overlap Count At Entry", "Overlapping Trade",
    "Detector Version", "Data Range", "Execution Basis", "Unrealized Price Change", "Unrealized R",
)
INVALID_COLUMNS = (
    "Signal ID", "Swing Type", "Trade Direction", "Confirmation Time UTC",
    "Confirmation Time Turkey", "Entry Candidate", "Stop Candidate", "Invalid Reason",
    "Swing Open Row", "Swing Close Row",
)
PRICE_COLUMNS = {22, 26, 27, 28, 29, 30, 35, 36, 42, 54, 56, 69}
PERCENT_COLUMNS = {31, 50}
R_COLUMNS = {37, 48, 49, 55, 57, 70}


def shown(value: Decimal | None) -> Decimal | None:
    if value is None:
        return None
    return value.quantize(EIGHT, rounding=ROUND_HALF_UP)


def revision_from_names(names: list[str]) -> int:
    pattern = re.compile(REVISION_PATTERN, re.IGNORECASE)
    numbers = [int(match.group(1)) for name in names if (match := pattern.match(Path(name).name))]
    return max(numbers) + 1 if numbers else 0


def _iso(moment: datetime | None) -> str | None:
    if moment is None:
        return None
    value = moment.astimezone(timezone.utc)
    if value.microsecond:
        return value.strftime("%Y-%m-%dT%H:%M:%S.") + f"{value.microsecond // 1000:03d}Z"
    return value.strftime("%Y-%m-%dT%H:%M:%SZ")


def _hours(trade: Trade) -> Decimal | None:
    if trade.duration_minutes is None:
        return None
    return Decimal(trade.duration_minutes) / Decimal(60)


def trade_row(trade: Trade, detector_version: str) -> list:
    hours = _hours(trade)
    reference = "Peak 2 High" if trade.direction == "SHORT" else "Dip 2 Low"
    return [
        trade.trade_id, STRATEGY_NAME, CONFIGURATION_VERSION, "BTCUSDT", "USD-M Perpetual",
        "30m", "1m", trade.swing_id, trade.swing_type, trade.direction, "TRUE", "FALSE", "FALSE",
        _iso(trade.swing_open_time), turkey_text(trade.swing_open_time),
        _iso(trade.swing_close_open_time), turkey_text(trade.swing_close_open_time),
        _iso(trade.signal_time), turkey_text(trade.signal_time),
        _iso(trade.entry_time), None if trade.entry_time is None else turkey_text(trade.entry_time),
        shown(trade.entry_price), reference,
        _iso(trade.stop_reference_time), None if trade.stop_reference_time is None else turkey_text(trade.stop_reference_time),
        shown(trade.stop_price), shown(trade.peak_2_high), shown(trade.dip_2_low), shown(trade.stop_price),
        shown(trade.risk_price), shown(trade.risk_percent), shown(Decimal("1.5")), shown(Decimal(1)),
        RISK_REWARD_RATIO, shown(trade.reward_distance), shown(trade.target_price), shown(trade.reward_ratio),
        _iso(trade.first_eligible_time), None if trade.first_eligible_time is None else turkey_text(trade.first_eligible_time),
        _iso(trade.exit_time), None if trade.exit_time is None else turkey_text(trade.exit_time),
        shown(trade.exit_price), trade.exit_reason, trade.outcome,
        "TRUE" if trade.resolved else "FALSE", "TRUE" if trade.censored else "FALSE",
        trade.censoring_reason or None, shown(trade.intended_r), shown(trade.realized_r), shown(trade.gross_percent),
        trade.duration_minutes, shown(hours) if hours is not None else None, trade.bars_to_exit,
        shown(trade.mfe_price), shown(trade.mfe_r), shown(trade.mae_price), shown(trade.mae_r),
        "TRUE" if trade.same_minute_both else "FALSE",
        SAME_BAR_POLICY if trade.same_minute_both else None,
        trade.gap_type or None, "TRUE" if trade.gap_fill else "FALSE",
        "TRUE" if trade.slippage_beyond_stop else "FALSE",
        trade.concurrent_at_entry, "TRUE" if trade.overlapping else "FALSE",
        detector_version, DATA_RANGE, PERFORMANCE_LABEL,
        shown(trade.unrealized_change), shown(trade.unrealized_r),
    ]


def _style_sheet(sheet, rows: list[list], display_name: str) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, row in enumerate(rows, start=2):
        outcome = row[43]
        fill = OPEN_FILL
        if outcome == "TAKE_PROFIT":
            fill = WIN_FILL
        elif row[57] == "TRUE":
            fill = AMBIGUOUS_FILL
        elif row[59] not in (None, "", "FALSE"):
            fill = GAP_FILL
        elif outcome == "STOP_LOSS":
            fill = LOSS_FILL
        for cell in sheet[index]:
            cell.fill = fill
        for column in PRICE_COLUMNS | PERCENT_COLUMNS | R_COLUMNS:
            sheet.cell(index, column).number_format = "0.00000000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TRADE_COLUMNS))}{max(1, sheet.max_row)}"
    sheet.row_dimensions[1].height = 32
    for column in range(1, len(TRADE_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    table = Table(displayName=display_name, ref=f"A1:{get_column_letter(len(TRADE_COLUMNS))}{max(1, sheet.max_row)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)
    if sheet.max_row > 1:
        outcome = get_column_letter(44)
        sheet.conditional_formatting.add(f"{outcome}2:{outcome}{sheet.max_row}", FormulaRule(formula=[f'{outcome}2="TAKE_PROFIT"'], fill=WIN_FILL))
        sheet.conditional_formatting.add(f"{outcome}2:{outcome}{sheet.max_row}", FormulaRule(formula=[f'{outcome}2="STOP_LOSS"'], fill=LOSS_FILL))


def _metrics(summary: dict) -> list[tuple[str, object]]:
    labels = (
        ("Eligible Signals", "eligible_signal_count"),
        ("Valid Trades", "valid_trade_count"),
        ("Invalid Signals", "invalid_signal_count"),
        ("Resolved Trades", "resolved_trade_count"),
        ("Unresolved Trades", "unresolved_trade_count"),
        ("Take Profits", "take_profit_count"),
        ("Stop Losses", "stop_loss_count"),
        ("Resolved Win Rate Percent", "win_rate_percent"),
        ("Gross Break-Even Win Rate Percent", "break_even_win_rate_percent"),
        ("Win Rate Edge Percentage Points", "win_rate_edge_percentage_points"),
        ("Average R", "average_r"),
        ("Median R", "median_r"),
        ("Total Cumulative R", "total_r"),
        ("Maximum Cumulative R Drawdown", "maximum_drawdown_r"),
        ("Profit Factor", "profit_factor"),
        ("Expectancy R", "expectancy_r"),
        ("Ideal Expectancy R", "ideal_expectancy_r"),
        ("Average Minutes To Take Profit", "average_minutes_to_take_profit"),
        ("Median Minutes To Take Profit", "median_minutes_to_take_profit"),
        ("Average Minutes To Stop Loss", "average_minutes_to_stop_loss"),
        ("Median Minutes To Stop Loss", "median_minutes_to_stop_loss"),
        ("Minimum Minutes To Resolution", "minimum_minutes_to_resolution"),
        ("Maximum Minutes To Resolution", "maximum_minutes_to_resolution"),
        ("Same Minute Ambiguous", "same_minute_ambiguous_count"),
        ("Adverse Gap Stops", "adverse_gap_stop_count"),
        ("Favorable Gap Targets", "favorable_gap_target_count"),
    )
    rows = []
    for label, key in labels:
        value = summary[key]
        rows.append((label, None if value is None else (format(value, "f") if isinstance(value, Decimal) else value)))
    return rows


def _add_table(sheet, headers: list[str], rows: list[list], display_name: str, start_row: int = 1) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(start_row, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            sheet.cell(start_row + offset, column, value)
    last = start_row + max(len(rows), 0)
    if not rows:
        last = start_row
    ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    table = Table(displayName=display_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)


def write_workbook(path: Path, trades: list[Trade], summaries: dict[str, dict], months: list[dict], diagnostics: list[tuple[str, object]], detector_version: str) -> None:
    book = Workbook()
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=False)
    book.remove(book.active)
    summary = book.create_sheet("Performance Summary")
    summary["A1"] = "Risk : Reward = 1 : 1.5"
    summary["A1"].font = NOTE_FONT
    summary["A2"] = "Take Profit = +1.5R. Stop Loss = -1.0R. Gross theoretical break-even win rate = 40.0000%."
    summary["A3"] = f"{PERFORMANCE_LABEL}. Historical simulation only. Not investment advice."
    summary["A4"] = STRATEGY_NAME
    summary["A5"] = CONFIGURATION_VERSION
    row_number = 7
    for title, key, table_name in (
        ("ALL TRADES", "ALL TRADES", "SummaryAll"),
        ("SHORT TRADES", "SHORT", "SummaryShort"),
        ("LONG TRADES", "LONG", "SummaryLong"),
    ):
        summary.cell(row_number, 1, title).font = Font(bold=True, size=14, name="Calibri")
        metrics = _metrics(summaries[key])
        _add_table(summary, ["Metric", "Value"], [[label, value] for label, value in metrics], table_name, row_number + 1)
        row_number += len(metrics) + 4
    summary.cell(row_number, 1, "MONTHLY RESULTS").font = Font(bold=True, size=14, name="Calibri")
    month_rows = [[item["month"], item["trades"], item["take_profit_count"], item["stop_loss_count"], item["unresolved_trade_count"], None if item["win_rate_percent"] is None else format(item["win_rate_percent"], "f"), None if item["total_r"] is None else format(item["total_r"], "f")] for item in months]
    _add_table(summary, ["Month", "Trades", "Take Profits", "Stop Losses", "Unresolved", "Win Rate Percent", "Total R"], month_rows, "SummaryMonths", row_number + 1)
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 28
    groups = {
        "All Trades": ("AllTrades", [trade for trade in trades if trade.outcome != "INVALID"]),
        "Short Trades": ("ShortTrades", [trade for trade in trades if trade.direction == "SHORT" and trade.outcome != "INVALID"]),
        "Long Trades": ("LongTrades", [trade for trade in trades if trade.direction == "LONG" and trade.outcome != "INVALID"]),
        "Open and Censored": ("OpenCensored", [trade for trade in trades if trade.outcome == "OPEN_AT_DATASET_END"]),
    }
    for title, (table_name, selected) in groups.items():
        sheet = book.create_sheet(title)
        sheet.append(list(TRADE_COLUMNS))
        rows = [trade_row(trade, detector_version) for trade in selected]
        for row in rows:
            if len(row) != len(TRADE_COLUMNS):
                raise RuntimeError(f"{title} width {len(row)}")
            sheet.append(row)
        _style_sheet(sheet, rows, table_name)
    invalid = book.create_sheet("Invalid Signals")
    invalid.append(list(INVALID_COLUMNS))
    invalid_rows = []
    for trade in trades:
        if trade.outcome != "INVALID":
            continue
        invalid_rows.append([
            trade.swing_id, trade.swing_type, trade.direction, _iso(trade.signal_time), turkey_text(trade.signal_time),
            shown(trade.candidate_entry), shown(trade.candidate_stop), trade.exit_reason, trade.open_row, trade.close_row,
        ])
    for row in invalid_rows:
        invalid.append(row)
    for cell in invalid[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    invalid.freeze_panes = "A2"
    invalid.auto_filter.ref = f"A1:J{max(1, invalid.max_row)}"
    for column in range(1, 11):
        invalid.column_dimensions[get_column_letter(column)].width = 28
    table = Table(displayName="InvalidSignals", ref=f"A1:J{max(1, invalid.max_row)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    invalid.add_table(table)
    parameters = book.create_sheet("Parameters")
    parameters.append(["Parameter", "Value"])
    for key, value in PARAMETERS:
        parameters.append([key, value])
    diagnostics_sheet = book.create_sheet("Diagnostics")
    diagnostics_sheet.append(["Item", "Value"])
    for key, value in diagnostics:
        diagnostics_sheet.append([key, "" if value is None else value])
    readme = book.create_sheet("README")
    readme.append(["Topic", "Explanation"])
    for key, value in (
        ("Experiment", "Independent historical outcome of each confirmed Primary 30m special swing at a fixed gross 1:1.5 reward-to-risk."),
        ("Signals", "The existing detector engine supplies the swings. Excel is not parsed and the detector is not modified."),
        ("Entry", "Entry is the completed Swing Close close. The trade does not exist during that candle."),
        ("Stops", "A swing high is short and stops at Peak 2 High. A swing low is long and stops at Dip 2 Low."),
        ("Target", "The target is exactly 1.5 times the initial stop distance. A normal winner is +1.5R and a normal loser is -1.0R."),
        ("Break-even", "With those fixed outcomes, the gross theoretical break-even win rate is 1 / (1 + 1.5) = 40.0000%, before fees, slippage, and funding."),
        ("Resolution", "Canonical 1-minute candles after confirmation decide which level is touched first."),
        ("Ambiguity", "If one minute touches both levels, the ordering is unknown and the stop is applied."),
        ("Unresolved", "Trades still open at the dataset end stay censored and are excluded from the resolved win rate."),
        ("Overlap", "Every eligible signal is simulated independently. Overlap is measured, not blocked."),
        ("Costs", "Results are gross. They are not a live-trading result and not investment advice."),
        ("Protected packages", "The detector and the existing 1:1 and 1:2 strategy packages were not modified."),
    ):
        readme.append([key, value])
    for sheet, name in ((parameters, "StrategyParameters"), (diagnostics_sheet, "BacktestDiagnostics"), (readme, "BacktestReadme")):
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 110
        last = max(1, sheet.max_row)
        table = Table(displayName=name, ref=f"A1:B{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)
    if book.sheetnames != list(SHEETS):
        raise RuntimeError(f"sheet order {book.sheetnames}")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def validate_workbook(path: Path, trades: list[Trade]) -> dict:
    normal = load_workbook(path, read_only=False, data_only=False)
    readonly = load_workbook(path, read_only=True, data_only=False)
    data_only = load_workbook(path, read_only=True, data_only=True)
    try:
        if normal.sheetnames != list(SHEETS) or readonly.sheetnames != list(SHEETS) or data_only.sheetnames != list(SHEETS):
            raise RuntimeError(f"sheets {normal.sheetnames}")
        if normal.vba_archive is not None or list(normal._external_links):
            raise RuntimeError("workbook contains macros or external links")
        valid = [trade for trade in trades if trade.outcome != "INVALID"]
        expected = {
            "All Trades": valid,
            "Short Trades": [trade for trade in valid if trade.direction == "SHORT"],
            "Long Trades": [trade for trade in valid if trade.direction == "LONG"],
            "Open and Censored": [trade for trade in valid if trade.outcome == "OPEN_AT_DATASET_END"],
            "Invalid Signals": [trade for trade in trades if trade.outcome == "INVALID"],
        }
        counts = {}
        for name, selected in expected.items():
            rows = [row for row in normal[name].iter_rows(min_row=2, values_only=True) if row[0]]
            counts[name] = len(rows)
            if len(rows) != len(selected):
                raise RuntimeError(f"{name} count {len(rows)} != {len(selected)}")
            if not normal[name].tables:
                raise RuntimeError(f"{name} has no table")
        if normal["Performance Summary"]["A1"].value != "Risk : Reward = 1 : 1.5":
            raise RuntimeError("1:1.5 banner is missing")
        parameters = {row[0]: row[1] for row in normal["Parameters"].iter_rows(min_row=2, max_col=2, values_only=True)}
        if parameters["gross_break_even_win_rate"] != "40.0000%" or parameters["reward_multiple"] != "1.5":
            raise RuntimeError("parameters are not 1:1.5")
        if counts["All Trades"] != counts["Short Trades"] + counts["Long Trades"]:
            raise RuntimeError("direction sheet reconciliation failed")
        return {"sheets": list(normal.sheetnames), "rows": counts, "readable": True, "macros": False, "external_links": False}
    finally:
        normal.close()
        readonly.close()
        data_only.close()
